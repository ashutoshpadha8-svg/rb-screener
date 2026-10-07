#!/usr/bin/env python3
"""
ema_screener.py -- EMA 9/21 cross screener, DAILY candles (7 Oct 2026, RB).
SEPARATE from the master scan: info only, no orders, nothing in rb/rbtrack.
NOT backtested in this form (ema_backtest.py 29 Sep: the plain EMA 9/21
cross lost to Nifty in 2013-19) -> look at the charts, SL/exit decided later.

Rule (RB, all on COMPLETED daily closes, never the intraday price):
  1. EMA9 crossed above EMA21: cross day = first day EMA9 > EMA21 after a
     day with EMA9 <= EMA21.
  2. Close above EMA100 AND EMA200.
  3. "Cross ke baad 3-4 din stock upside rehna chahiye": at least HOLD_DAYS
     sessions have passed since the cross day, and on EVERY day from the
     cross to the last close:
       - EMA9 stayed above EMA21 (no cross back down),
       - close stayed above EMA100 and EMA200,
       - close never fell below the cross-day close,
     and the last close is above the cross-day close.
  4. Cross not older than MAX_AGE sessions (older = old news).
Crosses from the last 0..HOLD_DAYS-1 sessions that pass so far go to the
'EMA_Waiting' sheet (NOT in the watchlist) -- they confirm in a day or two.

Universe + data = daily_screener / momentum_screener (>= Rs 10,000 Cr, free
history + NSE / broker gap fill). LTP column = live price, display only.

Output (own folder ~/RB_Screener/EMA/, RB 7 Oct; old files in reports/
are moved there once):
  EMA/EMA_Screener_YYYY-MM-DD.xlsx  (EMA_Confirmed, EMA_Waiting)
  EMA/Watchlist_EMA.txt             TradingView 'Import list': ONLY the
                                    confirmed stocks (overwritten daily);
                                    section 1 = also in another screener
Other screeners (RB 7 Oct): a stock that is ALSO in today's master scan
(Momentum_Top20 / Swing / Investing BUY-FIT-LATE) gets a gold row and the
'Doosre screener mein' column, e.g. 'MOM #3 + SWING FIT' (+ SUPER-BUY).
Runs by itself inside `rb` (after the portfolio, once per completed
session; a failure never stops rb). By hand:
      python3 ema_screener.py            (skips if today's file exists)
      python3 ema_screener.py --force    (run again)
      python3 ema_screener.py --hold 4   (4 sessions up after the cross)
"""

import argparse
import os
import sys

import numpy as np
import pandas as pd

import daily_screener as ds

FAST, SLOW, MID, LONG = 9, 21, 100, 200
HOLD_DAYS = 3        # sessions after the cross the stock must stay up
MAX_AGE = 10         # cross at most this many sessions ago
EMA_DIR = os.path.join(ds.HERE, "EMA")
WATCH_FILE = os.path.join(EMA_DIR, "Watchlist_EMA.txt")
OTHER = "Doosre screener mein"
GOLD = "F7E3A5"


def move_old(src=ds.REPORTS, dst=EMA_DIR):
    """One-time: EMA files written into reports/ by v22/v23 -> EMA/."""
    import glob
    import shutil
    os.makedirs(dst, exist_ok=True)
    for f in glob.glob(os.path.join(src, "EMA_Screener_*.xlsx")) + \
            glob.glob(os.path.join(src, "Watchlist_EMA.txt")):
        t = os.path.join(dst, os.path.basename(f))
        try:
            if os.path.exists(t):
                os.remove(f)
            else:
                shutil.move(f, t)
        except OSError:
            pass


def other_lists(path=None):
    """{SYMBOL: 'MOM #3 + SWING FIT + INV BUY + SUPER-BUY'} from today's
    master scan (same sheets and labels as All_Picks)."""
    try:
        import momentum_screener as ms
        import signal_tracker as st
        path = path or ms.todays_report()
        if not path:
            return {}
        sw, iv, mo = (st._sheet(path, n) for n in ("Swing", "Investing",
                                                    "Momentum_Top20"))
    except Exception:
        return {}
    out = {}

    def add(sym, txt):
        out.setdefault(sym, []).append(txt)
    if "Mom Rank" in mo:
        for sym, rk in zip(mo["Symbol"], mo["Mom Rank"]):
            add(sym, "MOM #%d" % int(float(rk)))
    for d, lab in ((sw, "SWING"), (iv, "INV")):
        if "Action" in d:
            for sym, act in zip(d["Symbol"], d["Action"]):
                add(sym, "%s %s" % (lab, str(act).strip().upper()))
    res = {}
    for sym, xs in out.items():
        mom = any(x.startswith("MOM") for x in xs)
        wtt = any(x.split()[0] in ("SWING", "INV") and not x.endswith("LATE")
                  for x in xs)
        res[sym] = " + ".join(list(dict.fromkeys(xs)) +
                              (["SUPER-BUY"] if mom and wtt else []))
    return res


def ema(s, n):
    return s.ewm(span=n, adjust=False).mean()


def check_stock(close, hold=HOLD_DAYS, max_age=MAX_AGE):
    """One stock's daily closes (oldest first, completed bars only).
    Returns a dict for a cross that passes so far, else None.
    'status' = CONFIRMED (>= hold sessions up) or WAITING."""
    c = pd.Series(close, dtype=float).dropna()
    if len(c) < LONG + max_age + 5:
        return None
    e9, e21, e100, e200 = (ema(c, n) for n in (FAST, SLOW, MID, LONG))
    above = (e9 > e21).values
    n = len(c)
    cross = None
    for i in range(n - 1, max(n - 2 - max_age, 0), -1):
        if above[i] and not above[i - 1]:
            cross = i
            break
        if not above[i]:
            return None              # EMA9 below EMA21 now / after the cross
    if cross is None:
        return None
    seg = slice(cross, n)
    cv = c.values
    if not (above[seg].all() and (cv[seg] > e100.values[seg]).all() and
            (cv[seg] > e200.values[seg]).all()):
        return None
    base = cv[cross]
    age = n - 1 - cross
    if age > 0 and ((cv[cross + 1:] < base).any() or cv[-1] <= base):
        return None
    return {"cross_date": c.index[cross].date(), "age": age,
            "status": "CONFIRMED" if age >= hold else "WAITING",
            "cross_close": base, "close": cv[-1],
            "low_since": cv[cross:].min(),
            "e9": e9.iloc[-1], "e21": e21.iloc[-1], "e100": e100.iloc[-1],
            "e200": e200.iloc[-1], "date": c.index[-1].date()}


def scan(frames, want=None, hold=HOLD_DAYS, max_age=MAX_AGE):
    """frames {sym: OHLCV df}. Only bars up to `want` (last COMPLETED
    session) are used. Returns a DataFrame, one row per passing stock."""
    rows = []
    for s, f in frames.items():
        c = f["Close"]
        if want is not None:
            c = c[c.index.date <= want]
        r = check_stock(c, hold, max_age)
        if r:
            r["symbol"] = s.upper()
            rows.append(r)
    return pd.DataFrame(rows)


def momentum_ranks():
    try:
        import momentum_screener as ms
        r = pd.read_csv(ms.RANKS_FILE)
        return dict(zip(r["symbol"].astype(str).str.upper(), r["rank"]))
    except Exception:
        return {}


def table(res, live, ranks, other=None):
    if res.empty:
        return pd.DataFrame()
    other = other or {}
    res = res.copy()
    res["_o"] = [0 if s in other else 1 for s in res["symbol"]]
    res = res.sort_values(["_o", "age", "symbol"])
    out = pd.DataFrame({
        "Symbol": res["symbol"].values,
        OTHER: [other.get(s, "") for s in res["symbol"]],
        "Cross Date": [d.strftime("%d-%m-%Y") for d in res["cross_date"]],
        "Days Since Cross": res["age"].values,
        "Cross Close": res["cross_close"].round(2).values,
        "Last Close": res["close"].round(2).values,
        "Since Cross %": ((res["close"] / res["cross_close"] - 1) * 100)
        .round(2).values,
        "LTP (live)": [live.get(s, np.nan) for s in res["symbol"]],
        "EMA9": res["e9"].round(2).values,
        "EMA21": res["e21"].round(2).values,
        "EMA100": res["e100"].round(2).values,
        "EMA200": res["e200"].round(2).values,
        "% above EMA200": ((res["close"] / res["e200"] - 1) * 100)
        .round(2).values,
        "Mom Rank": [ranks.get(s, np.nan) for s in res["symbol"]]})
    return out


def write_book(path, conf, wait, banner):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for name, df, note in (
            ("EMA_Confirmed", conf, "In Watchlist_EMA.txt"),
            ("EMA_Waiting", wait, "Cross < %d sessions ago - NOT in the "
             "watchlist yet" % HOLD_DAYS)):
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=banner).font = Font(bold=True)
        ws.cell(row=2, column=1, value=note + "  |  GOLD row = also in "
                "another screener today (see '%s')" % OTHER)
        if df.empty:
            ws.cell(row=4, column=1, value="No stock today.")
            continue
        for j, h in enumerate(df.columns, 1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        hit = PatternFill("solid", fgColor=GOLD)
        for i, row in enumerate(df.itertuples(index=False), 5):
            both = OTHER in df.columns and bool(row[list(df.columns)
                                                    .index(OTHER)])
            for j, v in enumerate(row, 1):
                if isinstance(v, float) and v != v:
                    v = None
                c = ws.cell(row=i, column=j, value=v)
                if both:
                    c.fill = hit
                    if j <= 2:
                        c.font = Font(bold=True, color="79601E")
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = "A4:%s%d" % (ws.cell(row=4, column=len(
            df.columns)).column_letter, 4 + len(df))
    try:
        import cap_class
        cap_class.add_columns(wb)
    except Exception as e:
        print("  ! cap columns skipped (%s)" % type(e).__name__)
    try:
        import xl_fit
        xl_fit.fit_workbook(wb, skip=())
    except Exception:
        pass
    wb.save(path)


def write_watchlist(symbols, path=WATCH_FILE, other=None):
    """TradingView 'Import list' -- ONLY the found stocks; the ones also in
    another screener first, in their own section."""
    syms = list(dict.fromkeys(s for s in symbols if s))
    other = other or {}
    secs = [("EMA + other screener", [s for s in syms if s in other]),
            ("EMA 9-21 Cross", [s for s in syms if s not in other])]
    lines = ["###%s,%s" % (t, ",".join("NSE:" + s for s in xs))
             for t, xs in secs if xs]
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w") as f:
        f.write(",".join(lines) + "\n" if lines else "")
    return path


def main():
    global HOLD_DAYS
    ap = argparse.ArgumentParser(description="EMA 9/21 cross screener")
    ap.add_argument("--hold", type=int, default=HOLD_DAYS,
                    help="sessions the stock must stay up after the cross")
    ap.add_argument("--max-age", type=int, default=MAX_AGE)
    ap.add_argument("--force", action="store_true",
                    help="run even if this session's file exists")
    a = ap.parse_args()
    HOLD_DAYS = max(1, a.hold)
    move_old()
    done = os.path.join(EMA_DIR, "EMA_Screener_%s.xlsx"
                        % ds.last_expected_session())
    if os.path.exists(done) and not a.force and a.hold == ap.get_default("hold") and \
            a.max_age == MAX_AGE:
        print("EMA screener already done for this session: %s\n"
              "Watchlist: %s  (again: --force)" % (done, WATCH_FILE))
        return 0
    import momentum_screener as ms
    sess = ms.get_session()
    frames, bm, caps, live, warns, want = ms.load_data(sess)
    res = scan(frames, want, HOLD_DAYS, a.max_age)
    live = {s.upper(): v for s, v in live.items()}
    ranks = momentum_ranks()
    last = max((f.index[-1].date() for f in frames.values()), default=want)
    last = min(last, want) if want else last
    other = other_lists()
    if res.empty:
        conf = wait = pd.DataFrame()
    else:
        conf = table(res[res["status"] == "CONFIRMED"], live, ranks, other)
        wait = table(res[res["status"] == "WAITING"], live, ranks, other)
    os.makedirs(EMA_DIR, exist_ok=True)
    path = os.path.join(EMA_DIR, "EMA_Screener_%s.xlsx" % last)
    banner = ("EMA 9/21 cross (daily) + close > EMA100 & EMA200 + up %d+ "
              "sessions after the cross | last close %s | %d stocks scanned"
              " | NOT backtested - info only" % (HOLD_DAYS, last, len(frames)))
    write_book(path, conf, wait, banner)
    syms = list(conf["Symbol"]) if len(conf) else []
    write_watchlist(syms, other=other)
    try:
        import drive_copy
        drive_copy.push(WATCH_FILE, quiet=True)
        drive_copy.push(path, quiet=True)
    except Exception:
        pass
    for w in warns[:5]:
        print("  ! %s" % w)
    print("\nEMA 9/21 CROSS -- last completed close %s" % last)
    print("CONFIRMED (cross %d-%d sessions ago, stayed up): %d" % (
        HOLD_DAYS, a.max_age, len(conf)))
    if len(conf):
        print(conf[["Symbol", "Cross Date", "Days Since Cross",
                    "Since Cross %", OTHER]].to_string(index=False))
    print("WAITING (fresh cross, not %d days yet): %d%s" % (
        HOLD_DAYS, len(wait), ("  " + ", ".join(wait["Symbol"]))
        if len(wait) else ""))
    both = [s for s in syms if s in other]
    print("ALSO IN ANOTHER SCREENER (gold rows): %s" % (
        ", ".join("%s (%s)" % (s, other[s]) for s in both) or "none"))
    print("\nExcel: %s" % path)
    print("TradingView watchlist (ONLY the %d confirmed): %s" % (
        len(syms), WATCH_FILE))


if __name__ == "__main__":
    sys.exit(main())
