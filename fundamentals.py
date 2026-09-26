#!/usr/bin/env python3
"""
FUNDAMENTAL CHECK FOR THE SCREENER SHORTLIST  (NSE)   --  v1
=============================================================

Reads today's shortlist from reports/RB_Screener_YYYY-MM-DD.xlsx (made by
daily_screener.py), fetches each stock's public Screener.in company page
(free, no login) and applies the checklists in FUNDAMENTALS.md section 9.

OUTPUT  written INTO the same reports/RB_Screener_YYYY-MM-DD.xlsx
  Swing / Investing : new columns on the right of every stock (Fund Check,
                      P/E, ROCE, ROE, D/E, growth, pledge ...). No row is
                      removed or re-ordered.
  Fundamentals      : new sheet -- every non-LATE shortlisted stock, sorted
                      by RS rank, with the full fundamental detail.
  Last 3 columns everywhere: Promoter / FII / DII holding change
  (percentage points, latest quarter vs previous). INFO ONLY.
  Nothing here removes a stock: the 2018-26 backtest showed the
  fundamental gate did not help (see IMPORTANT).

FUND CHECK (information only)
  PASS  : every automated rule passed
  FAIL  : at least one rule failed (see FAILED_RULES)
  CHECK : nothing failed, but some rule could not be computed
          (missing data, loss base, bank CRAR, insurer solvency ...)
  Missing data is never a silent PASS.

IMPORTANT
  * BACKTESTED (backtest.py --fundamentals, 2018-2026, point-in-time):
    the swing gate did NOT help -- in 2018-21 PASS stocks did much worse
    than FAIL stocks, in 2022-26 no difference. Treat every column here
    as information.
  * Screener.in data is restated and has no result dates: fine for live
    screening, useless for backtests.
  * Pledge %, auditor issues and SEBI action are NOT on the free page.
    Pledge is caught only if Screener's own "Cons" list mentions it.
    Check those by hand before buying.
  * Screener.in's terms discourage scraping. This fetches only the
    shortlist (tens of pages), slowly, and caches each page for the day.

RUN (after rbscan)
    python3 ~/Desktop/RB_Screener/fundamentals.py
  options:
    --file PATH        use a specific RB_Screener_*.xlsx
    --symbols A,B,C    just check these symbols (writes a separate
                       RB_Fundamentals_<date>.xlsx, no screener report needed)
  Run it again on the same day: it replaces its own columns/sheet.
  If the report is open in Excel/Numbers, close it first.
"""

import warnings
warnings.filterwarnings("ignore")

import os
import re
import sys
import glob
import html
import time
import datetime as dt
from urllib.parse import quote

import numpy as np
import pandas as pd
import requests

# ------------------------------------------------------------------ config
HERE = os.path.join(os.path.expanduser("~"), "Desktop", "RB_Screener")
DATA = os.path.join(HERE, "data")
REPORTS = os.path.join(HERE, "reports")
TAG = ""               # broker in file names, e.g. "DHAN" (set by account.py)
CACHE = os.path.join(DATA, "screener_pages")
for _d in (HERE, DATA, REPORTS, CACHE):
    os.makedirs(_d, exist_ok=True)

SCREENER = "https://www.screener.in/company/%s/%s"
HDRS = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/124.0 Safari/537.36",
        "Accept-Language": "en-US,en;q=0.9"}
PAUSE = 2.5            # seconds between page requests -- be polite

# thresholds (FUNDAMENTALS.md section 9 -- NOT optimised, NOT backtested)
SW_QPROFIT = 20.0      # swing: latest qtr profit YoY >= %
SW_QSALES = 15.0       # swing: latest qtr sales YoY >= %
SW_ROE = 10.0          # swing: 3-yr avg ROE >= %
SW_DE = 1.5            # swing: debt/equity <= (non-financials)
PLEDGE_MAX = 20.0      # both: promoter pledge < %
IN_ROE = 15.0          # investing: 3-yr avg ROE >= %
IN_ROCE = 15.0         # investing: 3-yr avg ROCE >= % (non-financials)
IN_DE = 1.0            # investing: debt/equity <= (non-financials)
IN_ICR = 3.0           # investing: interest coverage >= x
IN_CFO_PAT = 0.7       # investing: 5-yr CFO / 5-yr PAT >= (non-financials)
IN_LOSS_YEARS = 6      # investing: no loss year in last N fiscal years
IN_PROFIT_CAGR = 15.0  # investing: 3-yr profit CAGR >= %
IN_SALES_CAGR = 12.0   # investing: 3-yr sales CAGR >= %
BANK_NNPA = 2.0        # banks: net NPA < %
FIN_GNPA = 3.0         # banks / NBFCs: gross NPA < %
BANK_ROA = 1.0         # banks: ROA >= %
NBFC_ROA = 2.0         # NBFCs: ROA >= %



# ================================================================== fetch
def fetch_page(sym, standalone=False):
    """Screener.in company page HTML, cached for the day. None if missing."""
    kind = "standalone" if standalone else "consolidated"
    day = dt.date.today().isoformat()
    path = os.path.join(CACHE, "%s_%s_%s.html"
                        % (re.sub(r"[^A-Z0-9-]", "_", sym), kind, day))
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return f.read()
    url = SCREENER % (quote(sym, safe="-"),
                      "" if standalone else "consolidated/")
    for attempt in range(3):
        try:
            r = requests.get(url, headers=HDRS, timeout=30)
        except requests.RequestException:
            time.sleep(5)
            continue
        finally:
            time.sleep(PAUSE)
        if r.status_code == 404:
            return None
        if r.status_code == 429:           # rate limited -- back off
            print("\n  Screener says slow down, waiting 60s ...")
            time.sleep(60)
            continue
        if r.ok:
            with open(path, "w", encoding="utf-8") as f:
                f.write(r.text)
            return r.text
    return None


def clean_old_cache():
    today = dt.date.today().isoformat()
    for p in glob.glob(os.path.join(CACHE, "*.html")):
        if not p.endswith(today + ".html"):
            try:
                os.remove(p)
            except OSError:
                pass


# ================================================================== parse
def _txt(x):
    x = re.sub(r"<[^>]+>", " ", x)
    x = html.unescape(x).replace("\xa0", " ")
    return re.sub(r"\s+", " ", x).strip(" +")


def _num(x):
    if x is None:
        return np.nan
    x = str(x).replace(",", "").replace("%", "").replace("−", "-").strip()
    try:
        return float(x)
    except ValueError:
        return np.nan


def _table(h, anchor):
    """First table after the section anchor -> DataFrame (rows = line items,
    columns = period labels), numbers as floats."""
    i = h.find(anchor)
    if i < 0:
        return pd.DataFrame()
    j = h.find("<table", i)
    k = h.find("</table>", j)
    if j < 0 or k < 0:
        return pd.DataFrame()
    rows = [[_txt(c) for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, re.S)]
            for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", h[j:k], re.S)]
    if len(rows) < 2 or len(rows[0]) < 2:
        return pd.DataFrame()
    # "Mar 2024 15m" (changed fiscal year) -> "Mar 2024" so tables line up
    head = [re.sub(r"\s+\d+m$", "", h) for h in rows[0][1:]]
    data = {}
    for r in rows[1:]:
        if r and r[0] and r[0] not in data and len(r) == len(head) + 1:
            data[r[0]] = [_num(v) for v in r[1:]]
    if not data:
        return pd.DataFrame()
    return pd.DataFrame(data, index=head).T


def parse_page(h):
    p = {"q": _table(h, 'id="quarters"'),
         "pl": _table(h, 'id="profit-loss"'),
         "bs": _table(h, 'id="balance-sheet"'),
         "cf": _table(h, 'id="cash-flow"'),
         "ratios": _table(h, 'id="ratios"'),
         "shp": _table(h, 'id="quarterly-shp"')}
    i = h.find('id="top-ratios"')
    top = {}
    if i >= 0:
        k = h.find("</ul>", i)
        for n, v in re.findall(r'<span class="name">(.*?)</span>\s*'
                               r'<span class="nowrap value">(.*?)</span>\s*</li>',
                               h[i:k], re.S):
            top[_txt(n)] = _txt(v)
    p["price"] = _num(re.sub(r"[^\d.,-]", "", top.get("Current Price", "")))
    p["top_roe"] = _num(re.sub(r"[^\d.,-]", "", top.get("ROE", "")))
    p["top_roce"] = _num(re.sub(r"[^\d.,-]", "", top.get("ROCE", "")))
    p["pe"] = _num(re.sub(r"[^\d.,-]", "", top.get("Stock P/E", "")))
    ind = dict((t, _txt(n)) for t, n in re.findall(
        r'title="(Broad Sector|Sector|Broad Industry|Industry)">(.*?)</a>', h))
    p["sector"] = ind.get("Broad Sector", "")
    p["broad_ind"] = ind.get("Broad Industry", "")
    p["industry"] = ind.get("Industry", "")
    cons = ""
    i = h.find('<div class="cons"')
    if i >= 0:
        k = h.find("</ul>", i)
        cons = [_txt(x) for x in re.findall(r"<li>(.*?)</li>", h[i:k], re.S)]
        cons = [c for c in cons if c]
    p["cons"] = cons or []
    return p


def n_years(p):
    pl = p["pl"]
    return 0 if pl.empty else len([c for c in pl.columns if c != "TTM"])


def load_company(sym):
    """Consolidated figures, falling back to standalone when the
    consolidated history is missing or short."""
    h = fetch_page(sym)
    p = parse_page(h) if h else None
    basis = "consolidated"
    if p is None or n_years(p) < 6 or p["q"].shape[1] < 8:
        h2 = fetch_page(sym, standalone=True)
        p2 = parse_page(h2) if h2 else None
        if p2 is not None and (p is None or n_years(p2) > n_years(p)):
            p, basis = p2, "standalone"
    if p is None:
        return None
    p["basis"] = basis
    # banks / NBFCs: Screener shows NPA % on the standalone page only
    p["npa_q"] = p["q"]
    if basis == "consolidated" and fin_type(p) in ("bank", "nbfc"):
        g = _row(p["q"], "Gross NPA %")
        if g is None or g.isna().all():
            h2 = fetch_page(sym, standalone=True)
            if h2:
                p["npa_q"] = parse_page(h2)["q"]
    return p


# ================================================================== metrics
def _row(df, *names):
    for n in names:
        if n in df.index:
            return df.loc[n]
    return None


def _yoy(s, lag_label):
    """% growth of the last value in s vs the same period one year back.
    None when the base is <= 0 (loss / low base -> not meaningful)."""
    if s is None:
        return None
    s = s.dropna()
    if len(s) < 5:
        return None
    last = s.index[-1]
    prev = lag_label(last)
    if prev not in s.index:
        return None
    b, a = s[prev], s.iloc[-1]
    if not b > 0:
        return None
    return (a / b - 1) * 100


def _year_back(label):
    m = re.match(r"([A-Za-z]{3}) (\d{4})$", str(label))
    return "%s %d" % (m.group(1), int(m.group(2)) - 1) if m else None


def _yoy_series(s, n):
    """YoY % for each of the last n quarters (oldest first)."""
    out = []
    if s is None:
        return out
    s = s.dropna()
    for k in range(n, 0, -1):
        if len(s) < k:
            continue
        sub = s.iloc[:len(s) - k + 1]
        out.append(_yoy(sub, _year_back))
    return out


def _cagr(a, b, years):
    if not (a > 0 and b > 0):
        return None
    return ((a / b) ** (1.0 / years) - 1) * 100


def fin_type(p):
    if p["sector"] != "Financial Services":
        return ""
    bi = p["broad_ind"].lower()
    if "bank" in bi:
        return "bank"
    if "insurance" in bi:
        return "insurance"
    if bi == "finance" or "nbfc" in p["industry"].lower() \
            or "housing finance" in p["industry"].lower():
        return "nbfc"
    return "fin_other"


def metrics(p):
    q, pl, bs, cf, ra, shp = p["q"], p["pl"], p["bs"], p["cf"], p["ratios"], p["shp"]
    m = {"basis": p["basis"], "industry": p["industry"] or p["broad_ind"],
         "fin": fin_type(p), "price": p["price"], "cons": p["cons"],
         "pe": p["pe"] if p["pe"] == p["pe"] else None,
         "roce_now": p["top_roce"] if p["top_roce"] == p["top_roce"] else None,
         "roe_now": p["top_roe"] if p["top_roe"] == p["top_roe"] else None}
    fy = [c for c in pl.columns if c != "TTM"] if not pl.empty else []

    # ---- quarterly
    qsales = _row(q, "Sales", "Revenue")
    qnp = _row(q, "Net Profit")
    qpbt = _row(q, "Profit before tax")
    qoi = _row(q, "Other Income")
    qeps = _row(q, "EPS in Rs")
    m["q_label"] = qnp.dropna().index[-1] if qnp is not None and \
        len(qnp.dropna()) else ""
    m["q_sales_yoy"] = _yoy(qsales, _year_back)
    m["q_np_yoy"] = _yoy(qnp, _year_back)
    core = None
    if qpbt is not None and qoi is not None:
        core = qpbt - qoi.fillna(0)
    # financials: other income is core business (fees) -> use net profit
    m["q_core_yoy"] = m["q_np_yoy"] if m["fin"] else _yoy(core, _year_back)
    m["q_profit_yoy"] = m["q_core_yoy"]
    ser = _yoy_series(core if not m["fin"] else qnp, 3)
    m["q_profit_yoy_3"] = ser
    m["decel2"] = (len(ser) == 3 and None not in ser
                   and ser[2] < ser[1] < ser[0])
    m["accel"] = (len(ser) == 3 and None not in ser
                  and ser[2] > ser[1] > ser[0])
    ser_s = _yoy_series(qsales, 2)
    if len(ser) >= 2 and None not in ser[-2:] \
            and len(ser_s) == 2 and None not in ser_s:
        m["q_last2_pos"] = min(ser[-2:]) > 0 and min(ser_s) > 0
    else:
        m["q_last2_pos"] = None
    m["sue"] = None
    if qeps is not None and p["price"] > 0:
        e = qeps.dropna()
        if len(e) >= 5 and _year_back(e.index[-1]) in e.index:
            m["sue"] = (e.iloc[-1] - e[_year_back(e.index[-1])]) / p["price"] * 100
    oi_share = None
    if qpbt is not None and qoi is not None and len(qpbt.dropna()):
        pb, oi = qpbt.dropna().iloc[-1], qoi.reindex(qpbt.dropna().index).iloc[-1]
        if pb > 0 and oi == oi:
            oi_share = oi / pb * 100
    m["oi_share"] = oi_share
    npa = p.get("npa_q", q)
    m["gnpa"] = _row(npa, "Gross NPA %")
    m["nnpa"] = _row(npa, "Net NPA %")
    m["gnpa"] = m["gnpa"].dropna().iloc[-1] if m["gnpa"] is not None and \
        len(m["gnpa"].dropna()) else None
    m["nnpa"] = m["nnpa"].dropna().iloc[-1] if m["nnpa"] is not None and \
        len(m["nnpa"].dropna()) else None

    # ---- annual
    np_a = _row(pl, "Net Profit")
    np_a = np_a[fy].dropna() if np_a is not None else pd.Series(dtype=float)
    sales_a = _row(pl, "Sales", "Revenue")
    sales_a = sales_a[fy].dropna() if sales_a is not None else pd.Series(dtype=float)
    eps_a = _row(pl, "EPS in Rs")
    eps_a = eps_a[fy].dropna() if eps_a is not None else pd.Series(dtype=float)
    last6 = np_a.iloc[-IN_LOSS_YEARS:]
    m["loss_years"] = int((last6 < 0).sum()) if len(last6) >= IN_LOSS_YEARS \
        else (int((last6 < 0).sum()) if (last6 < 0).any() else None)
    m["profit_cagr3"] = _cagr(np_a.iloc[-1], np_a.iloc[-4], 3) \
        if len(np_a) >= 4 else None
    m["sales_cagr3"] = _cagr(sales_a.iloc[-1], sales_a.iloc[-4], 3) \
        if len(sales_a) >= 4 else None
    e = eps_a.iloc[-6:]
    g = [(e.iloc[i] / e.iloc[i - 1] - 1) for i in range(1, len(e))
         if e.iloc[i - 1] > 0]
    m["eps_var"] = float(np.std(g)) * 100 if len(g) >= 3 else None

    # ---- balance sheet: equity, debt, ROE, ROA
    eq = None
    if not bs.empty:
        ec, rs = _row(bs, "Equity Capital"), _row(bs, "Reserves")
        if ec is not None and rs is not None:
            eq = (ec + rs).dropna()
    debt = _row(bs, "Borrowings", "Borrowing") if not bs.empty else None
    m["de"] = None
    if eq is not None and debt is not None and len(eq) and eq.iloc[-1] > 0:
        m["de"] = debt.reindex(eq.index).fillna(0).iloc[-1] / eq.iloc[-1]
    roes = []
    if eq is not None and len(np_a):
        for y in list(np_a.index)[-3:]:
            if y in eq.index:
                k = list(eq.index).index(y)
                base = eq.iloc[k] if k == 0 else (eq.iloc[k] + eq.iloc[k - 1]) / 2
                if base > 0:
                    roes.append(np_a[y] / base * 100)
    if len(roes) < 3 and not ra.empty and "ROE %" in ra.index:
        r = ra.loc["ROE %"].dropna()
        roes = list(r.iloc[-3:]) if len(r) >= 3 else roes
    m["roe3"] = float(np.mean(roes)) if len(roes) == 3 else None
    if m["roe3"] is None and p["top_roe"] == p["top_roe"]:
        m["roe3_note"] = "latest ROE only"
    m["roce3"] = None
    if not ra.empty and "ROCE %" in ra.index:
        r = ra.loc["ROCE %"].dropna()
        if len(r) >= 3:
            m["roce3"] = float(r.iloc[-3:].mean())
    m["roa"] = None
    ta = _row(bs, "Total Assets") if not bs.empty else None
    if ta is not None and len(np_a):
        ta = ta.dropna()
        y = np_a.index[-1]
        if y in ta.index:
            k = list(ta.index).index(y)
            base = ta.iloc[k] if k == 0 else (ta.iloc[k] + ta.iloc[k - 1]) / 2
            if base > 0:
                m["roa"] = np_a[y] / base * 100

    # ---- interest cover, cash quality
    m["icr"] = None
    pbt_a, int_a = _row(pl, "Profit before tax"), _row(pl, "Interest")
    if pbt_a is not None and int_a is not None and fy:
        pb, it = pbt_a[fy[-1]], int_a[fy[-1]]
        if pb == pb and it == it:
            m["icr"] = 99.0 if it <= 0.5 else (pb + it) / it
    m["cfo_pat"] = None
    cfo = _row(cf, "Cash from Operating Activity") if not cf.empty else None
    if cfo is not None and len(np_a):
        yrs = [y for y in list(np_a.index)[-5:] if y in cfo.dropna().index]
        if len(yrs) >= 3 and np_a[yrs].sum() > 0:
            m["cfo_pat"] = cfo[yrs].sum() / np_a[yrs].sum()

    # ---- pledge (only if Screener's cons mention it)
    m["pledge"] = None
    for c in p["cons"]:
        mm = re.search(r"pledged\s+([\d.]+)\s*%", c, re.I)
        if mm:
            m["pledge"] = float(mm.group(1))

    # ---- holdings: change in percentage points, latest vs previous quarter
    m["shp_q"] = ""
    for key, names in (("d_prom", ("Promoters",)), ("d_fii", ("FIIs",)),
                       ("d_dii", ("DIIs",))):
        r = _row(shp, *names) if not shp.empty else None
        r = r.dropna() if r is not None else None
        if r is not None and len(r) >= 2:
            m[key] = round(r.iloc[-1] - r.iloc[-2], 2)
            m["shp_q"] = "%s vs %s" % (r.index[-1], r.index[-2])
        else:
            m[key] = "n/a" if r is None or not len(r) else None

    # D/E, interest cover, CFO/PAT, ROCE mean nothing for lenders/insurers
    if m["fin"]:
        m["de"] = m["icr"] = m["cfo_pat"] = m["roce3"] = m["roce_now"] = None
    return m


# ================================================================== rules
def _rule(out, ok, fail_text, missing_text):
    """ok: True pass / False fail / None could not compute."""
    if ok is None:
        out["check"].append(missing_text)
    elif not ok:
        out["fail"].append(fail_text)


def _f(x, fmt="%.1f"):
    return "n/a" if x is None else fmt % x


def swing_rules(m):
    o = {"fail": [], "check": [], "warn": []}
    if m["pledge"] is not None:
        _rule(o, m["pledge"] < PLEDGE_MAX, "pledge %.0f%%" % m["pledge"], "")
    if m["fin"] == "insurance":
        o["check"].append("insurer: qtr profit/premium lumpy, check VNB/solvency")
    else:
        _rule(o, None if m["q_profit_yoy"] is None
              else m["q_profit_yoy"] >= SW_QPROFIT,
              "qtr profit YoY %s%% < %.0f%%" % (_f(m["q_profit_yoy"]), SW_QPROFIT),
              "qtr profit YoY n/a (loss/low base)")
        _rule(o, None if m["q_sales_yoy"] is None
              else m["q_sales_yoy"] >= SW_QSALES,
              "qtr sales YoY %s%% < %.0f%%" % (_f(m["q_sales_yoy"]), SW_QSALES),
              "qtr sales YoY n/a")
    _rule(o, None if m["roe3"] is None else m["roe3"] >= SW_ROE,
          "3y ROE %s%% < %.0f%%" % (_f(m["roe3"]), SW_ROE), "3y ROE n/a")
    if not m["fin"]:
        _rule(o, None if m["de"] is None else m["de"] <= SW_DE,
              "D/E %s > %.1f" % (_f(m["de"], "%.2f"), SW_DE), "D/E n/a")
    if m["decel2"]:
        o["warn"].append("profit growth slowed 2 qtrs")
    if m["oi_share"] is not None and m["oi_share"] > 25 and not m["fin"]:
        o["warn"].append("other income %.0f%% of PBT" % m["oi_share"])
    return o


def invest_rules(m):
    o = {"fail": [], "check": [], "warn": []}
    if m["pledge"] is not None:
        _rule(o, m["pledge"] < PLEDGE_MAX, "pledge %.0f%%" % m["pledge"], "")
        if 5 <= m["pledge"] < PLEDGE_MAX:
            o["warn"].append("pledge %.0f%% (ideal < 5%%)" % m["pledge"])
    _rule(o, None if m["roe3"] is None else m["roe3"] >= IN_ROE,
          "3y ROE %s%% < %.0f%%" % (_f(m["roe3"]), IN_ROE), "3y ROE n/a")
    if not m["fin"]:
        _rule(o, None if m["roce3"] is None else m["roce3"] >= IN_ROCE,
              "3y ROCE %s%% < %.0f%%" % (_f(m["roce3"]), IN_ROCE), "3y ROCE n/a")
        _rule(o, None if m["de"] is None else m["de"] <= IN_DE,
              "D/E %s > %.1f" % (_f(m["de"], "%.2f"), IN_DE), "D/E n/a")
        _rule(o, None if m["icr"] is None else m["icr"] >= IN_ICR,
              "interest cover %sx < %.0fx" % (_f(m["icr"], "%.1f"), IN_ICR),
              "interest cover n/a")
        _rule(o, None if m["cfo_pat"] is None else m["cfo_pat"] >= IN_CFO_PAT,
              "CFO/PAT %s < %.1f" % (_f(m["cfo_pat"], "%.2f"), IN_CFO_PAT),
              "CFO/PAT n/a")
    _rule(o, None if m["loss_years"] is None else m["loss_years"] == 0,
          "loss in %s of last %d yrs" % (m["loss_years"], IN_LOSS_YEARS),
          "loss-year history short")
    _rule(o, None if m["profit_cagr3"] is None
          else m["profit_cagr3"] >= IN_PROFIT_CAGR,
          "3y profit CAGR %s%% < %.0f%%" % (_f(m["profit_cagr3"]), IN_PROFIT_CAGR),
          "3y profit CAGR n/a")
    if m["fin"] != "insurance":
        _rule(o, None if m["sales_cagr3"] is None
              else m["sales_cagr3"] >= IN_SALES_CAGR,
              "3y sales CAGR %s%% < %.0f%%" % (_f(m["sales_cagr3"]), IN_SALES_CAGR),
              "3y sales CAGR n/a")
        _rule(o, m["q_last2_pos"],
              "profit or sales fell YoY in last 2 qtrs",
              "last 2 qtrs YoY n/a")
    if m["fin"] == "bank":
        _rule(o, None if m["nnpa"] is None else m["nnpa"] < BANK_NNPA,
              "net NPA %s%%" % _f(m["nnpa"], "%.2f"), "net NPA n/a")
        _rule(o, None if m["gnpa"] is None else m["gnpa"] < FIN_GNPA,
              "gross NPA %s%%" % _f(m["gnpa"], "%.2f"), "gross NPA n/a")
        _rule(o, None if m["roa"] is None else m["roa"] >= BANK_ROA,
              "ROA %s%% < %.1f%%" % (_f(m["roa"], "%.2f"), BANK_ROA), "ROA n/a")
        o["check"].append("CRAR >= 15%: check by hand")
    elif m["fin"] == "nbfc":
        if m["gnpa"] is not None:
            _rule(o, m["gnpa"] < FIN_GNPA, "gross NPA %.2f%%" % m["gnpa"], "")
        else:
            o["check"].append("GNPA n/a")
        _rule(o, None if m["roa"] is None else m["roa"] >= NBFC_ROA,
              "ROA %s%% < %.1f%%" % (_f(m["roa"], "%.2f"), NBFC_ROA), "ROA n/a")
        o["check"].append("CRAR / ALM: check by hand")
    elif m["fin"] == "insurance":
        o["check"].append("solvency >= 180%: check by hand")
    if m["decel2"]:
        o["warn"].append("profit growth slowed 2 qtrs")
    return o


def verdict(o):
    if o["fail"]:
        return "FAIL"
    return "CHECK" if o["check"] else "PASS"


def detail(o):
    parts = o["fail"] + ["? " + c for c in o["check"]] + \
        ["! " + w for w in o["warn"]]
    return "; ".join(parts)




# ================================================================== input
def latest_report():
    files = [f for f in glob.glob(os.path.join(REPORTS, "RB_Screener_*.xlsx"))
             if not os.path.basename(f).startswith("~$")]
    return max(files, key=os.path.getmtime) if files else None


def read_shortlist(path):
    out = []
    for sheet in ("Swing", "Investing", "Momentum_Top20"):
        try:
            d = pd.read_excel(path, sheet_name=sheet)
        except Exception:
            continue
        if "Symbol" not in d.columns or "RS Rank" not in d.columns:
            continue
        d = d[pd.to_numeric(d["RS Rank"], errors="coerce").notna()]
        for _, r in d.iterrows():
            out.append({"list": sheet, "symbol": str(r["Symbol"]).strip().upper(),
                        "action": r.get("Action", ""),
                        "rs_rank": float(r["RS Rank"])})
    return pd.DataFrame(out)


# ================================================================== excel
PP = "+0.00;-0.00;0.00"
# (header, key, width, number format)
COLS_SWING = [
    ("Fund Check (info)", "swing_pass", 9, None),
    ("Failed / Missing", "swing_detail", 36, None),
    ("P/E", "pe", 7, "0.0"), ("ROCE %", "roce_now", 7, "0.0"),
    ("ROE %", "roe_now", 7, "0.0"), ("3y ROE %", "roe3", 7, "0.0"),
    ("D/E", "de", 6, "0.00"),
    ("Qtr Profit YoY %", "q_profit_yoy", 9, "0.0"),
    ("Qtr Sales YoY %", "q_sales_yoy", 9, "0.0"),
    ("Pledge % (if flagged)", "pledge", 8, "0.0"),
    ("Industry", "industry", 20, None),
]
COLS_INVEST = [
    ("Fund Check (info)", "inv_pass", 9, None),
    ("Failed / Missing", "inv_detail", 36, None),
    ("P/E", "pe", 7, "0.0"), ("ROCE %", "roce_now", 7, "0.0"),
    ("3y ROCE %", "roce3", 7, "0.0"), ("ROE %", "roe_now", 7, "0.0"),
    ("3y ROE %", "roe3", 7, "0.0"), ("D/E", "de", 6, "0.00"),
    ("Int. Cover x", "icr", 7, "0.0"), ("CFO/PAT 5y", "cfo_pat", 7, "0.00"),
    ("3y Profit CAGR %", "profit_cagr3", 8, "0.0"),
    ("3y Sales CAGR %", "sales_cagr3", 8, "0.0"),
    ("Qtr Profit YoY %", "q_profit_yoy", 9, "0.0"),
    ("Qtr Sales YoY %", "q_sales_yoy", 9, "0.0"),
    ("Net NPA %", "nnpa", 7, "0.00"),
    ("Pledge % (if flagged)", "pledge", 8, "0.0"),
    ("Industry", "industry", 20, None),
]
COLS_HOLD = [("Promoter \u0394 (pp)", "d_prom", 9, PP),
             ("FII \u0394 (pp)", "d_fii", 9, PP),
             ("DII \u0394 (pp)", "d_dii", 9, PP)]
COLS_FUND = [
    ("Symbol", "symbol", 13, None), ("Action", "action", 8, None),
    ("RS Rank", "rs_rank", 7, "0"), ("On sheets", "lists", 13, None),
    ("Industry", "industry", 22, None),
    ("Swing Check (info)", "swing_pass", 9, None),
    ("Investing Check (info)", "inv_pass", 9, None),
    ("Swing: Failed / Missing", "swing_detail", 34, None),
    ("Investing: Failed / Missing", "inv_detail", 34, None),
    ("P/E", "pe", 7, "0.0"), ("ROCE %", "roce_now", 7, "0.0"),
    ("3y ROCE %", "roce3", 7, "0.0"), ("ROE %", "roe_now", 7, "0.0"),
    ("3y ROE %", "roe3", 7, "0.0"), ("D/E", "de", 6, "0.00"),
    ("Int. Cover x", "icr", 7, "0.0"), ("CFO/PAT 5y", "cfo_pat", 7, "0.00"),
    ("Latest Qtr", "q_label", 9, None),
    ("Qtr Profit YoY %", "q_profit_yoy", 9, "0.0"),
    ("Qtr Sales YoY %", "q_sales_yoy", 9, "0.0"),
    ("SUE % of price", "sue", 8, "0.00"),
    ("3y Profit CAGR %", "profit_cagr3", 8, "0.0"),
    ("3y Sales CAGR %", "sales_cagr3", 8, "0.0"),
    ("Net NPA %", "nnpa", 7, "0.00"), ("ROA %", "roa", 7, "0.00"),
    ("Pledge % (if flagged)", "pledge", 8, "0.0"),
    ("Screener Cons (machine generated)", "cons_txt", 45, None),
    ("Data", "basis", 11, None),
] + COLS_HOLD
FUND_MARK = "Fund Check (info)"      # first added header on Swing/Investing


def _styles():
    from openpyxl.styles import Font, PatternFill, Alignment
    F = "Arial"
    return {
        "F": F, "Font": Font, "Alignment": Alignment,
        "head": Font(name=F, bold=True, color="FFFFFF"),
        "fill_main": PatternFill("solid", fgColor="1F3864"),
        "fill_fund": PatternFill("solid", fgColor="375623"),
        "fill_hold": PatternFill("solid", fgColor="7F7F7F"),
        "status": {"PASS": PatternFill("solid", fgColor="C6EFCE"),
                   "FAIL": PatternFill("solid", fgColor="FFC7CE"),
                   "CHECK": PatternFill("solid", fgColor="FFEB9C")},
        "green": Font(name=F, color="006100"),
        "red": Font(name=F, color="9C0006"),
    }


def _put(ws, row, col, v, fmt, key, st):
    if v is None or (isinstance(v, float) and v != v):
        v = None
    elif isinstance(v, (np.floating, np.integer)):
        v = float(v)
    elif isinstance(v, np.bool_):
        v = bool(v)
    c = ws.cell(row=row, column=col, value=v)
    c.font = st["Font"](name=st["F"])
    if fmt:
        c.number_format = fmt
    if fmt == PP and isinstance(v, float) and v:
        c.font = st["green"] if v > 0 else st["red"]
    if isinstance(v, str) and v in st["status"] and key.endswith("_pass"):
        c.fill = st["status"][v]
    if key in ("cons_txt", "swing_detail", "inv_detail"):
        c.alignment = st["Alignment"](wrap_text=True, vertical="top")
    return c


def _header(ws, row, col, h, width, st, fill):
    from openpyxl.utils import get_column_letter
    c = ws.cell(row=row, column=col, value=h)
    c.font, c.fill = st["head"], fill
    c.alignment = st["Alignment"](horizontal="center", vertical="center",
                                  wrap_text=True)
    ws.column_dimensions[get_column_letter(col)].width = width


def _table_rows(ws):
    """Row numbers of the data table (column A symbols, stop at the first
    empty row -- the screener writes notes after one blank row)."""
    rows, r = [], 2
    while ws.cell(row=r, column=1).value not in (None, ""):
        rows.append(r)
        r += 1
    return rows


def add_columns(ws, df, cols, st):
    """Append fundamental columns to the right of an existing sheet.
    Re-running replaces the previously added block."""
    from openpyxl.utils import get_column_letter
    last = ws.max_column
    first = None
    for c in range(1, last + 1):
        if ws.cell(row=1, column=c).value == FUND_MARK:
            first = c
            break
    if first is not None:                      # remove the old block
        ws.delete_cols(first, last - first + 1)
    width = 0
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value not in (None, ""):
            width = c
    start = width + 1
    allc = cols + COLS_HOLD
    for i, (h, k, w, _) in enumerate(allc):
        fill = st["fill_hold"] if (h, k, w, _) in COLS_HOLD else st["fill_fund"]
        _header(ws, 1, start + i, h, w, st, fill)
    rows = _table_rows(ws)
    by_sym = df.drop_duplicates("symbol").set_index("symbol")
    for r in rows:
        sym = str(ws.cell(row=r, column=1).value).strip().upper()
        if sym not in by_sym.index:
            continue
        rec = by_sym.loc[sym]
        for i, (h, k, w, fmt) in enumerate(allc):
            _put(ws, r, start + i, rec.get(k), fmt, k, st)
    end = start + len(allc) - 1
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(end),
                                      max(rows[-1] if rows else 2, 2))
    ws.row_dimensions[1].height = 42


def fund_sheet(wb, df, banner, shp_q, st):
    if "Fundamentals" in wb.sheetnames:
        del wb["Fundamentals"]
    ws = wb.create_sheet("Fundamentals")
    for i, (h, k, w, _) in enumerate(COLS_FUND, 1):
        fill = st["fill_hold"] if (h, k, w, _) in COLS_HOLD else (
            st["fill_main"] if i <= 5 else st["fill_fund"])
        _header(ws, 1, i, h, w, st, fill)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"
    for n, (_, r) in enumerate(df.iterrows(), start=2):
        for i, (h, k, w, fmt) in enumerate(COLS_FUND, 1):
            _put(ws, n, i, r.get(k), fmt, k, st)
    from openpyxl.utils import get_column_letter
    last = ws.max_row
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLS_FUND)),
                                      max(last, 2))
    notes = [
        banner,
        "Every shortlisted stock that is not LATE, sorted by RS rank. "
        "Nothing is removed because of fundamentals.",
        "BACKTEST 2018-26 (backtest.py --fundamentals): the fundamental gate "
        "did NOT improve results -- in 2018-21 PASS stocks did much worse than "
        "FAIL stocks. Read these columns as information; the price signal "
        "decides.",
        "Check (info): PASS = all automated rules passed, FAIL = a rule failed, "
        "CHECK = nothing failed but data was missing. '?' = could not check, "
        "'!' = warning.",
        "Qtr Profit YoY for non-financials = profit before tax minus other "
        "income. Loss or negative base = n/a. P/E, ROCE %, ROE % = latest, "
        "from Screener.",
        "NOT checked automatically: pledge % (only if Screener's Cons mention "
        "it), auditor resignation/qualification, SEBI/forensic action, bank "
        "CRAR, insurer solvency.",
        "Grey columns = change in holding, percentage points, %s. Quarterly "
        "data, published up to 21 days after quarter-end."
        % (shp_q or "latest quarter vs previous"),
    ]
    for i, t in enumerate(notes):
        c = ws.cell(row=last + 2 + i, column=1, value=t)
        c.font = st["Font"](name=st["F"], italic=True, bold=(i == 0),
                            color="404040")
    return ws


def _note_on_sheet(ws, text, st):
    """One line under the screener's own notes."""
    r = ws.max_row + 2
    c = ws.cell(row=r, column=1, value=text)
    c.font = st["Font"](name=st["F"], italic=True, color="375623", bold=True)


def save_report(path, df, fund_df, banner, shp_q):
    from openpyxl import load_workbook
    st = _styles()
    wb = load_workbook(path)
    for sheet, cols, key in (("Swing", COLS_SWING, "Swing"),
                             ("Investing", COLS_INVEST, "Investing"),
                             ("Momentum_Top20", COLS_SWING, "Momentum_Top20")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        already = any(ws.cell(row=1, column=c).value == FUND_MARK
                      for c in range(1, ws.max_column + 1))
        add_columns(ws, df[df["list"] == key], cols, st)
        if not already:
            _note_on_sheet(ws, "Green columns = fundamentals from Screener.in "
                           "(information only, backtest 2018-26 showed no "
                           "benefit). Details: Fundamentals sheet.", st)
    fund_sheet(wb, fund_df, banner, shp_q, st)
    fill_comparison(wb, df, st)
    try:
        wb.save(path)
    except PermissionError:
        alt = path.replace(".xlsx", "_fund.xlsx")
        wb.save(alt)
        print("! %s is open in Excel -- saved a copy instead." %
              os.path.basename(path))
        path = alt
    return path


def fill_comparison(wb, df, st):
    """Write the swing fundamental verdict into Strategy_Comparison's
    'Fundamental Status' column (made by momentum_screener.py)."""
    if "Strategy_Comparison" not in wb.sheetnames:
        return
    ws = wb["Strategy_Comparison"]
    col = None
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == "Fundamental Status":
            col = c
            break
    if col is None:
        return
    status = dict(zip(df["symbol"], df["swing_pass"]))
    for r in _table_rows(ws):
        sym = str(ws.cell(row=r, column=1).value).strip().upper()
        v = status.get(sym, "n/a")
        c = ws.cell(row=r, column=col, value=v)
        c.font = st["Font"](name=st["F"])
        if v in st["status"]:
            c.fill = st["status"][v]


def save_standalone(path, fund_df, banner, shp_q):
    from openpyxl import Workbook
    st = _styles()
    wb = Workbook()
    wb.remove(wb.active)
    fund_sheet(wb, fund_df, banner, shp_q, st)
    wb.save(path)
    return path


# ================================================================== main
def main():
    import account
    acc = account.activate()
    args = sys.argv[1:]
    src, syms = None, None
    if "--file" in args:
        src = args[args.index("--file") + 1]
    if "--symbols" in args:
        syms = [s.strip().upper() for s in
                args[args.index("--symbols") + 1].split(",") if s.strip()]

    if syms:
        short = pd.DataFrame({"list": "Swing", "symbol": syms, "action": "",
                              "rs_rank": np.nan})
        stamp = dt.date.today().isoformat()
        banner = "Manual symbol list | fundamentals from Screener.in on %s" % stamp
    else:
        src = src or latest_report()
        if not src or not os.path.exists(src):
            print("! No RB_Screener_*.xlsx in %s -- run rbscan first." % REPORTS)
            sys.exit(1)
        print("Shortlist from %s" % os.path.basename(src))
        short = read_shortlist(src)
        banner = ("Fundamentals from Screener.in on %s for %s"
                  % (dt.date.today().isoformat(), os.path.basename(src)))
    if short.empty:
        print("Shortlist is empty -- nothing to check.")
        return

    clean_old_cache()
    uniq = list(dict.fromkeys(short["symbol"]))
    print("Fetching %d companies from Screener.in (about %d s) ..."
          % (len(uniq), int(len(uniq) * PAUSE * 1.3)))
    rows, missing = {}, []
    for i, s in enumerate(uniq, 1):
        sys.stdout.write("\r  %3d/%d  %-14s" % (i, len(uniq), s))
        sys.stdout.flush()
        try:
            p = load_company(s)
            if p is None or p["pl"].empty:
                missing.append(s)
                continue
            rows[s] = metrics(p)
        except Exception as e:          # one bad page must not stop the run
            missing.append("%s (%s)" % (s, type(e).__name__))
    print()
    if missing:
        print("! No usable Screener data for: " + ", ".join(missing))

    recs = []
    for _, r in short.iterrows():
        m = rows.get(r["symbol"])
        d = dict(r)
        if m is None:
            d.update({"swing_pass": "CHECK", "inv_pass": "CHECK",
                      "swing_detail": "? no Screener data",
                      "inv_detail": "? no Screener data"})
        else:
            d.update(m)
            so, io = swing_rules(m), invest_rules(m)
            d.update({"swing_pass": verdict(so), "swing_detail": detail(so),
                      "inv_pass": verdict(io), "inv_detail": detail(io),
                      "cons_txt": " | ".join(m["cons"])})
        recs.append(d)
    df = pd.DataFrame(recs)

    # Fundamentals sheet: one row per stock, not LATE, sorted by RS rank
    lists = df.groupby("symbol")["list"].apply(lambda x: " + ".join(x))
    one = df.drop_duplicates("symbol").copy()
    one["lists"] = one["symbol"].map(lists)
    one = one[one["action"] != "LATE"].sort_values("rs_rank", ascending=False)

    shp_q = ""
    if "shp_q" in df and df["shp_q"].notna().any():
        shp_q = df["shp_q"].dropna().mode().iloc[0]

    try:
        if syms:
            path = save_standalone(os.path.join(
                REPORTS, "RB_Fundamentals_%s%s.xlsx" % (TAG + "_" if TAG else "",
                                                        stamp)), one, banner, shp_q)
        else:
            path = save_report(src, df, one, banner, shp_q)
    except ImportError:
        path = os.path.join(REPORTS, "RB_Fundamentals_%s%s.csv"
                            % (TAG + "_" if TAG else "",
                               dt.date.today().isoformat()))
        one.to_csv(path, index=False)
        print("! openpyxl not installed (pip3 install openpyxl) -> CSV saved")

    # ---------------------------------------------------------- summary
    pd.set_option("display.width", 220)

    def fmt(x, f="%+.2f"):
        if isinstance(x, str):
            return x
        return "" if x is None or x != x else f % x

    print("\n==== FUNDAMENTALS (%d stocks, not LATE, by RS rank) -- info only ===="
          % len(one))
    if len(one):
        t = pd.DataFrame({
            "symbol": one["symbol"], "act": one["action"],
            "RS": one["rs_rank"].map(lambda x: fmt(x, "%.0f")),
            "swing": one["swing_pass"], "invest": one["inv_pass"],
            "P/E": one.get("pe", pd.Series(dtype=float)).map(
                lambda x: fmt(x, "%.0f")),
            "ROCE": one.get("roce_now", pd.Series(dtype=float)).map(
                lambda x: fmt(x, "%.0f")),
            "qtrPrft%": one.get("q_profit_yoy", pd.Series(dtype=float)).map(
                lambda x: fmt(x, "%.0f")),
            "Prom": one.get("d_prom", pd.Series(dtype=object)).map(fmt),
            "FII": one.get("d_fii", pd.Series(dtype=object)).map(fmt),
            "DII": one.get("d_dii", pd.Series(dtype=object)).map(fmt)})
        print(t.to_string(index=False))
    print("\nHolding change = %s, percentage points. Info only."
          % (shp_q or "latest vs previous quarter"))
    print("Backtest 2018-26: the fundamental check did NOT improve results -- "
          "nothing is removed.")
    print("Not checked: pledge (unless Screener flags it), auditor, SEBI.")
    print("\nExcel: %s" % path)
    account.banner(acc)


if __name__ == "__main__":
    main()
