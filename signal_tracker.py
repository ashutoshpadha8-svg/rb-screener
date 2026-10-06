#!/usr/bin/env python3
"""
SIGNAL TRACKER  --  how did every stock the screener found actually do? (30 Sep 2026)
===================================================================================

Reads EVERY master scan (reports/RB_Screener_YYYY-MM-DD.xlsx) -- so the days
before this file existed are included too -- and keeps a permanent log
(data/signals_log.csv; rows are only ever added / updated, never dropped,
so a deleted xlsx loses nothing).

What counts as a "find" (one row each):
  W+TT      a Swing-sheet row BUY / FIT / LATE, one row per signal date
            (a stock that signals again later = a new row)
  MOMENTUM  a Momentum Top 20 stock, one row per stretch in the list
            (out of the list for > 20 days and back = a new row)
  both      marked in the 'Both lists' column (= Super-Buy that day)

Per find: price the day it was first found (the scan's own price), today's
price, return, Nifty over the same days, best / worst point since (daily
high / low), whether the 20% stop would have hit, and a rule status
(W+TT: close below the 40-week MA = exit; momentum: today's rank > 40 = sell
at the next rebalance). Finds younger than 30 days are marked TOO EARLY --
a week of prices says nothing.

Writes sheet "Signal_Tracker" into today's master scan (portfolio.py copies it
into every Portfolio file) + data/signal_tracker_latest.csv. No orders.

python3 signal_tracker.py            (rb_scan.py runs it after every scan)
"""

import os
import re
import sys
import glob
import datetime as dt

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import daily_screener as ds                                    # noqa: E402

LOG = os.path.join(ds.DATA, "signals_log.csv")
OUT = os.path.join(ds.DATA, "signal_tracker_latest.csv")
REPORTS = os.path.join(ds.HERE, "reports")
SHEET = "Signal_Tracker"
SUMMARY_SHEET = "Signal_Summary"
PICKS_SHEET = "All_Picks"
MOM_GAP_DAYS = 20          # out of the top 20 longer than this = new find
TOO_EARLY = 30             # calendar days
STOP = 0.20
LOG_COLS = ["source", "symbol", "key", "first_found", "first_status",
            "price_found", "signal_date", "signal_price", "rank_found",
            "last_seen", "days_listed", "both_lists"]
SYM_OK = re.compile(r"^[A-Z0-9][A-Z0-9&\-_.]{0,19}$")


# ================================================================== scans
def scan_files():
    """{date: path} for every master scan; the plain file wins over a
    *_fund.xlsx copy of the same day."""
    out = {}
    for p in glob.glob(os.path.join(REPORTS, "RB_Screener_*.xlsx")):
        b = os.path.basename(p)
        m = re.match(r"RB_Screener_(\d{4}-\d{2}-\d{2})(_fund)?\.xlsx$", b)
        if not m:
            continue
        d = m.group(1)
        if d not in out or not m.group(2):
            out[d] = p
    return dict(sorted(out.items()))


def _num(v):
    try:
        f = float(v)
        return f if f == f else None
    except (TypeError, ValueError):
        return None


def read_scan(day, path):
    """Rows (source, symbol, status, price, rank, signal_date, signal_price)."""
    rows = []
    try:
        sw = pd.read_excel(path, sheet_name="Swing")
        for _, r in sw.iterrows():
            s = str(r.get("Symbol", "")).strip().upper()
            act = str(r.get("Action", "")).strip().upper()
            if not SYM_OK.match(s) or act not in ("BUY", "FIT", "LATE"):
                continue
            sd = str(r.get("Signal Date", ""))[:10]
            rows.append(dict(day=day, source="W+TT", symbol=s, status=act,
                             price=_num(r.get("Price")),
                             rank=_num(r.get("RS Rank")), signal_date=sd,
                             signal_price=_num(r.get("Signal Price"))))
    except Exception as e:
        print("  ! %s: Swing sheet not read (%s)" % (day, type(e).__name__))
    try:
        mo = pd.read_excel(path, sheet_name="Momentum_Top20")
        for _, r in mo.iterrows():
            s = str(r.get("Symbol", "")).strip().upper()
            rk = _num(r.get("Mom Rank"))
            if not SYM_OK.match(s) or rk is None:
                continue
            px = _num(r.get("LTP")) or _num(r.get("Last Close"))
            rows.append(dict(day=day, source="MOMENTUM", symbol=s,
                             status="TOP20", price=px, rank=rk,
                             signal_date="", signal_price=None))
    except Exception as e:
        print("  ! %s: Momentum sheet not read (%s)" % (day, type(e).__name__))
    return rows


def load_log():
    try:
        d = pd.read_csv(LOG, dtype=str).fillna("")
        for c in LOG_COLS:
            if c not in d:
                d[c] = ""
        return d[LOG_COLS]
    except (IOError, OSError, ValueError):
        return pd.DataFrame(columns=LOG_COLS)


def build_log(files):
    """All scans -> one row per find. Existing log rows are kept (a scan file
    that was deleted later still counts)."""
    raw = []
    for day, path in files.items():
        raw += read_scan(day, path)
    finds = {}
    if raw:
        R = pd.DataFrame(raw).sort_values(["day"])
        wtt_days = {(r.symbol, r.day) for r in R[(R.source == "W+TT") & (
            R.status != "LATE")].itertuples()}
        mom_days = {(r.symbol, r.day) for r in R[R.source == "MOMENTUM"]
                    .itertuples()}
        for r in R[R.source == "W+TT"].itertuples():
            key = r.signal_date or r.day
            k = ("W+TT", r.symbol, key)
            f = finds.get(k)
            if f is None:
                finds[k] = dict(source="W+TT", symbol=r.symbol, key=key,
                                first_found=r.day, first_status=r.status,
                                price_found=r.price, signal_date=r.signal_date,
                                signal_price=r.signal_price, rank_found=r.rank,
                                last_seen=r.day, days_listed=1,
                                both_lists="YES" if (r.symbol, r.day)
                                in mom_days else "")
            else:
                f["last_seen"] = r.day
                f["days_listed"] += 1
        for sym, g in R[R.source == "MOMENTUM"].groupby("symbol"):
            start = prev = None
            for r in g.itertuples():
                d = pd.Timestamp(r.day)
                if prev is None or (d - prev).days > MOM_GAP_DAYS:
                    start = r.day
                    finds[("MOMENTUM", sym, start)] = dict(
                        source="MOMENTUM", symbol=sym, key=start,
                        first_found=r.day, first_status="TOP20",
                        price_found=r.price, signal_date="",
                        signal_price=None, rank_found=r.rank,
                        last_seen=r.day, days_listed=1,
                        both_lists="YES" if (sym, r.day) in wtt_days else "")
                else:
                    f = finds[("MOMENTUM", sym, start)]
                    f["last_seen"] = r.day
                    f["days_listed"] += 1
                prev = d
    new = pd.DataFrame(list(finds.values()), columns=LOG_COLS).astype(str) \
        .replace({"None": "", "nan": ""})
    old = load_log()
    om = old[old.source == "MOMENTUM"]
    for i in new.index[new.source == "MOMENTUM"]:   # same stretch, earlier
        o = om[(om.symbol == new.at[i, "symbol"]) &      # start in the log
               (om.first_found <= new.at[i, "first_found"])]
        if len(o):
            o = o.sort_values("first_found").iloc[-1]
            gap = (pd.Timestamp(new.at[i, "first_found"]) -
                   pd.Timestamp(o.last_seen)).days
            if gap <= MOM_GAP_DAYS:
                new.at[i, "key"] = o.key
    ok_ = old.source + "|" + old.symbol + "|" + old.key
    nk = new.source + "|" + new.symbol + "|" + new.key
    keep = old[~ok_.isin(nk)]
    prev = old.set_index(ok_)
    for i, k in enumerate(nk):            # an older scan file was deleted:
        if k in prev.index:               # the log remembers the first find
            o = prev.loc[k]
            o = o.iloc[0] if isinstance(o, pd.DataFrame) else o
            if o.first_found and o.first_found < new.at[i, "first_found"]:
                for c in ("first_found", "first_status", "price_found",
                          "rank_found", "both_lists"):
                    new.at[i, c] = o[c]
            new.at[i, "last_seen"] = max(o.last_seen, new.at[i, "last_seen"])
            new.at[i, "days_listed"] = str(max(
                int(_num(o.days_listed) or 0),
                int(_num(new.at[i, "days_listed"]) or 0)))
    log = pd.concat([keep, new], ignore_index=True)
    log = log.sort_values(["first_found", "source", "symbol"])
    tmp = LOG + ".tmp"
    log.to_csv(tmp, index=False)
    os.replace(tmp, LOG)
    return log


# ================================================================ All_Picks
def _sheet(path, name):
    try:
        d = pd.read_excel(path, sheet_name=name)
        d["Symbol"] = [str(x).strip().upper() if x == x and x is not None
                       else "" for x in d["Symbol"]]
        d = d[d["Symbol"].map(lambda x: bool(SYM_OK.match(str(x))))]
        if "Mom Rank" in d:
            return d[pd.to_numeric(d["Mom Rank"], errors="coerce").notna()]
        return d[[str(x).strip().upper() in ("BUY", "FIT", "LATE")
                  for x in d["Action"]]]
    except Exception:
        return pd.DataFrame(columns=["Symbol"])


def all_picks(path, log, today):
    """6 Oct 2026 (RB: "momentum, swing, investing ek he sheet per + har
    stock ki age"): today's Momentum top 20 + Swing + Investing as ONE row
    per stock. Age = calendar days since the stock was first found in this
    stretch: W+TT = its signal date, momentum = first scan of the current
    top-20 stretch (signals_log.csv); in both lists = the older one."""
    sw, iv, mo = (_sheet(path, n) for n in ("Swing", "Investing",
                                             "Momentum_Top20"))
    syms = list(mo["Symbol"]) if "Mom Rank" in mo else []
    for d in (sw, iv):
        syms += [x for x in d["Symbol"] if x not in syms]
    if not syms:
        return pd.DataFrame()
    mlog = log[log.source == "MOMENTUM"] if len(log) else log
    t0 = pd.Timestamp(today)
    first_scan = min(log["first_found"]) if len(log) else ""
    row_of = lambda d, s: d[d["Symbol"] == s].iloc[0] \
        if s in set(d["Symbol"]) else None                          # noqa
    out = []
    for s in syms:
        m, w, i = row_of(mo, s), row_of(sw, s), row_of(iv, s)
        cands = []                      # (date, price then, what)
        if m is not None:
            g = mlog[mlog.symbol == s].sort_values("first_found")
            if len(g):
                r = g.iloc[-1]
                cands.append((r.first_found, _num(r.price_found), "MOM"))
            else:
                cands.append((today, _num(m.get("LTP")), "MOM"))
        for x in (w, i):
            if x is not None:
                sd = str(x.get("Signal Date", ""))[:10]
                if re.match(r"\d{4}-\d{2}-\d{2}$", sd):
                    cands.append((sd, _num(x.get("Signal Price")), "W+TT"))
        cands.sort(key=lambda c: c[0])
        found, p0, what = cands[0] if cands else (today, None, "")
        now = None
        for x, k in ((m, "LTP"), (w, "Price"), (i, "Price")):
            if x is not None and now is None:
                now = _num(x.get(k))
        lists = []
        if m is not None:
            lists.append("MOM #%d" % int(_num(m.get("Mom Rank")) or 0))
        if w is not None:
            lists.append("SWING " + str(w.get("Action", "")).upper())
        if i is not None:
            lists.append("INV " + str(i.get("Action", "")).upper())
        wtt = w if w is not None else i
        age = max(0, (t0 - pd.Timestamp(found)).days)
        out.append({
            "Symbol": s, "Age": age_label(age), "Age (din)": age,
            "Pehli baar mila": found,
            "Kahan se": " + ".join(lists),
            "Super-Buy": "YES" if m is not None and wtt is not None and
            str(wtt.get("Action", "")).upper() != "LATE" else "",
            "Mom Rank": _num(m.get("Mom Rank")) if m is not None else None,
            "Swing": str(w.get("Action", "")).upper() if w is not None
            else "",
            "Investing": str(i.get("Action", "")).upper() if i is not None
            else "",
            "RS Rank": _num((wtt if wtt is not None else m).get("RS Rank")),
            "Price then": p0, "Price now": now,
            "Since found %": 100 * (now / p0 - 1) if now and p0 else None,
            "Age from": "signal date (W+TT)" if what == "W+TT" else
            ("first scan in top 20" + (" (scans start %s)" % first_scan
                                       if found == first_scan else ""))})
    return pd.DataFrame(out)


# ================================================================== prices
def prices(symbols):
    """Daily history per symbol (free source + broker gap fill) and live
    prices when a broker token is active."""
    frames = {}
    for s in symbols:
        df = ds.fetch_eod(s.lower())
        if df is not None and len(df):
            frames[s.lower()] = df
    bm = ds.fetch_eod(ds.BENCH)
    live, note = {}, "free source only (no broker token)"
    try:
        import momentum_screener as ms
        import broker_api as ba
        sess = ms.get_session()
        if sess and frames:
            frames, bm, live = ba.refresh(sess, frames, bm,
                                          ds.last_expected_session(), [])
            note = "%s fill + live price" % sess.label
    except SystemExit:
        pass
    except Exception as e:
        note = "broker step skipped (%s)" % type(e).__name__
    return frames, bm, live, note


def _ranks_now():
    try:
        r = pd.read_csv(os.path.join(ds.DATA, "momentum_ranks_latest.csv"))
        return dict(zip(r["symbol"].astype(str).str.upper(), r["rank"]))
    except Exception:
        return {}


def evaluate(log, frames, bm, live, today):
    ranks = _ranks_now()
    bmc = bm["Close"] if bm is not None else None
    rows = []
    for f in log.itertuples():
        s = f.symbol.lower()
        df = frames.get(s)
        found = pd.Timestamp(f.first_found)
        p0 = _num(f.price_found)
        base = dict(Symbol=f.symbol, Source=f.source,
                    **{"First status": f.first_status,
                       "Found on": f.first_found, "Price then": p0,
                       "Both lists": f.both_lists,
                       "Rank then": _num(f.rank_found),
                       "Days listed": _num(f.days_listed),
                       "Last in list": f.last_seen})
        if df is None or not p0:
            rows.append(dict(base, Status="NO DATA"))
            continue
        after = df[df.index > found]
        now = live.get(s) or float(df["Close"].iloc[-1])
        hi = max([now] + list(after["High"].dropna()))
        lo = min([now] + list(after["Low"].dropna()))
        n0 = n1 = None
        if bmc is not None and len(bmc):
            b_then = bmc[bmc.index <= found]
            n0 = float(b_then.iloc[-1]) if len(b_then) else None
            n1 = float(bmc.iloc[-1])
        ret = now / p0 - 1
        nret = n1 / n0 - 1 if n0 and n1 else None
        age = (pd.Timestamp(today) - found).days
        # rule status
        c = df["Close"]
        if f.source == "W+TT":
            ma200 = c.rolling(200).mean()
            below = after.index[(c.reindex(after.index) <
                                 ma200.reindex(after.index))]
            stop_hit = lo <= p0 * (1 - STOP)
            if stop_hit:
                status = "STOP HIT (-20%)"
            elif len(below):
                status = "SWING EXIT (close < 40w MA %s)" % below[0].date()
            else:
                status = "rule: hold"
        else:
            rk = ranks.get(f.symbol)
            if not ranks:
                status = "rank unknown (no ranks file)"
            elif rk is None:
                status = "not ranked now -> SELL at rebalance"
            else:
                status = "rank %d" % rk + (" -> SELL at rebalance" if rk > 40
                                           else " (keep, <= 40)")
        if age < TOO_EARLY:
            status += " | TOO EARLY"
        rows.append(dict(base, **{
            "Price now": round(now, 2), "Return %": 100 * ret,
            "Nifty same days %": 100 * nret if nret is not None else None,
            "vs Nifty %": 100 * (ret - nret) if nret is not None else None,
            "Best since %": 100 * (hi / p0 - 1),
            "Worst since %": 100 * (lo / p0 - 1),
            "Days since found": age, "Status": status}))
    return pd.DataFrame(rows)


def summary(t):
    if t.empty or "Return %" not in t:
        return pd.DataFrame()
    t = t.dropna(subset=["Return %"])
    groups = [("W+TT BUY (found on signal day)",
               (t.Source == "W+TT") & (t["First status"] == "BUY")),
              ("W+TT FIT (found later)",
               (t.Source == "W+TT") & (t["First status"] == "FIT")),
              ("W+TT LATE (not tested)",
               (t.Source == "W+TT") & (t["First status"] == "LATE")),
              ("Momentum Top 20", t.Source == "MOMENTUM"),
              ("  momentum rank 1-5 when found",
               (t.Source == "MOMENTUM") & (t["Rank then"] <= 5)),
              ("  momentum rank 6-10",
               (t.Source == "MOMENTUM") & (t["Rank then"] > 5) &
               (t["Rank then"] <= 10)),
              ("  momentum rank 11+ (selected)",
               (t.Source == "MOMENTUM") & (t["Rank then"] > 10)),
              ("In BOTH lists (Super-Buy)",
               (t["Both lists"] == "YES") & (t.Source == "W+TT")),
              ("ALL finds", t.Source.notna())]
    out = []
    for name, m in groups:
        g = t[m]
        if not len(g):
            continue
        vs = g["vs Nifty %"].dropna()
        days = pd.to_numeric(g["Days since found"], errors="coerce").dropna()
        out.append({"Group": name, "Finds": len(g),
                    "Stocks": int(g["Symbol"].nunique()),
                    "Up now": int((g["Return %"] > 0).sum()),
                    "Down now": int((g["Return %"] <= 0).sum()),
                    "Avg return %": g["Return %"].mean(),
                    "Median %": g["Return %"].median(),
                    "Avg vs Nifty %": vs.mean() if len(vs) else None,
                    "Beat Nifty": "%d of %d" % ((vs > 0).sum(), len(vs)),
                    "Best now %": g["Return %"].max(),
                    "Worst now %": g["Return %"].min(),
                    "Avg worst dip %": g["Worst since %"].mean(),
                    "20% stop hit": int(g["Status"].str.startswith(
                        "STOP").sum()),
                    "Days tracked": ("-" if not len(days) else "%d" %
                                     days.min() if days.min() == days.max()
                                     else "%d-%d" % (days.min(), days.max())),
                    "Older than 30 days": int((g["Days since found"] >=
                                               TOO_EARLY).sum())})
    return pd.DataFrame(out)


# ================================================================== output
AGE_FILL = [("TODAY", "E2EFDA"), ("1-5 days", "DDEBF7"),
            ("6-29 days", "FBF3E4")]           # soft, not flashy; 30+ = white


def age_label(d):
    """Calendar days since found -> TODAY / 1-5 days / 6-29 days / 30+ days."""
    try:
        d = int(d)
    except (TypeError, ValueError):
        return ""
    if d <= 0:
        return "TODAY"
    if d <= 5:
        return "1-5 days"
    if d < TOO_EARLY:
        return "6-29 days"
    return "30+ days"


TODAY_CELL = "A9D08E"          # the days cell of today's finds: a bit stronger
OLD_FILLS = ("C6EFCE", "EDEDED")   # screener's old BUY green / tracker grey
DAY_COLS = {"Swing": "Days Since Signal", "Investing": "Days Since Signal",
            "Momentum_Top20": "Days in Top 20",
            SHEET: "Days since found", PICKS_SHEET: "Age (din)"}


def _rgb(cell):
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return ""
    return str(f.fgColor.rgb or "")[-6:].upper()


def paint_ages(wb):
    """Same soft age colours on every sheet with a days column: the whole
    row by age (cells that carry another colour -- LATE orange, fundamentals
    columns -- keep theirs) and the days cell of today's finds a bit darker
    + bold. Safe to run again (rbsig) -- it only re-paints."""
    from openpyxl.styles import Font, PatternFill
    ages = {lab: PatternFill("solid", fgColor=col) for lab, col in AGE_FILL}
    none = PatternFill(fill_type=None)
    strong = PatternFill("solid", fgColor=TODAY_CELL)
    ours = set(OLD_FILLS) | {c for _, c in AGE_FILL} | {TODAY_CELL}
    for name, col in DAY_COLS.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        hrow = dcol = None
        for r in range(1, 6):
            for c in range(1, ws.max_column + 1):
                if ws.cell(row=r, column=c).value == col:
                    hrow, dcol = r, c
            if dcol:
                break
        if not dcol:
            continue
        ncol = 0
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=hrow, column=c).value not in (None, ""):
                ncol = c
        r = hrow + 1
        while ws.cell(row=r, column=1).value not in (None, ""):
            try:
                lab = age_label(ws.cell(row=r, column=dcol).value)
            except Exception:
                lab = ""
            if lab:
                for c in range(1, ncol + 1):
                    cell = ws.cell(row=r, column=c)
                    if _rgb(cell) in ours or _rgb(cell) == "":
                        cell.fill = ages.get(lab, none)
                d = ws.cell(row=r, column=dcol)
                if lab == "TODAY":
                    d.fill = strong
                    d.font = Font(name=d.font.name, bold=True)
            r += 1


def write_sheet(path, summ, t, note, today, picks=None):
    """Two tabs: Signal_Tracker = every find (header + Symbol frozen, filter
    on, so sorting always moves whole rows); Signal_Summary = the groups."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L
    wb = load_workbook(path)
    for n in (SHEET, SUMMARY_SHEET, PICKS_SHEET):
        if n in wb.sheetnames:
            del wb[n]
    head = Font(bold=True, color="FFFFFF")
    navy = PatternFill("solid", fgColor="1F4E78")
    ages = {lab: PatternFill("solid", fgColor=col) for lab, col in AGE_FILL}
    note_txt = ("Row colour = Age: green = found TODAY (days cell darker), "
                "blue = 1-5 days, beige = 6-29 days (TOO EARLY to judge), "
                "white = 30+ days. Same colours on Swing / Investing / "
                "Momentum_Top20. "
                "Rank then = momentum rank (MOMENTUM rows) or RS "
                "rank 0-100 (W+TT rows). Sort ONLY with the header arrows "
                "(filter) -- sorting one selected column mixes the rows. "
                "Backtest (2013-26): W+TT ~39% winners, avg +11.8% per trade "
                "over months; momentum top 20 beat the universe by ~6% in 6 "
                "months. Info only -- no orders.")

    def table(ws, df, r0, widths):
        for j, h in enumerate(df.columns, 1):
            c = ws.cell(row=r0, column=j, value=h)
            c.font, c.fill = head, navy
            c.alignment = Alignment(wrap_text=True, vertical="center",
                                    horizontal="center")
            ws.column_dimensions[L(j)].width = widths.get(h, 12)
        ws.row_dimensions[r0].height = 32
        for i, row in enumerate(df.itertuples(index=False), r0 + 1):
            age = None
            for j, v in enumerate(row, 1):
                if isinstance(v, float):
                    v = None if v != v else round(v, 1)
                c = ws.cell(row=i, column=j, value=v)
                h = df.columns[j - 1]
                if isinstance(v, float) and ("%" in h):
                    c.number_format = "+0.0;-0.0;0.0"
                    c.font = Font(color="1E7B34" if v > 0 else "C00000"
                                  if v < 0 else "000000")
                if h == "Age":
                    age = v
            if age in ages:
                for j in range(1, len(df.columns) + 1):
                    ws.cell(row=i, column=j).fill = ages[age]
        return r0 + len(df)

    ws = wb.create_sheet(SHEET)
    ws.cell(row=1, column=1, value="SIGNAL TRACKER | %s | prices: %s | "
            "summary: tab %s" % (today, note, SUMMARY_SHEET)).font = \
        Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=note_txt).font = Font(italic=True, size=10)
    last = table(ws, t, 3, {"Symbol": 14, "Age": 10, "Source": 11,
                            "Status": 40, "Found on": 11,
                            "Last in list": 11})
    ws.freeze_panes = "C4"           # header row + Symbol + Age always seen
    if len(t):
        ws.auto_filter.ref = "A3:%s%d" % (L(len(t.columns)), last)

    if picks is not None and len(picks):     # ONE list: mom + swing + inv
        wp = wb.create_sheet(PICKS_SHEET, 0)
        wp.cell(row=1, column=1, value="ALL PICKS | %s | Momentum top 20 + "
                "Swing + Investing, ek row per stock" % today).font = \
            Font(bold=True, size=13)
        wp.cell(row=2, column=1, value=(
            "Age = kitne din pehle mila: W+TT = signal date, momentum = top 20 "
            "mein pehla scan (scans 26 Sep 2026 se, usse pehle ka pata nahi). "
            "Rang: green = aaj, blue = 1-5 din, beige = 6-29 din (judge karne "
            "ke liye jaldi), white = 30+ din. Super-Buy = momentum top 20 + "
            "W+TT dono. Detail: Swing / Investing / Momentum_Top20 tabs. "
            "Info only -- koi order nahi.")).font = Font(italic=True, size=10)
        last = table(wp, picks, 3, {"Symbol": 14, "Age": 10, "Kahan se": 30,
                                    "Pehli baar mila": 12, "Age from": 30})
        wp.freeze_panes = "C4"
        wp.auto_filter.ref = "A3:%s%d" % (L(len(picks.columns)), last)
        try:
            import momentum_focus as mf
            top5 = mf.symbols(mf.load()[0])
            for r in range(4, last + 1):
                c = wp.cell(row=r, column=1)
                if c.value in top5:
                    c.fill = PatternFill("solid", fgColor=mf.GOLD_FILL)
                    c.font = Font(bold=True, color=mf.GOLD_TEXT)
        except Exception:
            pass

    wsum = wb.create_sheet(SUMMARY_SHEET)
    wsum.cell(row=1, column=1, value="SIGNAL SUMMARY | %s | every stock: tab "
              "%s" % (today, SHEET)).font = Font(bold=True, size=13)
    wsum.cell(row=2, column=1, value=note_txt).font = Font(italic=True,
                                                           size=10)
    if len(summ):
        table(wsum, summ, 3, {"Group": 34})
        wsum.freeze_panes = "B4"
    try:                               # momentum rank 1-5: look here first
        import momentum_focus as mf
        gold = PatternFill("solid", fgColor=mf.GOLD_FILL)
        gfont = Font(bold=True, color=mf.GOLD_TEXT)
        for r in range(4, 4 + len(summ)):
            c = wsum.cell(row=r, column=1)
            if "rank 1-5" in str(c.value):
                c.fill, c.font = gold, gfont
        n = 4 + len(summ) + 1
        for txt in ("Beat Nifty = only finds with a Nifty comparison. Best / "
                    "Worst now % = today's return (Best since % in "
                    "Signal_Tracker = the peak). Rank 11+ (selected) = the "
                    "sector cap can pick ranks above 20.",
                    "Gold rank 1-5 group = rank WHEN FOUND. Gold symbols in "
                    "Signal_Tracker = in TODAY's momentum rank 1-5. " +
                    mf.NOTE):
            wsum.cell(row=n, column=1, value=txt).font = Font(italic=True,
                                                               size=10)
            n += 1
        now = mf.symbols(mf.load()[0])
        cols = list(t.columns)
        if now and "Symbol" in cols and "Source" in cols:
            for i, x in enumerate(t.itertuples(index=False), 4):
                if x[cols.index("Source")] == "MOMENTUM" and \
                        x[cols.index("Symbol")] in now:
                    c = ws.cell(row=i, column=1)
                    c.fill, c.font = gold, gfont
    except Exception:
        pass                           # highlight only -- never stop a report
    try:                               # Mcap + Large/Mid/Small, every sheet
        import cap_class
        cap_class.add_columns(wb)
    except Exception:
        pass
    paint_ages(wb)
    try:                               # columns fit their content (layout)
        import xl_fit
        xl_fit.fit_workbook(wb)
    except ImportError:
        pass
    wb.save(path)


def main():
    today = ds.now_ist().date().isoformat()
    files = scan_files()
    if not files:
        print("No master scans in %s yet -- run rb first." % REPORTS)
        return
    print("Signal tracker: %d master scan(s) %s .. %s"
          % (len(files), list(files)[0], list(files)[-1]))
    log = build_log(files)
    syms = sorted(set(log.symbol))
    frames, bm, live, note = prices(syms)
    t = evaluate(log, frames, bm, live, today)
    t["Age"] = t["Days since found"].map(age_label)
    cols = ["Symbol", "Age", "Source", "First status", "Rank then", "Found on",
            "Days since found",
            "Price then", "Price now", "Return %", "Nifty same days %",
            "vs Nifty %", "Best since %", "Worst since %", "Both lists",
            "Days listed", "Last in list", "Status"]
    t = t.reindex(columns=cols).sort_values(["Found on", "Return %"],
                                            ascending=[False, False])
    summ = summary(t)
    t.to_csv(OUT, index=False)
    latest = list(files.values())[-1]
    try:
        picks = all_picks(latest, log, today)
    except Exception as e:
        print("  ! All_Picks skipped (%s)" % type(e).__name__)
        picks = None
    try:
        write_sheet(latest, summ, t, note, today, picks)
        print("  sheet '%s' written into %s" % (SHEET, os.path.basename(latest)))
        try:                        # fresh copy in Google Drive (Sheets)
            import drive_copy
            drive_copy.push(latest, quiet=True)
        except Exception:
            pass
    except Exception as e:
        print("  ! could not write the sheet (%s) -- CSV: %s"
              % (type(e).__name__, OUT))
    print("\n==== SIGNAL TRACKER (%d finds, prices: %s) ====" % (len(t), note))
    if len(summ):
        pd.set_option("display.width", 200)
        print(summ[["Group", "Finds", "Up now", "Down now", "Avg return %",
                    "Avg vs Nifty %", "Beat Nifty", "Worst now %",
                    "Older than 30 days"]]
              .to_string(index=False, float_format=lambda x: "%+.1f" % x))
    young = int(t["Status"].astype(str).str.contains("TOO EARLY").sum())
    if young:
        print("  %d of %d finds are younger than %d days -> TOO EARLY to judge."
              % (young, len(t), TOO_EARLY))


if __name__ == "__main__":
    main()
