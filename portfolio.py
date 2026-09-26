#!/usr/bin/env python3
"""
PORTFOLIO  (rbport)  --  everything about YOUR holdings, per account
===================================================================

The market scan (rbscan) is one shared file a day. This script does the
account part, for the account in token.txt:

  Holdings   every stock you hold (demat + PAPER portfolio from split.csv):
             trend      40w MA, 30w MA + Weinstein stage, 50-DMA
             technical  RSI(14), % from 52w high, ATR %, 6m / 12m return
             strategies W+TT signal today?  swing rule (20% stop / 40w MA),
                        investing rule (Stage 4), momentum rank (keep <= 40)
             fundamentals (Screener.in, info only)  +  latest news
             RECOMMENDATION + WHY
  Rebalance  momentum: SELL / BUY / HOLD per LIVE and PAPER (1st trading day)
  Actions    the master Strategy_Comparison list + an Action dropdown
             (BUY / BUY MTF / PAPER / PAPER MTF / WATCH) -> rbtrack

RECOMMENDATION
  A stock tagged in split.csv is judged by ITS strategy's exit rule (the
  backtested rules; worst leg wins):
     EXIT        rule fired (stop / 40w MA / Stage 4)
     SELL@REBAL  momentum rank > 40 -> sell at the monthly rebalance
     WATCH       close to an exit level
     HOLD        rule says stay
  An untagged holding (bought outside the system) gets the combined check:
     SELL   Stage 4, or below the 40w MA AND momentum rank > 40
     WEAK   one warning       KEEP  above 40w MA and rank <= 40
  -> the single rules are backtested; the combined check on outside holdings
     is NOT. Fundamentals and news never change the verdict (the fundamental
     gate did not help in the 2018-26 backtest) -- they are for your eyes.

OUTPUT accounts/<BROKER>_<ID>/reports/Portfolio_<BROKER>_<Name>_<date>.xlsx
RUN    rbport               (rbscan first, once a day)
       rbport --no-news     faster
       rbport --no-fund     skip Screener.in for holdings not in the scan
"""

import warnings
warnings.filterwarnings("ignore")

import os
import sys
import glob

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import daily_screener as ds
import momentum_screener as ms
import position_tracker as pt

REPORTS = os.path.join(ds.HERE, "reports")     # routed to the account folder
TAG = ""                                       # e.g. DHAN_Ashutosh (routed)
KEEP_RANK = ms.BUFFER * ms.SLOTS
LEG_ORDER = {"EXIT": 0, "SELL@REBAL": 1, "SELL": 1, "WATCH": 2, "WEAK": 2,
             "HOLD": 3, "KEEP": 3}


def path_for(day=None):
    day = day or ds.now_ist().date().isoformat()
    return os.path.join(REPORTS, "Portfolio_%s%s.xlsx"
                        % (TAG + "_" if TAG else "", day))


def latest():
    """Today's portfolio file, else the newest one (for rbtrack)."""
    p = path_for()
    if os.path.exists(p):
        return p
    files = [f for f in glob.glob(os.path.join(REPORTS, "Portfolio_*.xlsx"))
             if not os.path.basename(f).startswith("~$")]
    return max(files, key=os.path.getmtime) if files else None


# ================================================================== inputs
def read_split():
    if not os.path.exists(ms.SPLIT_FILE):
        return pd.DataFrame()
    sp = pd.read_csv(ms.SPLIT_FILE, dtype={"order_id": str})
    sp["symbol"] = sp["symbol"].astype(str).str.upper().str.strip()
    for c in ("swing_qty", "investing_qty", "momentum_qty", "entry_price"):
        sp[c] = pd.to_numeric(sp.get(c, 0), errors="coerce").fillna(0)
    sp["mode"] = sp.get("mode", "LIVE")
    sp["mode"] = sp["mode"].fillna("").astype(str).str.upper().replace("", "LIVE")
    sp["strategy"] = sp.get("strategy", "").fillna("") if "strategy" in sp \
        else ""
    return sp[sp["symbol"].isin(["", "EXAMPLE", "NAN"]) == False]  # noqa: E712


def positions(broker_holdings, sp):
    """One row per (mode, symbol): qty, entry, entry_date, legs."""
    out = []
    live = sp[sp["mode"] == "LIVE"] if len(sp) else sp
    for h in broker_holdings:
        s = h["symbol"]
        rows = live[live["symbol"] == s] if len(live) else live
        legs = {}
        for _, r in rows.iterrows():
            for leg in ("swing", "investing", "momentum"):
                if r["%s_qty" % leg] > 0:
                    legs[leg] = legs.get(leg, 0) + r["%s_qty" % leg]
        tagged_qty = sum(legs.values())
        edate = rows["entry_date"].dropna().astype(str).min() if len(rows) \
            and "entry_date" in rows else None
        entry = float(rows["entry_price"][rows["entry_price"] > 0].mean()) \
            if len(rows) and (rows["entry_price"] > 0).any() else h["avg_price"]
        note = ""
        if legs and abs(tagged_qty - h["qty"]) > 0.5:
            note = "split.csv legs total %g, demat %g" % (tagged_qty, h["qty"])
        out.append({"mode": "LIVE", "symbol": s, "qty": h["qty"],
                    "entry": entry or h["avg_price"], "entry_date": edate,
                    "legs": legs, "note": note})
    held = {h["symbol"] for h in broker_holdings}
    for _, r in (live[~live["symbol"].isin(held)].iterrows() if len(live)
                 else []):
        q = r["swing_qty"] + r["investing_qty"] + r["momentum_qty"]
        legs = {leg: r["%s_qty" % leg] for leg in ("swing", "investing",
                                                   "momentum")
                if r["%s_qty" % leg] > 0}
        if q > 0:
            out.append({"mode": "LIVE", "symbol": r["symbol"], "qty": q,
                        "entry": r["entry_price"], "entry_date":
                        r.get("entry_date"), "legs": legs, "note":
                        "in split.csv but NOT in demat (AMO pending? rbsync)"})
    paper = sp[sp["mode"] == "PAPER"] if len(sp) else sp
    for s, g in (paper.groupby("symbol") if len(paper) else []):
        legs = {leg: g["%s_qty" % leg].sum() for leg in
                ("swing", "investing", "momentum") if g["%s_qty" % leg].sum()}
        q = sum(legs.values())
        if q > 0:
            out.append({"mode": "PAPER", "symbol": s, "qty": q,
                        "entry": float(g["entry_price"].mean()),
                        "entry_date": g["entry_date"].astype(str).min()
                        if "entry_date" in g else None,
                        "legs": legs, "note": ""})
    return out


# ================================================================== analysis
def technicals(c, lo, hi):
    px = float(c.iloc[-1])
    ma50, ma150, ma200 = (c.rolling(n).mean() for n in (50, 150, 200))
    rising = bool(ma150.iloc[-1] > ma150.iloc[-11])
    above = px > ma150.iloc[-1]
    stage = ("Stage 2 (uptrend)" if above and rising else
             "Stage 4 (downtrend)" if not above and not rising else
             "Stage 3 (topping?)" if above else "Stage 1 / pullback")
    d = c.diff()
    up = d.clip(lower=0).ewm(alpha=1 / 14, adjust=False).mean()
    dn = (-d.clip(upper=0)).ewm(alpha=1 / 14, adjust=False).mean()
    rsi = float(100 - 100 / (1 + up.iloc[-1] / dn.iloc[-1])) if dn.iloc[-1] \
        else 100.0
    tr = np.maximum(np.maximum(hi - lo, (hi - c.shift(1)).abs()),
                    (lo - c.shift(1)).abs())
    atr = float(tr.ewm(alpha=1 / 14, adjust=False).mean().iloc[-1])
    hi52 = float(c.iloc[-252:].max())
    ret = lambda n: (px / float(c.iloc[-n - 1]) - 1) * 100 if len(c) > n \
        else np.nan
    return {"px": px, "ma50": float(ma50.iloc[-1]),
            "ma150": float(ma150.iloc[-1]), "ma200": float(ma200.iloc[-1]),
            "stage": stage, "rsi": rsi, "atr_pct": atr / px * 100,
            "dist52": (px / hi52 - 1) * 100, "ret6": ret(126), "ret12": ret(252)}


def combined(t, st4_now, rk):
    """Untagged holding: SELL / WEAK / KEEP (see docstring)."""
    below = t["px"] < t["ma200"]
    weak_rank = rk is None or rk > KEEP_RANK
    if st4_now or (below and rk is not None and rk > KEEP_RANK):
        return "SELL"
    if below or (weak_rank and rk is not None):
        return "WEAK"
    return "KEEP"


def stage4_now(c):
    ma150 = c.rolling(150).mean()
    below = (c < ma150).astype(int)
    run = below.groupby((below == 0).cumsum()).cumsum()
    return bool(run.iloc[-1] >= 10 and ma150.iloc[-1] < ma150.iloc[-11])


def fundamentals_for(syms, fund_sheet, fetch):
    """{sym: dict} from the master Fundamentals sheet, else Screener.in."""
    out = {}
    if not fund_sheet.empty and "Symbol" in fund_sheet:
        for _, r in fund_sheet.iterrows():
            out[str(r["Symbol"]).upper()] = {
                "P/E": r.get("P/E"), "ROCE %": r.get("ROCE %"),
                "ROE %": r.get("ROE %"), "D/E": r.get("D/E"),
                "Qtr Profit YoY %": r.get("Qtr Profit YoY %"),
                "Qtr Sales YoY %": r.get("Qtr Sales YoY %"),
                "Fund Swing": r.get("Swing Check (info)"),
                "Fund Invest": r.get("Investing Check (info)"),
                "Promoter Δ": r.get("Promoter Δ (pp)"),
                "FII Δ": r.get("FII Δ (pp)"), "DII Δ": r.get("DII Δ (pp)")}
    todo = [s for s in syms if s not in out]
    if not (fetch and todo):
        return out
    import fundamentals as fu
    print("Fundamentals for %d holding(s) not in today's scan (Screener.in, "
          "~%d s) ..." % (len(todo), int(len(todo) * fu.PAUSE * 1.3)))
    for s in todo:
        try:
            p = fu.load_company(s)
            if p is None or p["pl"].empty:
                continue
            m = fu.metrics(p)
            r1 = lambda k: round(m[k], 1) if isinstance(m.get(k), (int, float)) \
                and m[k] == m[k] else None
            out[s] = {"P/E": r1("pe"), "ROCE %": r1("roce_now"),
                      "ROE %": r1("roe_now"), "D/E": r1("de"),
                      "Qtr Profit YoY %": r1("q_profit_yoy"),
                      "Qtr Sales YoY %": r1("q_sales_yoy"),
                      "Fund Swing": fu.verdict(fu.swing_rules(m)),
                      "Fund Invest": fu.verdict(fu.invest_rules(m)),
                      "Promoter Δ": m.get("d_prom"), "FII Δ": m.get("d_fii"),
                      "DII Δ": m.get("d_dii")}
        except Exception:
            continue
    return out


def analyse(pos, closes, lows, highs, prov, ranks, wtt, fund, news):
    s = pos["symbol"]
    c = closes.get(s)
    base = {"Mode": pos["mode"], "Symbol": s, "Qty": pos["qty"],
            "Entry": round(pos["entry"], 2) if pos["entry"] else None}
    if c is None or len(c) < 210:
        base.update({"Recommendation": "NO DATA", "Why":
                     "not enough price history (ETF/SGB, new listing or "
                     "renamed symbol) " + pos["note"]})
        return base
    lo, hi = lows[s], highs[s]
    t = technicals(c, lo, hi)
    rk = ranks.get(s)
    legs, why, verdicts = pos["legs"], [], []
    if "swing" in legs:
        v, w = pt.judge_swing(c, lo, pos["entry"], pos["entry_date"])[:2]
        verdicts.append(v)
        why.append("Swing rule: %s (%s)" % (v, w))
    if "investing" in legs:
        v, w = pt.judge_investing(c, pos["entry_date"])[:2]
        verdicts.append(v)
        why.append("Investing rule: %s (%s)" % (v, w))
    if "momentum" in legs:
        v, w = pt.judge_momentum(s, c, lo, hi, pos["entry"],
                                 pos["entry_date"], ranks)[:2]
        verdicts.append(v)
        why.append("Momentum rule: %s (%s)" % (v, w))
    if verdicts:
        rec = min(verdicts, key=lambda v: LEG_ORDER.get(v, 9))
        basis = "+".join(legs)
    else:
        rec = combined(t, stage4_now(c), rk)
        basis = "untagged (combined check)"
        why.append("%s: %s 40w MA %.0f; momentum rank %s"
                   % (rec, "above" if t["px"] > t["ma200"] else "BELOW",
                      t["ma200"], rk if rk is not None else "none (outside "
                      "the >= Rs 10k Cr universe)"))
    if prov.get(s) and rec not in ("HOLD", "KEEP"):
        why.append("live price during market hours -- counts only if it "
                   "CLOSES here")
    if pos["note"]:
        why.append(pos["note"])
    f = fund.get(s, {})
    n = news.get(s, [])
    base.update({
        "Recommendation": rec, "Basis": basis, "Why": " | ".join(why),
        "LTP": round(t["px"], 2),
        "P&L %": round((t["px"] / pos["entry"] - 1) * 100, 1)
        if pos["entry"] else None,
        "Value (Rs)": round(t["px"] * pos["qty"]),
        "Stage": t["stage"], "40w MA": round(t["ma200"], 1),
        "vs 40w %": round((t["px"] / t["ma200"] - 1) * 100, 1),
        "30w MA": round(t["ma150"], 1), "50 DMA": round(t["ma50"], 1),
        "RSI 14": round(t["rsi"], 0), "From 52w High %": round(t["dist52"], 1),
        "ATR %": round(t["atr_pct"], 1),
        "6m Ret %": round(t["ret6"], 1) if t["ret6"] == t["ret6"] else None,
        "12m Ret %": round(t["ret12"], 1) if t["ret12"] == t["ret12"] else None,
        "Mom Rank": rk, "W+TT today": wtt.get(s, "-"),
        "P/E": f.get("P/E"), "ROCE %": f.get("ROCE %"), "ROE %": f.get("ROE %"),
        "D/E": f.get("D/E"), "Qtr Profit YoY %": f.get("Qtr Profit YoY %"),
        "Qtr Sales YoY %": f.get("Qtr Sales YoY %"),
        "Fund (swing)": f.get("Fund Swing"), "Fund (invest)": f.get("Fund Invest"),
        "Promoter Δ": f.get("Promoter Δ"), "FII Δ": f.get("FII Δ"),
        "DII Δ": f.get("DII Δ"),
        "News (7 days)": " || ".join("%s: %s (%s)" % x for x in n)})
    return base


# ================================================================== excel
HCOLS = ["Recommendation", "Mode", "Symbol", "Qty", "Entry", "LTP", "P&L %",
         "Value (Rs)", "Basis", "Why", "Stage", "40w MA", "vs 40w %", "30w MA",
         "50 DMA", "RSI 14", "From 52w High %", "ATR %", "6m Ret %",
         "12m Ret %", "Mom Rank", "W+TT today", "P/E", "ROCE %", "ROE %", "D/E",
         "Qtr Profit YoY %", "Qtr Sales YoY %", "Fund (swing)", "Fund (invest)",
         "Promoter Δ", "FII Δ", "DII Δ", "News (7 days)"]
WIDTH = {"Why": 70, "News (7 days)": 90, "Basis": 16, "Stage": 16,
         "Recommendation": 13, "Symbol": 13, "Fund (swing)": 9,
         "Fund (invest)": 9}
REC_FILL = {"EXIT": "F8CBAD", "SELL": "F8CBAD", "SELL@REBAL": "FCE4D6",
            "WATCH": "FFE699", "WEAK": "FFE699", "HOLD": "C6EFCE",
            "KEEP": "C6EFCE", "NO DATA": "D9D9D9"}


def write_book(path, hold, rebal, comp, held_modes, old_actions, banner):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    head = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F4E78")
    wb = Workbook()

    def table(ws, cols, rows, fills=None, widths=None):
        for i, h in enumerate(cols, 1):
            c = ws.cell(row=1, column=i, value=h)
            c.font, c.fill = head, hfill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = \
                (widths or {}).get(h, 10)
        for r, x in enumerate(rows, 2):
            for i, h in enumerate(cols, 1):
                v = x.get(h)
                if isinstance(v, (np.floating, np.integer)):
                    v = float(v)
                if isinstance(v, float) and v != v:
                    v = None
                ws.cell(row=r, column=i, value=v)
            if fills:
                k, fmap = fills
                col = fmap.get(x.get(k))
                if col:
                    ws.cell(row=r, column=cols.index(k) + 1).fill = \
                        PatternFill("solid", fgColor=col)
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)),
                                          max(len(rows) + 1, 2))
        return len(rows) + 1

    ws = wb.active
    ws.title = "Holdings"
    last = table(ws, HCOLS, hold, ("Recommendation", REC_FILL), WIDTH)
    for i, t in enumerate([banner,
                           "Tagged (split.csv) -> its strategy's backtested "
                           "exit rule; untagged -> combined check (SELL/WEAK/"
                           "KEEP, not backtested as a whole).",
                           "Fundamentals + news are information only -- they "
                           "never change the recommendation.",
                           "Nothing is sold from code: place sells in your "
                           "broker app."], last + 2):
        ws.cell(row=i, column=1, value=t).font = Font(italic=True, bold=(
            i == last + 2))

    wr = wb.create_sheet("Rebalance")
    rcols = ["Section", "Mode", "Symbol", "Mom Rank", "Momentum Score",
             "Sector", "Qty Held", "Entry", "LTP", "P&L %", "Shares to Buy",
             "Amount (Rs)", "Note"]
    last = table(wr, rcols, rebal, ("Section", {"SELL": "FFC7CE",
                                                "BUY": "C6EFCE",
                                                "HOLD": "DDEBF7"}),
                 {"Note": 45, "Sector": 22, "Symbol": 13})
    wr.cell(row=last + 2, column=1, value=(
        "Momentum only (split.csv strategy = Momentum). Do it on the 1st "
        "trading day of the month: sell SELL rows at the open, buy BUY rows "
        "that fill a free slot. Keep while rank <= %d." % KEEP_RANK)).font = \
        Font(italic=True)

    wa = wb.create_sheet("Actions")
    acols = list(comp.columns) + ["Held here", "Action"] if len(comp) else \
        ["Ticker", "Held here", "Action"]
    rows = []
    for _, x in comp.iterrows():
        d = dict(x)
        t = str(d.get("Ticker", "")).upper()
        d["Held here"] = ", ".join(sorted(held_modes.get(t, [])))
        d["Action"] = old_actions.get(t, "")
        rows.append(d)
    widths = {"Ticker": 13, "Strategy Overlap": 14, "Sector / Industry": 22,
              "Regime": 18, "Action": 11, "Held here": 10}
    last = table(wa, acols, rows, ("Strategy Overlap", {"Super-Buy": "C6EFCE"}),
                 widths)
    col = get_column_letter(len(acols))
    for r in range(2, last + 1):
        wa.cell(row=r, column=len(acols)).fill = PatternFill("solid",
                                                             fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(ms.ACTIONS),
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Action", error="Pick from the list",
                        promptTitle="Action", showInputMessage=True,
                        prompt="BUY / BUY MTF = real AMO, PAPER = mock, "
                               "WATCH = no order")
    dv.add("%s2:%s%d" % (col, col, max(last, 2)))
    wa.add_data_validation(dv)
    wa.cell(row=last + 2, column=1, value=(
        "Pick an Action (yellow), save + close, then rbtrack after 15:30. "
        "Momentum buys only on the 1st trading day of the month.")).font = \
        Font(italic=True)
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.replace(".xlsx", "_%s.xlsx" % ds.now_ist().strftime("%H%M"))
        wb.save(alt)
        print("! %s is open in Excel -- saved as %s (your Action picks there "
              "are NOT read by rbtrack until you close + re-run)"
              % (os.path.basename(path), os.path.basename(alt)))
        return alt


# ================================================================== main
def main():
    import account
    acc = account.activate()
    a = sys.argv[1:]
    sess = acc.session if acc.token_ok else None
    warns = []
    master = ms.todays_report()
    today = ds.now_ist().date().isoformat()
    if not master:
        print("! No master scan yet -- run rbscan first.")
        sys.exit(1)
    if today not in os.path.basename(master):
        warns.append("master scan is %s, not today -- run rbscan"
                     % os.path.basename(master))
    if not os.path.exists(pt.RANKS_FILE):
        print("! No momentum ranking -- run rbscan first.")
        sys.exit(1)
    rk = pd.read_csv(pt.RANKS_FILE)
    rk["symbol"] = rk["symbol"].astype(str).str.upper()
    ranks = dict(zip(rk["symbol"], rk["rank"].astype(int)))
    regime_red = bool(rk["regime_red"].iloc[0]) if "regime_red" in rk and \
        len(rk) else False

    # ---- holdings
    broker = []
    if sess:
        print("Fetching holdings from %s ..." % sess.label)
        broker = pt.get_holdings(sess)
    else:
        warns.append("token missing/expired -> demat holdings NOT read, only "
                     "the PAPER portfolio")
    sp = read_split()
    pos = positions(broker, sp)
    syms = sorted({p["symbol"] for p in pos})
    print("  %d demat holding(s), %d position row(s)" % (len(broker), len(pos)))

    # ---- prices
    closes, lows, highs, prov = {}, {}, {}, {}
    note = "no holdings"
    if syms:
        print("Loading prices ...")
        if sess:
            closes, lows, highs, prov, note = pt.load_prices(sess, syms, warns)
        else:
            for s in syms:
                df = ds.fetch_eod(s.lower())
                if df is not None:
                    closes[s], lows[s], highs[s] = df["Close"], df["Low"], df["High"]
            note = "free source only (no token)"

    # ---- master scan inputs
    sw = ms._read_sheet(master, "Swing")
    wtt = dict(zip(sw["Symbol"].astype(str).str.upper(), sw["Action"])) \
        if "Symbol" in sw and "Action" in sw else {}
    fund = fundamentals_for(syms, ms._read_sheet(master, "Fundamentals"),
                            "--no-fund" not in a)
    news = {}
    if syms and "--no-news" not in a:
        import news_feed as nf
        print("News for %d stock(s) ..." % len(syms))
        news = {s: nf.latest(s, n=2) for s in syms}

    hold = [analyse(p, closes, lows, highs, prov, ranks, wtt, fund, news)
            for p in pos]
    hold.sort(key=lambda x: (x["Mode"] != "LIVE",
                             LEG_ORDER.get(x["Recommendation"], 8),
                             -(x.get("Value (Rs)") or 0)))

    # ---- rebalance (momentum) + actions
    full = rk.rename(columns={})
    top = rk[rk.get("in_top", False) == True] if "in_top" in rk else \
        rk[rk["rank"] <= ms.SLOTS]                                  # noqa
    rebal = ms.rebalance_plan(top, full) if len(top) else []
    demat = {h["symbol"] for h in broker}
    for x in rebal:        # already owned, just not tagged Momentum in split.csv
        if x["Section"] == "BUY" and x["Mode"] == "LIVE" and x["Symbol"] in demat:
            x["Note"] = ("ALREADY IN DEMAT (not tagged Momentum) -- don't buy "
                         "again; set momentum_qty/strategy in split.csv")
    comp = ms._read_sheet(master, "Strategy_Comparison")
    if len(comp) and "Ticker" in comp:
        comp = comp[comp["Ticker"].notna() & comp["Strategy Overlap"].isin(
            ["Super-Buy", "Momentum only", "W+TT only"])]
    held_modes = {}
    for p in pos:
        held_modes.setdefault(p["symbol"], set()).add(p["mode"])
    old_actions = {}
    prev = path_for()
    if os.path.exists(prev):
        o = ms._read_sheet(prev, "Actions")
        if "Ticker" in o and "Action" in o:
            old_actions = {str(t).upper(): str(v) for t, v in
                           zip(o["Ticker"], o["Action"])
                           if isinstance(v, str) and v.strip()}

    # ---- terminal
    live_rows = [h for h in hold if h["Mode"] == "LIVE"]
    val = sum(h.get("Value (Rs)") or 0 for h in live_rows)
    cost = sum((h.get("Entry") or 0) * h["Qty"] for h in live_rows
               if h.get("Value (Rs)"))
    print("\n" + "=" * 78)
    print(" PORTFOLIO %s | master scan %s" % (acc.label,
                                              os.path.basename(master)))
    print(" prices: %s" % note)
    if val:
        print(" LIVE value ~Rs %s | cost ~Rs %s | P&L %+.1f%%"
              % (format(int(val), ","), format(int(cost), ","),
                 (val / cost - 1) * 100 if cost else 0))
    print("=" * 78)
    for h in hold:
        print("\n %-10s %-5s %-12s qty %-6g P&L %s  %s | RSI %s | rank %s | "
              "W+TT %s" % (h["Recommendation"], h["Mode"], h["Symbol"],
                           h["Qty"], "%+.1f%%" % h["P&L %"]
                           if h.get("P&L %") is not None else "  n/a",
                           h.get("Stage", ""), h.get("RSI 14", "-"),
                           h.get("Mom Rank") if h.get("Mom Rank") is not None
                           else "-", h.get("W+TT today", "-")))
        print("    why: %s" % h["Why"])
        if h.get("Fund (swing)") or h.get("P/E"):
            f = lambda v: "n/a" if v is None or v != v else (
                "%.0f" % v if isinstance(v, float) else str(v))
            print("    fundamentals (info): P/E %s, ROCE %s%%, qtr profit YoY "
                  "%s%%, check swing %s / invest %s"
                  % (f(h.get("P/E")), f(h.get("ROCE %")),
                     f(h.get("Qtr Profit YoY %")), h.get("Fund (swing)"),
                     h.get("Fund (invest)")))
        if h.get("News (7 days)"):
            print("    news: %s" % h["News (7 days)"][:220])
    if not hold:
        print("\n (no holdings -- demat empty and no PAPER rows)")
    for m in ("LIVE", "PAPER"):
        rr = [x for x in rebal if x["Mode"] == m]
        if rr:
            print("\n REBALANCE %s (1st trading day): SELL %s | BUY %s | HOLD %d"
                  % (m, ", ".join(x["Symbol"] for x in rr
                                  if x["Section"] == "SELL") or "-",
                     ", ".join(x["Symbol"] for x in rr if x["Section"] == "BUY"
                               and "fills" in x["Note"]) or "-",
                     sum(1 for x in rr if x["Section"] == "HOLD")))
    if regime_red:
        print("\n!!! MARKET RED (Nifty < 200-DMA): new buys did worse in the "
              "backtest. Paper first.")
    for w in warns:
        print("  ! " + w)
    path = write_book(path_for(), hold, rebal, comp, held_modes, old_actions,
                      "Portfolio %s | %s | master %s | prices: %s"
                      % (acc.label, today, os.path.basename(master), note))
    print("\nExcel: %s  (sheets Holdings, Rebalance, Actions)" % path)
    if old_actions:
        print("  kept your Action picks: %s" % ", ".join(
            "%s=%s" % kv for kv in old_actions.items()))
    print("Next: pick Actions in the Actions sheet -> save + close -> rbtrack "
          "(after 15:30)")
    account.banner(acc)


if __name__ == "__main__":
    main()
