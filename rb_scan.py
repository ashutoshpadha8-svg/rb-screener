#!/usr/bin/env python3
"""
MASTER SCAN  (rbscan)  --  ONE market scan per day, shared by every account
===========================================================================

The strategy is the same for everybody, so the stock lists are the same:
one file a day, reports/RB_Screener_YYYY-MM-DD.xlsx, with the sheets
  Swing, Investing          daily_screener.py   (W+TT signals)
  Momentum_Top20,           momentum_screener.py (RAMOM top 20, ranking)
  Strategy_Comparison
  Fundamentals + columns    fundamentals.py     (Screener.in, info only)

Runs the three steps in order (each only if the previous one worked).
Already scanned today -> it does NOT scan again, except:
  * the last scan ran during market hours (09:15-15:30) and the market is
    now closed -> one more scan for the closing prices
  * rbscan --force

Your own holdings, rebalance and Action picks: rbport (portfolio.py).
"""

import os
import sys
import subprocess
import datetime as dt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import daily_screener as ds

STEPS = ["daily_screener.py", "momentum_screener.py", "fundamentals.py"]
NEEDED = ("Swing", "Momentum_Top20", "Fundamentals")


def master_path(day=None):
    day = day or ds.now_ist().date().isoformat()
    return os.path.join(ds.HERE, "reports", "RB_Screener_%s.xlsx" % day)


def complete(path):
    """Scan file exists and all three steps wrote their sheets."""
    if not os.path.exists(path):
        return False
    try:
        from openpyxl import load_workbook
        names = load_workbook(path, read_only=True).sheetnames
    except Exception:
        return False
    return all(n in names for n in NEEDED)


def made_in_market_hours(path):
    t = dt.datetime.fromtimestamp(os.path.getmtime(path), ds.IST)
    return (t.date() == ds.now_ist().date() and t.weekday() < 5 and
            (9, 15) <= (t.hour, t.minute) < (15, 30))


def main():
    force = "--force" in sys.argv[1:]
    path = master_path()
    if complete(path) and not force:
        if made_in_market_hours(path) and not ds.market_open():
            print("Today's scan was made during market hours -- re-scanning "
                  "once for the closing prices.")
        else:
            print("Today's master scan is already done:\n  %s" % path)
            print("Nothing to scan again (same strategy, same stocks).")
            print("  Your portfolio:  rbport      Re-scan anyway:  rbscan --force")
            return
    for step in STEPS:
        print("\n" + "#" * 70 + "\n# %s\n" % step + "#" * 70)
        r = subprocess.call([sys.executable, os.path.join(HERE, step)])
        if r != 0:
            print("\n! %s stopped (exit %d) -- scan not complete." % (step, r))
            sys.exit(r)
    print("\nMASTER SCAN DONE: %s" % path)
    import gdrive_sync
    gdrive_sync.push(path)
    print("Next: rbport  (your holdings, rebalance, Action picks)")


if __name__ == "__main__":
    main()
