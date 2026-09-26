#!/usr/bin/env python3
"""
MOMENTUM SCREENER (NSE-style, monthly)  --  cash market only
============================================================

Strategy (strategy_lab.py "RAMOM", backtested 2013-2026 on the point-in-time
>= Rs 10,000 Cr universe, Rs 2 lakh, real Dhan costs + tax):
  score = 0.5 x z(6-month return / 1-year vol) + 0.5 x z(12-month return / 1-year vol)
          (z = cross-sectional z-score among eligible stocks)
  eligible = >= Rs 10,000 Cr today, 60-day median turnover > Rs 5 Cr,
             >= ~1 year of history
  hold the top 20, equal slots, rebalance on the 1st trading day of the month;
  a holding is kept while its rank stays inside the top 40 (2N buffer).

RISK SWITCHES (each one backtested, see CLAUDE.md "Momentum enhancements")
  SECTOR_CAP = 4       ON : max 4 picks per NSE industry. Better in BOTH halves,
                            max DD -43% -> -37%.
  ATR_SIZING = False   OFF: slot x (median ATR% / stock ATR%), 0.5x-2x. Lower
                            DD (-36%) but ~3%/yr less in 2020-26.
  Market regime        warning + "High Risk - Market Red" label only. Blocking
                            new buys when Nifty < 200DMA cost ~2.5%/yr.
  Trailing / breakeven stops: not here -- they destroyed the backtest
                            (18% -> 2-17%). position_tracker.py has an opt-in
                            switch (MOMENTUM_SMART_SL) if you insist.

OUTPUT (inside today's reports/RB_Screener_YYYY-MM-DD.xlsx; other sheets kept)
  Momentum_Top20      top 20 with Shares to Buy (Rs 2 lakh / 20 slots)
  Strategy_Comparison W+TT Swing list vs Momentum list, "Super-Buy" when in
                      both, and an empty "Action" column for your BUY
  data/momentum_ranks_latest.csv  full ranking (position_tracker.py uses it
                      for the monthly "rank > 40 -> sell" rule)

RUN (normally via rbscan, after daily_screener.py)
    python3 ~/Desktop/RB_Screener/momentum_screener.py
"""

import warnings
warnings.filterwarnings("ignore")

import os
import sys
import glob
import math
import datetime as dt

import numpy as np
import pandas as pd
import requests

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import daily_screener as ds

# ------------------------------------------------------------------ config
CAPITAL = 200000          # Rs, total momentum capital
SLOTS = 20                # top N
BUFFER = 2                # keep a holding while rank <= BUFFER x SLOTS
SECTOR_CAP = 4            # max picks per NSE industry (None = off)
ATR_SIZING = False        # True = volatility-scaled slots (see docstring)
RANKS_FILE = os.path.join(ds.DATA, "momentum_ranks_latest.csv")
INDUSTRY_FILE = os.path.join(ds.DATA, "_nse_industry.csv")
INDUSTRY_URL = ("https://nsearchives.nseindia.com/content/indices/"
                "ind_niftytotalmarket_list.csv")
SPLIT_FILE = os.path.join(ds.HERE, "split.csv")
RED = "High Risk - Market Red"


# ================================================================== helpers
def get_token():
    """Dhan token from dhan_token.txt (client id read from the token)."""
    tok = ds.read_token()
    if not tok:
        return None
    cid, exp = ds.token_info(tok)
    if exp and exp < ds.now_ist():
        print("! Token expired at %s -- running on the free source only."
              % exp.strftime("%d %b %H:%M"))
        return None
    if not cid:
        return None
    ds.CLIENT_ID = cid
    return tok


def industry_map():
    """NSE industry per symbol (Nifty Total Market list, refreshed weekly)."""
    fresh = os.path.exists(INDUSTRY_FILE) and \
        (dt.datetime.now().timestamp() - os.path.getmtime(INDUSTRY_FILE)) < 7 * 86400
    if not fresh:
        try:
            r = requests.get(INDUSTRY_URL, headers=ds.NSE_HDRS, timeout=30)
            if r.status_code == 200 and b"Industry" in r.content[:200]:
                open(INDUSTRY_FILE, "wb").write(r.content)
        except requests.RequestException:
            pass
    if not os.path.exists(INDUSTRY_FILE):
        return {}
    d = pd.read_csv(INDUSTRY_FILE)
    return dict(zip(d["Symbol"].astype(str).str.upper(), d["Industry"]))


def slot_amount(atr_pct, median_atr):
    """Rupees for one slot. Equal by default; volatility-scaled if ATR_SIZING."""
    base = CAPITAL / SLOTS
    if ATR_SIZING and atr_pct and median_atr and atr_pct > 0:
        base *= float(np.clip(median_atr / atr_pct, 0.5, 2.0))
    return base


def shares_for(amount, price):
    return int(math.floor(amount / price)) if price and price > 0 else 0


def momentum_holdings():
    """Symbols held under the Momentum strategy in split.csv."""
    if not os.path.exists(SPLIT_FILE):
        return set()
    sp = pd.read_csv(SPLIT_FILE)
    if "strategy" not in sp or "momentum_qty" not in sp:
        return set()
    sp = sp[(sp["strategy"].astype(str).str.lower() == "momentum") &
            (pd.to_numeric(sp["momentum_qty"], errors="coerce").fillna(0) > 0)]
    return set(sp["symbol"].astype(str).str.upper())


def todays_report():
    today = os.path.join(ds.REPORTS, "RB_Screener_%s.xlsx"
                         % ds.now_ist().date().isoformat())
    if os.path.exists(today):
        return today
    files = [f for f in glob.glob(os.path.join(ds.REPORTS, "RB_Screener_*.xlsx"))
             if not os.path.basename(f).startswith("~$")]
    return max(files, key=os.path.getmtime) if files else None


# ================================================================== data
def load_data(tok):
    """Universe (>= Rs 10k Cr today) + history, gap-filled from Dhan, and
    live LTP. Same sources and steps as daily_screener.py."""
    warns = []
    want = ds.last_expected_session()
    syms, caps = ds.build_universe(warns)
    frames = {}
    print("Loading history (%d stocks) ..." % len(syms))
    for i, s in enumerate(syms, 1):
        if i % 25 == 0 or i == len(syms):
            sys.stdout.write("\r  %3d/%d" % (i, len(syms)))
            sys.stdout.flush()
        df = ds.fetch_eod(s)
        if df is not None and len(df) >= 260:
            frames[s] = df
    print()
    bm = ds.fetch_eod(ds.BENCH)
    if bm is None:
        print("Could not load Nifty 50 history.")
        sys.exit(1)
    live = {}
    if tok:
        try:
            scrip = ds.load_scrip_map()
            last = max(f.index[-1] for f in frames.values()).date()
            if last < want:
                print("  free source is behind (%s) -- filling from Dhan ..." % last)
                bm = ds.gap_fill(bm, tok, "13", "IDX_I", "INDEX", want, warns,
                                 "NIFTY")
                for i, s in enumerate(list(frames), 1):
                    sid = scrip.get(s.upper())
                    if sid:
                        frames[s] = ds.gap_fill(frames[s], tok, sid, "NSE_EQ",
                                                "EQUITY", want, warns, s)
                    if i % 25 == 0 or i == len(frames):
                        sys.stdout.write("\r    %3d/%d" % (i, len(frames)))
                        sys.stdout.flush()
                print()
            ids = {s: scrip[s.upper()] for s in frames if s.upper() in scrip}
            print("  fetching today's prices from Dhan ...")
            ltp = ds.dhan_ltp(tok, list(ids.values()))
            live = {s: ltp[i] for s, i in ids.items() if i in ltp}
        except PermissionError as e:
            print("! Dhan refused (%s) -- continuing on the free source." % e)
        except Exception as e:
            print("! Dhan step failed (%s) -- continuing without it." % e)
    else:
        print("  (no usable Dhan token -- free source only, prices may be old)")
    return frames, bm, caps, live, warns, want


def panels(frames, bm):
    cal = bm.index
    for f in frames.values():
        cal = cal.union(f.index)
    cal = cal[cal >= bm.index[0]]
    P = {k: pd.DataFrame({s: f[k] for s, f in frames.items()}).reindex(cal)
         for k in ("Open", "High", "Low", "Close", "Volume")}
    P["BM"] = bm["Close"].reindex(cal).ffill(limit=5)
    return P


# ================================================================== scores
def compute(P):
    """Scores exactly as strategy_lab.scores() RAMOM + helper columns."""
    C = P["Close"].ffill(limit=5)
    V = P["Volume"]
    liq = (C * V).rolling(60).median() > 5e7
    r = C.pct_change().clip(-0.5, 0.5)
    vol = r.rolling(252, min_periods=200).std() * np.sqrt(252)
    m6 = C / C.shift(126) - 1
    m12 = C / C.shift(252) - 1
    ok = liq & vol.notna() & m12.notna()

    def z(x):
        x = x.where(ok)
        return x.sub(x.mean(axis=1), axis=0).div(x.std(axis=1), axis=0)

    score = 0.5 * z(m6 / vol) + 0.5 * z(m12 / vol)
    H, L = P["High"], P["Low"]
    tr = np.maximum(np.maximum(H - L, (H - C.shift(1)).abs()),
                    (L - C.shift(1)).abs())
    atr = tr.ewm(alpha=1 / 14, adjust=False, min_periods=14).mean() / C * 100
    BM = P["BM"]
    rs6 = (C / C.shift(126)).div(BM / BM.shift(126), axis=0)
    rs = rs6.rank(axis=1, pct=True) * 100
    hi52 = C.rolling(252).max()
    t = -1
    out = pd.DataFrame({
        "symbol": [s.upper() for s in C.columns],
        "score": score.iloc[t].values, "rs_rank": rs.iloc[t].values,
        "close": C.iloc[t].values, "atr_pct": atr.iloc[t].values,
        "ret6": 100 * m6.iloc[t].values, "ret12": 100 * m12.iloc[t].values,
        "vol": 100 * vol.iloc[t].values,
        "dist52": 100 * (C.iloc[t] / hi52.iloc[t] - 1).values,
        "eligible": ok.iloc[t].values})
    nifty = float(BM.iloc[-1])
    nifty200 = float(BM.rolling(200).mean().iloc[-1])
    return out, C.index[-1].date(), nifty, nifty200


def select(tab, sectors):
    """Rank eligible stocks; take the top SLOTS with the sector cap."""
    r = tab[tab["eligible"] & tab["score"].notna()].sort_values(
        "score", ascending=False).reset_index(drop=True)
    r["rank"] = np.arange(1, len(r) + 1)
    r["sector"] = r["symbol"].map(lambda s: sectors.get(s, "?"))
    picks, count, skipped = [], {}, []
    for _, x in r.iterrows():
        if len(picks) >= SLOTS:
            break
        sec = x["sector"]
        if SECTOR_CAP and sec != "?" and count.get(sec, 0) >= SECTOR_CAP:
            skipped.append(x["symbol"])
            continue
        picks.append(x)
        count[sec] = count.get(sec, 0) + 1
    return r, pd.DataFrame(picks).reset_index(drop=True), skipped


# ================================================================== excel
def _read_sheet(path, name):
    try:
        return pd.read_excel(path, sheet_name=name)
    except Exception:
        return pd.DataFrame()


def write_sheets(path, top, allrank, swing, fund_status, regime_red,
                 old_action, banner, sells):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    F = "Arial"
    head = Font(name=F, bold=True, color="FFFFFF")
    fill_h = PatternFill("solid", fgColor="7030A0")
    super_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    act_fill = PatternFill("solid", fgColor="FFF2CC")

    wb = load_workbook(path)
    for name in ("Momentum_Top20", "Strategy_Comparison"):
        if name in wb.sheetnames:
            del wb[name]

    def header(ws, cols):
        for i, (h, w) in enumerate(cols, 1):
            c = ws.cell(row=1, column=i, value=h)
            c.font, c.fill = head, fill_h
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 36
        ws.freeze_panes = "B2"

    def put(ws, r, c, v, fmt=None):
        if isinstance(v, (np.floating, np.integer)):
            v = float(v)
        if isinstance(v, float) and v != v:
            v = None
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=F)
        if fmt:
            cell.number_format = fmt
        return cell

    def notes(ws, start, lines):
        for i, t in enumerate(lines):
            c = ws.cell(row=start + i, column=1, value=t)
            c.font = Font(name=F, italic=True, bold=(i == 0),
                          color="C00000" if (i == 0 and regime_red) else "404040")

    # ---------------------------------------------------- Momentum_Top20
    ws = wb.create_sheet("Momentum_Top20")
    cols = [("Symbol", 13), ("Mom Rank", 7), ("Momentum Score", 9),
            ("RS Rank", 7), ("Sector / Industry", 22), ("Last Close", 10),
            ("LTP", 10), ("Price Source", 8), ("ATR %", 7), ("Slot (Rs)", 9),
            ("Shares to Buy", 8), ("Amount (Rs)", 10), ("6m Ret %", 8),
            ("12m Ret %", 8), ("1y Vol %", 8), ("Dist 52W High %", 8),
            ("Held?", 7), ("Suggested", 12), ("Regime", 20)]
    header(ws, cols)
    for n, (_, x) in enumerate(top.iterrows(), start=2):
        vals = [x["symbol"], int(x["rank"]), x["score"], x["rs_rank"],
                x["sector"], x["close"], x["price"], x["px_src"], x["atr_pct"],
                x["slot"], int(x["shares"]), x["amount"], x["ret6"], x["ret12"],
                x["vol"], x["dist52"], "YES" if x["held"] else "",
                "HOLD" if x["held"] else "BUY (rebalance)",
                RED if regime_red else "OK"]
        fmts = [None, "0", "0.00", "0", None, "#,##0.0", "#,##0.0", None,
                "0.0", "#,##0", "0", "#,##0", "0.0", "0.0", "0.0", "0.0",
                None, None, None]
        for i, (v, f) in enumerate(zip(vals, fmts), 1):
            c = put(ws, n, i, v, f)
            if regime_red and i == len(vals):
                c.fill = red_fill
    last = ws.max_row
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), max(last, 2))
    lines = [
        banner,
        "Score = 0.5 z(6m ret/1y vol) + 0.5 z(12m ret/1y vol) among >= Rs 10k "
        "Cr stocks with > Rs 5 Cr median turnover. Max %s per industry."
        % SECTOR_CAP,
        "Trade on the 1st trading day of the month at the open (ranks use the "
        "last close). Keep a holding while its rank <= %d; sell when it drops "
        "below." % (BUFFER * SLOTS),
        "Shares to Buy = floor(slot / LTP), slot = Rs %s / %d%s." %
        (format(CAPITAL, ","), SLOTS, " x volatility factor" if ATR_SIZING else ""),
        "Backtest 2013-26 (tax paid): ~18-20%/yr with sector cap vs Nifty ~10%, "
        "max DD ~-37%. 2013-19 was only ~Nifty. Realistic: 13-16%.",
    ]
    if regime_red:
        lines.insert(1, "MARKET REGIME RED: Nifty < 200-DMA. New buys NOT "
                     "recommended. (Backtest: blocking buys in red markets cost "
                     "~2.5%/yr -- your call.)")
    if sells:
        lines.append("Momentum holdings now OUTSIDE the top %d -> SELL at the "
                     "next rebalance: %s" % (BUFFER * SLOTS, ", ".join(sells)))
    notes(ws, last + 2, lines)

    # ---------------------------------------------------- Strategy_Comparison
    wc = wb.create_sheet("Strategy_Comparison")
    ccols = [("Ticker", 13), ("Strategy Overlap", 14), ("W+TT Status", 8),
             ("RS Rank", 7), ("Mom Rank", 7), ("Momentum Score", 9),
             ("ATR %", 7), ("Dist 52W High %", 9), ("Sector / Industry", 22),
             ("Fundamental Status", 11), ("Regime", 20), ("LTP", 10),
             ("Shares (Rs slot)", 9), ("Action", 10)]
    header(wc, ccols)
    wc.cell(row=1, column=len(ccols)).fill = PatternFill("solid",
                                                         fgColor="BF8F00")
    mom = set(top["symbol"])
    sw = {}
    if not swing.empty and "Symbol" in swing:
        s2 = swing[pd.to_numeric(swing.get("RS Rank"), errors="coerce").notna()]
        sw = {str(r["Symbol"]).upper(): r for _, r in s2.iterrows()}
    by = allrank.set_index("symbol")
    rows = []
    for s in list(dict.fromkeys(list(top["symbol"]) + list(sw))):
        a = by.loc[s] if s in by.index else None
        overlap = ("Super-Buy" if s in mom and s in sw else
                   "Momentum only" if s in mom else "W+TT only")
        t = top[top["symbol"] == s]
        price = float(t["price"].iloc[0]) if len(t) else \
            (a["price"] if a is not None else np.nan)
        slot = float(t["slot"].iloc[0]) if len(t) else CAPITAL / SLOTS
        rows.append({
            "Ticker": s, "Strategy Overlap": overlap,
            "W+TT Status": sw[s].get("Action", "") if s in sw else "",
            "RS Rank": (float(sw[s]["RS Rank"]) if s in sw else
                        (a["rs_rank"] if a is not None else np.nan)),
            "Mom Rank": a["rank"] if a is not None else np.nan,
            "Momentum Score": a["score"] if a is not None else np.nan,
            "ATR %": a["atr_pct"] if a is not None else np.nan,
            "Dist 52W High %": a["dist52"] if a is not None else np.nan,
            "Sector / Industry": a["sector"] if a is not None else "?",
            "Fundamental Status": fund_status.get(s, "(run fundamentals)"),
            "Regime": RED if regime_red else "OK", "LTP": price,
            "Shares (Rs slot)": shares_for(slot, price),
            "Action": old_action.get(s, "")})
    order = {"Super-Buy": 0, "Momentum only": 1, "W+TT only": 2}
    rows.sort(key=lambda x: (order[x["Strategy Overlap"]],
                             x["Mom Rank"] if x["Mom Rank"] == x["Mom Rank"]
                             else 9999))
    fmts = [None, None, None, "0", "0", "0.00", "0.0", "0.0", None, None, None,
            "#,##0.0", "0", None]
    for n, x in enumerate(rows, start=2):
        for i, ((h, _), f) in enumerate(zip(ccols, fmts), 1):
            c = put(wc, n, i, x[h], f)
            if h == "Strategy Overlap" and x[h] == "Super-Buy":
                c.fill = super_fill
            if h == "Regime" and regime_red:
                c.fill = red_fill
            if h == "Action":
                c.fill = act_fill
    last = wc.max_row
    wc.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(ccols)),
                                      max(last, 2))
    notes(wc, last + 2, [
        banner,
        "Type BUY in the yellow Action column for what you actually bought, "
        "save the file, then run: rbtrack",
        "Super-Buy = in BOTH the W+TT Swing list and the Momentum top %d. "
        "(Not backtested as a separate strategy.)" % SLOTS,
        "Momentum only -> tracked as Strategy=Momentum (exit: rank > %d at the "
        "monthly rebalance). W+TT only -> swing leg (20%% stop / 40w MA)."
        % (BUFFER * SLOTS),
        "Fundamental Status is filled by fundamentals.py (information only -- "
        "the fundamental gate did not help in the 2018-26 backtest)."])
    try:
        wb.save(path)
    except PermissionError:
        print("\n! %s is open in Excel. Close it and run momentum_screener.py "
              "again." % os.path.basename(path))
        sys.exit(1)


# ================================================================== main
def main():
    tok = get_token()
    print("MOMENTUM SCREENER -- NSE-style momentum, top %d, max %s per sector"
          % (SLOTS, SECTOR_CAP))
    frames, bm, caps, live, warns, want = load_data(tok)
    P = panels(frames, bm)
    tab, data_day, nifty, nifty200 = compute(P)
    sectors = industry_map()
    allrank, top, skipped = select(tab, sectors)
    if top.empty:
        print("No eligible stocks -- check data.")
        sys.exit(1)

    held = momentum_holdings()
    price = {s.upper(): v for s, v in live.items()}
    full = tab.merge(allrank[["symbol", "rank"]], on="symbol", how="left")
    full["sector"] = full["symbol"].map(lambda s: sectors.get(s, "?"))
    for d in (allrank, top, full):
        d["price"] = [price.get(s, c) for s, c in zip(d["symbol"], d["close"])]
        d["px_src"] = ["live" if s in price else "close" for s in d["symbol"]]
    med = float(np.nanmedian(top["atr_pct"]))
    top["slot"] = [slot_amount(a, med) for a in top["atr_pct"]]
    top["shares"] = [shares_for(a, p) for a, p in zip(top["slot"], top["price"])]
    top["amount"] = top["shares"] * top["price"]
    top["held"] = top["symbol"].isin(held)
    ranks = dict(zip(allrank["symbol"], allrank["rank"]))
    sells = sorted(s for s in held if ranks.get(s, 10 ** 9) > BUFFER * SLOTS)

    regime_red = nifty < nifty200
    allrank[["symbol", "rank", "score", "sector"]].assign(
        date=str(data_day)).to_csv(RANKS_FILE, index=False)

    # terminal
    print("\n" + "=" * 70)
    print(" DATA: last complete session %s | prices: %s"
          % (data_day, "Dhan live" if live else "last close"))
    if data_day < want and not live:
        print(" !!! DATA IS BEHIND (expected %s) -- do not trade from this." % want)
    print("=" * 70)
    if regime_red:
        print("\n\033[1m\033[91m" + "*" * 70)
        print(" MARKET REGIME RED: NIFTY < 200-DMA. NEW BUYS NOT RECOMMENDED.")
        print(" Nifty %.1f vs 200-DMA %.1f" % (nifty, nifty200))
        print(" (Backtest: skipping buys in red markets cut return ~2.5%/yr "
              "but also cut max DD.)")
        print("*" * 70 + "\033[0m")
    else:
        print("\n Market regime OK: Nifty %.1f > 200-DMA %.1f" % (nifty, nifty200))
    print("\n==== MOMENTUM TOP %d (eligible: %d) ====" % (SLOTS, len(allrank)))
    show = top[["rank", "symbol", "score", "sector", "price", "atr_pct",
                "shares", "amount"]].copy()
    show["held"] = np.where(top["held"], "HELD", "")
    print(show.to_string(index=False, float_format=lambda x: "%.2f" % x))
    print("  total to deploy: Rs %s" % format(int(top["amount"].sum()), ","))
    if skipped:
        print("  skipped by the %s-per-sector cap: %s"
              % (SECTOR_CAP, ", ".join(skipped[:10])))
    if sells:
        print("\n  SELL at the next rebalance (momentum holdings now ranked "
              "> %d): %s" % (BUFFER * SLOTS, ", ".join(sells)))
    today = ds.now_ist().date()
    first_td = today.day <= 3 and today.weekday() < 5
    print("\n  Rebalance = 1st trading day of the month (%s)." %
          ("maybe TODAY" if first_td else "not today"))

    path = todays_report()
    if not path:
        print("\n! No RB_Screener report found -- run daily_screener.py first.")
        sys.exit(1)
    swing = _read_sheet(path, "Swing")
    fund = _read_sheet(path, "Fundamentals")
    fund_status = {}
    if not fund.empty and "Symbol" in fund:
        col = next((c for c in fund.columns if str(c).startswith("Swing Check")),
                   None)
        if col:
            fund_status = dict(zip(fund["Symbol"].astype(str).str.upper(),
                                   fund[col].astype(str)))
    old = _read_sheet(path, "Strategy_Comparison")
    old_action = {}
    if not old.empty and "Ticker" in old and "Action" in old:
        old_action = {str(t).upper(): str(a) for t, a in
                      zip(old["Ticker"], old["Action"])
                      if isinstance(a, str) and a.strip()}
    banner = ("Momentum %s | data to %s | %s" %
              (today, data_day, RED if regime_red else "Market regime OK"))
    write_sheets(path, top, full, swing, fund_status, regime_red,
                 old_action, banner, sells)
    print("\nExcel: %s (sheets Momentum_Top20, Strategy_Comparison)" % path)
    if old_action:
        print("  kept your existing Action entries: %s"
              % ", ".join("%s=%s" % kv for kv in old_action.items()))


if __name__ == "__main__":
    main()
