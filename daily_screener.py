#!/usr/bin/env python3
"""
DAILY SWING + INVESTING SCREENER  (NSE)   --  v4
================================================

v4: results go to ONE Excel file with two sheets, "Swing" and
    "Investing" (reports/RB_Screener_YYYY-MM-DD.xlsx).
    Dhan client ID and token expiry are read from the token itself --
    nothing to configure except dhan_token.txt.

v3: UNIVERSE = every NSE-listed company with market cap >= Rs 10,000 Cr,
    read fresh each day from NSE's own MCAP file (inside the daily
    PR bhavcopy zip). All files live in ~/Desktop/RB_Screener.

WHAT CHANGED FROM v1
  * Signal rules now match the backtest EXACTLY (v1 used a looser
    Weinstein rule, so its list was not what we tested).
  * The free history source can lag by several days. v2 fills the
    missing days from Dhan and checks every stock at TODAY's price
    (live LTP during market hours, last close after hours).

RULES (straight from the backtest -- do not edit)
  Weinstein entry : close > 30-week MA, 30-week MA rising,
                    close = 30-day high, volume > 2x 50-day avg,
                    RS rank > 50, 60-day median turnover > Rs 5 Cr
  Minervini TT    : close > 50DMA > 150DMA > 200DMA, 200DMA rising,
                    >30% above 52w low, within 25% of 52w high, RS >= 70
  Entry           : Weinstein AND Trend Template on the same day,
                    buy next open
  Swing exit      : 20% below entry, or close below 40-week MA
  Tested result   : backtest.py, >= Rs 10,000 Cr universe point-in-time,
                    2013-2026: 1773 trades, win 39%, avg trade +11.8%,
                    PF 2.43; 20-slot portfolio ~12.5%/yr vs Nifty 10.5%
                    (price only), max drawdown ~-40% (see CLAUDE.md)

TODAY'S FIT CHECK (per stock)
  FIT    : signal is fresh or price is still near the signal price,
           AND at today's price the stock still passes the Trend
           Template, is above its 40-week MA, and the original 20%
           stop has not been hit.
  LATE   : still passes everything but price has run >10% past the
           signal. The backtest only tested next-day entry -- late
           entries are NOT backed by the numbers.
  NO FIT : fails at today's price (reason printed).

SETUP (put this file, position_tracker.py and dhan_token.txt in
       ~/Desktop/RB_Screener)
    pip3 install pandas numpy requests openpyxl
    dhan_token.txt : today's Dhan access token on line 1 (24h validity);
                     optional line 2 = a name for the account
  Without a token the screener still runs, but on possibly stale data,
  and it will say so loudly.

RUN
    python3 ~/Desktop/RB_Screener/daily_screener.py
"""

import warnings
warnings.filterwarnings("ignore")

import os
import io
import sys
import json
import re
import time
import base64
import zipfile
import datetime as dt

import numpy as np
import pandas as pd
import requests

# ------------------------------------------------------------------ config
CLIENT_ID = ""        # optional: normally read from the token itself

LOOKBACK = 20          # how many sessions back a signal may be
LATE_PCT = 10.0        # above this % past signal price -> LATE
INV_MAX_EXT = 25.0     # investing list: skip if >25% past signal
KEEP_ROWS = 900        # history rows kept per stock (need ~460)

MCAP_MIN_CR = 10000   # universe: market cap at or above this (Rs crore)

# everything lives here, wherever the script is run from
HERE = os.path.join(os.path.expanduser("~"), "Desktop", "RB_Screener")
DATA = os.path.join(HERE, "data")
REPORTS = os.path.join(HERE, "reports")
for _d in (HERE, DATA, REPORTS):
    os.makedirs(_d, exist_ok=True)
TOKEN_FILE = os.path.join(HERE, "dhan_token.txt")
MCAP_CACHE = os.path.join(DATA, "_nse_mcap_latest.csv")
SCRIP_FILE = os.path.join(DATA, "_dhan_scrip_master.csv")

EOD_BASE = "https://raw.githubusercontent.com/BennyThadikaran/eod2_data/main/daily/"
SCRIP_URL = "https://images.dhan.co/api-data/api-scrip-master.csv"
DHAN_HIST = "https://api.dhan.co/v2/charts/historical"
DHAN_LTP = "https://api.dhan.co/v2/marketfeed/ltp"
NSE_PR = "https://nsearchives.nseindia.com/archives/equities/bhavcopy/pr/PR%s.zip"
NSE_HDRS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; rv:109.0) "
                  "Gecko/20100101 Firefox/118.0",
    "Accept": "*/*",
    "Accept-Language": "en-US,en;q=0.5",
    "Referer": "https://www.nseindia.com/",
}

# fallback only: the 173-stock backtest universe, used if NSE's MCAP
# file cannot be downloaded and there is no cached copy
FALLBACK = """
aartiind abb acc adanient adanigreen adaniports adanipower alkem ambujacem
apollohosp apollotyre ashokley asianpaint astral atul aubank auropharma
axisbank bajaj-auto bajajfinsv bajajhldng bajfinance balkrisind bandhanbnk
bankbaroda bataindia bel bergepaint bharatforg bhartiartl bhel biocon
boschltd bpcl britannia canbk chamblfert cholafin cipla coalindia coforge
colpal concor coromandel crompton cumminsind dabur deepakntr divislab dixon
dlf dmart drreddy eichermot escorts exideind federalbnk fortis gail glenmark
gmrairport godrejcp godrejprop grasim hal havells hcltech hdfcbank hdfclife
heromotoco hindalco hindpetro hindunilvr hindzinc icicibank icicigi
icicipruli idfcfirstb igl indhotel indigo indusindbk infy ioc ipcalab irctc
itc jindalstel jkcement jswsteel jublfood kajariacer kotakbank kpittech
lalpathlab lauruslabs lichsgfin lodha lt ltm ltts lupin m&m manappuram
marico maruti maxhealth metropolis mgl motherson mphasis mrf muthootfin
nationalum naukri navinfluor nestleind nmdc ntpc oberoirlty ongc pageind
persistent petronet pfc phoenixltd pidilitind piind pnb polycab powergrid
prestige pvrinox ramcocem recltd relaxo reliance sail sbicard sbilife sbin
shreecem shriramfin siemens srf sundarmfin sunpharma suntv supremeind
syngene tatacomm tataconsum tataelxsi tatapower tatasteel tcs techm tiindia
titan tmpv torntpharm trent tvsmotor ubl ultracemco unitdspr upl vedl voltas
whirlpool wipro zeel zyduslife
""".split()
BENCH = "nifty 50"

IST = dt.timezone(dt.timedelta(hours=5, minutes=30))


# ================================================================== helpers
def now_ist():
    return dt.datetime.now(IST)


def last_expected_session():
    """Most recent weekday whose session should be complete."""
    n = now_ist()
    d = n.date()
    if n.weekday() >= 5 or (n.hour, n.minute) < (15, 45):
        d = d - dt.timedelta(days=1)
    while d.weekday() >= 5:
        d = d - dt.timedelta(days=1)
    return d


def market_open():
    n = now_ist()
    return n.weekday() < 5 and (9, 15) <= (n.hour, n.minute) < (15, 30)


def read_token_file():
    """dhan_token.txt -> (token, name, client_id_written).

    Either the plain token (one line, as before), or labelled lines:
        Client ID: 1100123456
        Name: Ashutosh Main
        Token: eyJ0eXAi...
    The token is recognised by its shape (a long x.y.z string) with or
    without the label; a plain 2nd line is taken as the name."""
    tok, name, cid, plain = None, "", "", []
    if not os.path.exists(TOKEN_FILE):
        return None, "", ""
    for line in open(TOKEN_FILE).read().splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(r"(?i)^(client\s*id|name|token)\s*[:=\-]\s*(.*)$", line)
        key, val = (m.group(1).lower(), m.group(2).strip()) if m else ("", line)
        if val in ("", "-"):
            continue
        if key.startswith("client"):
            cid = val
        elif key == "name":
            name = val
        elif tok is None and (key == "token" or
                              (val.count(".") == 2 and len(val) > 40
                               and " " not in val)):
            tok = val
        elif not key:
            plain.append(val)
    if not name and plain:
        name = plain[0]
    return tok, name[:40], cid


def read_token():
    return read_token_file()[0]


def token_info(tok):
    """Dhan tokens are JWTs: the client id and expiry time are inside.
    Returns (client_id, expiry_datetime_ist) or ("", None)."""
    try:
        part = tok.split(".")[1]
        part += "=" * (-len(part) % 4)
        d = json.loads(base64.urlsafe_b64decode(part))
        exp = d.get("exp")
        exp_dt = dt.datetime.fromtimestamp(int(exp), IST) if exp else None
        return str(d.get("dhanClientId") or ""), exp_dt
    except Exception:
        return "", None


def dhan_headers(tok):
    return {"access-token": tok, "client-id": CLIENT_ID,
            "Content-Type": "application/json", "Accept": "application/json"}


# ================================================================== universe
def _col(cols, *keys, avoid=()):
    for c in cols:
        lc = c.lower()
        if any(k in lc for k in keys) and not any(a in lc for a in avoid):
            return c
    return None


def parse_mcap(df):
    """NSE MCAP csv -> DataFrame(symbol, mcap_cr). Column names are
    matched loosely so small header changes do not break it."""
    df.columns = [str(c).strip() for c in df.columns]
    c_sym = _col(df.columns, "symbol")
    c_ser = _col(df.columns, "series")
    c_cat = _col(df.columns, "category")
    c_cap = _col(df.columns, "market cap", "mkt cap", "mcap", avoid=("date",))
    if not c_sym or not c_cap:
        return None
    d = pd.DataFrame({"symbol": df[c_sym].astype(str).str.strip()})
    cap = pd.to_numeric(df[c_cap].astype(str).str.replace(",", "")
                        .str.strip(), errors="coerce")
    if cap.max() > 1e8:          # file is in rupees -> convert to crore
        cap = cap / 1e7
    d["mcap_cr"] = cap
    keep = pd.Series(True, index=d.index)
    if c_ser:
        keep &= df[c_ser].astype(str).str.strip().isin(["EQ", "BE", "BZ"])
    if c_cat:
        keep &= df[c_cat].astype(str).str.strip().isin(["Listed", "Permitted"])
    keep &= ~d["symbol"].str.contains(r"-RE\d*$", na=False)
    d = d[keep].dropna()
    return (d.sort_values("mcap_cr", ascending=False)
            .drop_duplicates("symbol").reset_index(drop=True))


def fetch_mcap():
    """Latest NSE market-cap table. Tries the last ~10 weekdays, then the
    cached copy. Returns (df, date, source)."""
    day = now_ist().date()
    for back in range(0, 14):
        d = day - dt.timedelta(days=back)
        if d.weekday() >= 5:
            continue
        try:
            r = requests.get(NSE_PR % d.strftime("%d%m%y"), headers=NSE_HDRS,
                             timeout=30)
        except Exception:
            continue
        if r.status_code != 200 or r.content[:2] != b"PK":
            continue
        try:
            zf = zipfile.ZipFile(io.BytesIO(r.content))
            name = next((n for n in zf.namelist()
                         if n.lower().endswith(".csv") and "mcap" in n.lower()),
                        None)
            if not name:
                continue
            m = parse_mcap(pd.read_csv(zf.open(name)))
        except Exception:
            continue
        if m is None or m.empty:
            continue
        m["asof"] = d.isoformat()
        m.to_csv(MCAP_CACHE, index=False)
        return m, d, "NSE"
    if os.path.exists(MCAP_CACHE):
        m = pd.read_csv(MCAP_CACHE)
        return m, m["asof"].iloc[0], "cached copy"
    return None, None, None


def build_universe(warns):
    m, asof, src = fetch_mcap()
    if m is None:
        print("  !!! Could not get NSE market-cap file and no cached copy.")
        print("  !!! Falling back to the 173-stock backtest list.")
        return FALLBACK, {}
    big = m[m["mcap_cr"] >= MCAP_MIN_CR]
    print("  NSE market cap (%s, %s): %d companies >= Rs %s Cr"
          % (src, asof, len(big), format(MCAP_MIN_CR, ",")))
    if src != "NSE":
        warns.append("market-cap list is a cached copy from %s" % asof)
    caps = dict(zip(big["symbol"].str.lower(), big["mcap_cr"]))
    return list(caps), caps


# ================================================================== history
def fetch_eod(sym):
    """Adjusted daily history from the free eod2 repo, cached ~20h."""
    path = os.path.join(DATA, sym + ".csv")
    fresh = (os.path.exists(path) and
             time.time() - os.path.getmtime(path) < 20 * 3600)
    if not fresh:
        try:
            url = EOD_BASE + requests.utils.quote(sym + ".csv")
            r = requests.get(url, timeout=30)
            if r.status_code == 200:
                open(path, "wb").write(r.content)
        except Exception:
            pass
    if not os.path.exists(path):
        return None
    df = pd.read_csv(path, parse_dates=["Date"], index_col="Date")
    df = df[["Open", "High", "Low", "Close", "Volume"]]
    df = df[~df.index.duplicated(keep="last")].sort_index()
    return df.iloc[-KEEP_ROWS:]


# ================================================================== dhan
def load_scrip_map():
    """NSE equity symbol -> Dhan securityId. Cached for 7 days."""
    stale = (not os.path.exists(SCRIP_FILE) or
             time.time() - os.path.getmtime(SCRIP_FILE) > 7 * 86400)
    if stale:
        try:
            print("  downloading Dhan scrip master (once a week) ...")
            r = requests.get(SCRIP_URL, timeout=120)
            if r.status_code == 200:
                open(SCRIP_FILE, "wb").write(r.content)
        except Exception as e:
            print("  ! could not download scrip master:", e)
    if not os.path.exists(SCRIP_FILE):
        return {}
    sm = pd.read_csv(SCRIP_FILE, low_memory=False)
    c = {k.upper(): k for k in sm.columns}
    need = ["SEM_EXM_EXCH_ID", "SEM_SEGMENT", "SEM_TRADING_SYMBOL",
            "SEM_SMST_SECURITY_ID"]
    if not all(n in c for n in need):
        print("  ! scrip master format changed -- Dhan fill disabled")
        return {}
    m = sm[(sm[c["SEM_EXM_EXCH_ID"]] == "NSE") & (sm[c["SEM_SEGMENT"]] == "E")]
    if "SEM_SERIES" in c:
        m = m[m[c["SEM_SERIES"]].isin(["EQ", "BE"])]
    return dict(zip(m[c["SEM_TRADING_SYMBOL"]].astype(str).str.upper(),
                    m[c["SEM_SMST_SECURITY_ID"]].astype(str)))


def dhan_daily(tok, sec_id, seg, instr, frm, to):
    body = {"securityId": str(sec_id), "exchangeSegment": seg,
            "instrument": instr, "expiryCode": 0, "oi": False,
            "fromDate": frm.isoformat(), "toDate": to.isoformat()}
    r = requests.post(DHAN_HIST, headers=dhan_headers(tok), json=body,
                      timeout=30)
    if r.status_code in (401, 403):
        raise PermissionError("HTTP %d %s" % (r.status_code, r.text[:160]))
    j = r.json()
    if not isinstance(j, dict) or not j.get("timestamp"):
        return None
    idx = (pd.to_datetime(j["timestamp"], unit="s", utc=True)
           .tz_convert("Asia/Kolkata").normalize().tz_localize(None))
    df = pd.DataFrame({"Open": j["open"], "High": j["high"], "Low": j["low"],
                       "Close": j["close"], "Volume": j["volume"]},
                      index=idx)
    df.index.name = "Date"
    return df[~df.index.duplicated(keep="last")]


def dhan_ltp(tok, sec_ids):
    out = {}
    ids = [int(s) for s in sec_ids]
    for i in range(0, len(ids), 900):
        chunk = ids[i:i + 900]
        r = requests.post(DHAN_LTP, headers=dhan_headers(tok),
                          json={"NSE_EQ": chunk}, timeout=30)
        if r.status_code in (401, 403):
            raise PermissionError("HTTP %d %s" % (r.status_code, r.text[:160]))
        d = (r.json() or {}).get("data", {}).get("NSE_EQ", {})
        for k, v in d.items():
            p = v.get("last_price") if isinstance(v, dict) else None
            if p:
                out[str(k)] = float(p)
        time.sleep(1.1)
    return out


def gap_fill(df, tok, sec_id, seg, instr, want, warns, name):
    """Append missing sessions from Dhan. Skips if a split/bonus is
    suspected (Dhan prices are unadjusted)."""
    last = df.index[-1].date()
    if last >= want:
        return df
    frm = last + dt.timedelta(days=1)
    to = now_ist().date() + dt.timedelta(days=1)
    add = dhan_daily(tok, sec_id, seg, instr, frm, to)
    time.sleep(0.25)
    if add is None or add.empty:
        return df
    add = add[add.index > df.index[-1]]
    # today's bar during market hours is incomplete -> drop it
    if market_open():
        add = add[add.index.date < now_ist().date()]
    if add.empty:
        return df
    jump = add["Close"].iloc[0] / df["Close"].iloc[-1]
    if jump > 1.4 or jump < 0.6:
        warns.append("%s: %.0f%% gap between sources -- possible split/"
                     "bonus, Dhan fill skipped" % (name, (jump - 1) * 100))
        return df
    return pd.concat([df, add])


# ================================================================== excel
def write_excel(stamp, swing, inv, banner):
    """Two sheets: Swing, Investing. Derived columns are live formulas, so
    if you type your own price into the Price column they update."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.comments import Comment
    except ImportError:
        print("\n! openpyxl not installed -> run:  pip3 install openpyxl")
        print("  Saving CSV files instead.")
        swing.to_csv(os.path.join(REPORTS, "swing_%s.csv" % stamp), index=False)
        inv.to_csv(os.path.join(REPORTS, "investing_%s.csv" % stamp),
                   index=False)
        return None

    F = "Arial"
    head_font = Font(name=F, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    body = Font(name=F)
    fills = {"BUY": PatternFill("solid", fgColor="C6EFCE"),
             "LATE": PatternFill("solid", fgColor="FCE4D6")}
    order = {"BUY": 0, "FIT": 1, "LATE": 2}

    def action(r):
        if r["status"] == "LATE":
            return "LATE"
        return "BUY" if int(r["bars_ago"]) == 0 else "FIT"

    def ranked(df):
        d = df.copy()
        d["act"] = d.apply(action, axis=1)
        d["_o"] = d["act"].map(order)
        return d.sort_values(["_o", "rs_rank"], ascending=[True, False])

    def setup(ws, headers, widths):
        ws.append(headers)
        for i, c in enumerate(ws[1], 1):
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
            ws.column_dimensions[c.column_letter].width = widths[i - 1]
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"

    def add_notes(ws, start, lines):
        for i, t in enumerate(lines):
            c = ws.cell(row=start + i, column=1, value=t)
            warn = i == 0 and "!!!" in t
            c.font = Font(name=F, italic=True, bold=(i == 0),
                          color="C00000" if warn else "404040")

    wb = Workbook()

    # ---------------------------------------------------------- Swing
    ws = wb.active
    ws.title = "Swing"
    setup(ws, ["Symbol", "Action", "RS Rank", "Mcap (Rs Cr)", "Signal Date",
               "Days Since Signal", "Signal Price", "Price", "Price Source",
               "% vs Signal", "Stop (below Price)", "Exit: 40w MA",
               "Effective Exit (higher of the two)", "Risk to Exit %"],
          [13, 8, 8, 12, 12, 10, 11, 11, 9, 10, 12, 12, 15, 10])
    # Stop % input sits BELOW the table (not in row 1), so fundamentals.py
    # can add columns on the right and sorting/filtering stays safe.
    stop_row = len(swing) + 3
    stop_ref = "$F$%d" % stop_row
    for n, (_, r) in enumerate(ranked(swing).iterrows(), start=2):
        ws.append([r["symbol"], r["act"], int(r["rs_rank"]),
                   r["mcap_cr"], r["signal_date"], int(r["bars_ago"]),
                   float(r["signal_px"]), float(r["price"]), r["px_src"],
                   "=H%d/G%d-1" % (n, n),
                   "=H%d*(1-%s)" % (n, stop_ref),
                   float(r["exit_40w"]),
                   "=MAX(K%d,L%d)" % (n, n),
                   "=(H%d-M%d)/H%d" % (n, n, n)])
        for c in ws[n]:
            c.font = body
        ws["E%d" % n].number_format = "yyyy-mm-dd"
        ws["D%d" % n].number_format = "#,##0"
        for col in "GHKLM":
            ws["%s%d" % (col, n)].number_format = "#,##0.0"
        for col in "JN":
            ws["%s%d" % (col, n)].number_format = "0.0%"
        if r["act"] in fills:
            for c in ws[n][:14]:
                c.fill = fills[r["act"]]
    last = ws.max_row
    ws.auto_filter.ref = "A1:N%d" % max(last, 2)
    ws["A%d" % stop_row] = "Stop % below entry (edit to change the Stop column):"
    ws["A%d" % stop_row].font = Font(name=F, bold=True)
    c = ws["F%d" % stop_row]
    c.value, c.number_format = 0.20, "0%"
    c.font = Font(name=F, bold=True, color="0000FF")
    c.comment = Comment("Backtest rule: stop 20% below entry. Changing this "
                        "changes the Stop column.", "RB")
    add_notes(ws, stop_row + 2, [
        banner,
        "BUY = signal on the last session; the backtest bought at the NEXT "
        "day's open.",
        "FIT = older signal that still passes at today's price. The backtest "
        "never entered late -- the older, the weaker.",
        "LATE = ran >10% past the signal. Not tested. Skip.",
        "Exit = whichever comes first: the Stop, or a close below the 40-week "
        "MA. The 40w MA rises over time, so it is your trailing stop.",
        "NO profit target. In the earlier (173-stock) test a 25% target cut "
        "the average trade from +20.8% to +3.9%.",
        "Already bought? Type YOUR buy price into Price and the Stop and "
        "Risk columns recalculate.",
    ])

    # ---------------------------------------------------------- Investing
    wi = wb.create_sheet("Investing")
    setup(wi, ["Symbol", "Action", "RS Rank", "Mcap (Rs Cr)", "Signal Date",
               "Signal Price", "Price", "% vs Signal", "30w MA",
               "% above 30w MA"],
          [13, 8, 8, 12, 12, 11, 11, 10, 11, 11])
    for n, (_, r) in enumerate(ranked(inv).iterrows(), start=2):
        wi.append([r["symbol"], r["act"], int(r["rs_rank"]), r["mcap_cr"],
                   r["signal_date"], float(r["signal_px"]), float(r["price"]),
                   "=G%d/F%d-1" % (n, n), float(r["ma30w"]),
                   "=G%d/I%d-1" % (n, n)])
        for c in wi[n]:
            c.font = body
        wi["E%d" % n].number_format = "yyyy-mm-dd"
        wi["D%d" % n].number_format = "#,##0"
        for col in "FGI":
            wi["%s%d" % (col, n)].number_format = "#,##0.0"
        for col in "HJ":
            wi["%s%d" % (col, n)].number_format = "0.0%"
        if r["act"] in fills:
            for c in wi[n][:10]:
                c.fill = fills[r["act"]]
    last = wi.max_row
    wi.auto_filter.ref = "A1:J%d" % max(last, 2)
    add_notes(wi, last + 2, [
        banner,
        "Same entry signal as Swing, but held for the long term.",
        "Only exit: 10 straight closes below a FALLING 30-week MA (Stage 4 "
        "breakdown). A 20% dip is NOT an exit here.",
        "Backtest (>= Rs 10,000 Cr, 2013-2026): avg trade +17%, win 47%, "
        "~7% of trades doubled; 20-slot portfolio ~14.8%/yr, max DD -40%.",
        "Most of that edge came from 2020-26; in 2013-19 it only matched "
        "Nifty. Stocks delisted since are missing -> results a bit too good.",
    ])

    path = os.path.join(REPORTS, "RB_Screener_%s.xlsx" % stamp)
    try:
        wb.save(path)
    except PermissionError:           # file already open in Excel/Numbers
        path = os.path.join(REPORTS, "RB_Screener_%s_%s.xlsx"
                            % (stamp, now_ist().strftime("%H%M")))
        wb.save(path)
    return path


# ================================================================== main
def main():
    global CLIENT_ID
    import account                     # per-account folders (accounts/<ID>/)
    acc = account.activate()
    want = last_expected_session()
    tok = read_token()
    warns = []
    if tok:
        cid, exp = token_info(tok)
        if cid:
            CLIENT_ID = cid
        if exp and exp < now_ist():
            print("! Token in dhan_token.txt expired at %s. Paste a fresh one."
                  % exp.strftime("%d %b %H:%M"))
            tok = None
        elif not CLIENT_ID:
            print("! Could not read the client id from the token -- is "
                  "dhan_token.txt holding the full token?")
            tok = None
        elif exp:
            print("  Dhan token OK, valid till %s" % exp.strftime("%d %b %H:%M"))

    print("Building universe ...")
    SYMBOLS, caps = build_universe(warns)
    print("Loading history (%d stocks) ..." % len(SYMBOLS))
    frames, missing, young = {}, [], []
    for i, s in enumerate(SYMBOLS, 1):
        if i % 10 == 0 or i == len(SYMBOLS):
            sys.stdout.write("\r  %3d/%d" % (i, len(SYMBOLS)))
            sys.stdout.flush()
        df = fetch_eod(s)
        if df is None:
            missing.append(s)
        elif len(df) < 300:
            young.append(s)
        else:
            frames[s] = df
    bm = fetch_eod(BENCH)
    print()
    print("  usable: %d | no history file: %d | listed < ~14 months: %d"
          % (len(frames), len(missing), len(young)))
    if missing:
        warns.append("no free history for: " + ", ".join(missing[:15]) +
                     (" ..." if len(missing) > 15 else ""))
    if young:
        warns.append("too new to judge (rules need 52w + 200DMA): " +
                     ", ".join(young[:15]) + (" ..." if len(young) > 15 else ""))
    if bm is None:
        print("Could not load Nifty 50 history. Check internet and retry.")
        sys.exit(1)

    src_last = max(f.index[-1] for f in frames.values()).date()
    print("  free source last date: %s   (expected: %s)" % (src_last, want))

    live = {}
    scrip = {}
    if tok:
        try:
            scrip = load_scrip_map()
            if src_last < want:
                print("  source is behind -- filling missing days from Dhan ...")
                bm = gap_fill(bm, tok, "13", "IDX_I", "INDEX", want, warns,
                              "NIFTY")
                for i, s in enumerate(list(frames), 1):
                    sid = scrip.get(s.upper())
                    if i % 10 == 0 or i == len(frames):
                        sys.stdout.write("\r    %3d/%d" % (i, len(frames)))
                        sys.stdout.flush()
                    if not sid:
                        warns.append("%s: not found in Dhan scrip master, "
                                     "not filled" % s)
                        continue
                    frames[s] = gap_fill(frames[s], tok, sid, "NSE_EQ",
                                         "EQUITY", want, warns, s)
                print()
            ids = {s: scrip[s.upper()] for s in frames if s.upper() in scrip}
            print("  fetching today's prices from Dhan ...")
            ltp = dhan_ltp(tok, list(ids.values()))
            live = {s: ltp[i] for s, i in ids.items() if i in ltp}
        except PermissionError as e:
            print("\n*** Dhan refused the request (%s)." % e)
            print("*** The token is not expired by its own clock, so check:")
            print("***   - token copied fully into dhan_token.txt")
            print("***   - Dhan Data API access is active on your account")
            print("*** Continuing WITHOUT Dhan -- prices below may be old. ***\n")
            live = {}
        except Exception as e:
            print("\n  ! Dhan step failed (%s). Continuing without it." % e)
            live = {}
    else:
        print("  (no usable Dhan token -- running on the free source only)")

    # ---------------------------------------------------------------- panel
    cal = bm.index
    for f in frames.values():
        cal = cal.union(f.index)
    cal = cal[cal >= bm.index[0]]
    CL = pd.DataFrame({s: f["Close"] for s, f in frames.items()}).reindex(cal)
    VO = pd.DataFrame({s: f["Volume"] for s, f in frames.items()}).reindex(cal)
    OP = pd.DataFrame({s: f["Open"] for s, f in frames.items()}).reindex(cal)
    CL = CL.ffill(limit=5)
    BM = bm["Close"].reindex(cal).ffill(limit=5)
    if bm.index[-1] < cal[-1]:
        warns.append("Nifty history ends %s, stocks end %s -- RS uses "
                     "last Nifty value" % (bm.index[-1].date(), cal[-1].date()))
    last_bar = cal[-1]
    data_date = last_bar.date()
    stale = data_date < want

    # ---------------------------------------------------------------- rules
    liq = (CL * VO).rolling(60).median() > 5e7
    ma50 = CL.rolling(50).mean()
    ma150 = CL.rolling(150).mean()
    ma200 = CL.rolling(200).mean()            # = 40-week MA
    ma30w = ma150                               # = 30-week MA
    hi52 = CL.rolling(252).max()
    lo52 = CL.rolling(252).min()
    rs6 = (CL / CL.shift(126)).div(BM / BM.shift(126), axis=0)
    rs = rs6.rank(axis=1, pct=True) * 100

    wein = ((CL > ma30w) & (ma30w > ma30w.shift(10)) & (rs > 50) &
            (CL >= CL.rolling(30).max()) &
            (VO > 2 * VO.rolling(50).mean()) & liq)
    tt = ((CL > ma50) & (ma50 > ma150) & (ma150 > ma200) &
          (ma200 > ma200.shift(22)) & (CL > 1.3 * lo52) &
          (CL > 0.75 * hi52) & (rs >= 70))
    entry = (wein & tt).fillna(False).astype(bool)

    # ---------------------------------------------------------------- judge
    recent = entry.iloc[-LOOKBACK:]
    rows, nofit = [], []
    for s in CL.columns:
        hits = recent.index[recent[s].values]
        if len(hits) == 0:
            continue
        sd = hits[-1]
        k = cal.get_loc(sd)
        sig_px = float(CL[s].iloc[k])
        bars_ago = len(cal) - 1 - k
        # backtest entry = next open; if not printed yet, use signal close
        e_px = float(OP[s].iloc[k + 1]) if k + 1 < len(cal) and \
            not np.isnan(OP[s].iloc[k + 1]) else sig_px
        orig_stop = e_px * 0.80

        px = live.get(s, float(CL[s].iloc[-1]))
        src = "live" if s in live else "close"
        m50, m150, m200 = ma50[s].iloc[-1], ma150[s].iloc[-1], ma200[s].iloc[-1]
        ext = (px / sig_px - 1) * 100

        fails = []
        if px <= orig_stop:
            fails.append("original 20%% stop (%.1f) already hit" % orig_stop)
        if px < m200:
            fails.append("below 40-week MA %.1f (swing exit level)" % m200)
        if not (px > m50):
            fails.append("below 50DMA %.1f" % m50)
        if not (m50 > m150 > m200):
            fails.append("MAs not stacked 50>150>200")
        if not (ma200[s].iloc[-1] > ma200[s].iloc[-23]):
            fails.append("200DMA not rising")
        if px < 0.75 * hi52[s].iloc[-1]:
            fails.append("more than 25% below 52w high")
        if px < 1.3 * lo52[s].iloc[-1]:
            fails.append("less than 30% above 52w low")
        if rs[s].iloc[-1] < 70:
            fails.append("RS rank fell to %.0f" % rs[s].iloc[-1])

        base = {"symbol": s.upper(),
                "mcap_cr": int(round(caps.get(s, np.nan))) if s in caps else None,
                "rs_rank": int(round(rs[s].iloc[-1])),
                "signal_date": sd.date(), "bars_ago": bars_ago,
                "signal_px": round(sig_px, 1), "price": round(px, 1),
                "px_src": src, "vs_signal_pct": round(ext, 1)}
        if fails:
            base["why"] = "; ".join(fails)
            nofit.append(base)
            continue
        base["status"] = "FIT" if ext <= LATE_PCT else "LATE"
        base["stop_if_buy_today"] = round(px * 0.80, 1)
        base["exit_40w"] = round(m200, 1)
        base["room_to_exit_pct"] = round((px / m200 - 1) * 100, 1)
        base["ma30w"] = round(m150, 1)
        rows.append(base)

    # ---------------------------------------------------------------- report
    stamp = str(now_ist().date())
    print("\n" + "=" * 70)
    print(" DATA: last complete session %s | prices: %s"
          % (data_date, "Dhan live/LTP" if live else "last close in data"))
    if stale and not live:
        print(" !!! DATA IS BEHIND (expected %s). Prices and stops below are"
              % want)
        print(" !!! OLD. Add today's Dhan token for a correct run.")
    elif stale:
        print(" ! Signals use data to %s (expected %s). Dhan may not have"
              % (data_date, want))
        print("   published today's daily candle yet. Prices ARE today's.")
    print("=" * 70)

    ok = pd.DataFrame(rows)
    if ok.empty:
        print("\nNo stock fits today. Nothing to buy -- that is a valid answer.")
    else:
        ok = ok.sort_values(["status", "rs_rank"], ascending=[True, False])
        fresh = ok[ok.bars_ago == 0]
        print("\n>>> AAJ KYA KARNA HAI")
        if len(fresh):
            print("  NEW signal on the last session (backtest entry = next open):")
            for _, x in fresh.iterrows():
                print("    BUY  %-12s ~%.1f  stop %.1f  exit below %.1f"
                      % (x.symbol, x.price, x.stop_if_buy_today, x.exit_40w))
        else:
            print("  No new signal on the last session.")
        fits = ok[(ok.status == "FIT") & (ok.bars_ago > 0)]
        if len(fits):
            print("  Still valid (FIT, % vs signal): " + ", ".join(
                "%s %+.0f%%" % (x.symbol, x.vs_signal_pct)
                for x in fits.itertuples()))
        lates = ok[ok.status == "LATE"]
        if len(lates):
            print("  LATE (ran >%d%% past signal, entry not tested): " % LATE_PCT
                  + ", ".join(lates.symbol))

        swing = ok[["symbol", "mcap_cr", "status", "rs_rank", "signal_date",
                    "bars_ago", "signal_px", "price", "px_src",
                    "vs_signal_pct", "stop_if_buy_today", "exit_40w",
                    "room_to_exit_pct"]]
        inv = ok[ok.vs_signal_pct < INV_MAX_EXT][
            ["symbol", "mcap_cr", "status", "rs_rank", "signal_date",
             "bars_ago", "signal_px", "price", "vs_signal_pct", "ma30w"]]
        print("\n==== SWING (%d) ====" % len(swing))
        print(swing.drop(columns=["signal_px"]).to_string(index=False))
        print("\n==== INVESTING (%d) ====" % len(inv))
        print(inv.drop(columns=["signal_px", "bars_ago"]).to_string(index=False)
              if len(inv) else "  (none)")
        xl = write_excel(stamp, swing, inv,
                         "Data to %s | prices: %s%s" % (
                             data_date,
                             "Dhan live/LTP" if live else "last close in data",
                             "  | !!! DATA IS BEHIND -- do not trade from this"
                             if (stale and not live) else ""))
        if xl:
            print("\nExcel: %s" % xl)

    nf = pd.DataFrame(nofit)
    if not nf.empty:
        print("\n==== SIGNALLED RECENTLY BUT NO FIT AT TODAY'S PRICE (%d) ===="
              % len(nf))
        for _, x in nf.sort_values("rs_rank", ascending=False).iterrows():
            print("  %-12s sig %s  px %8.1f  -> %s"
                  % (x.symbol, x.signal_date, x.price, x.why))

    if warns:
        print("\n---- warnings ----")
        for w in warns[:25]:
            print("  !", w)

    print("""
REMINDERS
  - Win rate ~39%, not 60-70%. Money comes from avg win (+51%) being
    ~4x avg loss (-14%). The MEDIAN trade loses ~6%. Many trades WILL lose.
  - Stop = 20% below YOUR entry. Exit = close below 40-week MA.
    No profit target, no tightening -- both lost money in testing.
  - LATE names were never tested. If you take them, size smaller.
  - Backtest of this universe (backtest.py, 2013-2026): 20-slot
    portfolio ~12.5%/yr swing, ~14.8%/yr investing exit, Nifty 10.5%
    (price only), drawdowns ~-40%. Bad years exist (2015, 2018, 2024-25).
""")
    print("Reports saved in: %s" % REPORTS)
    account.banner(acc)


if __name__ == "__main__":
    main()
