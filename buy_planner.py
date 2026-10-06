"""
buy_planner.py -- ONE sheet to buy from (5 Oct 2026, RB approved):

  1) ACCOUNT: available balance, holdings value + cost, unrealised and
     realised P&L, money added (you type it once), overall P&L.
  2) KITNA LAGANA HAI: budget -> split by cap class (MID 60 / LARGE 25 /
     SMALL 15, editable), equal inside a class; a class with no pick gives its
     share to the highest preference present (MID > LARGE > SMALL).
  3) SHORTLIST: momentum top 20 + W+TT BUY/FIT, best preference first.
     Pick BUY / MTF / WATCH (blank = no); 'Qty (you)' fixes a stock's qty and the rest of the budget is
     re-split over the other BUY / MTF rows. Whole shares; what is left buys +1
     share down the list (max +1 each).

The sheet works with live formulas (Excel / Google Sheets). rbtrack does
NOT trust the cached formula values: it reads the inputs and recomputes with
distribute() below -- the same arithmetic -- so the order = what you saw.
No broker call and no order here.
"""

import math

import pandas as pd
from order_inputs import quantity

SHEET = "Buy_Planner"
SHARES = {"M": 0.60, "L": 0.25, "S": 0.15}
NAMES = {"M": "Mid", "L": "Large", "S": "Small"}
CODE = {v: k for k, v in NAMES.items()}
ORDER = ("M", "L", "S")
R0, NROWS = 29, 40                 # first stock row, rows in the table
INPUT_YEL, NAVY = "FFF2CC", "1F4E78"
PICKS = ("BUY", "MTF", "WATCH")        # BUY = CNC, MTF = margin buy, WATCH =
BUYING = ("BUY", "MTF")                 # watchlist only; blank = not taken
MTF_LEV = None                          # legacy input ignored; per-stock broker value required
LBL_LEV = "MTF leverage (andaaza)"
LBL_ADDED = "Kul paise add kiye (deposit)"
LBL_BUDGET = "Is baar invest (Rs)"


# ================================================================== maths
def _num(v, default=None):
    try:
        x = float(str(v).replace(",", ""))
        return default if not math.isfinite(x) else x
    except (TypeError, ValueError):
        return default


def pick_of(v):
    """'buy' / ' Buy ' -> 'BUY' (old 'YES' / True too); blank, 'NO' or
    anything unknown -> '' (= not taken)."""
    if v is True:
        return "BUY"
    v = " ".join(str(v or "").upper().split())
    v = {"YES": "BUY", "BUY MTF": "MTF"}.get(v, v)
    return v if v in PICKS else ""


def distribute(rows, budget, shares=None, lev=None):
    """rows (priority order): dicts symbol, cls ('M'/'L'/'S'/other), price,
    pick (BUY / MTF / WATCH / blank), qty_you (int or None). MTF rows cost only
    price / row["mtf_lev"] of your own money per share. Returns (rows with base / extra
    / final / amount (= OWN money), summary). Same arithmetic as the sheet."""
    shares = dict(SHARES, **(shares or {}))
    if budget not in (None, "") and (_num(budget) is None or _num(budget) < 0):
        raise ValueError("Planner budget must be a finite non-negative amount")
    budget = max(0.0, _num(budget, 0.0))
    if any(not isinstance(shares[c], (int,float)) or not math.isfinite(shares[c]) or not 0 <= shares[c] <= 1 for c in ORDER) or abs(sum(shares[c] for c in ORDER)-1) > .00001:
        raise ValueError("Mid / Large / Small percentages must total 100%")
    rows = [dict(r) for r in rows]
    for r in rows:
        r["pick"] = pick_of(r.get("pick"))
        r["buy"] = r["pick"] in BUYING
        r["price"] = _num(r.get("price"))
        leverage = _num(r.get("mtf_lev"))
        r["mtf_enabled"] = bool(leverage and math.isfinite(leverage) and 1 < leverage <= 10)
        r["unit"] = (r["price"] / (leverage if r["pick"] == "MTF" else 1.0)) \
            if r["price"] and r["price"] > 0 and (r["pick"] != "MTF" or r["mtf_enabled"]) else None
        if r["pick"] == "MTF" and not r["mtf_enabled"]:
            r["buy"] = False
        r["qty_you"] = quantity(r.get("qty_you"), "%s BUY Qty" % r.get("symbol", ""))
        r["auto"] = r["buy"] and r["qty_you"] is None and bool(r["unit"])
    n = {c: sum(1 for r in rows if r["auto"] and r["cls"] == c)
         for c in ORDER}
    missing = sum(shares[c] for c in ORDER if n[c] == 0)
    top = next((c for c in ORDER if n[c] > 0), None)
    eff = {c: (shares[c] + (missing if c == top else 0)) if n[c] else 0.0
           for c in ORDER}
    fixed = sum(r["qty_you"] * (r["unit"] or 0) for r in rows
                if r["buy"] and r["qty_you"] is not None)
    left = max(0.0, budget - fixed)
    for r in rows:
        r["weight"] = eff[r["cls"]] / n[r["cls"]] if r["auto"] and \
            r["cls"] in n and n[r["cls"]] else 0.0
        if not r["buy"]:
            r["base"] = 0
        elif r["qty_you"] is not None:
            r["base"] = r["qty_you"]
        elif r["auto"]:
            r["base"] = int(math.floor(left * r["weight"] / r["unit"]))
        else:
            r["base"] = 0
    pool = budget - fixed - sum(r["base"] * r["unit"] for r in rows
                                if r["auto"])
    used = 0.0
    for r in rows:                     # leftover: +1 share, top row first
        r["extra"] = 1 if r["auto"] and r["weight"] > 0 and \
            r["unit"] <= pool - used else 0
        used += r["extra"] * (r["unit"] or 0)
        r["final"] = r["base"] + r["extra"]
        r["amount"] = r["final"] * (r["unit"] or 0)       # own money
        r["position"] = r["final"] * (r["price"] or 0)
    total = sum(r["amount"] for r in rows)
    return rows, {"budget": budget, "fixed": fixed, "auto_pool": left,
                  "eff": eff, "count": n, "total": total,
                  "left": budget - total}


def priority(cls, overlap, mom_rank, rs_rank):
    """Sort key: MID > LARGE > SMALL > ?, Super-Buy, momentum rank, RS."""
    ci = ORDER.index(cls) if cls in ORDER else 9
    sb = 0 if overlap == "Super-Buy" else 1
    mr = mom_rank if mom_rank == mom_rank and mom_rank is not None else 9999
    rs = -(rs_rank if rs_rank == rs_rank and rs_rank is not None else 0)
    return (ci, sb, mr, rs)


def shortlist(comp, cap_of):
    """Momentum top 20 + W+TT BUY/FIT from the Actions list (comp)."""
    out = []
    if comp is None or not len(comp):
        return out
    for _, x in comp.iterrows():
        ov = str(x.get("Strategy Overlap", ""))
        st = str(x.get("W+TT Status", "")).upper()
        mom = ov in ("Super-Buy", "Momentum only")
        if not mom and st not in ("BUY", "FIT"):
            continue
        s = str(x.get("Ticker", "")).upper()
        mr = _num(x.get("Mom Rank"))
        why = ("Super-Buy (momentum #%d + W+TT %s)" % (mr, st)
               if ov == "Super-Buy" and mr else
               "Momentum #%d" % mr if mom and mr else "W+TT %s" % st)
        out.append({"symbol": s, "why": why, "overlap": ov,
                    "mom_rank": mr, "rs_rank": _num(x.get("RS Rank")),
                    "cls": cap_of(s), "price": _num(x.get("LTP")),
                    "sector": x.get("Sector / Industry", "")})
    out.sort(key=lambda r: priority(r["cls"], r["overlap"], r["mom_rank"],
                                    r["rs_rank"]))
    return out[:NROWS]


# ================================================================== sheet
def write_sheet(wb, acct, rows, inputs, picks, title, gold=()):
    """acct: dict cash, hold_value, hold_cost, realised (None ok).
    inputs: money_added, budget, shares. picks: {SYM: (pick, qty_you)}."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.sheet_view.showGridLines = False
    f = lambda c: PatternFill("solid", fgColor=c)                 # noqa
    thin = Side(style="thin", color="D9D9D9")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = (5, 13, 30, 8, 9, 11, 8, 10, 9, 9, 9, 10, 12, 9, 12)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    last = R0 + NROWS - 1

    def band(r, text):
        ws.cell(row=r, column=1, value=text).font = Font(
            bold=True, color="FFFFFF", size=12)
        for k in range(1, 16):
            ws.cell(row=r, column=k).fill = f(NAVY)

    def kv(r, label, value, fmt="#,##0", inp=False, bold=False, note=None):
        ws.cell(row=r, column=2, value=label).font = Font(bold=True)
        b = ws.cell(row=r, column=5, value=value)
        b.number_format = fmt
        b.font = Font(bold=bold or inp, size=12 if bold else 11)
        b.alignment = Alignment(horizontal="right")
        if inp:
            b.fill = f(INPUT_YEL)
        if note:
            ws.cell(row=r, column=7, value=note).font = Font(
                italic=True, color="7F7F7F", size=9)

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=15)
    ws.cell(row=2, column=1, value="Peele cells tum bharo; baaki formula "
            "hai (Excel / Google Sheets turant badalta hai). Save + close, "
            "phir rbtrack -- woh BUY / MTF rows ki FINAL QTY bhejta hai.").font = \
        Font(italic=True, color="7F7F7F")
    band(4, "1) ACCOUNT")
    na = "n/a (token?)"
    kv(5, "Available balance (broker)", acct.get("cash")
       if acct.get("cash") is not None else na,
       note=acct.get("cash_note", "rb chalne ke waqt broker se"))
    kv(6, "Holdings value (aaj)", acct.get("hold_value", 0))
    kv(7, "Invested in holdings (cost)", acct.get("hold_cost", 0))
    kv(8, "Unrealised P&L", "=E6-E7", '+#,##0;-#,##0', note="khule stocks")
    kv(9, "Unrealised %", "=IF(E7,E6/E7-1,0)", '+0.0%;-0.0%')
    kv(10, "Realised P&L (Journal, fees ke baad)", acct.get("realised") or 0,
       '+#,##0;-#,##0', note="bech chuke trades")
    kv(11, LBL_ADDED, inputs.get("money_added"), inp=True,
       note="ek baar bharo (broker API deposit nahi batata)")
    kv(12, "Account abhi (balance + holdings)", "=N(E5)+E6", bold=True)
    kv(13, "OVERALL profit / loss", '=IF(N(E11)>0,E12-E11,"E11 bharo")',
       '+#,##0;-#,##0', bold=True)
    kv(14, "OVERALL %", '=IF(N(E11)>0,E12/E11-1,"")', '+0.0%;-0.0%',
       bold=True)

    band(16, "2) KITNA LAGANA HAI")
    kv(17, LBL_BUDGET, inputs.get("budget"), inp=True, bold=True)
    kv(18, "Check", '=IF(N(E17)<=0,"E17 mein amount likho",IF(AND(ISNUMBER('
       'E5),E17>E5),"!! balance se zyada",IF(J25>E17,"!! tumhari Qty budget '
       'se zyada","OK")))', "@")
    kv(19, "MTF leverage", "Har stock ka alag (column R)", "@",
       note="0 / unknown = MTF blocked; rbtrack broker se dobara verify karta hai")
    for k, h in enumerate(["Class", "Target %", "Auto stocks", "Effective %",
                           "Rs"], 2):
        ws.cell(row=20, column=k, value=h).font = Font(bold=True)
    sh = dict(SHARES, **(inputs.get("shares") or {}))
    rng = lambda c: "$%s$%d:$%s$%d" % (c, R0, c, last)              # noqa
    buy = '((%s="BUY")+(%s="MTF"))' % (rng("G"), rng("G"))
    for i, c in enumerate(ORDER):
        r = 21 + i
        ws.cell(row=r, column=2, value=NAMES[c])
        t = ws.cell(row=r, column=3, value=sh[c])
        t.number_format, t.fill = "0%", f(INPUT_YEL)
        ws.cell(row=r, column=4, value='=SUMPRODUCT(%s*(%s="")*'
                '(%s=B%d)*(%s>0))' % (buy, rng("H"), rng("E"), r,
                                      rng("Q")))
        ws.cell(row=r, column=5, value='=IF(D%d>0,C%d+IF($D$24=B%d,$F$24,0),'
                '0)' % (r, r, r)).number_format = "0%"
        ws.cell(row=r, column=6, value="=$J$26*E%d" % r).number_format = \
            "#,##0"
    ws["B24"], ws["B25"] = "Khaali class ka hissa", "Fixed (Qty you) Rs"
    ws["B26"] = "Auto stocks ke liye Rs"
    ws["D24"] = '=IF(D21>0,"Mid",IF(D22>0,"Large",IF(D23>0,"Small","-")))'
    ws["F24"] = "=C21*(D21=0)+C22*(D22=0)+C23*(D23=0)"
    ws["F24"].number_format = "0%"
    ws["J25"] = '=SUMPRODUCT(%s*(%s<>"")*N(+%s)*%s)' % (
        buy, rng("H"), rng("H"), rng("Q"))
    ws["J26"] = "=MAX(0,N(E17)-J25)"
    ws["J27"] = '=N(E17)-J25-SUMPRODUCT(%s*(%s="")*%s*%s)' % (
        buy, rng("H"), rng("J"), rng("Q"))
    for c in ("J25", "J26"):
        ws[c].number_format = "#,##0"
    ws["J27"].font = Font(color="FFFFFF")
    for r in (24, 25, 26):
        ws.cell(row=r, column=2).font = Font(bold=True)

    band(27, "3) SHORTLIST -> Pick: BUY / MTF (margin) / WATCH (sirf "
         "watchlist); khaali = nahi lena. Chaaho to apni Qty likho")
    heads = ["#", "Symbol", "Kyun list mein", "Mom Rank", "Cap Class",
             "Price", "Pick", "Qty (you)", "Weight", "Base Qty", "+1 (bacha)",
             "FINAL QTY", "Own money (Rs)", "% budget", "Position (Rs)"]
    for k, h in enumerate(heads, 1):
        c = ws.cell(row=28, column=k, value=h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), \
            f(NAVY), box
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[28].height = 30
    ws["P28"] = 0
    color = {"Large": "1F4E78", "Mid": "1E7B34", "Small": "C55A11"}
    for i in range(NROWS):
        r = R0 + i
        x = rows[i] if i < len(rows) else None
        pk, qy = picks.get(x["symbol"], ("", None)) if x else ("", None)
        vals = [i + 1 if x else None, x["symbol"] if x else None,
                x["why"] if x else None, x["mom_rank"] if x else None,
                NAMES.get(x["cls"], "?") if x else None,
                x["price"] if x else None, pk or None, qy,
                '=IF(AND(OR(G{r}="BUY",G{r}="MTF"),H{r}="",Q{r}>0),IFERROR('
                'INDEX($E$21:$E$23,MATCH(E{r},$B$21:$B$23,0))/INDEX($D$21:'
                '$D$23,MATCH(E{r},$B$21:$B$23,0)),0),0)'.format(r=r),
                '=IF(OR(AND(G{r}<>"BUY",G{r}<>"MTF"),AND(G{r}="MTF",N(R{r})<=1)),0,IF(H{r}<>"",N(+H{r}),'
                'IF(Q{r}>0,INT($J$26*I{r}/Q{r}),0)))'.format(r=r),
                '=IF(AND(OR(G{r}="BUY",G{r}="MTF"),H{r}="",Q{r}>0,I{r}>0,'
                'Q{r}<=$J$27-P{p}),1,0)'.format(r=r, p=r - 1),
                "=J%d+K%d" % (r, r), "=L%d*Q%d" % (r, r),
                "=IF(N($E$17)>0,M%d/$E$17,0)" % r,
                "=L%d*N(F%d)" % (r, r)]
        for k, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=k, value=v)
            cell.border = box
            if k in (7, 8):
                cell.fill = f(INPUT_YEL)
            if k == 6:
                cell.number_format = "#,##0.0"
            if k in (9, 14):
                cell.number_format = "0.0%"
            if k in (13, 15):
                cell.number_format = "#,##0"
            if k in (2, 12):
                cell.font = Font(bold=True)
            if k == 5 and x:
                cell.font = Font(bold=True, color=color.get(v, "7F7F7F"))
        ws.cell(row=r, column=16, value="=P%d+K%d*Q%d" % (r - 1, r, r))
        ws.cell(row=r, column=17, value='=IF(N(F{r})>0,IF(G{r}="MTF",'
                'IF(N(R{r})>1,F{r}/R{r},0),F{r}),0)'.format(r=r))
        ws.cell(row=r, column=18, value=(x.get("mtf_lev") if x.get("mtf_lev") is not None else "UNKNOWN") if x else None).number_format = '0.00"x"'
        ws.cell(row=r, column=19, value=x.get("mtf_asof") if x else None)
        ws.cell(row=r, column=20, value=x.get("mtf_note", "UNKNOWN") if x else None)
        if x and x["symbol"] in gold:
            for k in (2, 4):
                ws.cell(row=r, column=k).fill = f("F7E3A5")
    ws["R28"] = "MTF x (broker)"
    ws["R28"].font = Font(bold=True,color="FFFFFF")
    ws["R28"].fill = f(NAVY)
    ws.column_dimensions["R"].width = 17
    ws["S28"], ws["T28"] = "MTF updated IST", "MTF status (broker)"
    ws.column_dimensions["S"].width = 25
    ws.column_dimensions["T"].width = 42
    for cell in (ws["S28"], ws["T28"]):
        cell.font = Font(bold=True,color="FFFFFF");cell.fill=f(NAVY)
    ws.column_dimensions["P"].hidden = True
    ws.column_dimensions["Q"].hidden = True
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(PICKS),
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Pick", error="BUY / MTF / WATCH chuno "
                        "(khaali = nahi lena)", showInputMessage=True,
                        promptTitle="Pick", prompt="BUY = normal, MTF = margin, "
                        "WATCH = sirf watchlist, khaali = nahi")
    dv.add("G%d:G%d" % (R0, last))
    ws.add_data_validation(dv)
    T = last + 1
    ws.cell(row=T, column=2, value="TOTAL").font = Font(bold=True)
    ws.cell(row=T, column=13, value="=SUM(M%d:M%d)" % (R0, last))
    ws.cell(row=T + 1, column=2, value="Bacha (round off)").font = \
        Font(bold=True)
    ws.cell(row=T + 1, column=13, value="=N(E17)-M%d" % T)
    for r in (T, T + 1):
        ws.cell(row=r, column=13).number_format = "#,##0"
        ws.cell(row=r, column=13).font = Font(bold=True)
    ws.conditional_formatting.add("E18", FormulaRule(
        formula=['LEFT(E18,2)="!!"'], font=Font(bold=True, color="C00000")))
    notes = [
        "Budget pehle class mein (Mid / Large / Small %, C21:C23), phir us "
        "class ke BUY / MTF stocks mein barabar. Khaali class ka hissa upar ki "
        "preference (Mid > Large > Small) ko. Backtest 2013-26: is mix ne "
        "2013-19 mein +3 pts, 2020-26 barabar (cap_mix_study.py).",
        "Qty (you) bharoge to woh stock utna hi; baaki budget dusre BUY "
        "stocks mein dobara batta hai. Poore shares -> bacha paisa list mein "
        "upar se +1 share.",
        "Momentum recommendations monthly; manual BUY needs confirmation. 20 slots + 4 per "
        "industry rbtrack khud check karta hai. Gold = momentum rank 1-5 "
        "(pehle dekho, BUY signal nahi).",
        "BUY = normal buy. MTF: har stock ka broker leverage alag (R); "
        "4.5x par own money = position/4.5, 0/unknown par MTF blocked. "
        "rbtrack fresh leverage se quantity aur required cash dobara dikhata hai. "
        "Backtest: 4x MTF ne 1x se KAM kamaya, DD -97%.",
        "WATCH = koi order nahi, sirf watchlist (rbport analyse karta hai). "
        "TRADING OFF = koi order nahi."]
    for i, t in enumerate(notes):
        ws.cell(row=T + 3 + i, column=1, value=t).font = Font(
            italic=True, size=9, color="404040")
    ws.freeze_panes = "A4"
    return ws


def read_sheet(xlsx):
    """(inputs, rows) from a saved Portfolio file. inputs: money_added,
    budget, shares; rows: symbol, cls, price, pick, qty_you, why (priority
    order). ({}, []) if the sheet is missing."""
    try:
        from openpyxl import load_workbook
        ws = load_workbook(xlsx, data_only=True)[SHEET]
        raw = load_workbook(xlsx, data_only=False)[SHEET]
    except Exception:
        return {}, []
    labels = {str(ws.cell(row=r, column=2).value or "").strip(): r
              for r in range(1, 20)}
    inp = {"money_added": _num(ws.cell(row=labels.get(LBL_ADDED, 11),
                                       column=5).value),
           "budget": _num(ws.cell(row=labels.get(LBL_BUDGET, 17),
                                  column=5).value),
           "shares": {},
           "mtf_lev": _num(ws.cell(row=labels.get(LBL_LEV, 19),
                                   column=5).value)}
    budget_raw = raw.cell(row=labels.get(LBL_BUDGET, 17), column=5).value
    if budget_raw not in (None, "") and (_num(budget_raw) is None or _num(budget_raw) < 0):
        raise ValueError("Planner budget must be a literal finite non-negative amount")
    for r in range(21, 24):
        c = CODE.get(str(ws.cell(row=r, column=2).value or ""))
        v = _num(ws.cell(row=r, column=3).value)
        if c and v is not None:
            inp["shares"][c] = v / 100 if v > 1 else v
    rows = []
    for r in range(R0, R0 + NROWS):
        s = ws.cell(row=r, column=2).value
        if not s:
            continue
        pick = pick_of(ws.cell(row=r, column=7).value)
        rows.append({"symbol": str(s).upper().strip(),
                     "why": ws.cell(row=r, column=3).value or "",
                     "cls": CODE.get(str(ws.cell(row=r, column=5).value), "?"),
                     "price": _num(ws.cell(row=r, column=6).value),
                     "pick": pick,
                     "mtf_lev": _num(ws.cell(row=r, column=18).value),
                     # Manual Qty must be literal; a formula cache may be absent/stale.
                     "qty_you": quantity(raw.cell(row=r, column=8).value,
                                         "%s BUY Qty" % s) if pick in BUYING else None})
    return inp, rows


def frame(rows):
    return pd.DataFrame(rows)
