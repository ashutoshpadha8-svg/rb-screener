"""Literal spreadsheet inputs and order flows, with external calls replaced."""
import ast
import contextlib
import datetime as dt
import io
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace as NS
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import buy_planner as bp
import order_inputs
import portfolio as pf
import execution_safety as safety


def function(name, env):
    tree = ast.parse((ROOT/'auto_tracker_update.py').read_text())
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == name)
    exec(compile(ast.Module(body=[node], type_ignores=[]),str(ROOT/'auto_tracker_update.py'),'exec'),env)
    return env[name]


class ManualInputs(unittest.TestCase):
    def test_quantity_valid_and_blank(self):
        for v, expected in [(None,None),('',None),(' ',None),(0,0),(3,3),('3',3),('1,000',1000)]:
            self.assertEqual(order_inputs.quantity(v),expected)

    def test_invalid_quantity_never_defaults(self):
        for v in ['abc',-1,2.5,'=1+2',True,float('nan'),float('inf')]:
            with self.assertRaises(ValueError):order_inputs.quantity(v)

    def test_invalid_buy_quantity_never_auto_allocates(self):
        for v in ['abc',-1,2.5,'=1+2']:
            with self.assertRaises(ValueError):
                bp.distribute([dict(symbol='ABC',cls='M',price=100,pick='BUY',qty_you=v)],1000)

    def test_planner_formula_quantity_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'p.xlsx';w=Workbook();s=w.active;s.title=bp.SHEET
            s.cell(bp.R0,2,'ABC');s.cell(bp.R0,7,'BUY');s.cell(bp.R0,8,'=1+2');w.save(p)
            with self.assertRaises(ValueError):bp.read_sheet(p)

    def sell_book(self,tmp,value):
        p=Path(tmp)/'s.xlsx';w=Workbook();s=w.active;s.title='Holdings'
        s['A1']=pf.SELL_CAP;s['B1']='ABC sell qty (held 40, CNC)'
        s['A2']='SELL';s['B2']=value;w.save(p)
        return p

    def test_invalid_sell_quantity_never_sells_all(self):
        read=function('read_sells',dict(pf=pf))
        with tempfile.TemporaryDirectory() as tmp:
            for v in ['abc',-1,2.5,'=1+2']:
                with self.assertRaises(ValueError):read(self.sell_book(tmp,v))

    def test_explicit_zero_sell_and_blank_all(self):
        read=function('read_sells',dict(pf=pf))
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(read(self.sell_book(tmp,0)),[])
            self.assertEqual(read(self.sell_book(tmp,None))[0]['qty'],40)
            self.assertEqual(read(self.sell_book(tmp,3))[0]['qty'],3)

    def run_main(self,before,after):
        calls=[];empty=pd.DataFrame(columns=['Ticker','act'])
        acc=NS(broker='DUMMY',cid='TEST');date=[before]
        def pull(path):date[0]=after
        modules=dict(account=NS(activate=lambda:acc,banner=lambda a:None),
                     drive_copy=NS(pull=pull),sip=NS(read_sheet=lambda path:[],due=lambda *a,**k:[]),
                     settings=NS(read_dashboard=lambda path,**k:'ON'))
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'p.xlsx';p.touch()
            env=dict(sys=sys,pd=pd,os=os,take_lock=lambda:True,
                     safety=NS(validate_report=lambda *a:None,validate_batch=safety.validate_batch),
                     ba=NS(open_intents=lambda:pd.DataFrame()),ms=NS(get_session=lambda:NS()),
                     recover_buys=lambda sess:None,ds=NS(now_ist=lambda:dt.datetime(2026,10,6,20)),
                     pf=NS(report_date=lambda path:date[0]),action_rows=lambda path:empty.copy(),
                     planner_rows=lambda *a:(pd.DataFrame(),None),
                     read_sells=lambda path:[dict(symbol='ABC',qty=3,product='CNC')],
                     place_sells=lambda *a:calls.append('SELL'))
            with patch.dict(sys.modules,modules),patch.object(sys,'argv',['rbtrack','--file',str(p)]),contextlib.redirect_stdout(io.StringIO()):
                function('main',env)()
        return calls

    def test_stale_report_stops_orders(self):
        self.assertEqual(self.run_main('2026-10-05','2026-10-05'),[])

    def test_unknown_date_stops_orders(self):
        self.assertEqual(self.run_main(None,None),[])

    def test_date_checked_after_drive_pull(self):
        self.assertEqual(self.run_main('2026-10-06','2026-10-05'),[])

    def test_current_report_reaches_sell_flow(self):
        self.assertEqual(self.run_main('2026-10-05','2026-10-06'),['SELL'])

    def buy_flow(self,answer='YES',funds=1000,identity=True):
        posts=[];row=dict(symbol='ABC',shares=3,entry_price=100,product='CNC',note='test')
        env=dict(ds=NS(market_open=lambda:False,MCAP_CACHE='dummy',now_ist=lambda:dt.datetime(2026,10,6)),LIMIT_BUFFER=.02,
                 safety=NS(validate_batch=safety.validate_batch,validate_caps=lambda *a:None,finite_positive=safety.finite_positive),pd=pd,
                 ba=NS(symbol_map=lambda sess:{'ABC':{}},verify_identity=lambda sess:(identity,'dummy'),
                       available_funds=lambda sess:(funds,''),pending_orders=lambda sess:set(),
                       holdings=lambda sess:[],BrokerError=RuntimeError),
                 input=lambda prompt:answer,
                 send_one=lambda *a,**k:(posts.append(a) or True,'dummy-order','PENDING'))
        with contextlib.redirect_stdout(io.StringIO()):
            accepted=function('place_orders',env)(NS(label='Dummy',client_id='TEST',broker='DUMMY'),[row],False,False)
        return posts,accepted

    def test_buy_exact_quantity_with_confirmation(self):
        posts,accepted=self.buy_flow();self.assertEqual(len(posts),1)
        self.assertEqual(posts[0][1:5],('ABC',3,'CNC','BUY'));self.assertEqual(len(accepted),1)

    def test_buy_declined_sends_nothing(self):
        self.assertEqual(self.buy_flow(answer='NO'),([],[]))

    def test_buy_insufficient_funds_sends_nothing(self):
        self.assertEqual(self.buy_flow(funds=200),([],[]))

    def test_wrong_account_sends_nothing(self):
        self.assertEqual(self.buy_flow(identity=False),([],[]))

    def sell_flow(self,answer='YES SELL',have=5):
        posts=[];writes=[]
        env=dict(ds=NS(market_open=lambda:False),dt=dt,pd=pd,
                 ba=NS(holdings=lambda sess:[dict(symbol='ABC',qty=have)],
                       pending_orders=lambda sess:set(),
                       symbol_map=lambda sess:{'ABC':{}},sold_recently=lambda *a,**k:False,
                       ordered_today=lambda *a,**k:False,verify_identity=lambda sess:(True,'dummy'),
                       BrokerError=RuntimeError),input=lambda prompt:answer,
                 read_split=lambda:pd.DataFrame([dict(symbol='ABC',mode='LIVE',note='')]),
                 write_split=lambda sp:writes.append(sp.copy()),
                 send_one=lambda *a,**k:(posts.append(a) or True,'dummy-order','PENDING'))
        with contextlib.redirect_stdout(io.StringIO()):
            done=function('place_sells',env)(NS(label='Dummy',client_id='TEST'),
                    [dict(symbol='ABC',qty=7,product='CNC')],False)
        return posts,done,writes

    def test_sell_never_exceeds_fresh_holdings(self):
        posts,done,writes=self.sell_flow();self.assertEqual(len(posts),1)
        self.assertEqual(posts[0][1:5],('ABC',5,'CNC','SELL'))
        self.assertEqual(done[0]['qty'],5);self.assertEqual(len(writes),1)

    def test_sell_declined_sends_nothing(self):
        self.assertEqual(self.sell_flow(answer='NO'),([],[],[]))

    def test_stock_absent_from_demat_sends_nothing(self):
        self.assertEqual(self.sell_flow(have=0),([],[],[]))


if __name__=='__main__':unittest.main()
