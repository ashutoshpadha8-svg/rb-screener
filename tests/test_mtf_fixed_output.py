"""Stable account-scoped report, numbered output, and atomic refresh regressions."""
import contextlib,datetime as dt,io,sys,tempfile,types,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
import mtf_check as m
import mtf_breakeven as b
from test_mtf_breakeven import fixture,TODAY
def metric(ws,label,stock='TCS'):
 row=next(c.row for c in ws['A'] if c.value==label)
 header=next(c.row for c in ws['A'] if c.value=='Metric')
 column=next(c.column for c in ws[header] if c.value==stock)
 return ws.cell(row,column).value
class FixedOutputTests(unittest.TestCase):
 def run_cli(self,reports,label='DUMMY A',day=TODAY,price=2100,broker='DHAN',token_ok=True):
  reports.mkdir(parents=True,exist_ok=True);r=fixture();r['rows'][0]['Price']=price
  r['rows'][0]['Price source']='DUMMY quote refreshed '+day.isoformat()
  acc=types.SimpleNamespace(activate=lambda:types.SimpleNamespace(token_ok=token_ok,session=types.SimpleNamespace(broker=broker),label=label,reports=str(reports)))
  class Date(dt.date):
   @classmethod
   def today(cls):return cls(day.year,day.month,day.day)
  stream=io.StringIO()
  with patch.dict(sys.modules,{'account':acc}),patch.object(m,'build',lambda *a:r),patch.object(sys,'argv',['mtf_check.py']),patch.object(m.dt,'date',Date),contextlib.redirect_stdout(stream):code=m.main()
  return code,stream.getvalue()
 def test_first_run_fixed_name(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'A/reports';code,txt=self.run_cli(p);self.assertEqual(code,0);self.assertTrue((p/'MTF_Check.xlsx').is_file());self.assertEqual(len(list(p.glob('*.xlsx'))),1)
 def test_same_day_refresh_changes_value(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'A/reports';self.run_cli(p);before=metric(load_workbook(p/'MTF_Check.xlsx',data_only=True).active,'Net P/L including exit reserve')
   self.run_cli(p,price=3500);after=metric(load_workbook(p/'MTF_Check.xlsx',data_only=True).active,'Net P/L including exit reserve')
   self.assertLess(before,0);self.assertGreater(after,0);self.assertEqual(len(list(p.glob('*.xlsx'))),1)
 def test_new_day_updates_same_file_asof(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'A/reports';self.run_cli(p);self.run_cli(p,day=dt.date(2026,10,8))
   w=load_workbook(p/'MTF_Check.xlsx',data_only=True).active
   self.assertEqual(w['B4'].value.date(),dt.date(2026,10,8));self.assertEqual(len(list(p.glob('*.xlsx'))),1)
 def test_second_account_does_not_overwrite_first(self):
  with tempfile.TemporaryDirectory() as d:
   a=Path(d)/'A/reports';z=Path(d)/'B/reports';self.run_cli(a,label='DUMMY A');original=(a/'MTF_Check.xlsx').read_bytes()
   self.run_cli(z,label='DUMMY B',price=3500);self.assertEqual((a/'MTF_Check.xlsx').read_bytes(),original)
   self.assertIn('DUMMY B',load_workbook(z/'MTF_Check.xlsx',data_only=True).active['A3'].value)
 def test_account_name_change_keeps_same_filename(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'A/reports';self.run_cli(p,label='OLD NAME');self.run_cli(p,label='NEW NAME')
   self.assertEqual(len(list(p.glob('*.xlsx'))),1);self.assertIn('NEW NAME',load_workbook(p/'MTF_Check.xlsx',data_only=True).active['A3'].value)
 def test_numbered_claude_style_output(self):
  with tempfile.TemporaryDirectory() as d:
   _,txt=self.run_cli(Path(d)/'A/reports')
   for heading in ['1. AAPKA LAGAYA PAISA','2. PER-STOCK HOLDING','3. AAJ BECHO TOH KYA MILEGA','4. BREAKEVEN / 30-DAY TARGET']:self.assertIn(heading,txt)
   self.assertIn('UNKNOWN',txt);self.assertNotIn('ACTUAL LOSS',txt)
 def test_profit_prints_plus(self):
  with tempfile.TemporaryDirectory() as d:
   _,txt=self.run_cli(Path(d)/'A/reports',price=3500);self.assertIn('+Rs',txt)
 def test_missing_price_overwrites_with_unknown_not_old_profit(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'A/reports';self.run_cli(p,price=3500);code,txt=self.run_cli(p,price=None)
   self.assertEqual(code,2);self.assertEqual(metric(load_workbook(p/'MTF_Check.xlsx',data_only=True).active,'Net P/L including exit reserve'),'UNKNOWN');self.assertIn('INCOMPLETE',txt)
 def test_invalid_token_does_not_overwrite_existing_report(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'A/reports';self.run_cli(p);before=(p/'MTF_Check.xlsx').read_bytes();code,_=self.run_cli(p,token_ok=False)
   self.assertEqual(code,1);self.assertEqual((p/'MTF_Check.xlsx').read_bytes(),before)
 def test_unsupported_broker_not_written_as_dhan(self):
  with tempfile.TemporaryDirectory() as d:
   # 9 Oct: ANGEL/ZERODHA now supported (mtf_generic); an UNKNOWN broker still refuses.
   p=Path(d)/'UPSTOX/reports';code,txt=self.run_cli(p,broker='UPSTOX');self.assertEqual(code,1);self.assertFalse((p/'MTF_Check.xlsx').exists());self.assertIn('rate card nahi',txt)
   import broker_api as ba
   p=Path(d)/'ANGEL/reports'
   with patch.object(ba,'mtf_holdings',side_effect=ba.BrokerError('DUMMY down')):code,txt=self.run_cli(p,broker='ANGEL')
   self.assertEqual(code,2);self.assertIn('MISSING holdings',txt);self.assertNotIn('DHAN',load_workbook(p/'MTF_Check.xlsx').active['A3'].value or '')
 def test_atomic_failure_preserves_old_report(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'MTF_Check.xlsx';p.write_bytes(b'ORIGINAL REPORT')
   with patch.object(b,'cache_formula_values',side_effect=OSError('DUMMY failed cache')):
    with self.assertRaises(OSError):b.write_report(fixture(),TODAY,p)
   self.assertEqual(p.read_bytes(),b'ORIGINAL REPORT');self.assertEqual(list(Path(d).glob('.mtf-report-*')),[])
 def test_symlink_report_refused(self):
  with tempfile.TemporaryDirectory() as d:
   original=Path(d)/'other-account.xlsx';original.write_bytes(b'KEEP OTHER ACCOUNT');p=Path(d)/'MTF_Check.xlsx';p.symlink_to(original)
   with self.assertRaises(ValueError):b.write_report(fixture(),TODAY,p)
   self.assertEqual(original.read_bytes(),b'KEEP OTHER ACCOUNT')
if __name__=='__main__':unittest.main()
