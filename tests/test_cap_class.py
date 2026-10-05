#!/usr/bin/env python3
"""
CAP CLASS COLUMN TESTS (5 Oct 2026) -- temp files only, no network/account.
Run:  cd ~/RB_Screener && python3 tests/test_cap_class.py
"""
import os
import sys
import tempfile

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import cap_class as cc                                          # noqa: E402

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                          "" if cond else "  <- " + str(detail)))


def heads(ws, r=1):
    return [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]


def main():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "mcap.csv")
    n = 300
    pd.DataFrame({"symbol": ["S%03d" % i for i in range(1, n + 1)],
                  "mcap_cr": [100000 - i * 100 for i in range(1, n + 1)],
                  "asof": "2026-10-05"}).to_csv(p, index=False)
    cc.MCAP = p
    cc._cache.clear()
    tab, asof = cc.table(p)
    cc.table = lambda path=None: (tab, asof)
    print("1) AMFI-style rank classes")
    check("rank 100 Large, 101 Mid, 250 Mid, 251 Small",
          [tab["S%03d" % i][2] for i in (100, 101, 250, 251)] ==
          ["Large", "Mid", "Mid", "Small"])
    check("'SYM | Company' and lower case read", cc.of("s005 | X Ltd", tab)
          and cc.of("s005 | X Ltd", tab)[2] == "Large")

    print("2) columns appended at the END, nothing moves")
    wb = Workbook()
    ws = wb.active
    ws.title = "Actions"
    for i, h in enumerate(["Ticker", "Action", "Amount (Rs)"], 1):
        ws.cell(row=1, column=i, value=h).font = Font(bold=True)
    for r, s in enumerate(["S001", "S150", "S280", "NOPE"], 2):
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value="BUY")
    ws.auto_filter.ref = "A1:C5"
    cc.add_columns(wb)
    check("headers = old 3 + Mcap + Cap Class", heads(ws) ==
          ["Ticker", "Action", "Amount (Rs)", cc.HEAD_M, cc.HEAD_C], heads(ws))
    check("classes Large / Mid / Small / ?",
          [ws.cell(row=r, column=5).value for r in range(2, 6)] ==
          ["Large", "Mid", "Small", "?"])
    check("Action column untouched", ws.cell(row=2, column=2).value == "BUY")
    check("filter widened to the new columns", ws.auto_filter.ref == "A1:E5",
          ws.auto_filter.ref)
    cc.add_columns(wb)
    check("second run: no duplicate columns", heads(ws).count(cc.HEAD_C) == 1
          and ws.max_column == 5, heads(ws))

    print("3) sheet that already has Mcap (Swing) / header on row 3")
    wb = Workbook()
    ws = wb.active
    ws.title = "Swing"
    for i, h in enumerate(["Symbol", "Action", "Mcap (Rs Cr)"], 1):
        ws.cell(row=1, column=i, value=h)
    ws.cell(row=2, column=1, value="S200")
    ws.cell(row=2, column=3, value=12345)
    wt = wb.create_sheet("Signal_Tracker")
    for i, h in enumerate(["Symbol", "Age"], 1):
        wt.cell(row=3, column=i, value=h)
    wt.cell(row=4, column=1, value="S260")
    for name in ("Dashboard", "Holdings"):
        x = wb.create_sheet(name)
        x.cell(row=1, column=1, value="Symbol")
        x.cell(row=2, column=1, value="S001")
    cc.add_columns(wb)
    check("Swing: only Cap Class added, own Mcap kept", heads(ws) ==
          ["Symbol", "Action", "Mcap (Rs Cr)", cc.HEAD_C] and
          ws.cell(row=2, column=3).value == 12345, heads(ws))
    check("Swing class Mid", ws.cell(row=2, column=4).value == "Mid")
    check("tracker header row 3 handled", wt.cell(row=3, column=4).value ==
          cc.HEAD_C and wt.cell(row=4, column=4).value == "Small")
    check("Dashboard / Holdings skipped", wb["Dashboard"].max_column == 1 and
          wb["Holdings"].max_column == 1)

    print("4) no market-cap file -> nothing added, no crash")
    cc.table = lambda path=None: ({}, None)
    wb = Workbook()
    wb.active.cell(row=1, column=1, value="Symbol")
    check("returns 0", cc.add_columns(wb) == 0 and wb.active.max_column == 1)

    bad = [n for n, ok in RESULTS if not ok]
    print("\n%d / %d passed%s" % (len(RESULTS) - len(bad), len(RESULTS),
                                  "" if not bad else " -- FAILED: " +
                                  "; ".join(bad)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
