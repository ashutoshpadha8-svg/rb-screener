"""rbmtf for ANGEL / ZERODHA (mtf_generic, 9 Oct). Mocked holdings + quotes, no network."""
import datetime as dt,sys,tempfile,types,unittest,json,io,contextlib
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import mtf_check as m,mtf_generic as g,mtf_portfolio_report as p
TODAY=dt.date(2026,10,9)
def args(**k):
 d=dict(mtf=None,buy_date=None,loan=None,unpaid_interest=None,own_cash=None,rate=None)
 d.update(k);return types.SimpleNamespace(**d)
Q=lambda sess,syms:{s:({'TCS':2156.0,'INFY':1500.0}[s],m.VERIFIED,'NSE demo') for s in syms}
def run(broker,hold,data,**k):
 rate=p.set_broker(broker);a=args(**k);a.rate=a.rate or rate
 R=g.build_generic(a,types.SimpleNamespace(broker=broker),TODAY,data,holdings_fn=lambda s:dict(hold),quotes_fn=Q)
 return R,p.calculate(R,TODAY,rate=a.rate)
ZH={'TCS':dict(qty=52.0,avg_price=3082.09,own_margin=60000.0,source='kite')}
AH={'TCS':dict(qty=52.0,avg_price=3082.09,own_margin=None,source='angel')}
class Generic(unittest.TestCase):
 def setUp(self):self.d=tempfile.TemporaryDirectory();self.data=self.d.name
 def tearDown(self):self.d.cleanup();p.set_broker('DHAN')
 def test_zerodha_loan_from_initial_margin_and_card(self):
  R,rows=run('ZERODHA',ZH,self.data)
  self.assertAlmostEqual(R['loan'].value,52*3082.09-60000,places=2);self.assertEqual(R['loan'].status,m.ESTIMATED)
  self.assertEqual(p.BROKER,.003);self.assertEqual(p.PLEDGE,15)
  self.assertIsNone(rows[0]['past']);self.assertIn('buy_dates',R['src'])          # no buy date yet
  self.assertIsNotNone(rows[0]['cash'])                                             # cash still known
 def test_angel_without_loan_is_unknown_with_command(self):
  R,rows=run('ANGEL',AH,self.data)
  self.assertFalse(R['loan'].known);self.assertIn('--loan',R['loan'].note);self.assertIsNone(rows[0]['cash'])
  self.assertFalse(p.complete(R,rows))
 def test_angel_complete_with_loan_and_buy_date(self):
  R,rows=run('ANGEL',AH,self.data,loan=100000.0,buy_date='TCS:2025-11-24')
  r=rows[0];self.assertEqual(r['status'],'ESTIMATED');self.assertTrue(p.complete(R,rows))
  self.assertAlmostEqual(r['past'],100000*.1499*319/365,places=2)                 # Angel 14.99%, 319 days
  self.assertAlmostEqual(r['sale_fees']['brokerage'],20.0);self.assertEqual(r['sale_fees']['dp'],20.0)
  start=dt.date(2026,10,1);self.assertAlmostEqual(R['unpaid'].value,round(100000*.1499*(TODAY-start).days/365,2))  # fortnightly
 def test_buy_date_remembered_and_dropped_on_qty_change(self):
  run('ANGEL',AH,self.data,loan=1.0,buy_date='TCS:2025-11-24')
  self.assertEqual(json.load(open(Path(self.data)/g.INPUTS_FILE))['TCS']['date'],'2025-11-24')
  R,_=run('ANGEL',AH,self.data,loan=1.0);self.assertIn('TCS',R['current_lots'])
  R,_=run('ANGEL',{'TCS':dict(AH['TCS'],qty=60.0)},self.data,loan=1.0);self.assertNotIn('TCS',R['current_lots']);self.assertIn('buy_date_TCS',R['src'])
 def test_bad_inputs_refused(self):
  for k in (dict(buy_date='TCS:2030-01-01'),dict(buy_date='INFY:2025-01-01'),dict(buy_date='TCS-2025'),dict(mtf='TCS:1.5')):
   with self.assertRaises(ValueError):run('ANGEL',AH,self.data,**k)
 def test_manual_mtf_override(self):
  R,rows=run('ANGEL',{},self.data,mtf='INFY:10@1400',loan=7000.0,buy_date='INFY:2026-02-04')
  self.assertEqual(rows[0]['symbol'],'INFY');self.assertAlmostEqual(rows[0]['cost'],14000.0);self.assertEqual(R['src']['holdings'][0],m.ESTIMATED)
 def test_holdings_failure_never_zero(self):
  rate=p.set_broker('ANGEL')
  def boom(s):raise RuntimeError('down')
  R=g.build_generic(args(rate=rate),types.SimpleNamespace(broker='ANGEL'),TODAY,self.data,holdings_fn=boom,quotes_fn=Q)
  self.assertFalse(R['open_cost_num'].known);self.assertFalse(R['loan'].known);self.assertFalse(p.complete(R,p.calculate(R,TODAY,rate=rate)))
 def test_report_written_with_broker_card(self):
  R,rows=run('ZERODHA',ZH,self.data,buy_date='TCS:2025-11-24');R['report_account']='ZERODHA | X | demo'
  path=Path(self.data)/'MTF_Check.xlsx';p.write_report(R,TODAY,path,rate=.146)
  w=load_workbook(path).active;vals=[c.value for row in w.iter_rows() for c in row]
  self.assertIn(.003,vals);self.assertIn('ZERODHA | X | demo',' '.join(str(v) for v in vals))
 def test_dhan_card_restored_by_main(self):
  p.set_broker('ANGEL');p.set_broker('DHAN');self.assertEqual((p.BROKER,p.DP,p.PLEDGE),(.0003,12.5,15))
if __name__=='__main__':unittest.main()
