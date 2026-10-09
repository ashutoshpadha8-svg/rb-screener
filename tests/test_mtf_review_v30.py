"""Independent MTF v30 findings and current-only workflow regressions.
Synthetic accounts only; external sockets forbidden. No credentials read.
"""
import contextlib,copy,datetime as dt,io,socket,sys,tempfile,types,unittest
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT),str(ROOT/'tests')]
def blocked(*a,**k):raise RuntimeError('OFFLINE MTF7')
socket.socket.connect=socket.socket.connect_ex=socket.create_connection=blocked
import mtf_check as m,mtf_prices as q,mtf_breakeven as b
from openpyxl import load_workbook
from test_mtf_breakeven import fixture,TODAY
from test_mtf_current_inventory import args,tr
from test_mtf_prices import payload
SECRET='SYNTHETIC_TOKEN_01234567890123456789'
def inventory_model(history=None,loan=10000):
 h=[dict(symbol='AAA',qty=10,avg_price=1000),dict(symbol='BBB',qty=20,avg_price=550)]
 ts=[tr(symbol='AAA',qty=10,price=1000,product='',tid='A'),tr(symbol='BBB',qty=20,price=550,product='',tid='B'),tr(symbol='OLD',qty=5,price=200,product='',tid='OLD')]
 led=m.ledger_rows([dict(narration='Net MTF Funding by Dhan',debit=0,credit=10000,voucherdate='Oct 07, 2026')])
 ba=types.SimpleNamespace(holdings=lambda *x:h)
 with patch.dict(sys.modules,{'broker_api':ba}),patch.object(m,'dhan_trades',return_value=(ts if history is None else history,dict(complete=True,notes=[]))),patch.object(m,'mtf_positions',return_value={}),patch.object(m,'dhan_ledger',return_value=led),patch.object(m,'quotes',return_value={'AAA':(900,m.VERIFIED,'synthetic exchange time'),'BBB':(600,m.VERIFIED,'synthetic exchange time')}):
  return m.build(args(loan=loan,confirm_requests={'AAA':10,'BBB':20}),types.SimpleNamespace(token=SECRET),TODAY,dt.date(2024,1,1))
def many_fixture():
 r=fixture();r['rows'].append(dict(Symbol='THIRD',Qty=10,Price=120,**{'Open cost (Rs)':m.Num(1000)}));r['fifo']['lots']['THIRD']=[['2026-09-01',10,100,0]];r['open_cost_num']=m.Num(415598.3);return r
class ReviewTests(unittest.TestCase):
 def test_foreign_csv_rejected_even_after_rename(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'ledger.csv';p.write_text('Client ID,222\nDate,Narration,Debit,Credit\n')
   with self.assertRaises(ValueError):m.validate_statement_account(p,'DHAN_111')
 def test_missing_identity_rejected(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'trades.csv';p.write_text('Date,Time,Price\n')
   with self.assertRaises(ValueError):m.validate_statement_account(p,'DHAN_111',True)
 def test_correct_ragged_csv_identity_allowed(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'ledger.csv';p.write_text('Client ID,111\nDate,Narration,Debit,Credit\n')
   m.validate_statement_account(p,'DHAN_111')
 def test_conflicting_filename_header_rejected(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'TRADE_HISTORY_CSV_111_2024-01-01_2026-10-07_0_.csv';p.write_text('Client ID,222\n')
   with self.assertRaises(ValueError):m.validate_statement_account(p,'DHAN_111',True)
 def test_renamed_xlsx_identity_allowed(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'ledger.csv';m.pd.DataFrame([['Client ID',111]]).to_excel(str(p)+'.xlsx',header=False,index=False);Path(str(p)+'.xlsx').rename(p);m.validate_statement_account(p,'DHAN_111')
 def test_failed_history_exception_secret_removed(self):
  ba=types.SimpleNamespace(holdings=lambda *x:[])
  with patch.dict(sys.modules,{'broker_api':ba}),patch.object(m,'dhan_trades',side_effect=RuntimeError(SECRET)),patch.object(m,'mtf_positions',return_value={}),patch.object(m,'dhan_ledger',return_value=m.ledger_rows([])),patch.object(m,'quotes',return_value={}):
   r=m.build(args(),types.SimpleNamespace(token=SECRET),TODAY,dt.date(2024,1,1));self.assertNotIn(SECRET,str(r['src']))
 def test_external_model_text_redacted_without_losing_numbers(self):
  r=inventory_model();r['src']['synthetic']=(m.UNKNOWN,SECRET);r['led'].loc[0,'narration']=SECRET;r['loan'].note=SECRET
  out=m.sanitize_report(r,types.SimpleNamespace(token=SECRET));self.assertNotIn(SECRET,str(out));self.assertEqual(out['loan'].value,10000)
 def test_jwt_and_key_value_redacted_without_session(self):
  out=q.redact_text('access-token=ABCD; eyJabc.def.ghi',object());self.assertNotIn('ABCD',out);self.assertNotIn('eyJabc',out)
 def test_atomic_audit_rejects_existing_link(self):
  with tempfile.TemporaryDirectory() as d:
   target=Path(d)/'other.xlsx';target.write_bytes(b'KEEP');p=Path(d)/'audit.xlsx';p.symlink_to(target)
   with self.assertRaises(ValueError):
    with b.atomic_excel_writer(p):pass
   self.assertEqual(target.read_bytes(),b'KEEP')
 def test_atomic_audit_rejects_parent_link(self):
  with tempfile.TemporaryDirectory() as d:
   real=Path(d)/'real';real.mkdir();link=Path(d)/'link';link.symlink_to(real,target_is_directory=True)
   with self.assertRaises(ValueError):
    with b.atomic_excel_writer(link/'audit.xlsx'):pass
   self.assertEqual(list(real.iterdir()),[])
 def test_audit_failure_keeps_previous_workbook(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'audit.xlsx';p.write_bytes(b'PREVIOUS')
   with self.assertRaises(RuntimeError):
    with b.atomic_excel_writer(p) as w:
     m.pd.DataFrame({'x':[1]}).to_excel(w);raise RuntimeError('simulated interrupted report')
   self.assertEqual(p.read_bytes(),b'PREVIOUS');self.assertEqual(list(Path(d).glob('.mtf-*')),[])
 def test_afterclose_no_intraday_quote_at_1535(self):
  n=dt.datetime(2026,10,9,15,35,tzinfo=q.IST)
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='09-Oct-2026 12:00:00'),'TCS',n,n.date()-dt.timedelta(days=1),False,True)
 def test_afterclose_no_yesterday_during_eod_delay(self):
  n=dt.datetime(2026,10,9,15,35,tzinfo=q.IST)
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='08-Oct-2026 16:00:00'),'TCS',n,n.date()-dt.timedelta(days=1),False,True)
 def test_one_second_future_rejected(self):
  n=dt.datetime(2026,10,9,12,tzinfo=q.IST)
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='09-Oct-2026 12:00:01'),'TCS',n,n.date()-dt.timedelta(days=1),True,True)
 def test_holiday_current_day_quote_rejected(self):
  n=dt.datetime(2026,10,9,12,tzinfo=q.IST)
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='09-Oct-2026 11:59:00'),'TCS',n,n.date()-dt.timedelta(days=1),False,False)
 def test_preopen_prior_completed_session_allowed(self):
  n=dt.datetime(2026,10,9,9,tzinfo=q.IST)
  self.assertEqual(q.parse_quote(payload(stamp='08-Oct-2026 16:00:00'),'TCS',n,n.date()-dt.timedelta(days=1),False,True)[1],m.VERIFIED)
 def test_missing_lots_current_qty_value_preserved(self):
  r=fixture();r['fifo']['lots']['TCS']=[]
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'report.xlsx';b.write_report(r,TODAY,p);w=load_workbook(p,data_only=True).active;f=load_workbook(p).active
   self.assertEqual(w['B10'].value,52);self.assertEqual(w['C18'].value,109200);self.assertAlmostEqual(w['D18'].value,-51068.6);self.assertEqual(w['G18'].value,'UNKNOWN');self.assertNotIn('SUMIF',str(f['B25'].value))
 def test_inconsistent_total_loan_basis_blocks_targets(self):
  r=many_fixture();r['open_cost_num']=m.Num(1);self.assertTrue(all(x['full'] is None for x in b.calculate(r,TODAY)))
 def test_three_stocks_use_one_default_sheet(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'report.xlsx';rows=b.write_report(many_fixture(),TODAY,p);w=load_workbook(p,data_only=True)
   self.assertEqual(w.sheetnames,['Breakeven']);self.assertEqual(len(rows),3);self.assertTrue(all(x['full'] is not None for x in rows));self.assertAlmostEqual(w.active['J13'].value,261888.33)
 def test_three_stocks_missing_quote_total_net_unknown(self):
  r=many_fixture();r['rows'][2]['Price']=None
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'report.xlsx';b.write_report(r,TODAY,p);s=load_workbook(p,data_only=True).active
   self.assertEqual(s['G13'].value,'UNKNOWN');self.assertEqual(s['E12'].value,'UNKNOWN');self.assertGreater(s['I12'].value,0)
 def test_current_acquisitions_separate_from_history(self):
  r=inventory_model();self.assertEqual(set(r['current_lots']),{'AAA','BBB'});self.assertFalse(r['complete_net'].known);self.assertTrue(r['open_cost_num'].known);self.assertEqual(r['open_cost_num'].value,21000)
  rows=b.calculate(r,TODAY);self.assertTrue(all(x['status']==m.ESTIMATED for x in rows));self.assertAlmostEqual(sum(x['past'] for x in rows),10000*.1249*36/365)
 def test_product_declaration_does_not_retag_blank_history(self):
  r=inventory_model();self.assertTrue(all(t['product']=='' for t in r['mtf']));self.assertFalse(r['fifo']['realised'].known)
 def test_uncovered_earlier_sell_blocks_purchase_date_model(self):
  ts=[tr(symbol='AAA',qty=1,price=1000,side='SELL',date='2026-08-01',ts=m._ts('2026-08-01 10:00:00'),product='CNC'),tr(symbol='AAA',qty=10,price=1000,product='CNC')]
  r=inventory_model(ts);self.assertNotIn('AAA',r['current_lots']);self.assertIsNone(next(x for x in b.calculate(r,TODAY) if x['symbol']=='AAA')['past'])
 def test_wrong_cost_basis_does_not_invent_lots(self):
  r=inventory_model([tr(symbol='AAA',qty=10,price=999,product='CNC')]);self.assertNotIn('AAA',r['current_lots'])
 def test_missing_trade_time_blocks_model(self):
  r=inventory_model([tr(symbol='AAA',qty=10,price=1000,product='CNC',ts=None)]);self.assertNotIn('AAA',r['current_lots'])
 def test_intraday_lots_not_used_for_acquisition(self):
  r=inventory_model([tr(symbol='AAA',qty=10,price=1000,product='INTRADAY')]);self.assertNotIn('AAA',r['current_lots'])
 def test_partial_fifo_uses_remaining_buy_date(self):
  ts=[tr(symbol='AAA',qty=15,price=1000,product='CNC'),tr(symbol='AAA',qty=5,price=1100,side='SELL',product='CNC',date='2026-09-02',ts=m._ts('2026-09-02 10:00:00'))]
  r=inventory_model(ts);self.assertEqual(r['current_lots']['AAA'][0][:3],['2026-09-01',10,1000])
 def test_lifetime_interest_cannot_change_current_targets(self):
  r=inventory_model();a=b.calculate(r,TODAY);r['paid']=m.Num(999999);r['unpaid']=m.Num(999999);strip=lambda L:[{k:v for k,v in x.items() if k not in ('loan_share','unpaid_share','cash_now')} for x in L];self.assertEqual(strip(b.calculate(r,TODAY)),strip(a))
 def test_default_report_literal_provenance_no_formula(self):
  r=fixture();r['rows'][0]['Price source']='=HYPERLINK("https://example.invalid")'
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'report.xlsx';b.write_report(r,TODAY,p);w=load_workbook(p).active;self.assertEqual(w['C10'].data_type,'n');self.assertFalse(load_workbook(p)._external_links)
 def test_default_cli_unknown_exit_reports_incomplete(self):
  r=inventory_model();r['exit']=m.Num(None,m.UNKNOWN,'ledger missing');r['src']['ledger']=(m.UNKNOWN,'ledger unavailable')
  with tempfile.TemporaryDirectory() as d:
   acc=types.SimpleNamespace(activate=lambda:types.SimpleNamespace(token_ok=True,session=types.SimpleNamespace(broker='DHAN'),label='DUMMY',reports=d))
   with patch.dict(sys.modules,{'account':acc}),patch.object(m,'build',return_value=r),patch.object(sys,'argv',['mtf_check.py']),contextlib.redirect_stdout(io.StringIO()) as stream:
    self.assertEqual(m.main(),2)
   self.assertIn('INCOMPLETE',stream.getvalue());self.assertIn('ledger unavailable',stream.getvalue())
 def test_historical_gap_does_not_block_complete_current_model(self):
  r=inventory_model();r['exit']=m.Num(9000,m.ESTIMATED,'synthetic current settlement')
  with tempfile.TemporaryDirectory() as d:
   acc=types.SimpleNamespace(activate=lambda:types.SimpleNamespace(token_ok=True,session=types.SimpleNamespace(broker='DHAN'),label='DUMMY',reports=d))
   with patch.dict(sys.modules,{'account':acc}),patch.object(m,'build',return_value=r),patch.object(sys,'argv',['mtf_check.py']),contextlib.redirect_stdout(io.StringIO()):self.assertEqual(m.main(),0)
 def test_current_model_cli_sanitizes_all_outputs(self):
  r=inventory_model();r['src']['synthetic']=(m.UNKNOWN,SECRET);r['led'].loc[0,'narration']=SECRET
  for audit in [False,True]:
   with self.subTest(audit=audit),tempfile.TemporaryDirectory() as d:
    acc=types.SimpleNamespace(activate=lambda:types.SimpleNamespace(token_ok=True,session=types.SimpleNamespace(broker='DHAN',token=SECRET),label='DUMMY',reports=d))
    with patch.dict(sys.modules,{'account':acc}),patch.object(m,'build',return_value=r),patch.object(sys,'argv',['mtf_check.py']+(['--audit'] if audit else [])),contextlib.redirect_stdout(io.StringIO()) as stream:
     m.main()
    self.assertNotIn(SECRET,stream.getvalue());w=load_workbook(next(Path(d).glob('*.xlsx')))
    self.assertNotIn(SECRET,str([[c.value for row in s for c in row] for s in w]))
if __name__=='__main__':unittest.main()
