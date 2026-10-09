#!/usr/bin/env python3
"""MTF7.2 (9 Oct 2026, RB: 'Google Sheets mein upload nahi ho raha'):
cache_formula_values must keep SpreadsheetML as the DEFAULT namespace
(<worksheet xmlns=...>), never a prefixed <s:worksheet>. Offline only."""
import os
import sys
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import mtf_breakeven as mb                                        # noqa: E402

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                          "" if cond else "  <- " + str(detail)))


def main():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "t.xlsx")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["B1"] = "=A1*3"
    ws["C1"] = '="x"&A1'
    wb.save(p)
    mb.cache_formula_values(p, {"B1": 6.0, "C1": "x2", "D1": None})
    xml = zipfile.ZipFile(p).read("xl/worksheets/sheet1.xml").decode()
    check("root is <worksheet xmlns=...> (default namespace)",
          "<worksheet xmlns=\"http://schemas.openxmlformats.org/"
          "spreadsheetml/2006/main\"" in xml, xml[:200])
    check("no prefixed SpreadsheetML tags (<s:..> / <ns0:..>)",
          "<s:" not in xml and "<ns0:" not in xml, xml[:200])
    check("formula kept + cached value written",
          "<f>A1*3</f><v>6.0</v>" in xml, xml)
    v = load_workbook(p, data_only=True).active
    f = load_workbook(p).active
    check("openpyxl reads cache 6.0 and formula =A1*3",
          v["B1"].value == 6.0 and f["B1"].value == "=A1*3",
          (v["B1"].value, f["B1"].value))
    bad = [n for n, ok in RESULTS if not ok]
    print("\n%d / %d passed%s" % (len(RESULTS) - len(bad), len(RESULTS),
                                  "" if not bad else " -- FAILED: " +
                                  "; ".join(bad)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
