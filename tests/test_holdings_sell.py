#!/usr/bin/env python3
"""
SELL FROM THE HOLDINGS CARDS (6 Oct 2026, RB) -- temp files only: no
network, no token, NO order.  Run: python3 tests/test_holdings_sell.py
"""
import os
import sys
import tempfile

import pandas as pd
from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import broker_api as ba                                         # noqa: E402
import momentum_screener as ms                                  # noqa: E402
import portfolio as pf                                          # noqa: E402

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                          "" if cond else "  <- " + str(detail)))


def holding(sym, rec, qty, legv=None):
    return {"Symbol": sym, "Mode": "LIVE", "Recommendation": rec, "Qty": qty,
            "Entry": 100.0, "Value (Rs)": 90.0 * qty, "P&L %": -10.0,
            "_door": 8.0, "_legv": legv or {}, "Kyun": "test", "Why": "test",
            "Aaj %": -1.5, "Aaj (Rs)": -135 * qty // 10,
            "_day": pd.Timestamp("2026-10-05").date()}


def main():
    d = tempfile.mkdtemp()
    ms.SPLIT_FILE = os.path.join(d, "split.csv")
    ba.sold_recently = lambda *a, **k: False
    sp = pd.DataFrame([{"symbol": "AAA", "swing_qty": 0, "investing_qty": 6,
                        "momentum_qty": 0, "entry_price": 100.0,
                        "entry_date": "2026-09-01", "strategy": "W+TT",
                        "mode": "LIVE", "product": "CNC", "order_id": "",
                        "note": ""}])
    hold = [holding("AAA", "EXIT", 10, {"investing": "EXIT"}),
            holding("BBB", "HOLD", 40), holding("PPP", "HOLD", 5)]
    hold[2]["Mode"] = "PAPER"
    demat = [{"symbol": "AAA", "qty": 10}, {"symbol": "BBB", "qty": 40}]

    print("1) rows: rule stock + every other LIVE demat stock")
    rows = pf.sell_rows(hold, sp, demat, "ON", {}, pd.Timestamp("2026-10-06")
                        .date())
    by = {x["Symbol"]: x for x in rows}
    check("AAA EXIT: default SELL, qty = firing leg 6 of 10",
          by["AAA"]["Sell?"] == "SELL" and by["AAA"]["Qty"] == 6 and
          by["AAA"]["Held qty"] == 10, by.get("AAA"))
    check("BBB HOLD: listed, blank Sell?, qty = all 40",
          by["BBB"]["Sell?"] == "" and by["BBB"]["Qty"] == 40, by.get("BBB"))
    check("PAPER stock not sellable", "PPP" not in by)
    off = {x["Symbol"]: x for x in pf.sell_rows(
        hold, sp, demat, "OFF", {}, pd.Timestamp("2026-10-06").date())}
    check("TRADING OFF: nothing pre-picked", off["AAA"]["Sell?"] == "")
    kept = {x["Symbol"]: x for x in pf.sell_rows(
        hold, sp, demat, "ON", {("BBB", "CNC"): ("SELL", 15),
                                ("AAA", "CNC"): ("", None)},
        pd.Timestamp("2026-10-06").date())}
    check("same-day picks kept: BBB SELL 15, AAA un-picked",
          kept["BBB"]["Sell?"] == "SELL" and kept["BBB"]["Qty"] == 15 and
          kept["AAA"]["Sell?"] == "", kept)
    big = {x["Symbol"]: x for x in pf.sell_rows(
        hold, sp, demat, "ON", {("BBB", "CNC"): ("SELL", 99)},
        pd.Timestamp("2026-10-06").date())}
    check("qty above held is cut to held (40)", big["BBB"]["Qty"] == 40)

    print("2) cards carry Sell? + qty, read back, rbtrack")
    path = os.path.join(d, "Portfolio_T.xlsx")
    pf.write_book(path, hold, [], pd.DataFrame(), {}, {},
                  "Portfolio T | 2026-10-06 | master x", sells=rows,
                  trading="ON")
    wb = load_workbook(path)
    check("no separate Sell sheet", "Sell" not in wb.sheetnames,
          wb.sheetnames)
    ws = wb["Holdings"]
    caps = [(c.row, c.column) for row in ws.iter_rows() for c in row
            if c.value == pf.SELL_CAP]
    check("2 LIVE cards have the Sell? box (PAPER none)", len(caps) == 2,
          caps)
    dvs = [str(v.sqref) for v in ws.data_validations.dataValidation]
    check("dropdown + qty validation on the inputs", len(dvs) == 2, dvs)
    got = {x["Symbol"]: x for x in pf.read_sell_table(path)}
    check("read back: AAA SELL 6, BBB blank 40", got["AAA"]["Sell?"] ==
          "SELL" and got["AAA"]["Qty"] == 6 and got["BBB"]["Sell?"] == ""
          and got["BBB"]["Held qty"] == 40, got)
    r, c = [rc for rc in caps if ws.cell(row=rc[0] - 10, column=rc[1]).value
            == "BBB"][0]
    ws.cell(row=r + 1, column=c, value="sell")      # you pick it, lower case
    ws.cell(row=r + 1, column=c + 1, value=12)
    wb.save(path)
    import auto_tracker_update as at
    s = {x["symbol"]: x for x in at.read_sells(path)}
    check("rbtrack: AAA 6 + BBB 12 (your qty)", s.get("AAA", {}).get("qty")
          == 6 and s.get("BBB", {}).get("qty") == 12, s)
    ws.cell(row=r + 1, column=c + 1).value = None
    wb.save(path)
    s = {x["symbol"]: x for x in at.read_sells(path)}
    check("blank qty = all held (40)", s["BBB"]["qty"] == 40, s)

    print("2b) today's gain / loss on the card")
    txt = [ws.cell(row=rc[0] - 8, column=rc[1] + 1).value for rc in caps]
    lab = [ws.cell(row=rc[0] - 8, column=rc[1]).value for rc in caps]
    check("card line: -1.50%  (-Rs 135) for AAA (qty 10)",
          "-1.50%  (-Rs 135)" in txt, txt)
    check("label says the session date when it is not today",
          all(str(x).startswith("Last session 05 Oct") for x in lab), lab)
    check("day_txt +/- and no data", pf.day_txt({"Aaj %": 2.0,
                                                 "Aaj (Rs)": 1234}) ==
          "+2.00%  (+Rs 1,234)" and pf.day_txt({}) == "-")
    wt = wb["Holdings_Table"]
    heads = [c.value for c in wt[1]]
    check("Holdings_Table has Aaj % + Aaj (Rs)", "Aaj %" in heads and
          "Aaj (Rs)" in heads, heads)

    print("2c) day change uses the last EARLIER day (two bars for today)")
    import numpy as np
    idx = list(pd.bdate_range(end="2026-10-05", periods=259)) + [
        pd.Timestamp("2026-10-06"), pd.Timestamp("2026-10-06 15:30")]
    c = pd.Series(np.linspace(1, 3, 261), index=pd.DatetimeIndex(idx))
    c.iloc[-3], c.iloc[-2], c.iloc[-1] = 2.10, 2.15, 2.15
    pos = {"symbol": "X", "mode": "LIVE", "qty": 554, "entry": 21.97,
           "entry_date": "2025-01-01", "legs": {}, "note": ""}
    h = pf.analyse(pos, {"X": c}, {"X": c}, {"X": c}, {}, {}, {}, {}, {})
    check("2.10 -> 2.15 = +2.38% (+Rs 28), not 0.00%",
          pf.day_txt(h) == "+2.38%  (+Rs 28)", pf.day_txt(h))
    band = [str(ws.cell(row=r, column=1).value) for r in range(1, 12)]
    check("LIVE band shows the account's AAJ total", any(
        "AAJ -Rs" in b for b in band), band)

    print("3) old file with a Sell sheet still works")
    old = os.path.join(d, "old.xlsx")
    w = Workbook()
    x = w.active
    x.title = "Sell"
    for i, h in enumerate(["Symbol", "Product", "Qty", "Sell?"], 1):
        x.cell(row=3, column=i, value=h)
    x.append(["ZZZ", "CNC", 7, "YES"])
    w.save(old)
    s = at.read_sells(old)
    check("old YES row read as SELL 7", s == [{"symbol": "ZZZ", "qty": 7,
                                              "product": "CNC"}], s)

    bad = [n for n, ok in RESULTS if not ok]
    print("\n%d / %d passed%s" % (len(RESULTS) - len(bad), len(RESULTS),
                                  "" if not bad else " -- FAILED: " +
                                  "; ".join(bad)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
