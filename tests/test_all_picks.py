#!/usr/bin/env python3
"""
ALL_PICKS sheet (6 Oct 2026, RB: momentum + swing + investing on one sheet,
with each stock's age) -- temp files only: no network, no token, NO order.
Run: python3 tests/test_all_picks.py
"""
import os
import sys
import tempfile

import pandas as pd
from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import signal_tracker as st                                     # noqa: E402

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                          "" if cond else "  <- " + str(detail)))


def master(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Swing"
    ws.append(["Symbol", "Action", "RS Rank", "Mcap (Rs Cr)", "Signal Date",
               "Days Since Signal", "Signal Price", "Price"])
    ws.append(["AAA", "FIT", 90, 20000, "2026-09-28", 6, 100.0, 110.0])
    ws.append(["CCC", "BUY", 80, 15000, "2026-10-05", 0, 50.0, 51.0])
    ws.append([])
    ws.append(["Stop % below entry (edit):", None, None, None, None, 0.2])
    wi = wb.create_sheet("Investing")
    wi.append(["Symbol", "Action", "RS Rank", "Mcap (Rs Cr)", "Signal Date",
               "Days Since Signal", "Signal Price", "Price"])
    wi.append(["AAA", "FIT", 90, 20000, "2026-09-28", 6, 100.0, 110.0])
    wi.append(["DDD", "LATE", 75, 12000, "2026-09-20", 12, 10.0, 12.0])
    wm = wb.create_sheet("Momentum_Top20")
    wm.append(["Symbol", "Mom Rank", "Days in Top 20", "In Top 20 since",
               "Momentum Score", "RS Rank", "Sector / Industry",
               "Last Close", "LTP"])
    wm.append(["BBB", 1, 10, "2026-09-26", 2.1, 99, "IT", 200.0, 210.0])
    wm.append(["AAA", 2, 3, "2026-10-03", 1.9, 90, "Pharma", 109.0, 110.0])
    wm.append([])
    wm.append(["NOTE", "text"])
    wb.save(path)


def main():
    d = tempfile.mkdtemp()
    path = os.path.join(d, "RB_Screener_2026-10-06.xlsx")
    master(path)
    log = pd.DataFrame([
        {"source": "MOMENTUM", "symbol": "BBB", "key": "2026-09-26",
         "first_found": "2026-09-26", "price_found": "180"},
        {"source": "MOMENTUM", "symbol": "AAA", "key": "2026-10-03",
         "first_found": "2026-10-03", "price_found": "105"}],
        columns=st.LOG_COLS).fillna("")
    p = st.all_picks(path, log, "2026-10-06")
    by = {r.Symbol: r for r in p.itertuples()}
    print("1) one row per stock, momentum first")
    check("4 stocks, no duplicates, no note rows",
          sorted(by) == ["AAA", "BBB", "CCC", "DDD"], list(p.Symbol))
    check("order: momentum rank 1, 2, then W+TT", list(p.Symbol)[:2] ==
          ["BBB", "AAA"], list(p.Symbol))
    check("AAA lists: MOM #2 + SWING FIT + INV FIT",
          by["AAA"]._5 == "MOM #2 + SWING FIT + INV FIT", by["AAA"]._5)
    check("AAA Super-Buy, DDD (LATE only) not", by["AAA"]._6 == "YES" and
          by["DDD"]._6 == "", (by["AAA"]._6, by["DDD"]._6))
    print("2) age")
    check("AAA age = older of signal 28 Sep / top-20 3 Oct = 8 days",
          by["AAA"]._3 == 8 and by["AAA"]._4 == "2026-09-28", by["AAA"])
    check("BBB: momentum since 26 Sep = 10 days, 6-29 label",
          by["BBB"]._3 == 10 and by["BBB"].Age == "6-29 days", by["BBB"])
    check("CCC signal yesterday = 1 day", by["CCC"]._3 == 1)
    check("price then / since found: BBB 180 -> 210 = +16.7%",
          abs(by["BBB"]._13 - 16.67) < 0.1, by["BBB"])
    check("BBB at the first scan date is flagged", "scans start" in
          by["BBB"]._14, by["BBB"]._14)
    print("3) sheet written first in the master, age colours")
    st.write_sheet(path, pd.DataFrame(), pd.DataFrame(columns=["Symbol"]),
                   "test", "2026-10-06", p)
    wb = load_workbook(path)
    check("All_Picks is the first tab", wb.sheetnames[0] == "All_Picks",
          wb.sheetnames)
    ws = wb["All_Picks"]
    heads = [c.value for c in ws[3]]
    check("header has Age + Age (din) + Kahan se", {"Age", "Age (din)",
          "Kahan se"} <= set(heads), heads)
    col = heads.index("Age (din)") + 1
    rows = {ws.cell(row=r, column=1).value: r for r in range(4, 8)}
    fill = ws.cell(row=rows["CCC"], column=col).fill.fgColor.rgb
    check("CCC (1 day) row painted blue", str(fill)[-6:].upper() ==
          dict(st.AGE_FILL)["1-5 days"].upper(), fill)
    check("tab is safe to rebuild (no duplicate sheet)",
          (st.write_sheet(path, pd.DataFrame(), pd.DataFrame(
              columns=["Symbol"]), "t", "2026-10-06", p) or True) and
          load_workbook(path).sheetnames.count("All_Picks") == 1)
    check("empty lists -> no sheet data", st.all_picks(
        os.path.join(d, "missing.xlsx"), log, "2026-10-06").empty)

    bad = [n for n, ok in RESULTS if not ok]
    print("\n%d / %d passed%s" % (len(RESULTS) - len(bad), len(RESULTS),
                                  "" if not bad else " -- FAILED: " +
                                  "; ".join(bad)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
