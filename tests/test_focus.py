#!/usr/bin/env python3
"""
MOMENTUM TOP 5 HIGHLIGHT TESTS (1 Oct 2026, RB approved A+C) -- temp files
only: no network, no token, no account, NO order.
Run:  cd ~/RB_Screener && python3 tests/test_focus.py

  1) focus = selected (in_top) AND rank 1-5 only; rank 6 / NaN / cap-skipped
     never gold; fewer than 5 is shown as fewer
  2) cohort = MOMENTUM finds with Rank then 1-5 (not W+TT RS 3)
  3) Momentum_Top20: gold Symbol + Mom Rank only on rank 1-5 rows
  4) Signal_Summary: new columns, labels, values unchanged; gold group row;
     Signal_Tracker gold symbol survives paint_ages
  5) Dashboard panel lines + Actions sheet code path
"""
import os
import sys
import tempfile

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import momentum_focus as mf                                     # noqa: E402
import momentum_screener as ms                                  # noqa: E402
import signal_tracker as st                                     # noqa: E402

RESULTS = []
TMP = tempfile.mkdtemp()


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                          "" if cond else "  <- " + str(detail)))


def ranks_table():
    # raw rank 3 skipped by the sector cap (in_top False), rank 6 selected,
    # a NaN rank, a duplicate and string booleans
    return pd.DataFrame({
        "symbol": ["aaa", "BBB", "CCC", "DDD", "EEE", "FFF", "GGG", "AAA"],
        "rank": [1, 2, 3, 4, 5, 6, np.nan, 1],
        "in_top": ["True", True, "False", "true", "TRUE", True, True, True],
        "score": [9, 8, 7, 6, 5, 4, 3, 9], "sector": ["X"] * 8,
        "price": [100.0] * 8, "date": ["2026-09-30"] * 8})


def main():
    print("1) focus membership")
    f = mf.focus(ranks_table())
    check("rank 1,2,4,5 selected only (3 cap-skipped, 6 + NaN out, no dup)",
          [x["symbol"] for x in f] == ["AAA", "BBB", "DDD", "EEE"],
          [x["symbol"] for x in f])
    check("empty / missing table -> []", mf.focus(None) == [] and
          mf.focus(pd.DataFrame()) == [])
    p = os.path.join(TMP, "ranks.csv")
    ranks_table().to_csv(p, index=False)
    rows, day, msg = mf.load(p)
    check("load() reads the file + date", len(rows) == 4 and
          day == "2026-09-30" and msg == "", (len(rows), day, msg))
    rows, day, msg = mf.load(os.path.join(TMP, "missing.csv"))
    check("missing file -> explicit message, no crash", rows == [] and msg)

    print("2) cohort (rank WHEN FOUND)")
    t = pd.DataFrame({
        "Symbol": ["A", "B", "C", "D", "E", "F", "A", "G"],
        "Source": ["MOMENTUM"] * 4 + ["W+TT", "MOMENTUM", "MOMENTUM",
                                      "MOMENTUM"],
        "Rank then": [1, 3, 5, 8, 3, 2, 4, 23],
        "Return %": [10.0, -4.0, 6.0, 1.0, 50.0, np.nan, 2.0, 1.0],
        "vs Nifty %": [9.0, -5.0, np.nan, 0.0, 49.0, 1.0, 1.0, 0.5],
        "Days since found": [4, 4, 4, 4, 4, 4, 35, 4]})
    c = mf.cohort(t)
    check("W+TT RS 3, rank 8 and no-return rows excluded -> 4 signals",
          c and c["signals"] == 4, c)
    check("unique stocks 3 (A twice)", c and c["stocks"] == 3, c)
    check("beat Nifty counts only compared finds: 2 of 3",
          c and (c["beat"], c["compared"]) == (2, 3), c)
    check("avg / median right", c and abs(c["avg"] - 3.5) < 1e-9 and
          abs(c["median"] - 4.0) < 1e-9, c)
    check("days 4-35, one 30+ old", c and (c["days_min"], c["days_max"],
                                          c["old"]) == (4, 35, 1), c)
    line = mf.cohort_line(c)
    check("line says TOO EARLY + sample size", "TOO EARLY" in line and
          "4 signals (3 stocks)" in line, line)

    print("3) Momentum_Top20 gold cells")
    from openpyxl import Workbook, load_workbook
    path = os.path.join(TMP, "RB_Screener_test.xlsx")
    wb = Workbook()
    wb.active.title = "Swing"
    wb.save(path)
    n = 7
    top = pd.DataFrame({
        "symbol": ["S%d" % i for i in range(1, n + 1)],
        "rank": [1, 2, 4, 5, 6, 7, 8],       # raw 3 was cap-skipped
        "score": np.linspace(5, 1, n), "rs_rank": 90, "sector": "X",
        "close": 100.0, "price": 100.0, "px_src": "live", "atr_pct": 3.0,
        "slot": 10000.0, "shares": 100, "amount": 10000.0, "ret6": 1.0,
        "ret12": 2.0, "vol": 30.0, "dist52": -5.0})
    ms.days_in_top = lambda syms: {s: (0, "") for s in syms}
    ms.write_sheets(path, top, top, pd.DataFrame(), {}, False, "banner")
    ws = load_workbook(path)["Momentum_Top20"]
    gold = [ws.cell(row=r, column=1).value for r in range(2, n + 2)
            if ws.cell(row=r, column=1).fill.fgColor.rgb.endswith(mf.GOLD_FILL)]
    check("gold = ranks 1,2,4,5 (S1-S4), not rank 6+", gold ==
          ["S1", "S2", "S3", "S4"], gold)
    check("Mom Rank cell gold too", ws.cell(row=2, column=2).fill.fgColor
          .rgb.endswith(mf.GOLD_FILL))
    check("other columns keep their fill", not ws.cell(row=2, column=5).fill
          .fgColor.rgb.endswith(mf.GOLD_FILL))
    check("note explains gold != buy", any(
        "NOT a buy signal" in str(ws.cell(row=r, column=1).value)
        for r in range(n + 2, ws.max_row + 1)))

    print("4) Signal_Summary + Signal_Tracker")
    t2 = t.assign(**{"First status": "", "Both lists": "",
                     "Worst since %": -1.0, "Status": "TOO EARLY",
                     "Age": "1-5 days", "Found on": "2026-09-26"})
    summ = st.summary(t2)
    g = summ.set_index("Group")
    check("rank 1-5 group values unchanged by the new columns",
          abs(g.loc["  momentum rank 1-5 when found", "Avg return %"] - 3.5)
          < 1e-9 and g.loc["  momentum rank 1-5 when found", "Finds"] == 4)
    check("new columns Stocks / Days tracked / Best now % / Worst now %",
          all(k in summ.columns for k in ("Stocks", "Days tracked",
                                          "Best now %", "Worst now %")),
          list(summ.columns))
    check("label 'rank 11+ (selected)'", any("11+ (selected)" in x
                                              for x in summ["Group"]))
    mf.RANKS = p
    old_load = mf.load
    mf.load = lambda path=p: old_load(path)
    tcols = ["Symbol", "Age", "Source", "Rank then", "Days since found",
             "Return %"]
    tt = pd.DataFrame({"Symbol": ["AAA", "CCC", "AAA"],
                       "Age": "1-5 days",
                       "Source": ["MOMENTUM", "MOMENTUM", "W+TT"],
                       "Rank then": [1, 3, 80], "Days since found": [4, 4, 4],
                       "Return %": [1.0, 2.0, 3.0]})[tcols]
    st.write_sheet(path, summ, tt, "test", "2026-10-01")
    wb = load_workbook(path)
    wsum = wb["Signal_Summary"]
    grow = [r for r in range(4, 4 + len(summ))
            if "rank 1-5" in str(wsum.cell(row=r, column=1).value)]
    check("summary group cell gold", grow and wsum.cell(row=grow[0], column=1)
          .fill.fgColor.rgb.endswith(mf.GOLD_FILL))
    wst = wb["Signal_Tracker"]
    fills = {(wst.cell(row=r, column=1).value,
              wst.cell(row=r, column=3).value):
             wst.cell(row=r, column=1).fill.fgColor.rgb for r in (4, 5, 6)}
    check("AAA MOMENTUM (today rank 1) gold after paint_ages",
          fills[("AAA", "MOMENTUM")].endswith(mf.GOLD_FILL), fills)
    check("CCC (cap-skipped today) and W+TT AAA not gold",
          not fills[("CCC", "MOMENTUM")].endswith(mf.GOLD_FILL) and
          not fills[("AAA", "W+TT")].endswith(mf.GOLD_FILL), fills)
    check("age colour kept on the other cells", wst.cell(row=4, column=2)
          .fill.fgColor.rgb.endswith("DDEBF7"))

    print("5) Dashboard panel")
    import portfolio as pf
    mf.load_cohort = lambda path=None: mf.cohort(t)
    pf.SLOT_RS, pf.SLOT_NOTE = 10000, "test"

    class Acc(object):
        label = "TEST"
    hold = [{"Symbol": "BBB", "Mode": "LIVE", "Recommendation": "HOLD",
             "Qty": 1, "Entry": 1.0, "Value (Rs)": 1.0}]
    d = pf.dashboard_data(Acc(), "2026-10-01", "x.xlsx", "n", False, hold, [],
                          pd.DataFrame(), pd.DataFrame(), [])
    heads = [a for a, _, _ in d["focus"]]
    check("panel lists #1 AAA, #2 BBB, #4 DDD, #5 EEE", heads[:4] ==
          ["#1 AAA", "#2 BBB", "#4 DDD", "#5 EEE"], heads)
    check("held stock marked HELD", "HELD" in d["focus"][1][1])
    check("panel says fewer than 5 + sample-size line", any(
        "skipped" in b for _, b, _ in d["focus"]) and any(
        "TOO EARLY" in b for _, b, _ in d["focus"]))
    wb = Workbook()
    pf.dashboard(wb.active, d, {"Momentum_Top20", "Signal_Summary"})
    vals = [wb.active.cell(row=r, column=1).value
            for r in range(1, wb.active.max_row + 1)]
    check("Dashboard renders the TOP 5 section", any(
        "MOMENTUM TOP 5" in str(v) for v in vals))

    bad = [n for n, ok in RESULTS if not ok]
    print("\n%d / %d passed%s" % (len(RESULTS) - len(bad), len(RESULTS),
                                  "" if not bad else " -- FAILED: " +
                                  "; ".join(bad)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
