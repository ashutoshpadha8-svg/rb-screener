"""MTF8.1 layout (RB 9 Oct): header over its numbers (same alignment), rupee money format."""
import datetime as dt,sys,tempfile,types,unittest,io,contextlib
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
import mtf_check as m
import mtf_portfolio_report as p
from test_mtf_breakeven import fixture
TODAY=dt.date(2026,10,9)
def model():
 r=fixture();r['loan']=m.Num(258386.43);r['current_lots']=r['fifo']['lots'];r['unpaid']=m.Num(176.32,m.ESTIMATED,'x');r['report_account']='DEMO'
 for x in r['rows']:x['Price']=2156 if x['Symbol']=='TCS' else 5770;x['Price source']='NSE demo'
 return r
class LayoutTests(unittest.TestCase):
 def setUp(self):
  self.d=tempfile.TemporaryDirectory();self.path=Path(self.d.name)/'MTF_Check.xlsx';self.r=model()
  self.rows=p.write_report(self.r,TODAY,self.path);self.s=load_workbook(self.path).active
 def tearDown(self):self.d.cleanup()
 def find(self,text):
  return next(c for row in self.s.iter_rows() for c in row if c.value==text)
 def test_headers_and_values_share_alignment(self):
  for head in ('Qty','Buy cost','Loan proxy','Interest MODEL','Pure BE now','Safe hold+exit+tax'):
   h=self.find(head);v=self.s.cell(h.row+1,h.column)
   self.assertEqual(h.alignment.horizontal,'center',head);self.assertEqual(v.alignment.horizontal,'center',head)
  m3=self.find('Metric')
  for c in range(2,2+len(self.rows)+1):
   self.assertEqual(self.s.cell(m3.row,c).alignment.horizontal,'center');self.assertEqual(self.s.cell(m3.row+1,c).alignment.horizontal,'center')
  self.assertEqual(self.find('Market value').alignment.horizontal,'left')
 def test_money_has_rupee_qty_and_days_do_not(self):
  self.assertIn('₹',self.find('Market value').offset(0,1).number_format)
  q=self.find('Qty');self.assertNotIn('₹',self.s.cell(q.row+1,q.column).number_format)
  self.assertNotIn('₹',self.s['F4'].number_format);self.assertIn('%',self.s['D4'].number_format)
  be=self.find('Pure BE now');self.assertIn('₹',self.s.cell(be.row+1,be.column).number_format)
 def test_terminal_rupee(self):
  buf=io.StringIO()
  with contextlib.redirect_stdout(buf):p.print_report(self.r,self.rows,types.SimpleNamespace(label='DEMO'),TODAY,types.SimpleNamespace(rate=.1249,buffer_days=30,settlement_buffer=3,tax_reserve=.208))
  out=buf.getvalue();self.assertTrue(p.RUPEE+'414,598.30' in out);self.assertIn('-'+p.RUPEE+'48,156.60',out)
if __name__=='__main__':unittest.main()
