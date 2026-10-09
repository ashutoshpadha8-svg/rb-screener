"""Official quote identity, freshness and fail-closed fallback regression tests."""
import copy,datetime as dt,sys,types,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import mtf_prices as q
import mtf_check as m
import broker_api as ba
import daily_screener as ds
NOW=dt.datetime(2026,10,9,16,10,tzinfo=q.IST)
DAY=NOW.date()
def payload(symbol='TCS',price=2156,stamp='09-Oct-2026 16:00:00'):
 return {'equityResponse':[{'metaData':{'symbol':symbol,'series':'EQ'},'orderBook':{'lastPrice':price},'lastUpdateTime':stamp}]}
class Response:
 def __init__(self,value):self.value=value
 def json(self):return self.value
 def raise_for_status(self):pass
class HTTP:
 def __init__(self,meta=None,value=None):
  self.headers={};self.calls=[];self.meta=meta or {'symbol':'TCS','activeSeries':['EQ'],'marketType':'N','isSuspended':'false'};self.value=value or payload()
 def get(self,url,params,timeout):
  self.calls.append((url,params));return Response(self.meta if params['functionName']=='getMetaData' else self.value)
class QuoteTests(unittest.TestCase):
 def test_official_price_has_exchange_time(self):
  p,status,note=q.parse_quote(payload(),'TCS',NOW,DAY,False)
  self.assertEqual(p,2156);self.assertEqual(status,'VERIFIED');self.assertIn('2026-10-09 16:00:00 IST',note)
 def test_intraday_fresh(self):
  n=NOW.replace(hour=12);self.assertEqual(q.parse_quote(payload(stamp='09-Oct-2026 12:05:00'),'TCS',n,DAY-dt.timedelta(days=1),True)[0],2156)
 def test_intraday_yesterday_rejected(self):
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='08-Oct-2026 16:00:00'),'TCS',NOW,DAY,True)
 def test_intraday_16_minutes_rejected(self):
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='09-Oct-2026 15:54:00'),'TCS',NOW,DAY,True)
 def test_weekend_last_session_accepted(self):
  self.assertEqual(q.parse_quote(payload(),'TCS',NOW+dt.timedelta(days=1),DAY,False)[0],2156)
 def test_prior_session_rejected(self):
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='08-Oct-2026 16:00:00'),'TCS',NOW,DAY,False)
 def test_future_rejected(self):
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='09-Oct-2026 16:13:00'),'TCS',NOW,DAY,False)
 def test_missing_time_rejected(self):
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp=None),'TCS',NOW,DAY,False)
 def test_invalid_prices_rejected(self):
  for p in [0,-1,None,'nan','inf','bad']:
   with self.subTest(p=p),self.assertRaises(ValueError):q.parse_quote(payload(price=p),'TCS',NOW,DAY,False)
 def test_wrong_stock_rejected(self):
  with self.assertRaises(ValueError):q.parse_quote(payload('OTHER'),'TCS',NOW,DAY,False)
 def test_wrong_series_rejected(self):
  p=payload();p['equityResponse'][0]['metaData']['series']='T0'
  with self.assertRaises(ValueError):q.parse_quote(p,'TCS',NOW,DAY,False)
 def test_ambiguous_rejected(self):
  p=payload();p['equityResponse']*=2
  with self.assertRaises(ValueError):q.parse_quote(p,'TCS',NOW,DAY,False)
 def test_malformed_nested_schema_rejected(self):
  for field in ['metaData','orderBook']:
   p=payload();p['equityResponse'][0][field]=None
   with self.subTest(field=field),self.assertRaises(ValueError):q.parse_quote(p,'TCS',NOW,DAY,False)
 def test_preclose_not_completed_session(self):
  with self.assertRaises(ValueError):q.parse_quote(payload(stamp='09-Oct-2026 15:20:00'),'TCS',NOW,DAY,False)
 def test_endpoint_requests_only_read_quotes(self):
  h=HTTP();self.assertEqual(q.nse_quotes(['TCS'],NOW,DAY,False,h)['TCS'][0],2156)
  self.assertEqual([p['functionName'] for _,p in h.calls],['getMetaData','getSymbolData'])
  self.assertTrue(all(url==q.URL for url,_ in h.calls));self.assertEqual(h.calls[1][1]['series'],'EQ')
 def test_metadata_mismatch_no_quote_request(self):
  h=HTTP(meta={'symbol':'WRONG','activeSeries':['EQ'],'marketType':'N'})
  self.assertEqual(q.nse_quotes(['TCS'],NOW,DAY,False,h),{});self.assertEqual(len(h.calls),1)
 def test_suspended_rejected(self):
  h=HTTP();h.meta['isSuspended']=True
  self.assertEqual(q.nse_quotes(['TCS'],NOW,DAY,False,h),{})
 def test_nse_success_skips_broker(self):
  s=types.SimpleNamespace(broker='DHAN')
  with patch.object(q,'nse_quotes',return_value={'TCS':(2156,'VERIFIED','NSE timestamp')}),patch.object(ba,'live_prices',side_effect=AssertionError('must not call broker')):
   self.assertEqual(m.quotes(s,['TCS'])['TCS'][0],2156)
 def test_failed_sources_dated_fallback(self):
  s=types.SimpleNamespace(broker='DHAN')
  with patch.object(q,'nse_quotes',return_value={}),patch.object(ba,'live_prices',side_effect=RuntimeError('down')),patch.object(q,'data_plan_note',return_value=''),patch.object(m,'last_closes',return_value={'TCS':(2076,DAY-dt.timedelta(days=1))}),patch.object(ds,'last_expected_session',return_value=DAY):
   p,status,note=m.quotes(s,['TCS'])['TCS'];self.assertEqual(status,'ESTIMATED');self.assertIn('NOT a current live quote',note)
 def test_no_source_means_missing_not_zero(self):
  with patch.object(q,'nse_quotes',return_value={}),patch.object(ba,'live_prices',return_value={}),patch.object(m,'last_closes',return_value={}):
   self.assertEqual(m.quotes(types.SimpleNamespace(broker='DHAN'),['TCS']),{})
 def test_data_plan_failure_is_explained(self):
  with patch.object(ba,'_call',return_value={'dataPlan':'Deactive'}):
   self.assertIn('Deactive',q.data_plan_note(ba,types.SimpleNamespace(broker='DHAN')))
 def test_secrets_redacted(self):
  s=types.SimpleNamespace(token='SECRET',_jwt='OTHER_SECRET')
  self.assertNotIn('SECRET',q.safe_error(RuntimeError('SECRET OTHER_SECRET'),s))
 def test_excel_recalculation_cannot_create_unknown_funding(self):
  import tempfile
  import mtf_breakeven as b
  from test_mtf_breakeven import fixture,TODAY
  from openpyxl import load_workbook
  r=fixture();r['open_cost_num']=m.Num(None,m.UNKNOWN,'unreconciled historical inventory')
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'report.xlsx';b.write_report(r,TODAY,p)
   w=load_workbook(p,data_only=False).active
   # Literal empty formulas survive Excel/LibreOffice recalculation; a SUM
   # of candidate row costs cannot create a verified funding denominator.
   self.assertEqual(w['G64'].value,'=""');self.assertEqual(w['H64'].value,'=""')
   self.assertEqual(load_workbook(p,data_only=True).active['G18'].value,'UNKNOWN')
 def test_workbook_price_source_visible_without_comments(self):
  import tempfile
  import mtf_breakeven as b
  from test_mtf_breakeven import fixture,TODAY
  from openpyxl import load_workbook
  r=fixture();r['rows'][0]['Price source']='NSE official LTP DUMMY timestamp'
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'report.xlsx';b.write_report(r,TODAY,p)
   self.assertIn('NSE official LTP DUMMY timestamp',load_workbook(p,data_only=True).active['A8'].value)
if __name__=='__main__':unittest.main()
