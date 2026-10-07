#!/usr/bin/env python3
"""
EMA 9/21 CROSS SCREENER (7 Oct 2026, RB) -- synthetic prices only, no
network, no token, NO order.  Run: python3 tests/test_ema_screener.py
"""
import os
import sys
import tempfile

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import ema_screener as es                                       # noqa: E402

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                          "" if cond else "  <- " + str(detail)))


def series(after, dip=12):
    """300 rising days, a `dip`-day pullback (EMA9 < EMA21, still far above
    EMA100/200), then `after` = list of daily % moves after the low."""
    up = 100 * np.cumprod(np.full(300, 1.004))
    down = up[-1] * np.cumprod(np.full(dip, 0.985))
    px = list(up) + list(down)
    for m in after:
        px.append(px[-1] * (1 + m / 100.0))
    idx = pd.bdate_range("2025-01-01", periods=len(px))
    return pd.Series(px, index=idx)


def first_cross(c):
    e9, e21 = es.ema(c, 9), es.ema(c, 21)
    a = (e9 > e21).values
    return [i for i in range(1, len(a)) if a[i] and not a[i - 1]]


def main():
    print("1) cross found, days counted")
    c = series([3.0] * 12)
    xs = first_cross(c)
    check("synthetic series has a recent cross", xs and xs[-1] > 300, xs)
    age = len(c) - 1 - xs[-1]
    r = es.check_stock(c)
    check("rising after the cross -> CONFIRMED, age right",
          r and r["status"] == "CONFIRMED" and r["age"] == age, (r, age))
    check("close above EMA100 and EMA200", r and r["close"] > r["e100"] and
          r["close"] > r["e200"], r)

    print("2) only 0-2 sessions after the cross -> WAITING")
    cut = c.iloc[:xs[-1] + 2]                  # cross day + 1 day
    r = es.check_stock(cut)
    check("age 1 -> WAITING", r and r["status"] == "WAITING" and
          r["age"] == 1, r)
    r = es.check_stock(cut, hold=1)
    check("--hold 1 -> CONFIRMED", r and r["status"] == "CONFIRMED", r)

    print("3) stock gave back the cross level -> out")
    d = c.iloc[:xs[-1] + 4].copy()
    d.iloc[-2] = d.iloc[xs[-1]] * 0.995      # below the cross-day close
    r = es.check_stock(d)
    check("close below cross-day close after the cross -> None", r is None, r)

    print("4) old cross -> out; EMA9 back below -> out")
    r = es.check_stock(c, max_age=age - 1)
    check("cross older than max_age -> None", r is None, r)
    e = c.copy()
    e = pd.concat([e, pd.Series(e.iloc[-1] * np.cumprod(np.full(15, 0.97)),
                                index=pd.bdate_range(e.index[-1] +
                                                     pd.Timedelta(days=1),
                                                     periods=15))])
    r = es.check_stock(e)
    check("falls after -> None", r is None, r)

    print("5) price below EMA200 -> out (downtrend cross)")
    dn = pd.Series(100 * np.cumprod(np.full(320, 0.997)),
                   index=pd.bdate_range("2025-01-01", periods=320))
    dn = pd.concat([dn, pd.Series(dn.iloc[-1] * np.cumprod(np.full(8, 1.02)),
                                  index=pd.bdate_range(dn.index[-1] +
                                                       pd.Timedelta(days=1),
                                                       periods=8))])
    xs2 = first_cross(dn)
    check("downtrend series has a cross", xs2 and xs2[-1] > 300, xs2)
    check("but below EMA200 -> None", es.check_stock(dn) is None)

    print("6) scan uses only completed bars + watchlist = found only")
    frames = {"AAA": pd.DataFrame({"Close": c}),
              "BBB": pd.DataFrame({"Close": dn})}
    res = es.scan(frames, want=c.index[-1].date())
    check("scan: only AAA", list(res["symbol"]) == ["AAA"], res)
    res2 = es.scan(frames, want=c.index[xs[-1] + 1].date())
    check("scan cut at `want` -> AAA WAITING", len(res2) == 1 and
          res2["status"].iloc[0] == "WAITING", res2)
    t = es.table(res, {"AAA": 123.4}, {"AAA": 7})
    check("table: LTP + Mom Rank + Since Cross %", t["LTP (live)"].iloc[0]
          == 123.4 and t["Mom Rank"].iloc[0] == 7 and
          t["Since Cross %"].iloc[0] > 0, t.T)
    d = tempfile.mkdtemp()
    wl = es.write_watchlist(["AAA", "CCC", "AAA"], os.path.join(d, "w.txt"))
    check("watchlist: only found stocks, no duplicates",
          open(wl).read() == "###EMA 9-21 Cross,NSE:AAA,NSE:CCC\n")
    es.write_watchlist([], wl)
    check("nothing found -> empty watchlist", open(wl).read() == "")
    path = os.path.join(d, "EMA.xlsx")
    es.write_book(path, t, es.table(res2, {}, {}), "banner")
    from openpyxl import load_workbook
    wb = load_workbook(path)
    check("Excel: EMA_Confirmed + EMA_Waiting", wb.sheetnames[:2] ==
          ["EMA_Confirmed", "EMA_Waiting"], wb.sheetnames)
    check("Excel row: AAA", wb["EMA_Confirmed"].cell(row=5, column=1).value
          == "AAA")

    bad = [n for n, ok in RESULTS if not ok]
    print("\n%d / %d passed%s" % (len(RESULTS) - len(bad), len(RESULTS),
                                  "" if not bad else " -- FAILED: " +
                                  "; ".join(bad)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
