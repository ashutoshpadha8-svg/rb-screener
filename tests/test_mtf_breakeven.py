"""Current-lot breakeven regression; dummy account, no broker/network."""
import contextlib,datetime as dt,io,sys,tempfile,types,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import mtf_check as m
import mtf_breakeven as b
TODAY=dt.date(2026,10,7)
def fixture():
 lots={'TCS':[['2025-11-24',32,3175.8,0],['2026-02-04',10,3006.3,0],['2026-02-12',10,2858,0]],'PERSISTENT':[['2025-12-10',5,6142,0],['2026-01-21',5,6115,0],['2026-01-21',3,6066,0],['2026-02-04',9,5977.5,0],['2026-02-12',9,5658.5,0],['2026-05-14',15,70122.7/15,0]]}
 rows=[{'Symbol':s,'Qty':sum(x[1] for x in v),'Open cost (Rs)':m.Num(round(sum(x[1]*x[2] for x in v),2),m.ESTIMATED),'Price':2100 if s=='TCS' else 5521} for s,v in lots.items()]
 return dict(loan=m.Num(261888.33),unpaid=m.Num(0, m.VERIFIED, 'synthetic no outstanding interest'),open_cost_num=m.Num(414598.3),fifo={'lots':lots},rows=rows)
class BreakevenTests(unittest.TestCase):
 def test_tcs_expected_broker_target(self):self.assertAlmostEqual(b.calculate(fixture(),TODAY)[0]['broker'],3307.35)
 def test_persistent_expected_broker_target(self):self.assertAlmostEqual(b.calculate(fixture(),TODAY)[1]['broker'],5857.2)
 def test_tax_targets(self):
  r=b.calculate(fixture(),TODAY);self.assertAlmostEqual(r[0]['full'],3366.6);self.assertAlmostEqual(r[1]['full'],5943.55)
 def test_net_positive_at_rounded_targets(self):
  for r in b.calculate(fixture(),TODAY):self.assertGreaterEqual(r['net'],2)
 def test_one_tick_below_fails_buffer(self):
  for r in b.calculate(fixture(),TODAY):
   v=r['qty']*(r['full']-.05);net=v-b.sell_cost(v)-max(v-r['cost'],0)*.208-r['cost']-r['past']-r['future']-r['buy'];self.assertLess(net,2)
 def test_loan_allocations_sum_to_current_loan(self):self.assertAlmostEqual(sum(r['funded'] for r in b.calculate(fixture(),TODAY)),261888.33)
 def test_interest_buffer_is_33_calendar_days(self):
  for r in b.calculate(fixture(),TODAY):self.assertAlmostEqual(r['future'],r['funded']*.1249*33/365)
 def test_more_holding_days_raise_target(self):self.assertGreater(b.calculate(fixture(),TODAY,days=60)[0]['full'],b.calculate(fixture(),TODAY,days=0)[0]['full'])
 def test_tax_disabled_equals_broker_target(self):
  for r in b.calculate(fixture(),TODAY,tax=0):self.assertEqual(r['full'],r['broker'])
 def test_no_ledger_double_count(self):
  r=fixture();original=b.calculate(r,TODAY);r['paid']=m.Num(999999);r['unpaid']=m.Num(999999);strip=lambda L:[{k:v for k,v in x.items() if k not in ('loan_share','unpaid_share','cash_now')} for x in L];self.assertEqual(strip(b.calculate(r,TODAY)),strip(original))
 def test_missing_loan_unknown(self):
  r=fixture();r['loan']=m.Num();self.assertTrue(all(x['broker'] is None for x in b.calculate(r,TODAY)))
 def test_excess_loan_unknown(self):
  r=fixture();r['loan']=m.Num(999999);self.assertEqual(b.calculate(r,TODAY)[0]['status'],'UNKNOWN')
 def test_zero_loan_no_interest(self):
  r=fixture();r['loan']=m.Num(0);a=b.calculate(r,TODAY)[0];self.assertEqual(a['past']+a['future'],0);self.assertIsNotNone(a['broker'])
 def test_missing_lot_unknown(self):
  r=fixture();r['fifo']['lots']['TCS']=[];self.assertIsNone(b.calculate(r,TODAY)[0]['full'])
 def test_future_lot_unknown(self):
  r=fixture();r['fifo']['lots']['TCS'][0][0]='2027-01-01';self.assertIsNone(b.calculate(r,TODAY)[0]['full'])
 def test_mismatched_cost_unknown(self):
  r=fixture();r['rows'][0]['Open cost (Rs)']=m.Num(1);self.assertIsNone(b.calculate(r,TODAY)[0]['full'])
 def test_invalid_scenario_unknown(self):self.assertIsNone(b.calculate(fixture(),TODAY,days=-1)[0]['full'])
 def test_uncapped_brokerage_solver(self):
  p=b.target(1,1000,10,5,3,tax=.208);v=p;self.assertGreaterEqual(v-b.sell_cost(v)-max(v-1000,0)*.208-1018,2)
 def test_capped_brokerage_solver(self):
  p=b.target(100,100000,1000,200,150,tax=.208);v=p*100;self.assertGreaterEqual(v-b.sell_cost(v)-max(v-100000,0)*.208-101350,2)
 def test_invalid_target_inputs(self):
  for q,c,t in [(0,100,.208),(1,0,.208),(1,100,float('nan'))]:self.assertIsNone(b.target(q,c,0,0,0,tax=t))
 def test_one_sheet_formulas_and_caches(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'test.xlsx';b.write_report(fixture(),TODAY,p);f=load_workbook(p);v=load_workbook(p,data_only=True)
   self.assertEqual(f.sheetnames,['Breakeven']);self.assertTrue(f.active['E10'].value.startswith('='));self.assertAlmostEqual(v.active['E10'].value,3366.6);self.assertEqual(v.active['B10'].value,52);self.assertFalse(f._external_links);self.assertEqual(len(f.active.data_validations.dataValidation),7)
 def test_unknown_sheet_keeps_unknown(self):
  r=fixture();r['loan']=m.Num()
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'test.xlsx';b.write_report(r,TODAY,p);self.assertEqual(load_workbook(p,data_only=True).active['E10'].value,'UNKNOWN')
 def test_default_cli_one_sheet(self):
  with tempfile.TemporaryDirectory() as d:
   acc=types.SimpleNamespace(activate=lambda:types.SimpleNamespace(token_ok=True,session=types.SimpleNamespace(broker='DHAN'),label='DUMMY',reports=d))
   with patch.dict(sys.modules,{'account':acc}),patch.object(m,'build',lambda *a:fixture()),patch.object(sys,'argv',['mtf_check.py']),contextlib.redirect_stdout(io.StringIO()):self.assertEqual(m.main(),0)
   files=list(Path(d).glob('*.xlsx'));self.assertEqual(len(files),1);self.assertEqual(load_workbook(files[0]).sheetnames,['Breakeven'])
 def test_invalid_new_inputs_before_account(self):
  acc=types.SimpleNamespace(activate=lambda:(_ for _ in ()).throw(AssertionError('must not activate')))
  for args in [['--buffer-days','-1'],['--settlement-buffer','99'],['--tax-reserve','nan']]:
   with patch.dict(sys.modules,{'account':acc}),patch.object(sys,'argv',['mtf_check.py']+args),contextlib.redirect_stdout(io.StringIO()):self.assertEqual(m.main(),1)
 def test_current_gross_loss(self):
  r=b.calculate(fixture(),TODAY);self.assertAlmostEqual(r[0]['price_pnl'],-51068.6);self.assertAlmostEqual(r[1]['price_pnl'],-363.7)
 def test_current_interest_and_charges_subtracted_once(self):
  for r in b.calculate(fixture(),TODAY):
   self.assertAlmostEqual(r['with_interest'],r['price_pnl']-r['past'])
   self.assertAlmostEqual(r['now_net_tax'],r['price_pnl']-r['past']-r['buy']-b.sell_cost(r['now_value'])-r['now_tax'])
 def test_loss_has_no_tax_credit(self):
  for r in b.calculate(fixture(),TODAY):self.assertEqual(r['now_tax'],0);self.assertEqual(r['now_net'],r['now_net_tax'])
 def test_price_profit_can_still_be_net_loss(self):
  f=fixture();f['rows'][0]['Price']=3100;r=b.calculate(f,TODAY)[0]
  self.assertGreater(r['price_pnl'],0);self.assertLess(r['with_interest'],0);self.assertLess(r['now_net_tax'],0)
 def test_true_net_profit_positive(self):
  f=fixture();f['rows'][0]['Price']=3500;r=b.calculate(f,TODAY)[0]
  self.assertGreater(r['price_pnl'],0);self.assertGreater(r['with_interest'],0);self.assertGreater(r['now_net_tax'],0)
  self.assertAlmostEqual(r['now_tax'],r['price_pnl']*.208)
 def test_profit_tax_disabled(self):
  f=fixture();f['rows'][0]['Price']=3500;r=b.calculate(f,TODAY,tax=0)[0]
  self.assertEqual(r['now_tax'],0);self.assertEqual(r['now_net'],r['now_net_tax'])
 def test_current_profit_excludes_future_interest(self):
  a=b.calculate(fixture(),TODAY,days=0,settlement=0);c=b.calculate(fixture(),TODAY,days=100,settlement=20)
  self.assertEqual([r['now_net_tax'] for r in a],[r['now_net_tax'] for r in c]);self.assertNotEqual(a[0]['full'],c[0]['full'])
 def test_missing_quote_net_unknown(self):
  f=fixture();f['rows'][0]['Price']=None;r=b.calculate(f,TODAY)[0]
  self.assertIsNone(r['price_pnl']);self.assertIsNone(r['now_net_tax']);self.assertIsNotNone(r['broker'])
 def test_zero_quote_not_treated_as_full_loss(self):
  f=fixture();f['rows'][0]['Price']=0;r=b.calculate(f,TODAY)[0];self.assertIsNone(r['price_pnl']);self.assertIsNone(r['now_net_tax'])
 def test_missing_loan_preserves_price_pnl(self):
  f=fixture();f['loan']=m.Num();r=b.calculate(f,TODAY)[0]
  self.assertAlmostEqual(r['price_pnl'],-51068.6);self.assertIsNone(r['with_interest']);self.assertIsNone(r['now_net_tax'])
 def test_loss_sheet_cache_and_formulas(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'loss.xlsx';rows=b.write_report(fixture(),TODAY,p);w=load_workbook(p,data_only=True).active;f=load_workbook(p).active
   self.assertAlmostEqual(w['G18'].value,rows[0]['now_net_tax']);self.assertAlmostEqual(w['G20'].value,sum(r['now_net_tax'] for r in rows));self.assertAlmostEqual(w['D20'].value,-51432.3);self.assertTrue(f['G18'].value.startswith('='))
   self.assertEqual(str(list(f.conditional_formatting)[0].sqref),'D18:D20 G18:G20')
 def test_profit_sheet_cache_positive(self):
  f=fixture();f['rows'][0]['Price']=3500;f['rows'][1]['Price']=6500
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'profit.xlsx';b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active
   self.assertGreater(w['G18'].value,0);self.assertGreater(w['G19'].value,0);self.assertGreater(w['G20'].value,0)
 def test_partial_quote_keeps_total_unknown(self):
  f=fixture();f['rows'][1]['Price']=None
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'unknown.xlsx';b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active
   self.assertEqual(w['G20'].value,'UNKNOWN');self.assertIsInstance(w['G18'].value,float)
 def test_missing_loan_sheet_gross_known_net_unknown(self):
  f=fixture();f['loan']=m.Num()
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'unknown.xlsx';b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active
   self.assertAlmostEqual(w['D18'].value,-51068.6);self.assertEqual(w['G18'].value,'UNKNOWN');self.assertEqual(w['G20'].value,'UNKNOWN')
 def test_days_held_plain_number_not_date_formula(self):
  # Apple Numbers turns date-date into a duration -> interest showed 0 (RB 9 Oct).
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'days.xlsx';b.write_report(fixture(),TODAY,p);f=load_workbook(p).active
   days=[f.cell(k,6).value for k in range(64,64+9)]
   self.assertEqual(days[0],(TODAY-dt.date(2025,11,24)).days);self.assertTrue(all(isinstance(x,int) for x in days))
   self.assertFalse(any('$B$4-' in str(c.value) for row in f.iter_rows() for c in row))
 def test_sold_now_net_is_price_minus_interest_minus_charges(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'net.xlsx';rows=b.write_report(fixture(),TODAY,p);w=load_workbook(p,data_only=True).active;f=load_workbook(p).active
   for k,r in zip((18,19),rows):
    self.assertGreater(w['E%d'%k].value,1000);self.assertAlmostEqual(w['F%d'%k].value,r['buy']+r['now_sale_fee']+r['now_tax'])
    self.assertAlmostEqual(w['G%d'%k].value,w['D%d'%k].value-w['E%d'%k].value-w['F%d'%k].value,places=6)
   self.assertEqual(f['G18'].value,'=IF(COUNT(D18:F18)=3,D18-E18-F18,"UNKNOWN")')
   self.assertAlmostEqual(w['H10'].value,rows[0]['full']-rows[0]['cmp']);self.assertIn('SELL ORDER PRICE',w['E9'].value)
 def test_formulas_recalculate_to_cached_values(self):
  try:from pycel import ExcelCompiler
  except ImportError:return
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'calc.xlsx';b.write_report(fixture(),TODAY,p);w=load_workbook(p,data_only=True).active;x=ExcelCompiler(filename=str(p))
   for ref in ('D10','E10','F10','H10','D11','E11','E18','F18','G18','E19','F19','G19','G20'):
    self.assertAlmostEqual(x.evaluate('Breakeven!'+ref),w[ref].value,places=4)
 def test_cash_in_hand_if_sold_today(self):
  f=fixture();f['unpaid']=m.Num(264.48,m.ESTIMATED,'x')
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'cash.xlsx';rows=b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active;fo=load_workbook(p).active
   want=sum(r['now_value']-r['now_sale_fee'] for r in rows)-261888.33-264.48
   self.assertAlmostEqual(w['D48'].value,want,places=6);self.assertEqual(fo['D48'].value,'=IF(COUNT(D44:D47)=4,D44-D45-D46-D47,"UNKNOWN")')
   # sell ONE stock: its own value - its fees - cost-ratio loan + unpaid share; parts add up to the total
   for j,r in enumerate(rows):
    c='BC'[j];share=r['cost']/414598.3
    self.assertEqual(w[c+'43'].value,r['symbol'])
    self.assertAlmostEqual(w[c+'48'].value,r['now_value']-r['now_sale_fee']-261888.33*share-264.48*share,places=4)
   self.assertAlmostEqual(w['B48'].value+w['C48'].value,w['D48'].value,places=4)
   f['unpaid']=m.Num();b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active
   self.assertEqual(w['D48'].value,'UNKNOWN');self.assertEqual(w['B48'].value,'UNKNOWN')
   f=fixture();f['unpaid']=m.Num(10);f['loan']=m.Num();b.write_report(f,TODAY,p);self.assertEqual(load_workbook(p,data_only=True).active['B48'].value,'UNKNOWN')
 def test_cash_formulas_recalculate(self):
  try:from pycel import ExcelCompiler
  except ImportError:return
  f=fixture();f['unpaid']=m.Num(264.48)
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'c.xlsx';b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active;x=ExcelCompiler(filename=str(p))
   for ref in ('B46','C46','B47','C47','B48','C48','D45','D48'):self.assertAlmostEqual(x.evaluate('Breakeven!'+ref),w[ref].value,places=4)
 def test_many_stocks_cash_per_stock_and_total(self):
  f=fixture();f['unpaid']=m.Num(300.0)
  lots=f['fifo']['lots'];lots['INFY']=[['2026-03-02',10,1500.0,0]]
  f['rows'].append({'Symbol':'INFY','Qty':10,'Open cost (Rs)':m.Num(15000.0,m.ESTIMATED),'Price':1400})
  f['open_cost_num']=m.Num(414598.3+15000);f['current_lots']=lots
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'many.xlsx';rows=b.write_report(f,TODAY,p);w=load_workbook(p,data_only=True).active
   self.assertEqual(len(rows),3);self.assertIn('HAATH MEIN',w['T9'].value)
   for k,r in zip(range(10,13),rows):
    self.assertEqual(w['A%d'%k].value,r['symbol']);self.assertAlmostEqual(w['T%d'%k].value,r['cash_now'],places=4)
    self.assertAlmostEqual(r['cash_now'],r['now_value']-r['now_sale_fee']-(261888.33+300)*r['cost']/429598.3,places=4)
   self.assertAlmostEqual(w['T13'].value,sum(r['now_value']-r['now_sale_fee'] for r in rows)-261888.33-300,places=4)
   try:from pycel import ExcelCompiler
   except ImportError:return
   x=ExcelCompiler(filename=str(p))
   for ref in ('R10','S11','T12','T13'):self.assertAlmostEqual(x.evaluate('Breakeven!'+ref),w[ref].value,places=4)
 def test_drive_copy_goes_to_account_folder(self):
  import os
  with tempfile.TemporaryDirectory() as d:
   drive=Path(d)/'MyDrive';drive.mkdir();src=Path(d)/'MTF_Check.xlsx';b.write_report(fixture(),TODAY,src)
   acc=types.SimpleNamespace(broker='DHAN',name='Tanya')
   with patch.dict(os.environ,{'RB_DRIVE_ROOT':str(drive)}),contextlib.redirect_stdout(io.StringIO()):dst=m.drive_mtf_copy(str(src),acc,any_path=True)
   self.assertIsNone(m.drive_mtf_copy(str(src),acc))  # temp/test file never copied
   self.assertEqual(Path(dst),drive/'RB_Reports'/'DHAN_Tanya'/'MTF_Check_DHAN_Tanya.xlsx');self.assertEqual(Path(dst).read_bytes(),src.read_bytes())
   with patch.dict(os.environ,{'RB_DRIVE_ROOT':str(Path(d)/'missing')}),contextlib.redirect_stdout(io.StringIO()):self.assertIsNone(m.drive_mtf_copy(str(src),acc,any_path=True))
if __name__=='__main__':unittest.main()
