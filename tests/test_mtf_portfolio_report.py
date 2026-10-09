"""Four-section MTF8 acceptance checks. Synthetic broker only; no API/order."""
import contextlib,copy,datetime as dt,io,os,sys,tempfile,types,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT),str(ROOT/'tests')]
import mtf_check as m,mtf_portfolio_report as p
TODAY=dt.date(2026,10,9)
def fixture():
 lots={'ALPHA':[['2026-09-01',6,1000,0],['2026-09-19',4,1000,0]],'BETA':[['2026-10-01',20,500,0]]}
 return dict(rows=[dict(Symbol=s,Qty=sum(z[1] for z in L),Price=1100 if s=='ALPHA' else 450,**{'Open cost (Rs)':m.Num(10000),'Price source':'synthetic exchange quote 09-Oct-2026 16:00 IST'}) for s,L in lots.items()],current_lots=lots,fifo={'lots':lots},loan=m.Num(15000),unpaid=m.Num(20),paid=m.Num(12345),open_cost_num=m.Num(20000),exit=m.Num(5000,m.ESTIMATED),report_account='SYNTHETIC DHAN A',src={'ledger':(m.VERIFIED,'synthetic ledger')})
def locate(s,label,stock='TOTAL'):
 row=next(c.row for c in s['A'] if c.value==label);head=next(c.row for c in s['A'] if c.value=='Metric');col=next(c.column for c in s[head] if c.value==stock);return s.cell(row,col)
class ReportTests(unittest.TestCase):
 def test_loan_and_own_split(self):
  rows=p.calculate(fixture(),TODAY);self.assertEqual(p.total(rows,'funded'),15000);self.assertEqual(p.total(rows,'own'),5000)
 def test_per_lot_ages_and_interest(self):
  rows=p.calculate(fixture(),TODAY);a=rows[0];self.assertEqual([z['days'] for z in a['lots']],[38,20]);self.assertAlmostEqual(a['past'],(6000*.75*38+4000*.75*20)*.1249/365)
 def test_daily_burn(self):self.assertAlmostEqual(p.total(p.calculate(fixture(),TODAY),'daily'),15000*.1249/365)
 def test_brokerage_cap_and_rounding(self):
  self.assertEqual(p.fees(1000)['brokerage'],.30);self.assertEqual(p.fees(100000)['brokerage'],20);self.assertEqual(p.fees(100500)['stt'],101);self.assertEqual(p.fees(10000,True)['stamp'],2)
 def test_dp_pledge_gst_not_flat_before_tax(self):
  f=p.fees(1000);self.assertEqual(f['dp']+f['pledge'],27.5);self.assertAlmostEqual(p.rounded(f['dp']*.18)+p.rounded(f['pledge']*.18),4.95)
 def test_no_sell_stamp(self):self.assertEqual(p.fees(10000)['stamp'],0)
 def test_buy_brokerage_min_not_fixed_twenty(self):self.assertEqual(p.fees(4000,True)['brokerage'],1.2)
 def test_paid_ledger_history_not_deducted_on_exit(self):
  f=fixture();old=p.calculate(f,TODAY);f['paid']=m.Num(999999);self.assertEqual(p.calculate(f,TODAY),old)
 def test_asof_unpaid_not_counted_twice(self):
  for r in p.calculate(fixture(),TODAY):self.assertAlmostEqual(r['cash']-r['pocket'],r['net_exit'])
 def test_full_sale_value_exit_reserve(self):
  for r in p.calculate(fixture(),TODAY):self.assertAlmostEqual(r['extra'],r['value']*.1249*3/365)
 def test_minimum_tick_targets(self):
  for r in p.calculate(fixture(),TODAY):
   for key,h,s,t,b in [('pure',0,0,0,0),('today_be',0,3,0,0),('full',30,3,.208,2)]:
    x=r[key];self.assertGreaterEqual(p.target_net(x,r['qty'],r['cost'],r['past'],r['buy'],r['funded'],.1249,h,s,t),b-1e-8);self.assertLess(p.target_net(x-.05,r['qty'],r['cost'],r['past'],r['buy'],r['funded'],.1249,h,s,t),b)
 def test_minimum_target_across_stt_rounding_jump(self):
  # Find the global first passing tick by a broad independent exhaustive scan.
  for cost in [1499,1499.45,1499.95,2498.7,2499.7,9998.7]:
   for tax in [0,.208,.7]:
    x=p.target(1,cost,10,18,750,.1249,30,3,tax,2)
    first=next(round(k*.05,8) for k in range(int(cost/.05),int(x/.05)+3) if p.target_net(round(k*.05,8),1,cost,10,18,750,.1249,30,3,tax)>=2)
    self.assertEqual(x,first)
 def test_small_qty_uncapped_target(self):
  x=p.target(1,1000,10,20,750,.1249,30,3,.208,2);self.assertIsNotNone(x);self.assertGreaterEqual(p.target_net(x,1,1000,10,20,750,.1249,30,3,.208),2)
 def test_profit_and_loss_both_work(self):
  r=p.calculate(fixture(),TODAY);self.assertGreater(r[0]['net_exit'],0);self.assertLess(r[1]['net_exit'],0)
 def test_tax_reserve_separate_from_cash(self):
  a=p.calculate(fixture(),TODAY,tax=0);b=p.calculate(fixture(),TODAY,tax=.208);self.assertEqual([r['cash'] for r in a],[r['cash'] for r in b]);self.assertGreater(b[0]['now_tax'],0);self.assertEqual(b[1]['now_tax'],0)
 def test_missing_quote_target_survives(self):
  f=fixture();f['rows'][0]['Price']=None;a=p.calculate(f,TODAY)[0];self.assertIsNone(a['cash']);self.assertIsNone(a['net_exit']);self.assertIsNotNone(a['full'])
 def test_missing_unpaid_only_cash_unknown(self):
  f=fixture();f['unpaid']=m.Num();a=p.calculate(f,TODAY)[0];self.assertIsNone(a['cash']);self.assertIsNotNone(a['net_exit']);self.assertIsNotNone(a['full'])
 def test_missing_loan(self):
  f=fixture();f['loan']=m.Num();a=p.calculate(f,TODAY)[0];self.assertIsNone(a['funded']);self.assertIsNone(a['past']);self.assertIsNone(a['full']);self.assertEqual(a['price_pnl'],1000)
 def test_bad_total_basis(self):
  f=fixture();f['open_cost_num']=m.Num(100);self.assertTrue(all(r['full'] is None for r in p.calculate(f,TODAY)))
 def test_excess_loan(self):
  f=fixture();f['loan']=m.Num(20001);self.assertTrue(all(r['cash'] is None for r in p.calculate(f,TODAY)))
 def test_missing_dates_not_invented(self):
  f=fixture();f['current_lots']['ALPHA']=[];a=p.calculate(f,TODAY)[0];self.assertIsNone(a['past']);self.assertIsNone(a['full']);self.assertEqual(a['value'],11000);self.assertIsNotNone(a['funded'])
 def test_fractional_equity_qty_unknown(self):
  f=fixture();f['rows'][0]['Qty']=.1;r=p.calculate(f,TODAY)[0];self.assertIsNone(r['qty']);self.assertIsNone(r['value']);self.assertIsNone(r['full'])
 def test_future_lot_blocks_interest(self):
  f=fixture();f['current_lots']['ALPHA'][0][0]='2027-01-01';self.assertIsNone(p.calculate(f,TODAY)[0]['past'])
 def test_failed_ledger_blocks_payout(self):
  f=fixture();f['src']['ledger']=(m.UNKNOWN,'synthetic ledger unavailable');self.assertIsNone(p.calculate(f,TODAY)[0]['cash'])
 def test_unpaid_over_model_does_not_fake_paid(self):
  f=fixture();f['unpaid']=m.Num(99999);self.assertTrue(all(r['paid_proxy'] is None for r in p.calculate(f,TODAY)))
 def test_zero_rate_and_zero_loan(self):
  f=fixture();f['loan']=m.Num(0);a=p.calculate(f,TODAY,rate=0)[0];self.assertEqual(a['past']+a['daily']+a['extra'],0);self.assertIsNotNone(a['full'])
 def test_any_current_stock_count_one_sheet(self):
  f=fixture();f['rows'].append(dict(Symbol='GAMMA',Qty=5,Price=210,**{'Open cost (Rs)':m.Num(1000)}));f['current_lots']['GAMMA']=[['2026-09-30',5,200,0]];f['open_cost_num']=m.Num(21000)
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'x.xlsx';r=p.write_report(f,TODAY,out);w=load_workbook(out,data_only=True);self.assertEqual(w.sheetnames,['Breakeven']);self.assertEqual(len(r),3);self.assertTrue(all(x['full'] for x in r))
 def test_all_four_sections_and_cached_numbers(self):
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'x.xlsx';r=p.write_report(fixture(),TODAY,out);w=load_workbook(out,data_only=True);s=w.active;self.assertTrue(all(any('SECTION '+str(n) in str(c.value) for c in s['A']) for n in [1,2,3,4]));self.assertAlmostEqual(locate(s,'NET CASH CREDIT SCENARIO').value,p.total(r,'cash'));self.assertIsNone(s['B17'].value if s['B17'].value!=m.UNKNOWN else None);self.assertFalse(w._external_links)
 def test_no_false_paid_interest_claim(self):
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'x.xlsx';p.write_report(fixture(),TODAY,out);w=load_workbook(out);txt=str([[c.value for row in s for c in row] for s in w]);self.assertNotIn('pehle hi ledger se kaat chuka',txt);self.assertEqual(w.active['B17'].value,m.UNKNOWN)
 def test_formula_injection_stock_source_literal(self):
  f=fixture();f['rows'][0]['Price source']='=HYPERLINK("https://invalid")';f['report_account']='=2+2'
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'x.xlsx';p.write_report(f,TODAY,out);w=load_workbook(out);self.assertFalse(w._external_links);self.assertNotEqual(w.active['K30'].data_type,'f')
 def test_atomic_failure_and_symlink(self):
  with tempfile.TemporaryDirectory() as d:
   out=Path(d)/'x.xlsx';out.write_bytes(b'KEEP')
   with patch.object(p,'cache_formula_values',side_effect=OSError('synthetic failed cache')):
    with self.assertRaises(OSError):p.write_report(fixture(),TODAY,out)
   self.assertEqual(out.read_bytes(),b'KEEP');link=Path(d)/'link.xlsx';link.symlink_to(out)
   with self.assertRaises(ValueError):p.write_report(fixture(),TODAY,link)
   self.assertEqual(out.read_bytes(),b'KEEP');self.assertFalse(list(Path(d).glob('.mtf-report-*')))
 def test_current_unpaid_estimator_has_no_forward_day(self):
  from test_mtf_current_inventory import args,tr
  ts=[tr(symbol='ALPHA',qty=10,price=1000,product='MTF')];ba=types.SimpleNamespace(holdings=lambda *x:[dict(symbol='ALPHA',qty=10,avg_price=1000)])
  led=m.ledger_rows([dict(narration='MTF Interest for Period 01/10/2026 To 07/10/2026',debit=70,credit=0,voucherdate='Oct 08, 2026')])
  with patch.dict(sys.modules,{'broker_api':ba}),patch.object(m,'dhan_trades',return_value=(ts,dict(complete=True,notes=[]))),patch.object(m,'mtf_positions',return_value={'ALPHA':10}),patch.object(m,'dhan_ledger',return_value=led),patch.object(m,'quotes',return_value={'ALPHA':(1000,m.VERIFIED,'dummy quote')}):
   r=m.build(args(loan=7500,unpaid_interest=None),object(),TODAY,dt.date(2024,1,1));self.assertEqual(r['unpaid'].value,20);self.assertIn('excludes future',r['unpaid'].note)
if __name__=='__main__':unittest.main()
