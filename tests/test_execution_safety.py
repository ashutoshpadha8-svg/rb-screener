"""Safety boundaries, actual broker fixtures and fresh per-stock MTF; no network."""
import datetime as dt
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace as NS
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import broker_api as ba
import broker_mtf as mtf
import buy_planner as bp
import execution_safety as safe
import journal
import momentum_screener as ms
import pandas as pd
from openpyxl import Workbook,load_workbook

class SafetyTests(unittest.TestCase):
    def test_mtf_45_actual_cash_and_qty(self):
        out,total=bp.distribute([dict(symbol='ABC',cls='M',price=900,pick='MTF',mtf_lev=4.5)],1000)
        self.assertEqual(out[0]['final'],5);self.assertEqual(total['total'],1000)
        self.assertEqual(out[0]['position'],4500)

    def test_zero_unknown_and_bad_leverage_block(self):
        for lev in [0,None,1,-1,float('inf'),float('nan'),11]:
            out,total=bp.distribute([dict(symbol='ABC',cls='M',price=900,pick='MTF',mtf_lev=lev,qty_you=5)],1000,lev=4)
            self.assertEqual(out[0]['final'],0);self.assertEqual(total['total'],0)

    def test_delivery_ignores_mtf_rate(self):
        out,_=bp.distribute([dict(symbol='ABC',cls='M',price=100,pick='BUY',mtf_lev=4.5)],1000)
        self.assertEqual(out[0]['final'],10)

    def test_refresh_runs_again_not_old_snapshot(self):
        rates=iter([4.5,0,None]);calls=[]
        adapter=NS(mtf_leverage=lambda s,sym,p:(calls.append(sym) or next(rates),'fixture'))
        snapshots=[mtf.refresh(NS(),{'ABC':100},adapter,'time'+str(n)) for n in range(3)]
        self.assertEqual([s['ABC']['leverage'] for s in snapshots],[4.5,0,None])
        self.assertEqual(calls,['ABC']*3);self.assertIn('UNKNOWN',snapshots[-1]['ABC']['note'])

    def test_refresh_exception_and_nan_are_unknown(self):
        def fail(*a):raise ba.BrokerError('offline')
        self.assertIsNone(mtf.refresh(NS(),{'ABC':100},NS(mtf_leverage=fail),'now')['ABC']['leverage'])
        self.assertIsNone(mtf.refresh(NS(),{'ABC':100},NS(mtf_leverage=lambda *a:(float('nan'),'bad')),'now')['ABC']['leverage'])

    def test_sheet_status_columns_no_input_shift(self):
        w=Workbook();s=w.active;s.title='Holdings_Table';s.append(['Symbol','Qty']);s.append(['ABC',3])
        snaps={'ABC':dict(leverage=4.5,asof='now',note='broker fixture')}
        mtf.add_columns(w,snaps);mtf.add_columns(w,snaps)
        self.assertEqual(s['B2'].value,3);self.assertEqual(s['C2'].value,4.5);self.assertEqual(s.max_column,5)
        self.assertEqual(w['MTF_Rates']['B2'].value,4.5)

    def test_planner_new_columns_literal_rate(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'test.xlsx';w=Workbook()
            bp.write_sheet(w,{},[dict(symbol='ABC',cls='M',price=900,why='test',mom_rank=1,mtf_lev=4.5,mtf_asof='now',mtf_note='fixture')],{'budget':1000},{'ABC':('MTF',None)},'fixture')
            w.save(p);_,rows=bp.read_sheet(p)
            self.assertEqual(rows[0]['mtf_lev'],4.5);self.assertEqual(w[bp.SHEET]['S29'].value,'now')
            self.assertIn('R29',w[bp.SHEET]['Q29'].value)
            self.assertNotIn('$E$19',w[bp.SHEET]['Q29'].value)

    def margin(self,broker,response):
        sess=NS(broker=broker,label=broker,client_id='dummy')
        with patch.object(ba,'symbol_map',return_value={'ABC':dict(id='1',tsym='ABC-EQ')}),patch.object(ba,'_call',return_value=response):
            return ba.mtf_leverage(sess,'ABC',100)

    def test_dhan_actual_and_explicit_zero(self):
        self.assertEqual(self.margin('DHAN',{'leverage':'1:4.5','totalMargin':22200})[0],4.5)
        self.assertEqual(self.margin('DHAN',{'leverage':0,'totalMargin':0})[0],0)

    def test_angel_ratio_and_kite_actual(self):
        self.assertEqual(self.margin('ANGEL',{'totalMarginRequired':100000/4.5})[0],4.5)
        self.assertEqual(self.margin('ZERODHA',[{'leverage':4.5,'total':22200}])[0],4.5)

    def test_invalid_margin_responses_unknown(self):
        for broker,response in [('DHAN',[]),('DHAN',{}),('ANGEL',None),('ZERODHA',[]),('ANGEL',{'totalMarginRequired':'bad'})]:
            self.assertIsNone(self.margin(broker,response)[0])

    def test_batch_duplicates_overlap_and_unverified_mtf(self):
        row=dict(symbol='ABC',shares=2,entry_price=100,lev=1,product='CNC')
        for buys,sells in [([row,row],[]),([row],[dict(symbol='ABC')]),([dict(row,product='MTF',lev=1)],[])]:
            with self.assertRaises(ValueError):safe.validate_batch(buys,sells)

    def test_mcap_floor_date_and_duplicate(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'caps.csv'
            for cap,date in [(9999,'2026-10-05'),(20000,'2026-09-25'),(20000,'2026-10-07')]:
                pd.DataFrame([dict(symbol='ABC',mcap_cr=cap,asof=date)]).to_csv(p,index=False)
                with self.assertRaises(ValueError):safe.validate_caps({'ABC'},p,dt.date(2026,10,6))
            pd.DataFrame([dict(symbol='ABC',mcap_cr=10000,asof='2026-10-05')]).to_csv(p,index=False)
            safe.validate_caps({'ABC'},p,dt.date(2026,10,6))
            d=pd.read_csv(p);pd.concat([d,d]).to_csv(p,index=False)
            with self.assertRaises(ValueError):safe.validate_caps({'ABC'},p,dt.date(2026,10,6))

    def test_rank_integrity_freshness_and_missing_manifest(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'rank.csv';d=pd.DataFrame([dict(symbol='ABC',rank=1,date='2026-10-05',regime_red=False)])
            safe.publish_ranks(p,d);self.assertEqual(len(safe.read_ranks(p,'2026-10-05')),1)
            with self.assertRaises(ValueError):safe.read_ranks(p,'2026-10-06')
            p.write_text(p.read_text()+'corrupt')
            with self.assertRaises(ValueError):safe.read_ranks(p,'2026-10-05')
            safe.publish_ranks(p,d);Path(str(p)+'.meta.json').unlink()
            with self.assertRaises(ValueError):safe.read_ranks(p,'2026-10-05')

    def test_intraday_ranking_day_accepted(self):
        # RB runs rb in market hours: the ranking is dated today (live bar)
        import daily_screener as ds
        IST=dt.timezone(dt.timedelta(hours=5,minutes=30))
        mon=dt.datetime(2026,10,5,11,0,tzinfo=IST); sun=dt.datetime(2026,10,4,11,0,tzinfo=IST)
        early=dt.datetime(2026,10,5,8,30,tzinfo=IST)
        old=ds.now_ist
        try:
            ds.now_ist=lambda:mon
            self.assertIn(dt.date(2026,10,5),safe.ranking_days(mon))
            ds.now_ist=lambda:early
            self.assertNotIn(dt.date(2026,10,5),safe.ranking_days(early))
            ds.now_ist=lambda:sun
            self.assertNotIn(dt.date(2026,10,4),safe.ranking_days(sun))
        finally:
            ds.now_ist=old
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'rank.csv'
            safe.publish_ranks(p,pd.DataFrame([dict(symbol='ABC',rank=1,date='2026-10-05',regime_red=False)]))
            self.assertEqual(len(safe.read_ranks(p,{dt.date(2026,10,2),dt.date(2026,10,5)})),1)
            with self.assertRaises(ValueError):safe.read_ranks(p,{dt.date(2026,10,1),dt.date(2026,10,2)})

    def test_lock_excludes_another_process(self):
        with tempfile.TemporaryDirectory() as t:
            self.assertTrue(safe.workflow_lock(t));self.assertTrue(safe.workflow_lock(t))
            code='import execution_safety,sys;sys.exit(0 if not execution_safety.workflow_lock(sys.argv[1]) else 1)'
            proc=subprocess.run([sys.executable,'-c',code,t],cwd=ROOT,capture_output=True)
            self.assertEqual(proc.returncode,0,proc.stderr)

    def hold(self,broker,rows,positions=()):
        with patch.object(ba,'_call',side_effect=[rows,list(positions)] if broker!='ZERODHA' else [rows]):
            return ba.holdings(NS(broker=broker))

    def test_dhan_available_zero_not_total(self):
        h=self.hold('DHAN',[dict(tradingSymbol='ABC',totalQty=10,availableQty=0)])[0]
        self.assertEqual(h['qty'],10);self.assertEqual(h['sellable_by_product']['CNC'],0)

    def test_angel_delivery_and_mtf_separate(self):
        h=self.hold('ANGEL',[dict(tradingsymbol='ABC-EQ',quantity=5,product='DELIVERY')],[dict(tradingsymbol='ABC-EQ',netqty=3,producttype='MARGIN',exchange='NSE',buyavgprice=100)])[0]
        self.assertEqual(h['qty'],8);self.assertEqual(h['sellable_by_product'],{'CNC':5,'MTF':3})

    def test_angel_mtf_seen_twice_not_double_counted(self):
        h=self.hold('ANGEL',[dict(tradingsymbol='ABC-EQ',quantity=3,product='MARGIN')],[dict(tradingsymbol='ABC-EQ',netqty=3,producttype='MARGIN',exchange='NSE')])[0]
        self.assertEqual(h['qty'],3);self.assertEqual(h['sellable_by_product']['MTF'],3)

    def test_angel_no_open_positions_null_ok(self):
        # RB 6 Oct: Angel getPosition = status true + data null when nothing
        # is open -> holdings still come from getHolding (rbport stopped)
        with patch.object(ba,'_call',side_effect=[[dict(tradingsymbol='ABC-EQ',quantity=5,product='DELIVERY')],None]):
            h=ba.holdings(NS(broker='ANGEL'))
        self.assertEqual([(x['symbol'],x['qty']) for x in h],[('ABC',5)])
        with patch.object(ba,'_call',side_effect=[[dict(tradingsymbol='ABC-EQ',quantity=5)],{}]):
            with self.assertRaises(ba.BrokerError):ba.holdings(NS(broker='ANGEL'))

    def test_angel_empty_order_book_null_ok(self):
        with patch.object(ba,'_call',return_value=None),patch.object(ba,'symbol_map',return_value={}):
            self.assertEqual(ba.pending_orders(NS(broker='ANGEL')),set())
        with patch.object(ba,'_call',return_value={}),patch.object(ba,'symbol_map',return_value={}):
            with self.assertRaises(ba.BrokerError):ba.pending_orders(NS(broker='ANGEL'))

    def test_kite_nested_mtf_and_used_qty(self):
        h=self.hold('ZERODHA',[dict(tradingsymbol='ABC',quantity=5,used_quantity=2,mtf=dict(quantity=4,used_quantity=1))])[0]
        self.assertEqual(h['qty'],9);self.assertEqual(h['sellable_by_product'],{'CNC':3,'MTF':3})

    def test_invalid_holdings_fail_closed(self):
        for rows in [None,{},[None],[dict(quantity='nan',tradingsymbol='ABC')]]:
            with self.assertRaises(ba.BrokerError):self.hold('ANGEL',rows)

    def test_journal_pending_buy_does_not_count_unfilled_qty(self):
        self.assertEqual(journal._qty(dict(momentum_qty=10,order_id='1',filled_qty_confirmed=0)),0)
        self.assertEqual(journal._qty(dict(momentum_qty=10,order_id='1',filled_qty_confirmed=3)),3)

    def test_journal_partial_sells_do_not_reuse_old_price(self):
        with tempfile.TemporaryDirectory() as t,patch.object(ms,'SPLIT_FILE',str(Path(t)/'split.csv')),patch.object(journal,'dividends_of',return_value=[]),patch.object(journal,'_split_after_close'):
            sp=pd.DataFrame([dict(symbol='ABC',mode='LIVE',entry_date='2026-10-01',entry_price=100,momentum_qty=10,swing_qty=0,investing_qty=0,product='CNC')])
            journal.sync(sp,[dict(symbol='ABC',qty=10)],{}, {},today='2026-10-01')
            trades=[dict(symbol='ABC',side='SELL',qty=3,price=110,date='2026-10-02')]
            with patch.object(journal,'ba_trades',return_value=trades):
                journal.sync(sp.assign(momentum_qty=7),[dict(symbol='ABC',qty=7)],{}, {},today='2026-10-02')
            trades.append(dict(symbol='ABC',side='SELL',qty=2,price=120,date='2026-10-03'))
            with patch.object(journal,'ba_trades',return_value=trades):
                j,_=journal.sync(sp.assign(momentum_qty=5),[dict(symbol='ABC',qty=5)],{}, {},today='2026-10-03')
            closed=j[j.status=='CLOSED']
            self.assertEqual(list(closed.sell_price),['110.00','120.00'])
            self.assertEqual(sum(float(q) for q in j[j.status=='OPEN'].qty),5)

    def test_unshortlisted_stocks_not_populated(self):
        w=Workbook();s=w.active;s.title='Holdings_Table';s.append(['Symbol','Qty']);s.append(['ABC',3]);s.append(['UNRELATED',4])
        mtf.add_columns(w,{'ABC':dict(leverage=4.5,asof='now',note='fixture')})
        self.assertIsNone(s['C3'].value);self.assertEqual(w['MTF_Rates'].max_row,2)

    def test_limit_uses_actual_tick_not_global_guess(self):
        self.assertEqual(safe.limit_price(101.02,.1),101.1)
        self.assertEqual(safe.limit_price(101.02,.01),101.02)
        with self.assertRaises(ValueError):safe.limit_price(100,None)

    def test_invalid_budget_and_percentages_fail(self):
        for budget in ['bad',-1,float('inf')]:
            with self.assertRaises(ValueError):bp.distribute([],budget)
        with self.assertRaises(ValueError):bp.distribute([],100,shares={'M':.8,'L':.25,'S':.15})

    def test_mtf_confirmation_changed_or_unknown_stops_order(self):
        import test_manual_order_inputs as manual
        posts=[]
        row=dict(symbol='ABC',shares=3,entry_price=100,product='MTF',lev=4.5,note='fixture')
        for new_rate in [4.0,0,None]:
            env=dict(ds=NS(market_open=lambda:False,MCAP_CACHE='dummy',now_ist=lambda:dt.datetime(2026,10,6)),LIMIT_BUFFER=.02,
                     safety=NS(validate_batch=safe.validate_batch,validate_caps=lambda *a:None,finite_positive=safe.finite_positive),pd=pd,
                     ba=NS(symbol_map=lambda sess:{'ABC':{}},verify_identity=lambda sess:(True,'dummy'),
                           available_funds=lambda sess:(1000,''),pending_orders=lambda sess:set(),holdings=lambda sess:[],
                           mtf_leverage=lambda *a:(new_rate,'fixture'),BrokerError=RuntimeError),
                     input=lambda prompt:'YES MTF',send_one=lambda *a,**k:(posts.append(a) or True,'dummy','PENDING'))
            import contextlib,io
            with contextlib.redirect_stdout(io.StringIO()):
                accepted=manual.function('place_orders',env)(NS(label='Dummy',client_id='TEST',broker='DUMMY'),[row],False,False)
            self.assertEqual(accepted,[])
        self.assertEqual(posts,[])

    def test_pending_or_existing_holding_prevents_buy(self):
        import test_manual_order_inputs as manual
        import contextlib,io
        row=dict(symbol='ABC',shares=3,entry_price=100,product='CNC',lev=1,note='fixture')
        for pending,hold in [({'ABC'},[]),(set(),[dict(symbol='ABC',qty=3)])]:
            env=dict(ds=NS(MCAP_CACHE='dummy',now_ist=lambda:dt.datetime(2026,10,6)),
                     safety=NS(validate_batch=safe.validate_batch,validate_caps=lambda *a:None),pd=pd,
                     ba=NS(pending_orders=lambda sess:pending,holdings=lambda sess:hold,BrokerError=RuntimeError))
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(manual.function('place_orders',env)(NS(),[row],False,False),[])

if __name__=='__main__':unittest.main()
