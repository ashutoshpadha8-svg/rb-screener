#!/usr/bin/env python3
"""
Per-account switches, set from the Dashboard of the Portfolio file.

  TRADING (buy + sell):  ON / OFF   (default OFF)
     OFF -> rbtrack sends NO order at all: no BUY, no BUY MTF, no SIP, no SELL
            (WATCH still works -- it is no order)
     ON  -> BUY / SIP orders as picked, and SELL for the Sell sheet rows
            marked YES (you still type YES / YES SELL before anything goes)

File: accounts/<BROKER>_<ID>/data/settings.json
"""

import os
import json

import momentum_screener as ms

LABEL = "TRADING (buy + sell)"
DEFAULT = {"trading": "OFF"}


def _path():
    return os.path.join(os.path.dirname(ms.SPLIT_FILE), "settings.json")


def load():
    try:
        with open(_path()) as f:
            d = json.load(f)
        return dict(DEFAULT, **d)
    except (IOError, OSError, ValueError):
        return dict(DEFAULT)


def save(d):
    with open(_path(), "w") as f:
        json.dump(d, f)


def trading_on():
    return load().get("trading") == "ON"


def read_dashboard(xlsx):
    """Take the switch the user set on the Dashboard (if the file has one).
    Returns the value now in force ('ON' / 'OFF')."""
    d = load()
    try:
        from openpyxl import load_workbook
        ws = load_workbook(xlsx, read_only=True)["Dashboard"]
        for row in ws.iter_rows(min_row=1, max_row=12, max_col=2,
                                values_only=True):
            if row and str(row[0] or "").strip() == LABEL:
                v = str(row[1] or "").upper().strip()
                if v in ("ON", "OFF"):
                    d["trading"] = v
                    save(d)
                break
    except Exception:
        pass
    return d["trading"]
