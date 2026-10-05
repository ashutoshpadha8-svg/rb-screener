"""
cap_class.py -- market cap + Large / Mid / Small on every stock sheet
(5 Oct 2026, RB: "har sheet per stocks ki marketcap chahiye ki midcap hai
small ya large"). DISPLAY ONLY.

Class like AMFI (SEBI): rank by full market cap among ALL NSE companies in
NSE's latest market-cap file (data/_nse_mcap_latest.csv, ~2,600 names):
1-100 Large, 101-250 Mid, 251+ Small. AMFI's own list is published twice a
year, so a stock near rank 100 / 250 can differ from AMFI's label.

add_columns(wb) appends 'Mcap (Rs Cr)' + 'Cap Class' at the END of every
table sheet with a Symbol / Ticker column (no existing column moves, so
rbtrack and formulas are unaffected); running it again only refreshes them.
"""

import os

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
MCAP = os.path.join(HERE, "data", "_nse_mcap_latest.csv")
LARGE, MID = 100, 250
HEAD_M, HEAD_C = "Mcap (Rs Cr)", "Cap Class"
COLOR = {"Large": "1F4E78", "Mid": "1E7B34", "Small": "C55A11"}
SKIP = ("Dashboard", "Holdings", "Journal", "SIP", "Symbols",
        "Signal_Summary", "Buy_Planner")
_cache = {}


def label(rank):
    return "Large" if rank <= LARGE else "Mid" if rank <= MID else "Small"


def table(path=MCAP):
    """{SYMBOL: (mcap_cr, rank, 'Large'|'Mid'|'Small')}, as-of date."""
    if path in _cache:
        return _cache[path]
    try:
        d = pd.read_csv(path)
        d["symbol"] = d["symbol"].astype(str).str.upper().str.strip()
        d["mcap_cr"] = pd.to_numeric(d["mcap_cr"], errors="coerce")
        d = d.dropna(subset=["mcap_cr"]).sort_values("mcap_cr",
                                                     ascending=False)
        d = d.drop_duplicates("symbol").reset_index(drop=True)
        out = {s: (float(m), i + 1, label(i + 1))
               for i, (s, m) in enumerate(zip(d["symbol"], d["mcap_cr"]))}
        asof = str(d["asof"].iloc[0]) if "asof" in d and len(d) else None
    except Exception:
        out, asof = {}, None
    _cache[path] = (out, asof)
    return out, asof


def of(symbol, tab=None):
    tab = tab if tab is not None else table()[0]
    s = str(symbol or "").split(" | ")[0].upper().strip()
    return tab.get(s)


def _header_row(ws):
    for r in range(1, 9):
        vals = [ws.cell(row=r, column=c).value
                for c in range(1, min(ws.max_column, 60) + 1)]
        for k in ("Symbol", "Ticker"):
            if k in vals:
                return r, vals.index(k) + 1
    return None, None


def add_sheet(ws, tab):
    from copy import copy
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as L
    h, scol = _header_row(ws)
    if not h:
        return False
    heads = [ws.cell(row=h, column=c).value
             for c in range(1, ws.max_column + 1)]
    own_m = HEAD_M not in heads        # Swing / Investing already have it
    if HEAD_C in heads:
        cc = heads.index(HEAD_C) + 1
        cm = heads.index(HEAD_M) + 1 if HEAD_M in heads else None
        own_m = cm == cc - 1           # our pair from an earlier run
    else:
        last = max(c for c in range(1, ws.max_column + 1)
                   if ws.cell(row=h, column=c).value not in (None, ""))
        cm = last + 1 if own_m else heads.index(HEAD_M) + 1
        cc = last + 2 if own_m else last + 1
        src = ws.cell(row=h, column=scol)
        new = ((cm, HEAD_M), (cc, HEAD_C)) if own_m else ((cc, HEAD_C),)
        for c, t in new:
            cell = ws.cell(row=h, column=c, value=t)
            cell.font, cell.fill = copy(src.font), copy(src.fill)
            cell.alignment, cell.border = copy(src.alignment), \
                copy(src.border)
    r = h + 1
    while ws.cell(row=r, column=scol).value not in (None, ""):
        x = of(ws.cell(row=r, column=scol).value, tab)
        if cm and (own_m or ws.cell(row=r, column=cm).value is None):
            a = ws.cell(row=r, column=cm, value=round(x[0]) if x else None)
            a.number_format = "#,##0"
        b = ws.cell(row=r, column=cc, value=x[2] if x else "?")
        b.font = Font(bold=True, color=COLOR.get(x[2]) if x else "7F7F7F")
        r += 1
    if ws.auto_filter.ref and r > h + 1:    # filter covers the new columns
        ws.auto_filter.ref = "%s%d:%s%d" % (L(1), h, L(cc), r - 1)
    return True


def add_columns(wb, skip=SKIP):
    tab, _ = table()
    if not tab:
        return 0
    n = 0
    for ws in wb.worksheets:
        if ws.title in skip or ws.sheet_state != "visible":
            continue
        try:
            n += bool(add_sheet(ws, tab))
        except Exception:
            pass                       # display only -- never stop a report
    return n
