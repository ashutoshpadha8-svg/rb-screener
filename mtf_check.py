#!/usr/bin/env python3
"""
mtf_check.py -- MTF hisaab (v30-Codex-MTF8, 9 Oct 2026).
READ-ONLY: no order, nothing changed in split.csv / journal.

v29 = rewrite after Codex's v27 review (8 findings, 17 synthetic cases):
every number carries a STATUS:
  VERIFIED   from the broker (trade history, ledger) or typed by you
  ESTIMATED  a model (last close instead of live price, interest by lot
             days, sell fees from the rate card, product guessed ...)
  UNKNOWN    a required input is missing -> the number is NOT shown as a
             value and every total that needs it is UNKNOWN too
             (a 'partial' subtotal of the known parts is shown separately)
Nothing missing is ever counted as 0.

Scopes are kept apart:
  MTF_Status         current vs historical result and source status
  MTF_Open           open MTF lots: cost, value, gross price P&L, interest
                     estimate by LOT (each lot its own days)
  MTF_Exit_Estimate  if sold now: gross - current loan - unpaid interest -
                     sell fees  (needs --loan; a hypothetical sale, never
                     'actual')
  MTF_Period         since --from: matched realised P&L (UNKNOWN when a
                     sell has no buy in the history) + open P&L - ALL MTF
                     interest debited - MTF trade charges
  Reconciliation     FIFO MTF qty vs demat qty, uncovered sells, demat
                     stocks with no MTF trade, unallocated account items
                     (DP / pledge / dividends / money in-out: account-wide,
                     NOT put into the MTF result)

Current loan: read explicit TODAY-dated Net MTF Funding snapshot from a
downloaded ledger, or type --loan. Never infer it from old interest.
Type it from the Dhan app (MTF / pledge section, 'funded amount'):
      python3 mtf_check.py --loan 272000
Without it loan, own money, exit cash and open-position interest are
UNKNOWN. 'Last interest debit / days x 365 / rate' is shown only as the
AVERAGE interest-bearing balance of that past period (a diagnostic).

Run:  python3 mtf_check.py                 (history from 1 Apr 2024)
      python3 mtf_check.py --loan 272000 --own-cash 140000
      python3 mtf_check.py --ledger ~/Downloads/ledger.xlsx   (API fails)
Output: terminal + accounts/<..>/reports/MTF_Check.xlsx (same file refreshed)
Exit code 2 = an essential source failed (report marked INCOMPLETE).
"""

import argparse
import datetime as dt
import math
import json
import tempfile
import os
import re
import sys

import pandas as pd
from mtf_prices import safe_error,redact_text

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

MTF_RATE = 0.1249          # Dhan's published base slab; ESTIMATE only
VERIFIED, ESTIMATED, UNKNOWN = "VERIFIED", "ESTIMATED", "UNKNOWN"
PAGE_LIMIT = 50


# ================================================================== values
class Num(object):
    """A money value with a status. value None <=> UNKNOWN."""

    def __init__(self, value=None, status=None, note="", partial=None):
        st = status or VERIFIED
        valid = value is not None and _finite(value)
        ok = valid and st in (VERIFIED, ESTIMATED)
        self.value = float(value) if ok else None
        self.status = st if ok else UNKNOWN
        self.note = note
        if not ok and partial is None and valid:
            partial = float(value)
        self.partial = float(partial) if partial is not None and _finite(partial) else None

    @property
    def known(self):
        return self.value is not None and self.status in (VERIFIED, ESTIMATED)

    def __neg__(self):
        return Num(-self.value if self.known else None, self.status,
                   self.note, -self.partial if self.partial is not None
                   else None)

    def __repr__(self):
        return "Num(%r, %s)" % (self.value, self.status)


def _finite(v):
    try:
        return math.isfinite(float(v))
    except (TypeError, ValueError):
        return False


def total(parts, note=""):
    """Sum of Num parts. Any UNKNOWN -> UNKNOWN (partial = known sum);
    any ESTIMATED -> ESTIMATED."""
    parts = list(parts)
    known = [p for p in parts if p.known]
    s = sum(p.value for p in known)
    if len(known) < len(parts):
        partial = s + sum(p.partial for p in parts if not p.known and p.partial is not None)
        return Num(None, UNKNOWN, note, partial=partial)
    st = ESTIMATED if any(p.status == ESTIMATED for p in parts) else VERIFIED
    return Num(s, st, note)


def rs(x):
    """'Rs 1,234' / '-Rs 1,234' / 'UNKNOWN'. Takes a number or a Num."""
    if isinstance(x, Num):
        if not x.known:
            return "UNKNOWN" + (" (known part %s)" % rs(x.partial)
                                if x.partial is not None else "")
        x = x.value
    if x is None or not _finite(x):
        return "-"
    return ("-Rs " if x < 0 else "Rs ") + format(abs(int(round(x))), ",")


def _num(v):
    """float, or nan for anything that is not a finite number."""
    try:
        if v is None or not str(v).strip():
            return float("nan")
        f = float(str(v).replace(",", "").strip())
        return f if math.isfinite(f) else float("nan")
    except ValueError:
        return float("nan")


# ================================================================== ledger
def classify(narration):
    """Bucket of one ledger row by its narration. '_' buckets are not
    charges: _balance (opening/closing), _trade (buy/sell bills),
    _money (money you added / took out)."""
    n = str(narration or "").lower()
    if "net mtf funding by dhan" in n:
        return "_funding"
    if "opening balance" in n or "closing balance" in n:
        return "_balance"
    if "mtf" in n and ("int" in n or "interest" in n):
        return "MTF interest"
    if "dividend" in n:
        return "Dividend (income)"
    if "interest" in n or "dpc" in n or "delayed payment" in n:
        return "Other interest"
    if "pledge" in n:
        return "Pledge charges"
    if re.search(r"\bdp\b|depository|cdsl|nsdl", n):
        return "DP charges"
    if re.search(r"bill|settlement|trade|stt", n):
        return "_trade"
    if re.search(r"payin|pay-in|payout|pay-out|fund|transfer|upi|neft|imps|"
                 r"rtgs|bank|withdraw|deposit", n):
        return "_money"
    return "Other charges"


DATE_FMTS = ("%b %d, %Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y",
             "%d-%m-%Y", "%d-%b-%Y", "%d %b %Y")


def parse_date(v):
    """date or None. Explicit formats; slashes/dashes are DAY first (Indian
    statements), never month first (Codex F8: 05/10/2026 = 5 Oct)."""
    if isinstance(v, (dt.datetime, pd.Timestamp)):
        return v.date() if not pd.isna(v) else None
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip()
    if not s or s.lower() in ("nan", "nat", "none"):
        return None
    for f in DATE_FMTS:
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    m = re.match(r"^(\d{4}-\d{2}-\d{2})(?:[ T].*)?$", s)
    if m:
        try:
            return dt.date.fromisoformat(m.group(1))
        except ValueError:
            return None
    return None


def ledger_rows(raw):
    """Rows (dicts) from the Dhan API or a downloaded file ->
    DataFrame date, narration, debit, credit, bucket, voucher, bad.
    Duplicate vouchers (same number/date/narration/amounts) are dropped."""
    cols_out = ["date", "narration", "debit", "credit", "bucket", "voucher",
                "bad"]
    d = pd.DataFrame(raw)
    if d.empty:
        return pd.DataFrame(columns=cols_out)
    cols = {c: str(c).lower().replace(" ", "") for c in d.columns}

    def pick(*keys):
        for c, lc in cols.items():
            if any(k == lc or (len(k) > 4 and k in lc) for k in keys):
                return d[c]
        return pd.Series([""] * len(d))
    def exact(*keys):
        for c, lc in cols.items():
            if lc in keys:
                return d[c]
        return pd.Series([""] * len(d))
    nar = pick("narration", "description", "particular", "particulars",
               "remark", "remarks")
    vd = pick("voucherdesc", "vouchertype")
    out = pd.DataFrame({
        "date": [parse_date(x) for x in pick("voucherdate", "date",
                                             "txndate", "transactiondate")],
        "narration": [("%s %s" % (a, b)).strip() if str(b).lower() not in
                      str(a).lower() else str(a)
                      for a, b in zip(nar.fillna(""), vd.fillna(""))],
        "debit": [_num(x) for x in pick("debit")],
        "credit": [_num(x) for x in pick("credit")],
        "voucher": [str(x) for x in exact("vouchernumber", "voucherno",
                                          "voucher_no", "voucherid")]})
    out["bucket"] = [classify(x) for x in out["narration"]]
    out["bad"] = (out["debit"].isna() | out["credit"].isna()
                  | out["date"].isna() | (out["debit"] < 0) | (out["credit"] < 0))
    key = out[["voucher", "date", "narration", "debit", "credit"]].astype(str)
    has_v = out["voucher"].str.strip().ne("") & out["voucher"].ne("nan")
    dup = key.duplicated() & has_v
    return out[~dup][cols_out].reset_index(drop=True)


def read_statement_table(path, required):
    """Detect Excel by file signature even when its extension says .csv.
    Downloaded statements have title/metadata rows above their real header."""
    with open(path, "rb") as f:
        magic = f.read(8)
    if magic.startswith(b"PK") or magic.startswith(bytes.fromhex("D0CF11E0")):
        raw = pd.read_excel(path, header=None)
    else:
        try:
            raw = pd.read_csv(path, header=None, encoding="utf-8-sig")
        except UnicodeDecodeError:
            raw = pd.read_csv(path, header=None, encoding="cp1252")
    def norm(v): return re.sub(r"[^a-z0-9]", "", str(v).lower())
    for i in range(min(50, len(raw))):
        cols = [norm(v) for v in raw.iloc[i]]
        if all(norm(c) in cols for c in required):
            out = raw.iloc[i+1:].copy()
            out.columns = [str(v).strip() for v in raw.iloc[i]]
            # Only data columns; remove sidebar advertisements/blank columns.
            out = out.loc[:, [not c.lower().startswith("nan") for c in out.columns]]
            out = out.dropna(how="all")
            return out
    raise ValueError("statement header not found: " + ", ".join(required))


def read_ledger_file(path):
    raw = read_statement_table(path, ("Date", "Debit", "Credit"))
    # Dhan download uses negative debits; API uses positive debits. Normalize
    # only the explicit downloaded debit convention, never arbitrary columns.
    debit = raw["Debit"].map(_num)
    credit = raw["Credit"].map(_num)
    finite = debit.dropna()
    if len(finite) and (finite < 0).any() and (finite > 0).any():
        raise ValueError("mixed debit sign convention; statement needs review")
    raw["Debit"] = debit.abs()
    raw["Credit"] = credit
    raw = raw.rename(columns={"Transaction ID":"vouchernumber", "Voucher Descriptions":"voucherdesc"})
    return ledger_rows(raw.to_dict("records"))


def symbol_name_map(path):
    """Exact normalized official names only; ambiguous names are rejected."""
    master = pd.read_csv(path, low_memory=False)
    master = master[(master["SEM_EXM_EXCH_ID"] == "NSE") &
                    (master["SEM_INSTRUMENT_NAME"] == "EQUITY")]
    def norm(v):
        x = re.sub(r"\b(limited|ltd)\.?$", "", str(v).lower()).strip()
        return re.sub(r"[^a-z0-9]", "", x)
    candidates = {}
    for _, r in master.iterrows():
        sym = re.sub(r"-(EQ|BE)$", "", str(r["SEM_TRADING_SYMBOL"]).upper())
        for col in ("SEM_CUSTOM_SYMBOL", "SM_SYMBOL_NAME", "SEM_TRADING_SYMBOL"):
            key = norm(r.get(col, ""))
            if key and key != "nan": candidates.setdefault(key,set()).add(sym)
    result = {k: next(iter(v)) for k,v in candidates.items() if len(v)==1}
    # Known historical display-name changes. These are company-name aliases,
    # never product classifications or invented trades.
    aliases = {"GMR Airports Infrastructure":"GMRAIRPORT", "Zomato":"ETERNAL",
               "National Securities Depository (NSDL)":"NSDL",
               "National Securities Depos":"NSDL"}
    for name,sym in aliases.items(): result.setdefault(norm(name),sym)
    return result, norm


def read_trade_file(path, instruments, frm, today):
    raw = read_statement_table(path, ("Date", "Time", "Name", "Buy/Sell", "Quantity/Lot", "Trade Price"))
    coverage = re.search(r"TRADE_HISTORY_CSV_\d+_(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})_",os.path.basename(path),re.I)
    if coverage and (parse_date(coverage.group(1)) > frm or parse_date(coverage.group(2)) < today):
        raise ValueError("downloaded trade history does not cover requested dates; refresh export / adjust --from")
    names, norm = symbol_name_map(instruments)
    out, notes = [], []
    bad = 0
    for i, r in raw.iterrows():
        if str(r.get("Status", "")).strip().lower() != "traded":
            continue
        if str(r.get("Segment", "")).strip().lower() not in ("equity", "cash"):
            continue
        d = parse_date(r.get("Date"))
        if not d: bad += 1; continue
        if not frm <= d <= today: continue
        sym = names.get(norm(r.get("Name")))
        if not sym:
            sym = "UNMAPPED:" + str(r.get("Name"))
            bad += 1; notes.append("unmapped exact company name: " + str(r.get("Name")))
        prod = str(r.get("Order", "")).strip().upper()
        prod = "CNC" if prod == "DELIVERY" else prod
        stamp = _ts(d.isoformat()+" "+str(r.get("Time", "")))
        q, price = _num(r.get("Quantity/Lot")), _num(r.get("Trade Price"))
        reported = _num(r.get("Trade Value"))
        valid_value = _finite(reported) and _finite(q) and _finite(price) and abs(q*price-reported) <= .02 + .00501*abs(q)
        if not valid_value:
            bad += 1; notes.append("trade value inconsistent at file row %s" % i)
        # Trade Price is displayed to 2 decimals; Trade Value preserves more
        # precision. Only use it if consistent with the displayed rounding.
        if _finite(reported) and _finite(q) and q > 0 and abs(q*price-reported) <= .02+.00501*abs(q):
            price = reported / q
        # No execution reference in this export: retain same-second fills.
        out.append(dict(symbol=sym, side=str(r.get("Buy/Sell", "")).upper(),qty=q,price=price,
            ts=stamp,date=d.isoformat(),product=prod,exch=str(r.get("Exchange"))+"_EQ",
            isin="",security_id="",tid="",oid="",charges=0,charges_ok=False,
            source="BROKER_DOWNLOADED_TRADE_STATEMENT", file_row=int(i)+1, invalid_source_row=not valid_value or sym.startswith("UNMAPPED:")))
    if not out: raise ValueError("no executed equity trades in requested file/date range")
    notes.insert(0,"downloaded statement: %d trades; no execution IDs, no automatic row deduplication; trade fees not supplied" % len(out))
    return out, dict(complete=not bad,notes=notes)


def validate_statement_account(path, account_key, trade_file=False):
    """Fail closed if a downloaded statement has no account identity.
    Dhan's trade export embeds client ID only in its original filename;
    ledger Excel/CSV must carry Client ID metadata. Contradictions reject.
    Filename is a user-supplied export identity, not cryptographic proof.
    """
    import csv
    cid = str(account_key).removeprefix("DHAN_")
    if not cid.isdigit():
        raise ValueError("statement validation needs a numeric active Dhan client ID")
    found = set()
    if trade_file:
        match = re.search(r"TRADE_HISTORY_CSV_(\d+)_", os.path.basename(path), re.I)
        if match:
            found.add(match.group(1))
            if match.group(1)!=cid:raise ValueError("statement belongs to another account")
    with open(path,"rb") as f:magic=f.read(8)
    if magic.startswith(b"PK") or magic.startswith(bytes.fromhex("D0CF11E0")):
        rows=pd.read_excel(path,header=None,nrows=20).fillna("").values.tolist()
    else:
        with open(path,encoding="utf-8-sig") as f:
            rows=[]
            for i,row in enumerate(csv.reader(f)):
                if i>=20:break
                rows.append(row)
    for row in rows:
        for i,v in enumerate(row):
            text=str(v).strip()
            if text.lower() in ("client id","clientid","dhanclientid","client code") and i+1<len(row):
                found.add(str(row[i+1]).strip().removesuffix(".0"))
            match=re.fullmatch(r"(?:Client\s*ID|Client\s*Code)\s*[:=]\s*(\d+)",text,re.I)
            if match:found.add(match.group(1))
    if not found:raise ValueError("statement account identity missing; keep original Dhan export name / Client ID header")
    if found!={cid}:raise ValueError("statement belongs to another account or conflicting account identities")


def statement_loan(led, today):
    """An explicit dated 'Net MTF Funding by Dhan' snapshot, not inferred interest."""
    x = led[led["bucket"] == "_funding"]
    if not len(x): return Num(None, UNKNOWN, "no explicit current funding snapshot")
    x = x[x.date == today]
    if len(x) != 1 or bool(x.bad.iloc[0]):
        return Num(None, UNKNOWN, "funding snapshot stale/ambiguous; --loan required")
    amount = float(x.credit.iloc[0])-float(x.debit.iloc[0])
    return Num(amount, VERIFIED, "explicit broker statement Net MTF Funding snapshot " + today.isoformat()) \
        if amount >= 0 else Num(None, UNKNOWN, "invalid funding snapshot")


def dhan_ledger(sess, frm, to):
    import broker_api as ba
    d = ba._call(sess, "GET", "/ledger", params={
        "from-date": frm.isoformat(), "to-date": to.isoformat()})
    d = d.get("data", d) if isinstance(d, dict) else d
    if not isinstance(d, list):
        raise ValueError("ledger reply is not a list")
    return ledger_rows(d)


def charges(led):
    """{bucket: Num(debit - credit, VERIFIED, 'n rows')} for the charge
    buckets; a bucket with an unreadable amount is UNKNOWN."""
    x = led[~led["bucket"].astype(str).str.startswith("_")]
    out = {}
    for b, g in x.groupby("bucket"):
        if g["bad"].any():
            ok = g[~g["bad"]]
            out[b] = Num(None, UNKNOWN, "%d unreadable rows" % g["bad"].sum(),
                         partial=round(ok["debit"].sum() - ok["credit"].sum(),
                                       2))
        else:
            out[b] = Num(round(g["debit"].sum() - g["credit"].sum(), 2),
                         VERIFIED, "%d rows" % len(g))
    return out


def closing_balance(led):
    """Ledger closing balance (+ = cash with Dhan), or None."""
    x = led[[("closing balance" in str(n).lower()) for n in led["narration"]]]
    if not len(x):
        return None
    r = x.iloc[-1]
    return round(r["credit"] - r["debit"], 2)


def money_in_out(led):
    """(added, withdrawn) from the _money rows (credit = added)."""
    x = led[led["bucket"] == "_money"]
    return round(x["credit"].sum(), 2), round(x["debit"].sum(), 2)


def interest_periods(led):
    """MTF interest rows with their period: [(start, end, debit, credit)]."""
    out = []
    for _, r in led[led["bucket"] == "MTF interest"].iterrows():
        m = re.findall(r"(\d{2})/(\d{2})/(\d{4})", str(r["narration"]))
        a = b = None
        if len(m) >= 2:
            try:
                a, b = (dt.date(int(y), int(mo), int(d)) for d, mo, y in m[:2])
            except ValueError:
                a = b = None
        out.append((a, b, r["debit"], r["credit"]))
    return out


def avg_interest_balance(led, rate=MTF_RATE):
    """DIAGNOSTIC ONLY (Codex F1): last MTF interest debit / its days x 365
    / rate = the AVERAGE interest-bearing balance of that past period at
    that assumed rate -- not today's loan. (avg Rs, per day Rs, period text,
    period end) or (None, None, '', None)."""
    if not rate or rate <= 0:
        return None, None, "", None
    per = [p for p in interest_periods(led) if p[0] and p[1] and p[2] > 0]
    if not per:
        return None, None, "", None
    a, b, deb, _ = max(per, key=lambda p: p[1])
    days = (b - a).days + 1
    if days <= 0:
        return None, None, "", None
    day = deb / days
    return round(day * 365 / rate, 2), round(day, 2), "%s -> %s" % (a, b), b


# ================================================================== trades
def _ts(v):
    """datetime of an execution, or None (never invented)."""
    s = str(v or "").strip()
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M",
              "%d-%m-%Y %H:%M:%S", "%d-%b-%Y %H:%M:%S"):
        for c in (s, s[:19]):
            try:
                return dt.datetime.strptime(c, f)
            except ValueError:
                pass
    return None


def parse_trade(x, ids=None):
    """Dhan trade-history row -> dict, keeping identity + exact time."""
    ids = ids or {}
    sid = str(x.get("securityId") or "")
    sym = ids.get(sid) or re.sub(r"-(EQ|BE)$", "", str(
        x.get("tradingSymbol") or x.get("customSymbol") or "").upper())
    ts = _ts(x.get("exchangeTime") or x.get("createTime") or
             x.get("updateTime"))
    fee_fields = ("sebiTax", "stt", "brokerageCharges", "serviceTax",
                  "exchangeTransactionCharges", "stampDuty")
    chg = [_num(x.get(k)) for k in fee_fields]
    return {
        "symbol": sym, "side": str(x.get("transactionType") or "").upper(),
        "qty": _num(x.get("tradedQuantity")),
        "price": _num(x.get("tradedPrice")),
        "ts": ts, "date": ts.date().isoformat() if ts else None,
        "product": str(x.get("productType") or "").upper(),
        "exch": str(x.get("exchangeSegment") or ""),
        "isin": str(x.get("isin") or ""), "security_id": sid,
        "tid": str(x.get("exchangeTradeId") or x.get("tradeId") or ""),
        "oid": str(x.get("orderId") or x.get("exchangeOrderId") or ""),
        "charges": sum(c for c in chg if c == c),
        "charges_ok": all(_finite(c) and c >= 0 for c in chg)}


def execution_key(t):
    # Exchange IDs are not assumed to be unique across all dates/stocks.
    return (t.get("exch"), t.get("date"), t.get("tid"), t.get("oid"),
            t.get("security_id") or t.get("symbol"), t.get("ts"))


def dedupe(trades):
    """Only collapse identical, identifiable execution replays.
    Reused IDs on different dates/orders/stocks are retained. With no ID
    and no order reference there is insufficient evidence to drop a fill.
    Conflicting payloads are retained and audited by dhan_trades()."""
    seen, out, drop = set(), [], 0
    for t in trades:
        identifiable = bool(t.get("ts") and str(t.get("tid") or "").strip().upper() not in ("","0","NA","N/A","NONE","NULL"))
        key = execution_key(t) + tuple(t.get(k) for k in
              ("symbol", "side", "qty", "price", "product", "charges", "charges_ok"))
        if identifiable and key in seen:
            drop += 1
            continue
        if identifiable:
            seen.add(key)
        out.append(t)
    return out, drop


def dhan_trades(sess, frm, to, chunk=90):
    """Every executed trade (Dhan GET /v2/trades/{from}/{to}/{page}, asked
    90 days at a time). Returns (trades, meta); meta['complete'] is False
    when a chunk still had data at the page limit."""
    import broker_api as ba
    ids = {}
    try:
        ids = {str(v["id"]): k for k, v in ba.symbol_map(sess).items()}
    except Exception:
        pass
    out, notes, complete = [], [], True
    a = frm
    while a <= to:
        b = min(to, a + dt.timedelta(days=chunk - 1))
        for page in range(PAGE_LIMIT):
            d = ba._call(sess, "GET", "/trades/%s/%s/%d" % (a, b, page))
            d = d.get("data", d) if isinstance(d, dict) else d
            if not isinstance(d, list) or any(not isinstance(x, dict) for x in d):
                raise ValueError("trade history reply must be a list of rows; null is unknown")
            if not d:
                break
            out += [parse_trade(x, ids) for x in d
                    if "EQ" in str(x.get("exchangeSegment", "EQ"))]
            if page == PAGE_LIMIT - 1:
                complete = False
                notes.append("%s..%s: still data at page %d" % (a, b, page))
        a = b + dt.timedelta(days=1)
    payloads = {}
    for t in out:
        if t.get("tid") and t.get("ts"):
            payloads.setdefault(execution_key(t), set()).add(tuple(t.get(k) for k in
                ("side", "qty", "price", "product", "charges", "charges_ok")))
    conflicts = sum(len(v) > 1 for v in payloads.values())
    if conflicts:
        complete = False
        notes.append("%d execution identities have conflicting payloads; statement required" % conflicts)
    out, drop = dedupe(out)
    if drop:
        notes.append("%d duplicate executions dropped" % drop)
    return out, {"complete": complete, "notes": notes}


def valid_trade(t, today=None):
    today = today or dt.date.today()
    return (not t.get("invalid_source_row",False) and t["side"] in ("BUY", "SELL") and _finite(t["qty"]) and
            t["qty"] > 0 and _finite(t["price"]) and t["price"] > 0 and
            t["ts"] is not None and t["ts"].date() <= today and
            bool(t["symbol"]))


def fifo(trades, today=None):
    """MTF trades -> FIFO by EXACT execution time (Codex F5: never re-sort a
    same-day SELL after a BUY). A SELL without enough earlier BUYs in the
    history = 'uncovered' -> realised P&L is UNKNOWN (the matched part is
    kept as 'matched').
    Returns dict: bought, sold, matched (Rs), realised (Num), charges (Num),
    lots {sym: [[date, qty, price, charge_left]]}, uncovered {sym: qty},
    invalid (count of rows without a valid time/qty/price)."""
    good = [t for t in trades if valid_trade(t, today)]
    invalid = len(trades) - len(good)
    order = sorted(range(len(good)), key=lambda i: (good[i]["ts"], i))
    lots, bought, sold, matched, chg, unc = {}, 0.0, 0.0, 0.0, 0.0, {}
    chg_ok = not invalid
    bad_fee_symbols = set()
    for i in order:
        t = good[i]
        L = lots.setdefault(t["symbol"], [])
        chg += t.get("charges", 0)
        fee_ok = t.get("charges_ok", False) and _finite(t.get("charges"))
        chg_ok = chg_ok and fee_ok
        if not fee_ok:
            bad_fee_symbols.add(t["symbol"])
        if t["side"] == "BUY":
            L.append([t["date"], t["qty"], t["price"], t.get("charges", 0)])
            bought += t["qty"] * t["price"]
        else:
            sold += t["qty"] * t["price"]
            left = t["qty"]
            while left > 1e-9 and L:
                take = min(left, L[0][1])
                matched += take * (t["price"] - L[0][2])
                L[0][3] -= L[0][3] * take / L[0][1]
                L[0][1] -= take
                left -= take
                if L[0][1] <= 1e-9:
                    L.pop(0)
            if left > 1e-9:
                unc[t["symbol"]] = unc.get(t["symbol"], 0) + left
    lots = {s: L for s, L in lots.items() if sum(x[1] for x in L) > 1e-9}
    realised = Num(round(matched, 2), VERIFIED, "") if not unc and \
        not invalid else Num(None, UNKNOWN, "sells without a buy in the "
                             "history / invalid rows", partial=round(matched,
                                                                     2))
    return {"bought": round(bought, 2), "sold": round(sold, 2),
            "matched": round(matched, 2), "realised": realised,
            "charges": Num(round(chg, 2), VERIFIED if chg_ok else UNKNOWN,
                           "per-trade fees in the history"),
            "lots": lots, "uncovered": unc, "invalid": invalid,
            "bad_fee_symbols": bad_fee_symbols}


def fifo_by_buy_origin(trades, today):
    """Economic FIFO across CNC/MTF equity acquisitions (intraday excluded).
    A DELIVERY sell can dispose of an earlier MTF-origin lot. This links
    economic price P&L, not the date/amount of a broker funding conversion.
    Every cross-product match is exported; attribution is ESTIMATED.
    Missing/unclassified acquisitions or uncovered sells taint the result."""
    cash = [t for t in trades if t["product"] in ("CNC","MTF","MARGIN")]
    good = sorted([t for t in cash if valid_trade(t,today)], key=lambda t:t["ts"])
    invalid = len(cash)-len(good)
    inventory, unc, matches = {}, {}, []
    bought=sold=matched=fees=0.0
    fee_ok=True
    bad_fee_symbols=set()
    for t in good:
        sym=t["symbol"];L=inventory.setdefault(sym,[])
        is_mtf=t["product"] in ("MTF","MARGIN")
        if t["side"]=="BUY":
            L.append([t["date"],t["qty"],t["price"],t.get("charges",0),t["product"]])
            if is_mtf:
                bought+=t["qty"]*t["price"]
                fees+=t.get("charges",0)
                ok=t.get("charges_ok",False)
                fee_ok=fee_ok and ok
                if not ok:bad_fee_symbols.add(sym)
        else:
            left=t["qty"];origin_qty=0.0
            while left>1e-9 and L:
                take=min(left,L[0][1]);lot=L[0]
                origin=lot[4]
                if origin in ("MTF","MARGIN"):
                    matched+=take*(t["price"]-lot[2]);sold+=take*t["price"];origin_qty+=take
                if origin!=t["product"]:
                    matches.append({"Symbol":sym,"Buy product":origin,"Buy date":lot[0],"Sell product":t["product"],"Sell time":t["ts"].isoformat(sep=" "),"Qty":take,"Buy price":lot[2],"Sell price":t["price"],"Price P&L":round(take*(t["price"]-lot[2]),2),"Status":ESTIMATED,"Note":"economic FIFO buy-origin attribution; not proof of broker conversion date or funded balance"})
                lot[3]-=lot[3]*take/lot[1];lot[1]-=take;left-=take
                if lot[1]<=1e-9:L.pop(0)
            if origin_qty:
                fees+=t.get("charges",0)*origin_qty/t["qty"]
                ok=t.get("charges_ok",False);fee_ok=fee_ok and ok
                if not ok:bad_fee_symbols.add(sym)
            if left>1e-9:unc[sym]=unc.get(sym,0)+left
    lots={sym:[x[:4] for x in L if x[4] in ("MTF","MARGIN")] for sym,L in inventory.items()}
    lots={sym:L for sym,L in lots.items() if L}
    st=ESTIMATED if matches else VERIFIED
    return {"bought":round(bought,2),"sold":round(sold,2),"matched":round(matched,2),
        "realised":Num(round(matched,2),st,"economic FIFO by buy product") if not unc and not invalid else Num(None,UNKNOWN,"economic FIFO has uncovered sells/invalid rows",partial=round(matched,2)),
        "charges":Num(round(fees,2),st if fee_ok else UNKNOWN,"MTF-origin fees incl. matched delivery exits"),
        "lots":lots,"uncovered":unc,"invalid":invalid,"bad_fee_symbols":bad_fee_symbols,
        "cross_matches":matches,
        "all_open_qty":{sym:sum(x[1] for x in L) for sym,L in inventory.items() if L}}


def lot_interest(lots, share, rate, today):
    """Interest ESTIMATE, each lot its own calendar days (Codex F6):
    sum(qty x price x share x rate x days/365). Broker billing (settlement
    days, slabs, daily balance) differs -> always ESTIMATED."""
    if not (_finite(share) and 0 <= share <= 1 and _finite(rate) and 0 <= rate < 1):
        return Num(None, UNKNOWN, "needs valid funded share and explicit rate (zero allowed)")
    s = 0.0
    for d, q, p, _ in lots:
        days = (today - dt.date.fromisoformat(d)).days
        if days < 0:
            return Num(None, UNKNOWN, "lot dated in the future")
        s += q * p * share * rate * days / 365
    return Num(round(s, 2), ESTIMATED, "lot days x %.0f%% funded x %.2f%%"
               % (100 * share, 100 * rate))


# ================================================================== prices
def last_closes(symbols):
    """{SYM: (close, date)} from the free history + NSE bhavcopy fill."""
    import broker_api as ba
    import daily_screener as ds
    fr = {}
    for x in symbols:
        try:
            f = ds.fetch_eod(x.lower())
        except Exception:
            f = None
        if f is not None and len(f):
            fr[x.lower()] = f
    try:
        ba.bhav_fill(fr, ds.last_expected_session())
    except Exception:
        pass
    return {x.upper(): (float(f["Close"].iloc[-1]), f.index[-1].date())
            for x, f in fr.items()}


def quotes(sess, symbols):
    """Official NSE primary; broker secondary; dated close last resort.
    Prices never select/change the holdings account. Stale data stays labelled.
    """
    import broker_api as ba
    import daily_screener as ds
    import nse_calendar as nc
    from mtf_prices import nse_quotes,safe_error,data_plan_note
    out, live = {}, {}
    # Synthetic legacy tests use object() sessions. Actual broker sessions
    # always have a broker attribute and use the official NSE primary source.
    if getattr(sess,'broker',None):
        now=ds.now_ist();trading=nc.is_trading_day(now.date())
        out=nse_quotes(symbols,now,ds.last_expected_session(),trading and ds.market_open(),trading_day=trading)
    missing=[s for s in symbols if s not in out]
    try:
        live = ba.live_prices(sess, missing) if missing else {}
    except Exception as e:
        print("! Broker quote unavailable (%s): %s" % (type(e).__name__,safe_error(e,sess)))
        note=data_plan_note(ba,sess)
        if note:print('! '+note)
    for s in symbols:
        if s not in out and _finite(live.get(s)) and live[s] > 0:
            out[s] = (float(live[s]), ESTIMATED, "%s broker LTP fetched %s IST (exchange timestamp unavailable)"%(getattr(sess,'broker','test'),ds.now_ist().strftime('%Y-%m-%d %H:%M:%S')))
    miss = [s for s in symbols if s not in out]
    if miss:
        for s, (c, d) in last_closes(miss).items():
            if s in miss and _finite(c) and c > 0:
                if d>ds.last_expected_session():continue
                out[s] = (c, ESTIMATED, "last close %s; NOT a current live quote" % d)
                print('! %s: live quote unavailable; using dated closing snapshot %s'%(s,d))
    return out


# ================================================================== build
def mtf_positions(sess):
    """{SYM: qty} of MTF rows in Dhan's positions API (product-tagged)."""
    import broker_api as ba
    pos = ba._call(sess, "GET", "/positions")
    pos = pos.get("data", pos) if isinstance(pos, dict) else pos
    if not isinstance(pos, list) or any(not isinstance(p, dict) for p in pos):
        raise ValueError("current positions reply invalid; inventory unknown")
    out = {}
    for p in pos:
        prod = str(p.get("productType") or "").upper()
        q = _num(p.get("netQty"))
        if prod in ("MTF", "MARGIN") and not _finite(q):
            raise ValueError("MTF position quantity missing/invalid")
        if prod in ("MTF", "MARGIN") and q > 0:
            sym = re.sub(r"-(EQ|BE)$", "", str(p.get("tradingSymbol") or
                                              "").upper())
            if not sym:
                raise ValueError("MTF position identity missing")
            out[sym] = out.get(sym, 0) + q
    return out


def reconcile(f, demat, demat_ok, current=None):
    """Open MTF lots vs demat qty. Dhan's demat qty is ALL products, so it
    can confirm 'not more than', never 'this much is MTF' (Codex F4).
    Returns (rows, use {sym: lots}) -- a stock is valued only when its MTF
    qty is consistent with the demat."""
    rows, use = [], {}
    for s in sorted(set(f["lots"]) | set(f["uncovered"])):
        lots = f["lots"].get(s, [])
        q = sum(x[1] for x in lots)
        dq = demat.get(s, {}).get("qty") if demat_ok else None
        if s in f["uncovered"]:
            st, note = UNKNOWN, ("sold %g more than bought since --from: "
                                 "older buys missing" % f["uncovered"][s])
        elif not demat_ok:
            st, note = UNKNOWN, "demat / current positions / history could not be read"
        elif s in (current or {}) and abs(current[s] - q) > 1e-9:
            st, note = UNKNOWN, "current broker MTF %g != history MTF %g; reconcile product/events" % (current[s], q)
        elif not dq:
            st, note = UNKNOWN, ("history says %g open but NOT in demat "
                                 "(sold / converted / transferred?)" % q)
        elif dq + 1e-9 < q:
            st, note = UNKNOWN, ("demat %g < MTF lots %g (partial sale / "
                                 "conversion not in history)" % (dq, q))
        elif dq > q + 1e-9:
            st, note = ESTIMATED, ("demat %g > MTF lots %g: rest = CNC or "
                                   "bought before --from" % (dq, q))
        else:
            st = VERIFIED if s in (current or {}) else ESTIMATED
            note = "MTF lots = demat qty" + ("; current MTF quantity confirmed" if st == VERIFIED
                                            else "; current product inferred from history")
        rows.append({"Symbol": s, "MTF qty (history)": q, "Demat qty": dq,
                     "Status": st, "Note": note})
        if q > 0 and st != UNKNOWN:
            use[s] = lots
    for s, q in sorted((current or {}).items()):
        if s not in f["lots"] and s not in f["uncovered"]:
            rows.append({"Symbol": s, "MTF qty (history)": 0,
                         "Demat qty": demat.get(s, {}).get("qty"),
                         "Status": UNKNOWN,
                         "Note": "broker shows %g MTF now but no MTF buy in "
                                 "the history (bought before --from?)" % q})
    if demat_ok:
        for s in sorted(demat):
            if s not in f["lots"] and s not in f["uncovered"] and \
                    s not in (current or {}):
                rows.append({"Symbol": s, "MTF qty (history)": 0,
                             "Demat qty": demat[s].get("qty"),
                             "Status": UNKNOWN,
                             "Note": "MTF status unconfirmed: held in demat but "
                                     "no matching MTF history; missing history does not prove CNC"})
    return rows, use


def parse_confirmations(value):
    out = {}
    for part in (value or "").split(","):
        if not part.strip():
            continue
        bits = part.strip().upper().split(":")
        sym = bits[0]
        if not re.fullmatch(r"[A-Z0-9&_.-]+", sym) or len(bits) > 2:
            raise ValueError("--confirm-mtf use SYMBOL or SYMBOL:QTY, comma separated")
        q = None if len(bits) == 1 else _num(bits[1])
        if q is not None and (not _finite(q) or q <= 0 or q != int(q)):
            raise ValueError("confirmed MTF quantity must be a positive whole number")
        if sym in out:
            raise ValueError("duplicate confirmed stock: " + sym)
        out[sym] = q
    return out


def confirmed_inventory(args, demat, today):
    """Account-local user product declarations, pinned to current qty/cost.
    Seven-day expiry and quantity/cost checks prevent silent reuse after a
    position change. This confirms product, not purchase dates/trade fees."""
    path = getattr(args, "confirmation_path", None)
    key = getattr(args, "account_key", "")
    saved, messages = {}, []
    if path and os.path.exists(path):
        try:
            with open(path) as file:
                d = json.load(file)
            if d.get("account_key") != key:
                raise ValueError("account mismatch")
            saved = d.get("stocks", {})
            if not isinstance(saved, dict):
                raise ValueError("invalid stock map")
        except (ValueError, TypeError, OSError) as e:
            raise ValueError("MTF confirmation file unreadable: " + type(e).__name__)
    requests = getattr(args, "confirm_requests", {}) or {}
    active, new = {}, dict(saved)
    for sym, r in saved.items():
        h = demat.get(sym, {})
        d = parse_date(r.get("as_of"))
        if d and 0 <= (today-d).days <= 7 and h and \
                _finite(h.get("qty")) and abs(h["qty"]-r["demat_qty"]) < 1e-9 and \
                _finite(h.get("avg_price")) and abs(h["avg_price"]-r["avg_price"]) < .005:
            active[sym] = r
        else:
            messages.append(sym + ": saved MTF confirmation expired or holdings changed; re-confirm")
    for sym, q in requests.items():
        h = demat.get(sym)
        if not h or not _finite(h.get("qty")) or h["qty"] <= 0:
            raise ValueError(sym + ": no positive broker demat holding; cannot confirm MTF")
        q = h["qty"] if q is None else q
        if q > h["qty"] or not _finite(h.get("avg_price")) or h["avg_price"] <= 0:
            raise ValueError(sym + ": invalid MTF quantity or missing broker average cost")
        # A subset has no MTF-specific average cost; never use combined average.
        r = dict(qty=q, demat_qty=h["qty"], avg_price=h["avg_price"], as_of=today.isoformat(),
                 provenance="USER_CONFIRMED", cost_valid=abs(q-h["qty"]) < 1e-9)
        active[sym] = new[sym] = r
    if requests and path:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=".mtf-confirm-", dir=os.path.dirname(path))
        try:
            with os.fdopen(fd, "w") as f:
                json.dump(dict(account_key=key, stocks=new), f, indent=2)
                f.flush(); os.fsync(f.fileno())
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
    return active, messages


def mtf_sell_fees(value, broker="DHAN"):
    """Single MTF sale model shared by summary and audit (never an order)."""
    if broker != "DHAN":raise ValueError("MTF fee model supports Dhan only")
    from mtf_breakeven import sell_cost
    return round(sell_cost(value),2)


def holding_row(sym, h, qty, px, product_note, cost_valid=True):
    """Current MTF stock independent of an incomplete historical FIFO.
    Average cost is a broker holdings basis, not a reconstructed MTF lot.
    Missing lot dates, fees and attributable interest remain unknown."""
    avg = h.get("avg_price")
    cost = Num(qty*avg, ESTIMATED, "broker holdings average; individual lots unavailable") \
        if cost_valid and _finite(avg) and avg > 0 else Num(None, UNKNOWN, "MTF-specific cost missing (mixed products?)")
    p = px.get(sym)
    value = Num(qty*p[0], ESTIMATED, "current qty; " + p[2] + "; " + product_note) \
        if p else Num(None, UNKNOWN, "price missing")
    gross = total([value, -cost])
    import journal
    fees = Num(mtf_sell_fees(value.value), ESTIMATED, "MTF rate card: assumes one sell order + one unpledge per stock") \
        if value.known else Num(None, UNKNOWN, "price missing")
    return dict(Symbol=sym, Qty=qty, **{"Avg cost":avg if cost.known else None,
        "Lots":None, "First buy":None, "Last buy":None,
        "Inventory status":ESTIMATED, "Product source":product_note,
        "Cost basis":"BROKER_HOLDINGS_AVERAGE" if cost.known else "UNKNOWN",
        "Open cost (Rs)":cost, "Price":p[0] if p else None,
        "Price source":p[2] if p else "UNKNOWN", "Value (Rs)":value,
        "Gross P&L (Rs)":gross,
        "Gross P&L %":round(100*gross.value/cost.value,2) if gross.known and cost.known else None,
        "Interest est. (Rs)":Num(None, UNKNOWN, "purchase lot dates missing; use actual ledger/statement"),
        "Sell fees est. (Rs)":fees, "Buy charges (Rs)":Num(None, UNKNOWN, "MTF lot fees missing")})


def open_table(use, px, today, share, rate, broker="DHAN", statuses=None,
               bad_fee_symbols=None):
    """One row per open MTF stock, every money column with its status."""
    import journal
    rows = []
    for s, lots in sorted(use.items()):
        q = sum(x[1] for x in lots)
        cost = sum(x[1] * x[2] for x in lots)
        p = px.get(s)
        inv_status = (statuses or {}).get(s, VERIFIED)
        value_status = (UNKNOWN if p and UNKNOWN in (p[1], inv_status) else
                        ESTIMATED if p and ESTIMATED in (p[1], inv_status) else inv_status)
        val = Num(q * p[0], value_status, p[2]) if p else Num(None, UNKNOWN, "no price")
        pnl = total([val, Num(-cost, inv_status)])
        buy_fee_status = UNKNOWN if s in (bad_fee_symbols or set()) else inv_status
        intr = lot_interest(lots, share, rate, today)
        fee = Num(mtf_sell_fees(val.value, broker), ESTIMATED,
                  "MTF rate card: one order + one unpledge per stock") if val.known else Num(None, UNKNOWN)
        rows.append({
            "Symbol": s, "Qty": q, "Avg cost": round(cost / q, 2),
            "Lots": len(lots), "First buy": lots[0][0],
            "Last buy": lots[-1][0], "Inventory status": inv_status,
            "Open cost (Rs)": Num(round(cost, 2), inv_status),
            "Price": p[0] if p else None, "Price source": p[2] if p else
            "UNKNOWN", "Value (Rs)": val, "Gross P&L (Rs)": pnl,
            "Gross P&L %": round(100 * pnl.value / cost, 2) if pnl.known
            and cost else None,
            "Interest est. (Rs)": intr, "Sell fees est. (Rs)": fee,
            "Buy charges (Rs)": Num(round(sum(x[3] for x in lots), 2), buy_fee_status)})
    return rows


def money_cell(n):
    return n.value if isinstance(n, Num) and n.known else None


def flatten(rows):
    """Num cells -> value column + '<col> status' column for Excel."""
    out = []
    for r in rows:
        o = {}
        for k, v in r.items():
            if isinstance(v, Num):
                o[k] = money_cell(v)
                o[k + " status"] = v.status
                o[k + " known part"] = v.partial if not v.known else None
                o[k + " note"] = v.note
            else:
                o[k] = v
        out.append(o)
    return pd.DataFrame(out)


def current_acquisition_lots(trades, demat, declared, today):
    """All-product economic acquisition dates for explicitly CURRENT MTF qty.
    Blank historical products never become historical MTF trades. Only a
    full qty/cost reconciliation to current holdings permits this model.
    Does NOT prove MTF funding began on acquisition date.
    """
    cash=[t for t in trades if t.get("product") != "INTRADAY"]
    lots={};bad=set()
    for t in cash:
        if not valid_trade(t,today):bad.add(t.get("symbol"))
    for t in sorted([t for t in cash if valid_trade(t,today)],key=lambda t:t["ts"]):
        sym=t["symbol"];L=lots.setdefault(sym,[])
        if t["side"]=="BUY":L.append([t["date"],t["qty"],t["price"],t.get("charges",0)])
        else:
            left=t["qty"]
            while left>1e-9 and L:
                take=min(left,L[0][1]);L[0][3]-=L[0][3]*take/L[0][1];L[0][1]-=take;left-=take
                if L[0][1]<=1e-9:L.pop(0)
            if left>1e-9:bad.add(sym)
    out={}
    for sym,q in declared.items():
        h=demat.get(sym,{});L=lots.get(sym,[])
        qty=sum(x[1] for x in L);cost=sum(x[1]*x[2] for x in L);avg=h.get("avg_price")
        if sym not in bad and L and _finite(avg) and avg>0 and abs(q-h.get("qty",0))<1e-9 and abs(qty-q)<1e-9 and abs(cost-q*avg)<=.02+.0051*q:
            out[sym]=L
    return out


def build(args, sess, today, frm):
    """All the numbers (no printing). Returns a dict of sections."""
    import broker_api as ba
    src = {}
    # 1. trade history (MTF product; blank product guessed only if no MTF)
    try:
        if getattr(args, "trades", None):
            instruments = getattr(args,"instruments",None) or ba.DHAN_SCRIP_FILE
            validate_statement_account(os.path.expanduser(args.trades),getattr(args,"account_key",""),trade_file=True)
            allt, meta = read_trade_file(os.path.expanduser(args.trades), instruments, frm, today)
        else:
            allt, meta = dhan_trades(sess, frm, today)
        src["trades"] = (VERIFIED if meta["complete"] else UNKNOWN,
                         "; ".join(meta["notes"]) or "%d trades" % len(allt))
    except Exception as e:
        allt = []
        src["trades"] = (UNKNOWN, "%s: %s" % (type(e).__name__, safe_error(e, sess)))
    prods = {}
    for t in allt:
        prods[t["product"] or "?"] = prods.get(t["product"] or "?", 0) + 1
    if prods.get("?"):
        src["historical_products"] = (UNKNOWN,
            "%d/%d trades have no broker product tag; MTF/CNC conversion records or complete product-tagged statement required" % (prods["?"],len(allt)))
    mtf = [t for t in allt if t["product"] in ("MTF", "MARGIN")]
    guessed = False
    if not mtf and any(not t["product"] for t in allt):
        mtf, guessed = [t for t in allt if not t["product"]], True
    f = fifo(mtf, today)
    cross_matches = []
    if getattr(args,"trades",None) and src["trades"][0] != UNKNOWN and not guessed and all(t["product"] in ("CNC","MTF","MARGIN","INTRADAY") for t in allt):
        f = fifo_by_buy_origin(allt,today)
        cross_matches=f["cross_matches"]
        src["economic_fifo"]=(ESTIMATED if cross_matches else VERIFIED,
            "%d cross-product lot matches; acquisition-origin price P&L, not funding conversion proof" % len(cross_matches))
    if f["invalid"]:
        src["trades"] = (UNKNOWN, src["trades"][1] + "; %d MTF rows without "
                         "a valid time/qty/price" % f["invalid"])
    # 2. demat (all products) for reconciliation
    try:
        demat = {h["symbol"]: h for h in ba.holdings(sess)}
        demat_ok = True
        src["demat"] = (VERIFIED, "broker holdings response")
    except Exception as e:
        demat, demat_ok = {}, False
        src["demat"] = (UNKNOWN, type(e).__name__)
    try:
        current = mtf_positions(sess)
        current_ok = True
        src["positions"] = (VERIFIED, "product-tagged current position response; empty does not prove carried inventory absent")
    except Exception as e:
        current, current_ok = {}, False
        src["positions"] = (UNKNOWN, "%s: %s" % (type(e).__name__, safe_error(e, sess)))
    confirmed, confirmation_notes = confirmed_inventory(args, demat, today) if demat_ok else ({}, [])
    if getattr(args, "confirm_requests", {}) and not demat_ok:
        raise ValueError("Cannot save MTF confirmation without broker holdings")
    src["product_confirmation"] = (ESTIMATED if confirmed else UNKNOWN if confirmation_notes else VERIFIED,
        "; ".join([s + " MTF USER_CONFIRMED " + r["as_of"] for s,r in sorted(confirmed.items())] + confirmation_notes)
        or "no user overrides; missing MTF trade is not proof of CNC")
    history_ok = src["trades"][0] != UNKNOWN
    inventory_ok = history_ok and demat_ok and current_ok
    rec, use = reconcile(f, demat, demat_ok and current_ok and (history_ok or bool(getattr(args,"trades",None))), current)
    mixed_blank = {t["symbol"] for t in allt if not t["product"]} if not guessed else set()
    for sym in sorted(mixed_blank):
        use.pop(sym, None)
        existing = next((r for r in rec if r["Symbol"] == sym), None)
        if existing is None:
            existing = {"Symbol": sym, "Demat qty": demat.get(sym, {}).get("qty")}
            rec.append(existing)
        existing.update(Status=UNKNOWN, Note="unclassified product in mixed history; resolve MTF/CNC from statement")
    if guessed:
        for key in ("realised", "charges"):
            n = f[key]
            if n.known:
                f[key] = Num(n.value, ESTIMATED, "product inferred from blank history")
    if not history_ok:
        f["realised"] = Num(None, UNKNOWN, "trade history incomplete", partial=f["matched"])
        f["charges"] = Num(None, UNKNOWN, "trade history incomplete", partial=f["charges"].value)
    src["trade_fees"] = (f["charges"].status, f["charges"].note)
    f["bought_num"] = Num(f["bought"], UNKNOWN if not history_ok else ESTIMATED if guessed else VERIFIED)
    f["sold_num"] = Num(f["sold"], UNKNOWN if not history_ok else ESTIMATED if guessed or cross_matches else VERIFIED)
    # an unreconciled MTF stock makes every open total UNKNOWN (no silent
    # omission, Codex F4)
    history_gap = [r["Symbol"] for r in rec if r["Status"] == UNKNOWN and r["Symbol"] not in confirmed]
    fallback = {}
    for sym, r in confirmed.items():
        # User confirmation applies only to CURRENT holdings, never all past trades.
        histqty = sum(x[1] for x in f["lots"].get(sym,[]))
        exact_history = sym in use and not guessed and sym not in mixed_blank and sym not in f["uncovered"] and abs(histqty-r["qty"]) < 1e-9
        if not exact_history:
            use.pop(sym, None)
            fallback[sym] = (demat[sym], r["qty"], "USER_CONFIRMED " + r["as_of"], r["cost_valid"])
        old = next((x for x in rec if x["Symbol"] == sym), None)
        row = {"Symbol":sym, "MTF qty (history)":sum(x[1] for x in f["lots"].get(sym, [])),
               "Demat qty":demat[sym]["qty"], "Current MTF qty":r["qty"],
               "Status":ESTIMATED, "Scope":"CURRENT", "Note":"MTF USER_CONFIRMED; " + ("explicit MTF trade lots match current quantity" if exact_history else "broker holdings qty/cost; purchase lots not invented")}
        if old is None: rec.append(row)
        else: old.update(row)
    declared={s:r["qty"] for s,r in confirmed.items() if r["cost_valid"]}
    declared.update({s:q for s,q in current.items() if s in demat and abs(q-demat[s]["qty"])<1e-9})
    acquisition=current_acquisition_lots(allt,demat,declared,today) if history_ok else {}
    for sym,L in acquisition.items():
        use[sym]=L;fallback.pop(sym,None)
        for r in rec:
            if r["Symbol"]==sym:
                r.update(Status=ESTIMATED,Scope="CURRENT",Note="current MTF declared; all-product acquisition qty/cost match broker. Funding since buy is an ESTIMATED assumption, not historical product/funding proof")
    src["current_acquisitions"]=(ESTIMATED if acquisition else UNKNOWN if declared else VERIFIED,
        "all-product dated acquisition lots reconciled for: "+", ".join(sorted(acquisition)) if acquisition else "no complete current acquisition lot basis; interest model remains UNKNOWN" if declared else "no current product declarations")
    # Historical orphan lots are not current positions. Only separate their
    # scope when ALL held stock quantities have an explicit current product basis.
    current_authoritative = bool(demat) and demat_ok and current_ok and all(
        (s in confirmed and abs(confirmed[s]["qty"]-h["qty"]) < 1e-9) or
        (s in current and abs(current[s]-h["qty"])<1e-9) or
        (getattr(args,"trades",None) and not guessed and
         abs(sum(x[1] for x in f["lots"].get(s,[]))-h["qty"]) < 1e-9)
        for s,h in demat.items())
    gap = []
    for r in rec:
        if r["Status"] == UNKNOWN:
            historical_only = current_authoritative and r["Symbol"] not in demat and r["Symbol"] not in current
            r["Scope"] = "HISTORY" if historical_only else "CURRENT/HISTORY"
            if not historical_only: gap.append(r["Symbol"])
    for sym, q in current.items():
        if sym in confirmed and abs(confirmed[sym]["qty"]-q) > 1e-9:
            gap.append(sym)
            fallback.pop(sym, None)
            rec.append({"Symbol":sym, "Status":UNKNOWN, "Scope":"CURRENT", "Note":"confirmed quantity differs from broker MTF position; reconcile before valuing"})
    # Without classified products, historical selection is only a candidate
    # subtotal; blank tags never become a complete MTF period result.
    if guessed or mixed_blank:
        f["realised"] = Num(None, UNKNOWN, "unclassified historical product; matched candidate part only", partial=f["matched"])
        f["charges"] = Num(None, UNKNOWN, "unclassified historical product; candidate fees only", partial=f["charges"].value if f["charges"].known else f["charges"].partial)
    if any(abs(sum(x[1] for x in f["lots"].get(sym,[]))-r["qty"]) > 1e-9 for sym,r in confirmed.items()):
        history_gap.append("confirmed current stocks have missing/inconsistent historical lots")
    if history_gap:
        f["realised"] = Num(None, UNKNOWN, "historical holdings/sales not reconciled", partial=f["matched"])
    if guessed:
        for r in rec:
            if r["Status"] != UNKNOWN and r["Status"] != "NOT MTF":
                r["Status"] = ESTIMATED
                r["Note"] += "; product blank in Dhan history, taken as MTF"
    for r in rec:
        if any(x["Symbol"]==r["Symbol"] for x in cross_matches) and r["Status"] != UNKNOWN:
            r["Status"]=ESTIMATED
            r["Note"]+="; includes cross-product economic FIFO matches"
    # 3. ledger
    led = ledger_rows([])
    try:
        if args.ledger:
            validate_statement_account(os.path.expanduser(args.ledger),getattr(args,"account_key",""))
            led = read_ledger_file(os.path.expanduser(args.ledger))
            src["ledger"] = (VERIFIED, "file " + os.path.basename(args.ledger))
        else:
            led = dhan_ledger(sess, frm, today)
            src["ledger"] = (VERIFIED, "Dhan ledger API %s -> %s" % (frm,
                                                                     today))
    except Exception as e:
        src["ledger"] = (UNKNOWN, "%s: %s" % (type(e).__name__, safe_error(e, sess)))
    if src["ledger"][0] != UNKNOWN and len(led):
        invalid_range = led["date"].map(lambda d: d is None or pd.isna(d) or not frm <= d <= today)
        if led["bad"].any() or invalid_range.any():
            src["ledger"] = (UNKNOWN, "statement has unreadable/missing amounts or invalid/out-of-range dates; verify full coverage")
    led_ok = src["ledger"][0] != UNKNOWN
    # ledger still charging MTF interest lately but no open MTF lot found
    # -> the inventory is incomplete, never 'no MTF position'
    if led_ok and not use and not fallback and not gap:
        ends = [p[1] for p in interest_periods(led) if p[1]]
        if ends and (today - max(ends)).days <= 10:
            gap.append("(ledger)")
            rec.append({"Symbol": "?", "MTF qty (history)": 0,
                        "Demat qty": None, "Status": UNKNOWN,
                        "Note": "MTF interest charged till %s but no open "
                                "MTF lot found in the trade history"
                                % max(ends)})
    ch = charges(led) if led_ok else {}
    paid = ch.get("MTF interest", Num(0.0, VERIFIED, "no MTF interest rows")) \
        if led_ok else Num(None, UNKNOWN, "ledger not read")
    avg_bal, per_day, per_txt, per_end = avg_interest_balance(led, args.rate) \
        if led_ok else (None, None, "", None)
    # 4. prices + open table
    price_symbols = sorted(set(use) | set(fallback))
    px = quotes(sess, price_symbols)
    missing_px = [s for s in price_symbols if s not in px]
    src["prices"] = (UNKNOWN if missing_px else ESTIMATED if any(p[1] == ESTIMATED for p in px.values()) else VERIFIED,
                     "missing quotes: " + ", ".join(missing_px) if missing_px else "quotes with per-stock provenance")
    open_cost = sum(x[1] * x[2] for L in use.values() for x in L)
    open_cost += sum(h.get("avg_price",0)*q for h,q,n,valid in fallback.values() if valid)
    loan = Num(args.loan, VERIFIED, "typed (--loan)") if args.loan is not None \
        else statement_loan(led, today) if led_ok else Num(None, UNKNOWN, "ledger missing / --loan required")
    share = (loan.value / open_cost) if loan.known and open_cost > 0 and \
        not gap else None
    if share is not None and share > 1:
        share = None
        loan = Num(None, UNKNOWN, "--loan is more than the open MTF cost")
    status_map = {r["Symbol"]: r["Status"] for r in rec}
    rows = open_table(use, px, today, share, args.rate, statuses=status_map,
                      bad_fee_symbols=f["bad_fee_symbols"])
    rows.extend(holding_row(sym, h, q, px, note, valid)
                for sym,(h,q,note,valid) in sorted(fallback.items()))
    rows.sort(key=lambda r:r["Symbol"])
    # 5. totals per scope
    hole = [Num(None, UNKNOWN, "unreconciled: " + ", ".join(gap))] if gap else []
    open_inventory_ok = demat_ok and current_ok and (history_ok or current_authoritative)
    if not open_inventory_ok:
        hole.append(Num(None, UNKNOWN, "essential inventory/history source failed"))
    inventory_status = UNKNOWN if hole else ESTIMATED if any(r.get("Inventory status") == ESTIMATED for r in rows) else VERIFIED
    open_cost_num = total([r["Open cost (Rs)"] for r in rows] + hole, "current holdings / reconciled lot cost")
    val = total([r["Value (Rs)"] for r in rows] + hole)
    gross = total([r["Gross P&L (Rs)"] for r in rows] + hole)
    fees = total([r["Sell fees est. (Rs)"] for r in rows] + hole)
    intr_open = total([r["Interest est. (Rs)"] for r in rows] + hole)
    buy_chg = total([r["Buy charges (Rs)"] for r in rows] + hole)
    if args.unpaid_interest is not None:
        unpaid = Num(args.unpaid_interest, VERIFIED, "typed")
    elif open_inventory_ok and not rows and not gap:
        unpaid = Num(0.0, VERIFIED, "no open MTF position")
    elif per_day is not None and per_end is not None:
        days = max(0, (today - per_end).days)   # accrued through valuation only; forward reserve is separate
        unpaid = Num(round(per_day * days, 2), ESTIMATED,
                     "%d as-of days x last Rs %.2f/day after %s; excludes future exit reserve; actual billing may differ" % (days, per_day, per_end))
    else:
        unpaid = Num(None, UNKNOWN, "no interest rate history")
    own = total([open_cost_num, -loan], "open cost - current loan; principal repayments are not expenses")
    exit_cash = total([val, -loan, -unpaid, -fees],
                      "value - loan - unpaid interest - sell fees")
    open_pnl = total([gross, -intr_open, -fees, -buy_chg],
                     "open lots: price - lot interest est. - sell fees - "
                     "buy charges")
    period = total([f["realised"], gross, -paid, -unpaid, -f["charges"],
                    -fees], "realised + open price - ALL MTF interest - "
                   "unpaid est. - MTF trade fees - sell fees")
    init_cash = Num(args.own_cash, VERIFIED, "typed (--own-cash)") \
        if args.own_cash is not None else Num(None, UNKNOWN,
                                              "needs --own-cash")
    roi = Num(round(100 * open_pnl.value / init_cash.value, 2),
              open_pnl.status, "% of initial own cash") \
        if open_pnl.known and init_cash.known and init_cash.value > 0 \
        else Num(None, UNKNOWN, "needs open P&L + --own-cash")
    unalloc = {k: v for k, v in ch.items() if k != "MTF interest"}
    unresolved_items = any(not v.known or abs(v.value) > 1e-9 for v in unalloc.values())
    unresolved_costs = any(k != "Dividend (income)" and (not v.known or abs(v.value) > 1e-9)
                           for k, v in unalloc.items())
    complete_net = Num(None, UNKNOWN, "unallocated holding costs/income need position mapping",
                       partial=period.value if period.known else period.partial) if unresolved_items else period
    if unresolved_items:
        period.note = "PARTIAL SUBTOTAL before unallocated holding costs/income; not comprehensive net"
        open_pnl.note = "PARTIAL model subtotal before unallocated holding costs/income"
        roi = Num(None, UNKNOWN, "comprehensive own-cash ROI needs cost/income allocation")
    if not led_ok:
        exit_cash = Num(None, UNKNOWN, "ledger unavailable; unpaid settlement obligations unknown",
                        partial=exit_cash.value if exit_cash.known else exit_cash.partial)
    if exit_cash.known:
        exit_cash.note += "; ESTIMATE: historical debited charges already paid, not deducted twice; future order/DP/unpledge charges estimated"
    add, wd = money_in_out(led) if led_ok else (None, None)
    essential_ok = all(src.get(k, (UNKNOWN,))[0] != UNKNOWN
                       for k in ("trades", "trade_fees", "ledger", "demat", "positions", "prices")) and not gap and not history_gap
    essential_ok = essential_ok and all(n.known for n in (open_cost_num, val, gross, loan, unpaid, exit_cash, complete_net))
    return {"src": src, "prods": prods, "guessed": guessed, "mtf": mtf,
            "fifo": f, "rec": rec, "rows": rows, "led": led, "ch": ch,
            "paid": paid, "avg": (avg_bal, per_day, per_txt), "loan": loan,
            "own": own, "val": val, "gross": gross, "fees": fees,
            "intr_open": intr_open, "unpaid": unpaid, "exit": exit_cash,
            "open_pnl": open_pnl, "period": period, "init_cash": init_cash,
            "roi": roi, "unalloc": unalloc, "money": (add, wd),
            "closing": closing_balance(led) if led_ok else None,
            "open_cost": open_cost, "open_cost_num": open_cost_num,
            "complete_net": complete_net, "ok": essential_ok, "confirmed": confirmed, "cross_matches":cross_matches, "current_lots":use}


def drive_mtf_copy(path, acc, any_path=False):
    """RB 9 Oct: MTF report also in Google Sheets. Copies ONLY MTF_Check.xlsx into the
    Google Drive for desktop folder My Drive/RB_Reports/<BROKER>_<Name>/ (same folder as
    the Portfolio file). No Drive app / any error -> local file is untouched."""
    try:
        import shutil, drive_copy, account as ac
        # only the real account report (tests / temp copies never reach the real Drive)
        accounts = os.path.realpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "accounts"))
        if not any_path and not os.path.realpath(path).startswith(accounts + os.sep):
            return None
        root = drive_copy.drive_root()
        if not root or not os.path.isfile(path):
            print("  Google Drive copy: Drive for desktop folder nahi mila -- local file hi hai.")
            return None
        tag = ac.file_tag(acc)
        dst = os.path.join(root, drive_copy.TOP, tag, "MTF_Check_%s.xlsx" % tag)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        shutil.copyfile(path, dst)
        print("  Google Drive copy: My Drive/%s" % os.path.relpath(dst, root))
        return dst
    except Exception as e:
        print("  ! Google Drive copy failed (%s) -- local file is fine." % type(e).__name__)
        return None


def main_generic(a, acc, sess, today):
    """ANGEL / ZERODHA: no trade-history / ledger API -> mtf_generic (RB 9 Oct)."""
    from mtf_generic import build_generic
    from mtf_portfolio_report import write_report, print_report, complete
    if a.audit or a.ledger or a.trades or getattr(a, "confirm_mtf", None):
        print("--audit / --ledger / --trades / --confirm-mtf are Dhan-only; for %s use --mtf, --buy-date, --loan, --unpaid-interest." % sess.broker)
        return 1
    try:
        R = build_generic(a, sess, today, getattr(acc, "data", None))
    except ValueError as e:
        print("MTF input error: " + safe_error(e, sess)); return 1
    R = sanitize_report(R, sess)
    R["report_account"] = redact_text(acc.label, sess)
    path = os.path.join(acc.reports, "MTF_Check.xlsx")
    out = write_report(R, today, path, days=a.buffer_days, settlement=a.settlement_buffer, tax=a.tax_reserve, rate=a.rate)
    print_report(R, out, acc, today, a)
    print("\n%s: no trade-history / ledger API -> rate %.2f%% (broker card), loan / buy dates from you." % (sess.broker, a.rate * 100))
    if not R["rows"] and R["src"].get("holdings", ("",))[0] != UNKNOWN:
        print("  Is account mein koi MTF stock nahi mila. Agar hai to: rbmtf --mtf SYMBOL:QTY")
    for k, (st, note) in sorted(R["src"].items()):
        if st == UNKNOWN:
            print("  MISSING %s: %s" % (k, note))
    if not R["loan"].known:
        print("  -> " + R["loan"].note)
    print("\nExcel (same account file refreshed):", path)
    drive_mtf_copy(path, acc)
    print("READ-ONLY broker access: koi order nahi gaya. Missing inputs remain UNKNOWN.")
    return 0 if complete(R, out) else 2


def sanitize_report(value,sess):
    """Redact external text in terminal/audit/default report, without altering numbers."""
    if isinstance(value,Num):return Num(value.value,value.status,redact_text(value.note,sess),value.partial)
    if isinstance(value,str):return redact_text(value,sess)
    if isinstance(value,pd.DataFrame):
        cellwise=getattr(value,"map",None) or value.applymap   # pandas 3 removed applymap
        return cellwise(lambda x:redact_text(x,sess) if isinstance(x,str) else x)
    if isinstance(value,dict):return {sanitize_report(k,sess):sanitize_report(v,sess) for k,v in value.items()}
    if isinstance(value,list):return [sanitize_report(x,sess) for x in value]
    if isinstance(value,tuple):return tuple(sanitize_report(x,sess) for x in value)
    return value


# ================================================================== main
def line(label, n, note=""):
    st = n.status if isinstance(n, Num) else ""
    nt = note or (n.note if isinstance(n, Num) else "")
    shown = ("%.2f%%" % n.value if n.known else rs(n)) if \
        label.endswith("%") and isinstance(n, Num) else rs(n)
    print("%-44s %26s  %-9s %s" % (label, shown, st, nt))
    return {"Item": label.strip(), "Rs": money_cell(n) if isinstance(n, Num)
            else n, "Status": st, "Note": nt, "Known part (Rs)": n.partial if isinstance(n, Num) and not n.known else None}


def print_current_report(R, rows, acc, today, args):
    """Numbered human-readable summary; unknown inputs never become zero."""
    def total_field(key):
        vals = [r.get(key) for r in rows]
        return sum(vals) if vals and all(_finite(v) for v in vals) else None
    def amount(v, signed=False):
        if not _finite(v): return "UNKNOWN"
        return ("+Rs {:,.2f}" if signed and v>0 else "Rs {:,.2f}").format(v)
    def item(label, value, note="", signed=False):
        print("  %-47s %20s%s" % (label,amount(value,signed),"  "+note if note else ""))
    def num_value(key):
        n=R.get(key)
        return n.value if isinstance(n,Num) and n.known else None
    print("\n==== MTF HISAAB | %s | %s ====" % (acc.label,today))
    print("Current holdings only. Interest/fees ESTIMATED; source prices shown below.")
    print("\n1. ABHI KE MTF STOCKS")
    print("  %-13s %5s %12s %16s %16s %18s" % ("Stock","Qty","Price","Price P/L","P/L + interest","Net P/L now"))
    for r in rows:
        print("  %-13s %5g %12s %16s %16s %18s" % (r['symbol'],r['qty'],amount(r['cmp']),amount(r['price_pnl'],True),amount(r['with_interest'],True),amount(r['now_net_tax'],True)))
        print("    Price source: %s" % r.get('price_source','UNKNOWN'))
    if not rows: print("  Current MTF inventory reconcile nahi hua; amounts UNKNOWN.")
    print("\n2. CAPITAL KA BREAKUP")
    cost=total_field('cost');loan=num_value('loan')
    item("Abhi ke stocks ki purchase cost",cost)
    item("Broker ka CURRENT funded amount",loan,"not inferred from old daily interest")
    item("Purchase cost ka unfunded hissa (model)",cost-loan if cost is not None and loan is not None and 0<=loan<=cost else None,"not verified lifetime own-cash contribution")
    print("\n3. CURRENT VALUE AUR INTEREST")
    item("Quote/snapshot par current value",total_field('now_value'))
    item("Price profit/loss",total_field('price_pnl'),signed=True)
    item("In CURRENT lots ka interest till today (model)",total_field('past'))
    item("Interest jodne ke baad profit/loss",total_field('with_interest'),signed=True)
    item("Current loan par daily interest (model)",loan*args.rate/365 if loan is not None else None,"rate %.2f%% p.a." % (args.rate*100))
    print("\n4. AAJ KE PRICE PAR BECHO TOH (ESTIMATE)")
    item("Sale proceeds",total_field('now_value'))
    item("Sell brokerage/taxes/DP/unpledge estimate",total_field('now_sale_fee'))
    item("Broker loan repayment",loan)
    item("Interest jo abhi debit hona baaki hai",num_value('unpaid'),"already paid interest is not deducted twice")
    item("Settlement mein cash released (estimate)",num_value('exit'),R.get('exit').note if isinstance(R.get('exit'),Num) and not R['exit'].known else "loan, unpaid interest and modeled sell charges deducted")
    if rows:
        print("  Agar SIRF ek stock becho (loan/unpaid interest stock-wise = cost ratio, MODEL):")
        print("  %-13s %15s %13s %15s %13s %17s" % ("Stock","Value","Sell chg","Loan wapas","Unpaid int","HAATH MEIN"))
        for r in rows:
            print("  %-13s %15s %13s %15s %13s %17s" % (r['symbol'],amount(r.get('now_value')),amount(r.get('now_sale_fee')),amount(r.get('loan_share')),amount(r.get('unpaid_share')),amount(r.get('cash_now'))))
    print("\n5. CURRENT NET PROFIT / LOSS")
    item("Buy brokerage/taxes/pledge estimate",total_field('buy'))
    item("Optional positive-gain tax reserve",total_field('now_tax'),"reserve, not confirmed personal tax")
    item("NET P/L: interest + modeled buy/sell costs",total_field('now_net_tax'),"ESTIMATED",signed=True)
    print("  Future holding/settlement interest is included in the target only.")
    print("\n6. %d-DAY BREAKEVEN + %d EXTRA INTEREST DAYS" % (args.buffer_days,args.settlement_buffer))
    print("  %-13s %20s %22s" % ("Stock","Broker-cost BE/share","BE + tax reserve/share"))
    for r in rows: print("  %-13s %20s %22s" % (r['symbol'],amount(r['broker']),amount(r['full'])))
    if any(r['now_net_tax'] is None or r['status']==UNKNOWN for r in rows) or not rows or (isinstance(R.get('exit'),Num) and not R['exit'].known):
        print("\nINCOMPLETE: see missing-input details below. UNKNOWN is not zero; known loan/price does not repair history.")
        print('Missing-input detail:')
        for key,(status,note) in sorted(R.get('src',{}).items()):
            if status==UNKNOWN:print('  %s: %s'%(key,note))
        for row in R.get('rows',[]):
            value=row.get('Interest est. (Rs)')
            if isinstance(value,Num) and not value.known:
                print('  %s interest: %s'%(row.get('Symbol','?'),value.note))
        for row in R.get('rec',[]):
            if row.get('Status')==UNKNOWN:
                print('  %s reconciliation: %s'%(row.get('Symbol','?'),row.get('Note','UNKNOWN')))
    print("\nSources: active account broker APIs, or explicitly supplied account-matched statements. Current funded-share interest is a MODEL, not stock-wise billed interest.")
    print("Detailed reconciliation: rbmtf --audit")


def main():
    ap = argparse.ArgumentParser(description="MTF hisaab with status")
    ap.add_argument("--from", dest="frm", default="2024-04-01",
                    help="start: before your FIRST MTF buy (YYYY-MM-DD)")
    ap.add_argument("--ledger", help="downloaded Dhan ledger (xlsx/csv; detected by file content)")
    ap.add_argument("--trades", help="downloaded Dhan trade history CSV; explicit MTF products")
    ap.add_argument("--instruments", help="optional local Dhan scrip master CSV for exact company-name mapping")
    ap.add_argument("--rate", type=float, default=None,
                    help="MTF interest per year as a fraction; default = the broker's published rate (Dhan 0.1249)")
    ap.add_argument("--loan", type=float, help="current MTF funded amount "
                    "from the Dhan app (Rs)")
    ap.add_argument("--unpaid-interest", type=float,
                    help="interest accrued but not yet debited (Rs)")
    ap.add_argument("--own-cash", type=float,
                    help="your own cash put into the open MTF lots (Rs)")
    ap.add_argument("--confirm-mtf", help="CURRENT stocks confirmed MTF: TCS,PERSISTENT or TCS:52; remembers this account qty/cost for 7 days")
    ap.add_argument("--mtf", help="ANGEL/ZERODHA: MTF stocks if the API does not tag them: TCS:52 or TCS:52@3082.5")
    ap.add_argument("--buy-date", help="ANGEL/ZERODHA: first buy date per MTF stock, remembered: TCS:2025-11-24,INFY:2026-02-04")
    ap.add_argument("--audit", action="store_true", help="optional detailed multi-sheet reconciliation")
    ap.add_argument("--buffer-days", type=int, default=30, help="future calendar holding days (default 30)")
    ap.add_argument("--settlement-buffer", type=int, default=3, help="extra calendar interest days after exit (scenario reserve)")
    ap.add_argument("--tax-reserve", type=float, default=.208, help="optional conservative reserve fraction on gross gain; 0 disables it")
    a = ap.parse_args()
    if not (0 <= a.buffer_days <= 3650 and 0 <= a.settlement_buffer <= 90 and _finite(a.tax_reserve) and 0 <= a.tax_reserve < .8):
        print("Invalid breakeven scenario: buffer 0..3650, settlement 0..90, tax reserve fraction 0..<0.8")
        return 1
    try:
        a.confirm_requests = parse_confirmations(a.confirm_mtf)
    except ValueError as e:
        print(str(e)); return 1
    rate_given = a.rate is not None
    if a.rate is None:
        a.rate = MTF_RATE        # replaced by the active broker's card below
    if 1 < a.rate <= 100:
        print("  --rate %.2f read as %.2f%% -> %.4f" % (a.rate, a.rate,
                                                       a.rate / 100))
        a.rate = a.rate / 100
    if not (_finite(a.rate) and 0 <= a.rate < 1):
        print("--rate must be a fraction like 0.1249")
        return 1
    for k in ("loan", "unpaid_interest", "own_cash"):
        v = getattr(a, k)
        if v is not None and (not _finite(v) or v < 0):
            print("--%s must be a number >= 0" % k.replace("_", "-"))
            return 1
    today = dt.date.today()
    frm = parse_date(a.frm)
    if not frm or frm > today:
        print("--from must be a past date YYYY-MM-DD")
        return 1
    import account
    acc = account.activate()
    if not acc.token_ok:
        print("Token expire hai -- token.txt mein naya token daalo.")
        return 1
    sess = acc.session
    import mtf_portfolio_report as mpr
    try:
        card_rate = mpr.set_broker(sess.broker)       # broker's own brokerage/DP/pledge/rate
    except ValueError:
        print("MTF rate card nahi hai is broker ke liye: %s" % sess.broker)
        return 1
    if not rate_given:
        a.rate = card_rate
    if sess.broker != "DHAN":
        try:
            return main_generic(a, acc, sess, today)
        finally:
            mpr.set_broker("DHAN")       # never leave another broker's card loaded
    a.account_key = getattr(acc, "key", acc.label)
    a.confirmation_path = os.path.join(acc.data, "mtf_product_confirmation.json") if getattr(acc, "data", None) else None
    try:
        R = build(a, sess, today, frm)
    except ValueError as e:
        print("MTF input error: " + safe_error(e, sess)); return 1
    R=sanitize_report(R,sess)
    if not a.audit:
        from mtf_portfolio_report import write_report, print_report, complete
        path = os.path.join(acc.reports, "MTF_Check.xlsx")
        R["report_account"] = redact_text(acc.label,sess)
        out = write_report(R, today, path, days=a.buffer_days, settlement=a.settlement_buffer, tax=a.tax_reserve, rate=a.rate)
        print_report(R, out, acc, today, a)
        print("\nExcel (same account file refreshed):", path)
        drive_mtf_copy(path, acc)
        print("READ-ONLY broker access: koi order nahi gaya. Missing inputs remain UNKNOWN.")
        return 0 if complete(R, out) else 2
    f = R["fifo"]

    print("\n==== MTF HISAAB  %s  %s  %s ====" % (
        acc.label, today, "" if R["ok"] else "** INCOMPLETE **"))
    for k, (st, note) in sorted(R["src"].items()):
        print("  source %-7s %-9s %s" % (k, st, note))
    print("  trades by product: %s%s" % (", ".join(
        "%s %d" % kv for kv in sorted(R["prods"].items())) or "none",
        "  -> blank-product rows are UNCLASSIFIED candidates; not confirmed MTF" if R["guessed"] else ""))
    if R["rows"]:
        t = flatten(R["rows"])
        print(t[["Symbol", "Qty", "Avg cost", "First buy", "Last buy",
                 "Price", "Price source", "Open cost (Rs)", "Value (Rs)",
                 "Gross P&L (Rs)", "Gross P&L %", "Interest est. (Rs)"]]
              .to_string(index=False))
    else:
        print("  (koi open MTF stock reconcile nahi hua -- Reconciliation "
              "dekho)")
    for r in R["rec"]:
        if r["Status"] != VERIFIED:
            print("  %-12s %-9s %s" % (r["Symbol"], r["Status"], r["Note"]))

    out_open, out_exit, out_period, out_acc = [], [], [], []
    print("\n-- OPEN MTF POSITIONS --")
    out_open.append(line("Open cost (lots)", R["open_cost_num"]))
    out_open.append(line("Value now", R["val"]))
    out_open.append(line("Gross price P&L", R["gross"]))
    out_open.append(line("Interest est. on these lots", -R["intr_open"]))
    out_open.append(line("Open P&L model subtotal", R["open_pnl"]))
    out_open.append(line("Own cash put in", R["init_cash"]))
    out_open.append(line("Return on initial own cash %", R["roi"],
                         "not annual; top-ups change it"))
    print("\n-- AGAR AAJ SAB BECHO (hypothetical, settlement nahi) --")
    out_exit.append(line("Value now", R["val"]))
    out_exit.append(line("Current MTF loan (Dhan ka paisa)", -R["loan"]))
    out_exit.append(line("Unpaid accrued interest", -R["unpaid"]))
    out_exit.append(line("Sell fees", -R["fees"]))
    out_exit.append(line("= Cash released", R["exit"]))
    out_exit.append(line("Own money in open lots", R["own"]))
    avg_bal, per_day, per_txt = R["avg"]
    print("\n-- MTF PERIOD (%s se aaj tak) --" % a.frm)
    out_period.append(line("MTF bought (history)", f["bought_num"]))
    out_period.append(line("MTF sold (history)", f["sold_num"], "includes FIFO-matched DELIVERY exits of MTF acquisitions" if R["cross_matches"] else ""))
    out_period.append(line("Realised price P&L", f["realised"],
                           "matched part %s" % rs(f["matched"])))
    out_period.append(line("Open price P&L", R["gross"]))
    out_period.append(line("MTF interest debited (ALL, incl. closed)",
                           -R["paid"]))
    out_period.append(line("Unpaid accrued interest", -R["unpaid"]))
    out_period.append(line("MTF trade fees (history)", -f["charges"]))
    out_period.append(line("Sell fees (open lots)", -R["fees"]))
    out_period.append(line("= MTF period subtotal", R["period"]))
    out_period.append(line("Comprehensive MTF net P&L", R["complete_net"]))
    if per_day is not None:
        print("  diagnostic: last interest %s = Rs %.2f/day -> AVERAGE "
              "interest-bearing balance ~%s at %.2f%% (NOT today's loan)"
              % (per_txt, per_day, rs(avg_bal), 100 * a.rate))
    print("\n-- ACCOUNT-WIDE / UNALLOCATED (MTF result mein NAHI jode) --")
    for k, v in sorted(R["unalloc"].items()):
        out_acc.append(line(k, -v))
    add, wd = R["money"]
    if add is not None:
        out_acc.append(line("Money added (ledger rows)", Num(add)))
        out_acc.append(line("Money withdrawn (ledger rows)", Num(-wd)))
    if R["closing"] is not None:
        out_acc.append(line("Ledger closing balance", Num(R["closing"])))
    need = []
    if not R["loan"].known:
        need.append("--loan <Dhan app ka MTF funded amount>")
    if not R["init_cash"].known:
        need.append("--own-cash <apna paisa jo in stocks mein laga>")
    if f["uncovered"]:
        need.append("--from <pehle MTF buy se pehle ki date; conversions also need statement>")
    if need:
        print("\nOptional missing inputs: rbmtf " + " ".join(need))
    if any(r.get("Status") == UNKNOWN for r in R["rec"]):
        print("History/product gaps require broker MTF stock statement, conversion records and complete trades; --loan alone does not resolve them.")
    if R["confirmed"]:
        print("MTF product is USER_CONFIRMED for current quantities. Re-confirm after conversion/quantity/cost change or after 7 days. Missing purchase dates were not invented.")

    os.makedirs(acc.reports, exist_ok=True)
    out = os.path.join(acc.reports, "MTF_Audit.xlsx" if a.audit else "MTF_Check.xlsx")
    from mtf_breakeven import atomic_excel_writer
    with atomic_excel_writer(out) as w:
        pd.DataFrame([
            {"Item":"Report as of", "Value":today.isoformat(),"Status":"AS_OF"},
            {"Item":"Current MTF holdings", "Value":", ".join(r["Symbol"] for r in R["rows"]) or "unreconciled", "Status":R["val"].status},
            {"Item":"Current price P&L (Rs)", "Value":money_cell(R["gross"]),"Status":R["gross"].status},
            {"Item":"Cash released estimate (Rs)","Value":money_cell(R["exit"]),"Status":R["exit"].status},
            {"Item":"Historical complete net (Rs)","Value":money_cell(R["complete_net"]),"Status":R["complete_net"].status},
            {"Item":"Read this first", "Value":"Current holdings and old trade history are separate. Known part is not final P&L. Cash released is a hypothetical estimate, not settled cash.","Status":"INFO"}
        ]).to_excel(w,sheet_name="MTF_Status",index=False)
        pd.DataFrame(out_open).to_excel(w, sheet_name="MTF_Open_Summary",
                                        index=False)
        flatten(R["rows"]).to_excel(w, sheet_name="MTF_Open", index=False)
        pd.DataFrame(out_exit).to_excel(w, sheet_name="MTF_Exit_Estimate",
                                        index=False)
        pd.DataFrame(out_period).to_excel(w, sheet_name="MTF_Period",
                                          index=False)
        src = pd.DataFrame([{"Symbol": "source " + k, "Status": v[0],
                             "Note": v[1]} for k, v in sorted(R["src"].items())])
        pd.concat([src, pd.DataFrame(R["rec"]), pd.DataFrame(out_acc)],
                  ignore_index=True).to_excel(w, sheet_name="Reconciliation",
                                              index=False)
        pd.DataFrame([{k: (v.isoformat(sep=" ") if k == "ts" and v else v)
                       for k, v in t.items()}
                      for t in R["mtf"]]).to_excel(w, sheet_name="MTF_Trades",
                                                   index=False)
        if R["cross_matches"]:
            pd.DataFrame(R["cross_matches"]).to_excel(w,sheet_name="CrossProduct_Matches",index=False)
        if len(R["led"]):
            R["led"].to_excel(w, sheet_name="Ledger", index=False)
        from openpyxl.styles import Font, PatternFill, Alignment
        for sheet in w.book.worksheets:
            sheet.freeze_panes="A2"
            sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]:
                cell.font=Font(color="FFFFFF",bold=True)
                cell.fill=PatternFill("solid",fgColor="243B53")
            for col in sheet.columns:
                letter=col[0].column_letter
                sheet.column_dimensions[letter].width=min(58,max(14,max(len(str(c.value or "")) for c in col[:100])+2))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,(int,float)):cell.number_format="#,##0.00;[Red]-#,##0.00"
                    if cell.value in (UNKNOWN,ESTIMATED,VERIFIED) if isinstance(cell.value,str) else False:
                        cell.fill=PatternFill("solid",fgColor={UNKNOWN:"FFE0E0",ESTIMATED:"FFF0C2",VERIFIED:"D9EAD3"}[cell.value])
                    if isinstance(cell.value,str) and len(cell.value)>50:cell.alignment=Alignment(wrap_text=True,vertical="top")
        # This report contains computed values, never spreadsheet formulas.
        # Preserve literal labels/narrations beginning '=' instead of Excel evaluating them.
        for sheet in w.book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.data_type = "s"
    print("\nExcel: %s" % out)
    print("READ-ONLY: koi order nahi gaya. Exact settlement sirf asli sell + "
          "contract note + ledger se.")
    return 0 if R["ok"] else 2


if __name__ == "__main__":
    sys.exit(main())
