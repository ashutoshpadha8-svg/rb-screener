"""Per-account, per-symbol broker MTF snapshots. No guessed rate or orders."""
import math
from copy import copy

HEADS = ('MTF x (broker)', 'MTF updated IST', 'MTF status (broker)')


def refresh(sess, prices, adapter, now):
    """New snapshot each report run; failures never preserve yesterday's rate."""
    out = {}
    for symbol, price in sorted(prices.items()):
        symbol = str(symbol).split(' | ')[0].upper().strip()
        if not symbol: continue
        rate, note = None, 'No broker session'
        if sess is not None:
            try:
                rate, note = adapter.mtf_leverage(sess, symbol, price)
                if rate is not None and (not math.isfinite(float(rate)) or
                                         not (rate == 0 or 1 < rate <= 10)):
                    rate, note = None, 'Invalid broker leverage response'
            except Exception as error:
                rate, note = None, 'Margin check failed: '+type(error).__name__
        out[symbol] = dict(leverage=rate, asof=now,
                           note=note if rate is not None else 'UNKNOWN: '+note)
    return out


def display(snapshot):
    rate = snapshot.get('leverage')
    return 'UNKNOWN' if rate is None else '%.2fx' % rate


def add_columns(wb, snapshots):
    """Append display columns; never move an editable input or formula column."""
    from openpyxl.utils import get_column_letter
    for ws in list(wb.worksheets):
        if ws.title in ('Dashboard','Holdings','Buy_Planner','Journal','Symbols','Signal_Summary','MTF_Rates'):
            continue
        header = None
        for r in range(1, min(ws.max_row, 8)+1):
            values = [ws.cell(r,c).value for c in range(1,ws.max_column+1)]
            for name in ('Symbol','Ticker'):
                if name in values:
                    header = r,values.index(name)+1
                    break
            if header: break
        if not header: continue
        r, sc = header
        heads = [ws.cell(r,c).value for c in range(1,ws.max_column+1)]
        cols = []
        for name in HEADS:
            if name in heads: c = heads.index(name)+1
            else:
                c = len(heads)+1;heads.append(name)
            cols.append(c)
            cell=ws.cell(r,c,name);src=ws.cell(r,sc)
            cell.font,cell.fill=copy(src.font),copy(src.fill)
            cell.alignment=copy(src.alignment)
            ws.column_dimensions[get_column_letter(c)].width=18 if name==HEADS[0] else 25 if name==HEADS[1] else 45
        for row in range(r+1,ws.max_row+1):
            raw=ws.cell(row,sc).value
            sym=str(raw or '').split(' | ')[0].strip().upper()
            if not sym: continue
            if sym not in snapshots:
                continue
            snap=snapshots[sym]
            for c,value in zip(cols,(snap.get('leverage') if snap.get('leverage') is not None else 'UNKNOWN',snap.get('asof'),snap.get('note'))):
                ws.cell(row,c,value)
            ws.cell(row,cols[0]).number_format='0.00"x"'
    if 'MTF_Rates' in wb: del wb['MTF_Rates']
    ws=wb.create_sheet('MTF_Rates')
    ws.append(('Symbol',)+HEADS)
    for symbol,snap in sorted(snapshots.items()):
        ws.append((symbol,snap['leverage'] if snap['leverage'] is not None else 'UNKNOWN',snap['asof'],snap['note']))
        ws.cell(ws.max_row,2).number_format='0.00"x"'
    for col,width in zip('ABCD',(15,18,26,65)):ws.column_dimensions[col].width=width
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
